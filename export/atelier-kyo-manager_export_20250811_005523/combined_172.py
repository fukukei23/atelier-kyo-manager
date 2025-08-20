
# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pandas\tests\arithmetic\test_datetime64.py ===
# Arithmetic tests for DataFrame/Series/Index/Array classes that should
# behave identically.
# Specifically for datetime64 and datetime64tz dtypes
from datetime import (
    datetime,
    time,
    timedelta,
)
from itertools import (
    product,
    starmap,
)
import operator

import numpy as np
import pytest
import pytz

from pandas._libs.tslibs.conversion import localize_pydatetime
from pandas._libs.tslibs.offsets import shift_months
from pandas.errors import PerformanceWarning

import pandas as pd
from pandas import (
    DateOffset,
    DatetimeIndex,
    NaT,
    Period,
    Series,
    Timedelta,
    TimedeltaIndex,
    Timestamp,
    date_range,
)
import pandas._testing as tm
from pandas.core import roperator
from pandas.tests.arithmetic.common import (
    assert_cannot_add,
    assert_invalid_addsub_type,
    assert_invalid_comparison,
    get_upcast_box,
)

# ------------------------------------------------------------------
# Comparisons


class TestDatetime64ArrayLikeComparisons:
    # Comparison tests for datetime64 vectors fully parametrized over
    #  DataFrame/Series/DatetimeIndex/DatetimeArray.  Ideally all comparison
    #  tests will eventually end up here.

    def test_compare_zerodim(self, tz_naive_fixture, box_with_array):
        # Test comparison with zero-dimensional array is unboxed
        tz = tz_naive_fixture
        box = box_with_array
        dti = date_range("20130101", periods=3, tz=tz)

        other = np.array(dti.to_numpy()[0])

        dtarr = tm.box_expected(dti, box)
        xbox = get_upcast_box(dtarr, other, True)
        result = dtarr <= other
        expected = np.array([True, False, False])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(result, expected)

    @pytest.mark.parametrize(
        "other",
        [
            "foo",
            -1,
            99,
            4.0,
            object(),
            timedelta(days=2),
            # GH#19800, GH#19301 datetime.date comparison raises to
            #  match DatetimeIndex/Timestamp.  This also matches the behavior
            #  of stdlib datetime.datetime
            datetime(2001, 1, 1).date(),
            # GH#19301 None and NaN are *not* cast to NaT for comparisons
            None,
            np.nan,
        ],
    )
    def test_dt64arr_cmp_scalar_invalid(self, other, tz_naive_fixture, box_with_array):
        # GH#22074, GH#15966
        tz = tz_naive_fixture

        rng = date_range("1/1/2000", periods=10, tz=tz)
        dtarr = tm.box_expected(rng, box_with_array)
        assert_invalid_comparison(dtarr, other, box_with_array)

    @pytest.mark.parametrize(
        "other",
        [
            # GH#4968 invalid date/int comparisons
            list(range(10)),
            np.arange(10),
            np.arange(10).astype(np.float32),
            np.arange(10).astype(object),
            pd.timedelta_range("1ns", periods=10).array,
            np.array(pd.timedelta_range("1ns", periods=10)),
            list(pd.timedelta_range("1ns", periods=10)),
            pd.timedelta_range("1 Day", periods=10).astype(object),
            pd.period_range("1971-01-01", freq="D", periods=10).array,
            pd.period_range("1971-01-01", freq="D", periods=10).astype(object),
        ],
    )
    def test_dt64arr_cmp_arraylike_invalid(
        self, other, tz_naive_fixture, box_with_array
    ):
        tz = tz_naive_fixture

        dta = date_range("1970-01-01", freq="ns", periods=10, tz=tz)._data
        obj = tm.box_expected(dta, box_with_array)
        assert_invalid_comparison(obj, other, box_with_array)

    def test_dt64arr_cmp_mixed_invalid(self, tz_naive_fixture):
        tz = tz_naive_fixture

        dta = date_range("1970-01-01", freq="h", periods=5, tz=tz)._data

        other = np.array([0, 1, 2, dta[3], Timedelta(days=1)])
        result = dta == other
        expected = np.array([False, False, False, True, False])
        tm.assert_numpy_array_equal(result, expected)

        result = dta != other
        tm.assert_numpy_array_equal(result, ~expected)

        msg = "Invalid comparison between|Cannot compare type|not supported between"
        with pytest.raises(TypeError, match=msg):
            dta < other
        with pytest.raises(TypeError, match=msg):
            dta > other
        with pytest.raises(TypeError, match=msg):
            dta <= other
        with pytest.raises(TypeError, match=msg):
            dta >= other

    def test_dt64arr_nat_comparison(self, tz_naive_fixture, box_with_array):
        # GH#22242, GH#22163 DataFrame considered NaT == ts incorrectly
        tz = tz_naive_fixture
        box = box_with_array

        ts = Timestamp("2021-01-01", tz=tz)
        ser = Series([ts, NaT])

        obj = tm.box_expected(ser, box)
        xbox = get_upcast_box(obj, ts, True)

        expected = Series([True, False], dtype=np.bool_)
        expected = tm.box_expected(expected, xbox)

        result = obj == ts
        tm.assert_equal(result, expected)


class TestDatetime64SeriesComparison:
    # TODO: moved from tests.series.test_operators; needs cleanup

    @pytest.mark.parametrize(
        "pair",
        [
            (
                [Timestamp("2011-01-01"), NaT, Timestamp("2011-01-03")],
                [NaT, NaT, Timestamp("2011-01-03")],
            ),
            (
                [Timedelta("1 days"), NaT, Timedelta("3 days")],
                [NaT, NaT, Timedelta("3 days")],
            ),
            (
                [Period("2011-01", freq="M"), NaT, Period("2011-03", freq="M")],
                [NaT, NaT, Period("2011-03", freq="M")],
            ),
        ],
    )
    @pytest.mark.parametrize("reverse", [True, False])
    @pytest.mark.parametrize("dtype", [None, object])
    @pytest.mark.parametrize(
        "op, expected",
        [
            (operator.eq, Series([False, False, True])),
            (operator.ne, Series([True, True, False])),
            (operator.lt, Series([False, False, False])),
            (operator.gt, Series([False, False, False])),
            (operator.ge, Series([False, False, True])),
            (operator.le, Series([False, False, True])),
        ],
    )
    def test_nat_comparisons(
        self,
        dtype,
        index_or_series,
        reverse,
        pair,
        op,
        expected,
    ):
        box = index_or_series
        lhs, rhs = pair
        if reverse:
            # add lhs / rhs switched data
            lhs, rhs = rhs, lhs

        left = Series(lhs, dtype=dtype)
        right = box(rhs, dtype=dtype)

        result = op(left, right)

        tm.assert_series_equal(result, expected)

    @pytest.mark.parametrize(
        "data",
        [
            [Timestamp("2011-01-01"), NaT, Timestamp("2011-01-03")],
            [Timedelta("1 days"), NaT, Timedelta("3 days")],
            [Period("2011-01", freq="M"), NaT, Period("2011-03", freq="M")],
        ],
    )
    @pytest.mark.parametrize("dtype", [None, object])
    def test_nat_comparisons_scalar(self, dtype, data, box_with_array):
        box = box_with_array

        left = Series(data, dtype=dtype)
        left = tm.box_expected(left, box)
        xbox = get_upcast_box(left, NaT, True)

        expected = [False, False, False]
        expected = tm.box_expected(expected, xbox)
        if box is pd.array and dtype is object:
            expected = pd.array(expected, dtype="bool")

        tm.assert_equal(left == NaT, expected)
        tm.assert_equal(NaT == left, expected)

        expected = [True, True, True]
        expected = tm.box_expected(expected, xbox)
        if box is pd.array and dtype is object:
            expected = pd.array(expected, dtype="bool")
        tm.assert_equal(left != NaT, expected)
        tm.assert_equal(NaT != left, expected)

        expected = [False, False, False]
        expected = tm.box_expected(expected, xbox)
        if box is pd.array and dtype is object:
            expected = pd.array(expected, dtype="bool")
        tm.assert_equal(left < NaT, expected)
        tm.assert_equal(NaT > left, expected)
        tm.assert_equal(left <= NaT, expected)
        tm.assert_equal(NaT >= left, expected)

        tm.assert_equal(left > NaT, expected)
        tm.assert_equal(NaT < left, expected)
        tm.assert_equal(left >= NaT, expected)
        tm.assert_equal(NaT <= left, expected)

    @pytest.mark.parametrize("val", [datetime(2000, 1, 4), datetime(2000, 1, 5)])
    def test_series_comparison_scalars(self, val):
        series = Series(date_range("1/1/2000", periods=10))

        result = series > val
        expected = Series([x > val for x in series])
        tm.assert_series_equal(result, expected)

    @pytest.mark.parametrize(
        "left,right", [("lt", "gt"), ("le", "ge"), ("eq", "eq"), ("ne", "ne")]
    )
    def test_timestamp_compare_series(self, left, right):
        # see gh-4982
        # Make sure we can compare Timestamps on the right AND left hand side.
        ser = Series(date_range("20010101", periods=10), name="dates")
        s_nat = ser.copy(deep=True)

        ser[0] = Timestamp("nat")
        ser[3] = Timestamp("nat")

        left_f = getattr(operator, left)
        right_f = getattr(operator, right)

        # No NaT
        expected = left_f(ser, Timestamp("20010109"))
        result = right_f(Timestamp("20010109"), ser)
        tm.assert_series_equal(result, expected)

        # NaT
        expected = left_f(ser, Timestamp("nat"))
        result = right_f(Timestamp("nat"), ser)
        tm.assert_series_equal(result, expected)

        # Compare to Timestamp with series containing NaT
        expected = left_f(s_nat, Timestamp("20010109"))
        result = right_f(Timestamp("20010109"), s_nat)
        tm.assert_series_equal(result, expected)

        # Compare to NaT with series containing NaT
        expected = left_f(s_nat, NaT)
        result = right_f(NaT, s_nat)
        tm.assert_series_equal(result, expected)

    def test_dt64arr_timestamp_equality(self, box_with_array):
        # GH#11034
        box = box_with_array

        ser = Series([Timestamp("2000-01-29 01:59:00"), Timestamp("2000-01-30"), NaT])
        ser = tm.box_expected(ser, box)
        xbox = get_upcast_box(ser, ser, True)

        result = ser != ser
        expected = tm.box_expected([False, False, True], xbox)
        tm.assert_equal(result, expected)

        if box is pd.DataFrame:
            # alignment for frame vs series comparisons deprecated
            #  in GH#46795 enforced 2.0
            with pytest.raises(ValueError, match="not aligned"):
                ser != ser[0]

        else:
            result = ser != ser[0]
            expected = tm.box_expected([False, True, True], xbox)
            tm.assert_equal(result, expected)

        if box is pd.DataFrame:
            # alignment for frame vs series comparisons deprecated
            #  in GH#46795 enforced 2.0
            with pytest.raises(ValueError, match="not aligned"):
                ser != ser[2]
        else:
            result = ser != ser[2]
            expected = tm.box_expected([True, True, True], xbox)
            tm.assert_equal(result, expected)

        result = ser == ser
        expected = tm.box_expected([True, True, False], xbox)
        tm.assert_equal(result, expected)

        if box is pd.DataFrame:
            # alignment for frame vs series comparisons deprecated
            #  in GH#46795 enforced 2.0
            with pytest.raises(ValueError, match="not aligned"):
                ser == ser[0]
        else:
            result = ser == ser[0]
            expected = tm.box_expected([True, False, False], xbox)
            tm.assert_equal(result, expected)

        if box is pd.DataFrame:
            # alignment for frame vs series comparisons deprecated
            #  in GH#46795 enforced 2.0
            with pytest.raises(ValueError, match="not aligned"):
                ser == ser[2]
        else:
            result = ser == ser[2]
            expected = tm.box_expected([False, False, False], xbox)
            tm.assert_equal(result, expected)

    @pytest.mark.parametrize(
        "datetimelike",
        [
            Timestamp("20130101"),
            datetime(2013, 1, 1),
            np.datetime64("2013-01-01T00:00", "ns"),
        ],
    )
    @pytest.mark.parametrize(
        "op,expected",
        [
            (operator.lt, [True, False, False, False]),
            (operator.le, [True, True, False, False]),
            (operator.eq, [False, True, False, False]),
            (operator.gt, [False, False, False, True]),
        ],
    )
    def test_dt64_compare_datetime_scalar(self, datetimelike, op, expected):
        # GH#17965, test for ability to compare datetime64[ns] columns
        #  to datetimelike
        ser = Series(
            [
                Timestamp("20120101"),
                Timestamp("20130101"),
                np.nan,
                Timestamp("20130103"),
            ],
            name="A",
        )
        result = op(ser, datetimelike)
        expected = Series(expected, name="A")
        tm.assert_series_equal(result, expected)


class TestDatetimeIndexComparisons:
    # TODO: moved from tests.indexes.test_base; parametrize and de-duplicate
    def test_comparators(self, comparison_op):
        index = date_range("2020-01-01", periods=10)
        element = index[len(index) // 2]
        element = Timestamp(element).to_datetime64()

        arr = np.array(index)
        arr_result = comparison_op(arr, element)
        index_result = comparison_op(index, element)

        assert isinstance(index_result, np.ndarray)
        tm.assert_numpy_array_equal(arr_result, index_result)

    @pytest.mark.parametrize(
        "other",
        [datetime(2016, 1, 1), Timestamp("2016-01-01"), np.datetime64("2016-01-01")],
    )
    def test_dti_cmp_datetimelike(self, other, tz_naive_fixture):
        tz = tz_naive_fixture
        dti = date_range("2016-01-01", periods=2, tz=tz)
        if tz is not None:
            if isinstance(other, np.datetime64):
                pytest.skip(f"{type(other).__name__} is not tz aware")
            other = localize_pydatetime(other, dti.tzinfo)

        result = dti == other
        expected = np.array([True, False])
        tm.assert_numpy_array_equal(result, expected)

        result = dti > other
        expected = np.array([False, True])
        tm.assert_numpy_array_equal(result, expected)

        result = dti >= other
        expected = np.array([True, True])
        tm.assert_numpy_array_equal(result, expected)

        result = dti < other
        expected = np.array([False, False])
        tm.assert_numpy_array_equal(result, expected)

        result = dti <= other
        expected = np.array([True, False])
        tm.assert_numpy_array_equal(result, expected)

    @pytest.mark.parametrize("dtype", [None, object])
    def test_dti_cmp_nat(self, dtype, box_with_array):
        left = DatetimeIndex([Timestamp("2011-01-01"), NaT, Timestamp("2011-01-03")])
        right = DatetimeIndex([NaT, NaT, Timestamp("2011-01-03")])

        left = tm.box_expected(left, box_with_array)
        right = tm.box_expected(right, box_with_array)
        xbox = get_upcast_box(left, right, True)

        lhs, rhs = left, right
        if dtype is object:
            lhs, rhs = left.astype(object), right.astype(object)

        result = rhs == lhs
        expected = np.array([False, False, True])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(result, expected)

        result = lhs != rhs
        expected = np.array([True, True, False])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(result, expected)

        expected = np.array([False, False, False])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(lhs == NaT, expected)
        tm.assert_equal(NaT == rhs, expected)

        expected = np.array([True, True, True])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(lhs != NaT, expected)
        tm.assert_equal(NaT != lhs, expected)

        expected = np.array([False, False, False])
        expected = tm.box_expected(expected, xbox)
        tm.assert_equal(lhs < NaT, expected)
        tm.assert_equal(NaT > lhs, expected)

    def test_dti_cmp_nat_behaves_like_float_cmp_nan(self):
        fidx1 = pd.Index([1.0, np.nan, 3.0, np.nan, 5.0, 7.0])
        fidx2 = pd.Index([2.0, 3.0, np.nan, np.nan, 6.0, 7.0])

        didx1 = DatetimeIndex(
            ["2014-01-01", NaT, "2014-03-01", NaT, "2014-05-01", "2014-07-01"]
        )
        didx2 = DatetimeIndex(
            ["2014-02-01", "2014-03-01", NaT, NaT, "2014-06-01", "2014-07-01"]
        )
        darr = np.array(
            [
                np.datetime64("2014-02-01 00:00"),
                np.datetime64("2014-03-01 00:00"),
                np.datetime64("nat"),
                np.datetime64("nat"),
                np.datetime64("2014-06-01 00:00"),
                np.datetime64("2014-07-01 00:00"),
            ]
        )

        cases = [(fidx1, fidx2), (didx1, didx2), (didx1, darr)]

        # Check pd.NaT is handles as the same as np.nan
        with tm.assert_produces_warning(None):
            for idx1, idx2 in cases:
                result = idx1 < idx2
                expected = np.array([True, False, False, False, True, False])
                tm.assert_numpy_array_equal(result, expected)

                result = idx2 > idx1
                expected = np.array([True, False, False, False, True, False])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 <= idx2
                expected = np.array([True, False, False, False, True, True])
                tm.assert_numpy_array_equal(result, expected)

                result = idx2 >= idx1
                expected = np.array([True, False, False, False, True, True])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 == idx2
                expected = np.array([False, False, False, False, False, True])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 != idx2
                expected = np.array([True, True, True, True, True, False])
                tm.assert_numpy_array_equal(result, expected)

        with tm.assert_produces_warning(None):
            for idx1, val in [(fidx1, np.nan), (didx1, NaT)]:
                result = idx1 < val
                expected = np.array([False, False, False, False, False, False])
                tm.assert_numpy_array_equal(result, expected)
                result = idx1 > val
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 <= val
                tm.assert_numpy_array_equal(result, expected)
                result = idx1 >= val
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 == val
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 != val
                expected = np.array([True, True, True, True, True, True])
                tm.assert_numpy_array_equal(result, expected)

        # Check pd.NaT is handles as the same as np.nan
        with tm.assert_produces_warning(None):
            for idx1, val in [(fidx1, 3), (didx1, datetime(2014, 3, 1))]:
                result = idx1 < val
                expected = np.array([True, False, False, False, False, False])
                tm.assert_numpy_array_equal(result, expected)
                result = idx1 > val
                expected = np.array([False, False, False, False, True, True])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 <= val
                expected = np.array([True, False, True, False, False, False])
                tm.assert_numpy_array_equal(result, expected)
                result = idx1 >= val
                expected = np.array([False, False, True, False, True, True])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 == val
                expected = np.array([False, False, True, False, False, False])
                tm.assert_numpy_array_equal(result, expected)

                result = idx1 != val
                expected = np.array([True, True, False, True, True, True])
                tm.assert_numpy_array_equal(result, expected)

    def test_comparison_tzawareness_compat(self, comparison_op, box_with_array):
        # GH#18162
        op = comparison_op
        box = box_with_array

        dr = date_range("2016-01-01", periods=6)
        dz = dr.tz_localize("US/Pacific")

        dr = tm.box_expected(dr, box)
        dz = tm.box_expected(dz, box)

        if box is pd.DataFrame:
            tolist = lambda x: x.astype(object).values.tolist()[0]
        else:
            tolist = list

        if op not in [operator.eq, operator.ne]:
            msg = (
                r"Invalid comparison between dtype=datetime64\[ns.*\] "
                "and (Timestamp|DatetimeArray|list|ndarray)"
            )
            with pytest.raises(TypeError, match=msg):
                op(dr, dz)

            with pytest.raises(TypeError, match=msg):
                op(dr, tolist(dz))
            with pytest.raises(TypeError, match=msg):
                op(dr, np.array(tolist(dz), dtype=object))
            with pytest.raises(TypeError, match=msg):
                op(dz, dr)

            with pytest.raises(TypeError, match=msg):
                op(dz, tolist(dr))
            with pytest.raises(TypeError, match=msg):
                op(dz, np.array(tolist(dr), dtype=object))

        # The aware==aware and naive==naive comparisons should *not* raise
        assert np.all(dr == dr)
        assert np.all(dr == tolist(dr))
        assert np.all(tolist(dr) == dr)
        assert np.all(np.array(tolist(dr), dtype=object) == dr)
        assert np.all(dr == np.array(tolist(dr), dtype=object))

        assert np.all(dz == dz)
        assert np.all(dz == tolist(dz))
        assert np.all(tolist(dz) == dz)
        assert np.all(np.array(tolist(dz), dtype=object) == dz)
        assert np.all(dz == np.array(tolist(dz), dtype=object))

    def test_comparison_tzawareness_compat_scalars(self, comparison_op, box_with_array):
        # GH#18162
        op = comparison_op

        dr = date_range("2016-01-01", periods=6)
        dz = dr.tz_localize("US/Pacific")

        dr = tm.box_expected(dr, box_with_array)
        dz = tm.box_expected(dz, box_with_array)

        # Check comparisons against scalar Timestamps
        ts = Timestamp("2000-03-14 01:59")
        ts_tz = Timestamp("2000-03-14 01:59", tz="Europe/Amsterdam")

        assert np.all(dr > ts)
        msg = r"Invalid comparison between dtype=datetime64\[ns.*\] and Timestamp"
        if op not in [operator.eq, operator.ne]:
            with pytest.raises(TypeError, match=msg):
                op(dr, ts_tz)

        assert np.all(dz > ts_tz)
        if op not in [operator.eq, operator.ne]:
            with pytest.raises(TypeError, match=msg):
                op(dz, ts)

        if op not in [operator.eq, operator.ne]:
            # GH#12601: Check comparison against Timestamps and DatetimeIndex
            with pytest.raises(TypeError, match=msg):
                op(ts, dz)

    @pytest.mark.parametrize(
        "other",
        [datetime(2016, 1, 1), Timestamp("2016-01-01"), np.datetime64("2016-01-01")],
    )
    # Bug in NumPy? https://github.com/numpy/numpy/issues/13841
    # Raising in __eq__ will fallback to NumPy, which warns, fails,
    # then re-raises the original exception. So we just need to ignore.
    @pytest.mark.filterwarnings("ignore:elementwise comp:DeprecationWarning")
    def test_scalar_comparison_tzawareness(
        self, comparison_op, other, tz_aware_fixture, box_with_array
    ):
        op = comparison_op
        tz = tz_aware_fixture
        dti = date_range("2016-01-01", periods=2, tz=tz)

        dtarr = tm.box_expected(dti, box_with_array)
        xbox = get_upcast_box(dtarr, other, True)
        if op in [operator.eq, operator.ne]:
            exbool = op is operator.ne
            expected = np.array([exbool, exbool], dtype=bool)
            expected = tm.box_expected(expected, xbox)

            result = op(dtarr, other)
            tm.assert_equal(result, expected)

            result = op(other, dtarr)
            tm.assert_equal(result, expected)
        else:
            msg = (
                r"Invalid comparison between dtype=datetime64\[ns, .*\] "
                f"and {type(other).__name__}"
            )
            with pytest.raises(TypeError, match=msg):
                op(dtarr, other)
            with pytest.raises(TypeError, match=msg):
                op(other, dtarr)

    def test_nat_comparison_tzawareness(self, comparison_op):
        # GH#19276
        # tzaware DatetimeIndex should not raise when compared to NaT
        op = comparison_op

        dti = DatetimeIndex(
            ["2014-01-01", NaT, "2014-03-01", NaT, "2014-05-01", "2014-07-01"]
        )
        expected = np.array([op == operator.ne] * len(dti))
        result = op(dti, NaT)
        tm.assert_numpy_array_equal(result, expected)

        result = op(dti.tz_localize("US/Pacific"), NaT)
        tm.assert_numpy_array_equal(result, expected)

    def test_dti_cmp_str(self, tz_naive_fixture):
        # GH#22074
        # regardless of tz, we expect these comparisons are valid
        tz = tz_naive_fixture
        rng = date_range("1/1/2000", periods=10, tz=tz)
        other = "1/1/2000"

        result = rng == other
        expected = np.array([True] + [False] * 9)
        tm.assert_numpy_array_equal(result, expected)

        result = rng != other
        expected = np.array([False] + [True] * 9)
        tm.assert_numpy_array_equal(result, expected)

        result = rng < other
        expected = np.array([False] * 10)
        tm.assert_numpy_array_equal(result, expected)

        result = rng <= other
        expected = np.array([True] + [False] * 9)
        tm.assert_numpy_array_equal(result, expected)

        result = rng > other
        expected = np.array([False] + [True] * 9)
        tm.assert_numpy_array_equal(result, expected)

        result = rng >= other
        expected = np.array([True] * 10)
        tm.assert_numpy_array_equal(result, expected)

    def test_dti_cmp_list(self):
        rng = date_range("1/1/2000", periods=10)

        result = rng == list(rng)
        expected = rng == rng
        tm.assert_numpy_array_equal(result, expected)

    @pytest.mark.parametrize(
        "other",
        [
            pd.timedelta_range("1D", periods=10),
            pd.timedelta_range("1D", periods=10).to_series(),
            pd.timedelta_range("1D", periods=10).asi8.view("m8[ns]"),
        ],
        ids=lambda x: type(x).__name__,
    )
    def test_dti_cmp_tdi_tzawareness(self, other):
        # GH#22074
        # reversion test that we _don't_ call _assert_tzawareness_compat
        # when comparing against TimedeltaIndex
        dti = date_range("2000-01-01", periods=10, tz="Asia/Tokyo")

        result = dti == other
        expected = np.array([False] * 10)
        tm.assert_numpy_array_equal(result, expected)

        result = dti != other
        expected = np.array([True] * 10)
        tm.assert_numpy_array_equal(result, expected)
        msg = "Invalid comparison between"
        with pytest.raises(TypeError, match=msg):
            dti < other
        with pytest.raises(TypeError, match=msg):
            dti <= other
        with pytest.raises(TypeError, match=msg):
            dti > other
        with pytest.raises(TypeError, match=msg):
            dti >= other

    def test_dti_cmp_object_dtype(self):
        # GH#22074
        dti = date_range("2000-01-01", periods=10, tz="Asia/Tokyo")

        other = dti.astype("O")

        result = dti == other
        expected = np.array([True] * 10)
        tm.assert_numpy_array_equal(result, expected)

        other = dti.tz_localize(None)
        result = dti != other
        tm.assert_numpy_array_equal(result, expected)

        other = np.array(list(dti[:5]) + [Timedelta(days=1)] * 5)
        result = dti == other
        expected = np.array([True] * 5 + [False] * 5)
        tm.assert_numpy_array_equal(result, expected)
        msg = ">=' not supported between instances of 'Timestamp' and 'Timedelta'"
        with pytest.raises(TypeError, match=msg):
            dti >= other


# ------------------------------------------------------------------
# Arithmetic


class TestDatetime64Arithmetic:
    # This class is intended for "finished" tests that are fully parametrized
    #  over DataFrame/Series/Index/DatetimeArray

    # -------------------------------------------------------------
    # Addition/Subtraction of timedelta-like

    @pytest.mark.arm_slow
    def test_dt64arr_add_timedeltalike_scalar(
        self, tz_naive_fixture, two_hours, box_with_array
    ):
        # GH#22005, GH#22163 check DataFrame doesn't raise TypeError
        tz = tz_naive_fixture

        rng = date_range("2000-01-01", "2000-02-01", tz=tz)
        expected = date_range("2000-01-01 02:00", "2000-02-01 02:00", tz=tz)

        rng = tm.box_expected(rng, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = rng + two_hours
        tm.assert_equal(result, expected)

        result = two_hours + rng
        tm.assert_equal(result, expected)

        rng += two_hours
        tm.assert_equal(rng, expected)

    def test_dt64arr_sub_timedeltalike_scalar(
        self, tz_naive_fixture, two_hours, box_with_array
    ):
        tz = tz_naive_fixture

        rng = date_range("2000-01-01", "2000-02-01", tz=tz)
        expected = date_range("1999-12-31 22:00", "2000-01-31 22:00", tz=tz)

        rng = tm.box_expected(rng, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = rng - two_hours
        tm.assert_equal(result, expected)

        rng -= two_hours
        tm.assert_equal(rng, expected)

    def test_dt64_array_sub_dt_with_different_timezone(self, box_with_array):
        t1 = date_range("20130101", periods=3).tz_localize("US/Eastern")
        t1 = tm.box_expected(t1, box_with_array)
        t2 = Timestamp("20130101").tz_localize("CET")
        tnaive = Timestamp(20130101)

        result = t1 - t2
        expected = TimedeltaIndex(
            ["0 days 06:00:00", "1 days 06:00:00", "2 days 06:00:00"]
        )
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

        result = t2 - t1
        expected = TimedeltaIndex(
            ["-1 days +18:00:00", "-2 days +18:00:00", "-3 days +18:00:00"]
        )
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

        msg = "Cannot subtract tz-naive and tz-aware datetime-like objects"
        with pytest.raises(TypeError, match=msg):
            t1 - tnaive

        with pytest.raises(TypeError, match=msg):
            tnaive - t1

    def test_dt64_array_sub_dt64_array_with_different_timezone(self, box_with_array):
        t1 = date_range("20130101", periods=3).tz_localize("US/Eastern")
        t1 = tm.box_expected(t1, box_with_array)
        t2 = date_range("20130101", periods=3).tz_localize("CET")
        t2 = tm.box_expected(t2, box_with_array)
        tnaive = date_range("20130101", periods=3)

        result = t1 - t2
        expected = TimedeltaIndex(
            ["0 days 06:00:00", "0 days 06:00:00", "0 days 06:00:00"]
        )
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

        result = t2 - t1
        expected = TimedeltaIndex(
            ["-1 days +18:00:00", "-1 days +18:00:00", "-1 days +18:00:00"]
        )
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

        msg = "Cannot subtract tz-naive and tz-aware datetime-like objects"
        with pytest.raises(TypeError, match=msg):
            t1 - tnaive

        with pytest.raises(TypeError, match=msg):
            tnaive - t1

    def test_dt64arr_add_sub_td64_nat(self, box_with_array, tz_naive_fixture):
        # GH#23320 special handling for timedelta64("NaT")
        tz = tz_naive_fixture

        dti = date_range("1994-04-01", periods=9, tz=tz, freq="QS")
        other = np.timedelta64("NaT")
        expected = DatetimeIndex(["NaT"] * 9, tz=tz).as_unit("ns")

        obj = tm.box_expected(dti, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = obj + other
        tm.assert_equal(result, expected)
        result = other + obj
        tm.assert_equal(result, expected)
        result = obj - other
        tm.assert_equal(result, expected)
        msg = "cannot subtract"
        with pytest.raises(TypeError, match=msg):
            other - obj

    def test_dt64arr_add_sub_td64ndarray(self, tz_naive_fixture, box_with_array):
        tz = tz_naive_fixture
        dti = date_range("2016-01-01", periods=3, tz=tz)
        tdi = TimedeltaIndex(["-1 Day", "-1 Day", "-1 Day"])
        tdarr = tdi.values

        expected = date_range("2015-12-31", "2016-01-02", periods=3, tz=tz)

        dtarr = tm.box_expected(dti, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = dtarr + tdarr
        tm.assert_equal(result, expected)
        result = tdarr + dtarr
        tm.assert_equal(result, expected)

        expected = date_range("2016-01-02", "2016-01-04", periods=3, tz=tz)
        expected = tm.box_expected(expected, box_with_array)

        result = dtarr - tdarr
        tm.assert_equal(result, expected)
        msg = "cannot subtract|(bad|unsupported) operand type for unary"
        with pytest.raises(TypeError, match=msg):
            tdarr - dtarr

    # -----------------------------------------------------------------
    # Subtraction of datetime-like scalars

    @pytest.mark.parametrize(
        "ts",
        [
            Timestamp("2013-01-01"),
            Timestamp("2013-01-01").to_pydatetime(),
            Timestamp("2013-01-01").to_datetime64(),
            # GH#7996, GH#22163 ensure non-nano datetime64 is converted to nano
            #  for DataFrame operation
            np.datetime64("2013-01-01", "D"),
        ],
    )
    def test_dt64arr_sub_dtscalar(self, box_with_array, ts):
        # GH#8554, GH#22163 DataFrame op should _not_ return dt64 dtype
        idx = date_range("2013-01-01", periods=3)._with_freq(None)
        idx = tm.box_expected(idx, box_with_array)

        expected = TimedeltaIndex(["0 Days", "1 Day", "2 Days"])
        expected = tm.box_expected(expected, box_with_array)

        result = idx - ts
        tm.assert_equal(result, expected)

        result = ts - idx
        tm.assert_equal(result, -expected)
        tm.assert_equal(result, -expected)

    def test_dt64arr_sub_timestamp_tzaware(self, box_with_array):
        ser = date_range("2014-03-17", periods=2, freq="D", tz="US/Eastern")
        ser = ser._with_freq(None)
        ts = ser[0]

        ser = tm.box_expected(ser, box_with_array)

        delta_series = Series([np.timedelta64(0, "D"), np.timedelta64(1, "D")])
        expected = tm.box_expected(delta_series, box_with_array)

        tm.assert_equal(ser - ts, expected)
        tm.assert_equal(ts - ser, -expected)

    def test_dt64arr_sub_NaT(self, box_with_array, unit):
        # GH#18808
        dti = DatetimeIndex([NaT, Timestamp("19900315")]).as_unit(unit)
        ser = tm.box_expected(dti, box_with_array)

        result = ser - NaT
        expected = Series([NaT, NaT], dtype=f"timedelta64[{unit}]")
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

        dti_tz = dti.tz_localize("Asia/Tokyo")
        ser_tz = tm.box_expected(dti_tz, box_with_array)

        result = ser_tz - NaT
        expected = Series([NaT, NaT], dtype=f"timedelta64[{unit}]")
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(result, expected)

    # -------------------------------------------------------------
    # Subtraction of datetime-like array-like

    def test_dt64arr_sub_dt64object_array(self, box_with_array, tz_naive_fixture):
        dti = date_range("2016-01-01", periods=3, tz=tz_naive_fixture)
        expected = dti - dti

        obj = tm.box_expected(dti, box_with_array)
        expected = tm.box_expected(expected, box_with_array).astype(object)

        with tm.assert_produces_warning(PerformanceWarning):
            result = obj - obj.astype(object)
        tm.assert_equal(result, expected)

    def test_dt64arr_naive_sub_dt64ndarray(self, box_with_array):
        dti = date_range("2016-01-01", periods=3, tz=None)
        dt64vals = dti.values

        dtarr = tm.box_expected(dti, box_with_array)

        expected = dtarr - dtarr
        result = dtarr - dt64vals
        tm.assert_equal(result, expected)
        result = dt64vals - dtarr
        tm.assert_equal(result, expected)

    def test_dt64arr_aware_sub_dt64ndarray_raises(
        self, tz_aware_fixture, box_with_array
    ):
        tz = tz_aware_fixture
        dti = date_range("2016-01-01", periods=3, tz=tz)
        dt64vals = dti.values

        dtarr = tm.box_expected(dti, box_with_array)
        msg = "Cannot subtract tz-naive and tz-aware datetime"
        with pytest.raises(TypeError, match=msg):
            dtarr - dt64vals
        with pytest.raises(TypeError, match=msg):
            dt64vals - dtarr

    # -------------------------------------------------------------
    # Addition of datetime-like others (invalid)

    def test_dt64arr_add_dtlike_raises(self, tz_naive_fixture, box_with_array):
        # GH#22163 ensure DataFrame doesn't cast Timestamp to i8
        # GH#9631
        tz = tz_naive_fixture

        dti = date_range("2016-01-01", periods=3, tz=tz)
        if tz is None:
            dti2 = dti.tz_localize("US/Eastern")
        else:
            dti2 = dti.tz_localize(None)
        dtarr = tm.box_expected(dti, box_with_array)

        assert_cannot_add(dtarr, dti.values)
        assert_cannot_add(dtarr, dti)
        assert_cannot_add(dtarr, dtarr)
        assert_cannot_add(dtarr, dti[0])
        assert_cannot_add(dtarr, dti[0].to_pydatetime())
        assert_cannot_add(dtarr, dti[0].to_datetime64())
        assert_cannot_add(dtarr, dti2[0])
        assert_cannot_add(dtarr, dti2[0].to_pydatetime())
        assert_cannot_add(dtarr, np.datetime64("2011-01-01", "D"))

    # -------------------------------------------------------------
    # Other Invalid Addition/Subtraction

    # Note: freq here includes both Tick and non-Tick offsets; this is
    #  relevant because historically integer-addition was allowed if we had
    #  a freq.
    @pytest.mark.parametrize("freq", ["h", "D", "W", "2ME", "MS", "QE", "B", None])
    @pytest.mark.parametrize("dtype", [None, "uint8"])
    def test_dt64arr_addsub_intlike(
        self, request, dtype, index_or_series_or_array, freq, tz_naive_fixture
    ):
        # GH#19959, GH#19123, GH#19012
        # GH#55860 use index_or_series_or_array instead of box_with_array
        #  bc DataFrame alignment makes it inapplicable
        tz = tz_naive_fixture

        if freq is None:
            dti = DatetimeIndex(["NaT", "2017-04-05 06:07:08"], tz=tz)
        else:
            dti = date_range("2016-01-01", periods=2, freq=freq, tz=tz)

        obj = index_or_series_or_array(dti)
        other = np.array([4, -1])
        if dtype is not None:
            other = other.astype(dtype)

        msg = "|".join(
            [
                "Addition/subtraction of integers",
                "cannot subtract DatetimeArray from",
                # IntegerArray
                "can only perform ops with numeric values",
                "unsupported operand type.*Categorical",
                r"unsupported operand type\(s\) for -: 'int' and 'Timestamp'",
            ]
        )
        assert_invalid_addsub_type(obj, 1, msg)
        assert_invalid_addsub_type(obj, np.int64(2), msg)
        assert_invalid_addsub_type(obj, np.array(3, dtype=np.int64), msg)
        assert_invalid_addsub_type(obj, other, msg)
        assert_invalid_addsub_type(obj, np.array(other), msg)
        assert_invalid_addsub_type(obj, pd.array(other), msg)
        assert_invalid_addsub_type(obj, pd.Categorical(other), msg)
        assert_invalid_addsub_type(obj, pd.Index(other), msg)
        assert_invalid_addsub_type(obj, Series(other), msg)

    @pytest.mark.parametrize(
        "other",
        [
            3.14,
            np.array([2.0, 3.0]),
            # GH#13078 datetime +/- Period is invalid
            Period("2011-01-01", freq="D"),
            # https://github.com/pandas-dev/pandas/issues/10329
            time(1, 2, 3),
        ],
    )
    @pytest.mark.parametrize("dti_freq", [None, "D"])
    def test_dt64arr_add_sub_invalid(self, dti_freq, other, box_with_array):
        dti = DatetimeIndex(["2011-01-01", "2011-01-02"], freq=dti_freq)
        dtarr = tm.box_expected(dti, box_with_array)
        msg = "|".join(
            [
                "unsupported operand type",
                "cannot (add|subtract)",
                "cannot use operands with types",
                "ufunc '?(add|subtract)'? cannot use operands with types",
                "Concatenation operation is not implemented for NumPy arrays",
            ]
        )
        assert_invalid_addsub_type(dtarr, other, msg)

    @pytest.mark.parametrize("pi_freq", ["D", "W", "Q", "h"])
    @pytest.mark.parametrize("dti_freq", [None, "D"])
    def test_dt64arr_add_sub_parr(
        self, dti_freq, pi_freq, box_with_array, box_with_array2
    ):
        # GH#20049 subtracting PeriodIndex should raise TypeError
        dti = DatetimeIndex(["2011-01-01", "2011-01-02"], freq=dti_freq)
        pi = dti.to_period(pi_freq)

        dtarr = tm.box_expected(dti, box_with_array)
        parr = tm.box_expected(pi, box_with_array2)
        msg = "|".join(
            [
                "cannot (add|subtract)",
                "unsupported operand",
                "descriptor.*requires",
                "ufunc.*cannot use operands",
            ]
        )
        assert_invalid_addsub_type(dtarr, parr, msg)

    @pytest.mark.filterwarnings("ignore::pandas.errors.PerformanceWarning")
    def test_dt64arr_addsub_time_objects_raises(self, box_with_array, tz_naive_fixture):
        # https://github.com/pandas-dev/pandas/issues/10329

        tz = tz_naive_fixture

        obj1 = date_range("2012-01-01", periods=3, tz=tz)
        obj2 = [time(i, i, i) for i in range(3)]

        obj1 = tm.box_expected(obj1, box_with_array)
        obj2 = tm.box_expected(obj2, box_with_array)

        msg = "|".join(
            [
                "unsupported operand",
                "cannot subtract DatetimeArray from ndarray",
            ]
        )
        # pandas.errors.PerformanceWarning: Non-vectorized DateOffset being
        # applied to Series or DatetimeIndex
        # we aren't testing that here, so ignore.
        assert_invalid_addsub_type(obj1, obj2, msg=msg)

    # -------------------------------------------------------------
    # Other invalid operations

    @pytest.mark.parametrize(
        "dt64_series",
        [
            Series([Timestamp("19900315"), Timestamp("19900315")]),
            Series([NaT, Timestamp("19900315")]),
            Series([NaT, NaT], dtype="datetime64[ns]"),
        ],
    )
    @pytest.mark.parametrize("one", [1, 1.0, np.array(1)])
    def test_dt64_mul_div_numeric_invalid(self, one, dt64_series, box_with_array):
        obj = tm.box_expected(dt64_series, box_with_array)

        msg = "cannot perform .* with this index type"

        # multiplication
        with pytest.raises(TypeError, match=msg):
            obj * one
        with pytest.raises(TypeError, match=msg):
            one * obj

        # division
        with pytest.raises(TypeError, match=msg):
            obj / one
        with pytest.raises(TypeError, match=msg):
            one / obj


class TestDatetime64DateOffsetArithmetic:
    # -------------------------------------------------------------
    # Tick DateOffsets

    # TODO: parametrize over timezone?
    @pytest.mark.parametrize("unit", ["s", "ms", "us", "ns"])
    def test_dt64arr_series_add_tick_DateOffset(self, box_with_array, unit):
        # GH#4532
        # operate with pd.offsets
        ser = Series(
            [Timestamp("20130101 9:01"), Timestamp("20130101 9:02")]
        ).dt.as_unit(unit)
        expected = Series(
            [Timestamp("20130101 9:01:05"), Timestamp("20130101 9:02:05")]
        ).dt.as_unit(unit)

        ser = tm.box_expected(ser, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = ser + pd.offsets.Second(5)
        tm.assert_equal(result, expected)

        result2 = pd.offsets.Second(5) + ser
        tm.assert_equal(result2, expected)

    def test_dt64arr_series_sub_tick_DateOffset(self, box_with_array):
        # GH#4532
        # operate with pd.offsets
        ser = Series([Timestamp("20130101 9:01"), Timestamp("20130101 9:02")])
        expected = Series(
            [Timestamp("20130101 9:00:55"), Timestamp("20130101 9:01:55")]
        )

        ser = tm.box_expected(ser, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        result = ser - pd.offsets.Second(5)
        tm.assert_equal(result, expected)

        result2 = -pd.offsets.Second(5) + ser
        tm.assert_equal(result2, expected)
        msg = "(bad|unsupported) operand type for unary"
        with pytest.raises(TypeError, match=msg):
            pd.offsets.Second(5) - ser

    @pytest.mark.parametrize(
        "cls_name", ["Day", "Hour", "Minute", "Second", "Milli", "Micro", "Nano"]
    )
    def test_dt64arr_add_sub_tick_DateOffset_smoke(self, cls_name, box_with_array):
        # GH#4532
        # smoke tests for valid DateOffsets
        ser = Series([Timestamp("20130101 9:01"), Timestamp("20130101 9:02")])
        ser = tm.box_expected(ser, box_with_array)

        offset_cls = getattr(pd.offsets, cls_name)
        ser + offset_cls(5)
        offset_cls(5) + ser
        ser - offset_cls(5)

    def test_dti_add_tick_tzaware(self, tz_aware_fixture, box_with_array):
        # GH#21610, GH#22163 ensure DataFrame doesn't return object-dtype
        tz = tz_aware_fixture
        if tz == "US/Pacific":
            dates = date_range("2012-11-01", periods=3, tz=tz)
            offset = dates + pd.offsets.Hour(5)
            assert dates[0] + pd.offsets.Hour(5) == offset[0]

        dates = date_range("2010-11-01 00:00", periods=3, tz=tz, freq="h")
        expected = DatetimeIndex(
            ["2010-11-01 05:00", "2010-11-01 06:00", "2010-11-01 07:00"],
            freq="h",
            tz=tz,
        ).as_unit("ns")

        dates = tm.box_expected(dates, box_with_array)
        expected = tm.box_expected(expected, box_with_array)

        for scalar in [pd.offsets.Hour(5), np.timedelta64(5, "h"), timedelta(hours=5)]:
            offset = dates + scalar
            tm.assert_equal(offset, expected)
            offset = scalar + dates
            tm.assert_equal(offset, expected)

            roundtrip = offset - scalar
            tm.assert_equal(roundtrip, dates)

            msg = "|".join(
                ["bad operand type for unary -", "cannot subtract DatetimeArray"]
            )
            with pytest.raises(TypeError, match=msg):
                scalar - dates

    # -------------------------------------------------------------
    # RelativeDelta DateOffsets

    @pytest.mark.parametrize("unit", ["s", "ms", "us", "ns"])
    def test_dt64arr_add_sub_relativedelta_offsets(self, box_with_array, unit):
        # GH#10699
        vec = DatetimeIndex(
            [
                Timestamp("2000-01-05 00:15:00"),
                Timestamp("2000-01-31 00:23:00"),
                Timestamp("2000-01-01"),
                Timestamp("2000-03-31"),
                Timestamp("2000-02-29"),
                Timestamp("2000-12-31"),
                Timestamp("2000-05-15"),
                Timestamp("2001-06-15"),
            ]
        ).as_unit(unit)
        vec = tm.box_expected(vec, box_with_array)
        vec_items = vec.iloc[0] if box_with_array is pd.DataFrame else vec

        # DateOffset relativedelta fastpath
        relative_kwargs = [
            ("years", 2),
            ("months", 5),
            ("days", 3),
            ("hours", 5),
            ("minutes", 10),
            ("seconds", 2),
            ("microseconds", 5),
        ]
        for i, (offset_unit, value) in enumerate(relative_kwargs):
            off = DateOffset(**{offset_unit: value})

            exp_unit = unit
            if offset_unit == "microseconds" and unit != "ns":
                exp_unit = "us"

            # TODO(GH#55564): as_unit will be unnecessary
            expected = DatetimeIndex([x + off for x in vec_items]).as_unit(exp_unit)
            expected = tm.box_expected(expected, box_with_array)
            tm.assert_equal(expected, vec + off)

            expected = DatetimeIndex([x - off for x in vec_items]).as_unit(exp_unit)
            expected = tm.box_expected(expected, box_with_array)
            tm.assert_equal(expected, vec - off)

            off = DateOffset(**dict(relative_kwargs[: i + 1]))

            expected = DatetimeIndex([x + off for x in vec_items]).as_unit(exp_unit)
            expected = tm.box_expected(expected, box_with_array)
            tm.assert_equal(expected, vec + off)

            expected = DatetimeIndex([x - off for x in vec_items]).as_unit(exp_unit)
            expected = tm.box_expected(expected, box_with_array)
            tm.assert_equal(expected, vec - off)
            msg = "(bad|unsupported) operand type for unary"
            with pytest.raises(TypeError, match=msg):
                off - vec

    # -------------------------------------------------------------
    # Non-Tick, Non-RelativeDelta DateOffsets

    # TODO: redundant with test_dt64arr_add_sub_DateOffset?  that includes
    #  tz-aware cases which this does not
    @pytest.mark.filterwarnings("ignore::pandas.errors.PerformanceWarning")
    @pytest.mark.parametrize(
        "cls_and_kwargs",
        [
            "YearBegin",
            ("YearBegin", {"month": 5}),
            "YearEnd",
            ("YearEnd", {"month": 5}),
            "MonthBegin",
            "MonthEnd",
            "SemiMonthEnd",
            "SemiMonthBegin",
            "Week",
            ("Week", {"weekday": 3}),
            "Week",
            ("Week", {"weekday": 6}),
            "BusinessDay",
            "BDay",
            "QuarterEnd",
            "QuarterBegin",
            "CustomBusinessDay",
            "CDay",
            "CBMonthEnd",
            "CBMonthBegin",
            "BMonthBegin",
            "BMonthEnd",
            "BusinessHour",
            "BYearBegin",
            "BYearEnd",
            "BQuarterBegin",
            ("LastWeekOfMonth", {"weekday": 2}),
            (
                "FY5253Quarter",
                {
                    "qtr_with_extra_week": 1,
                    "startingMonth": 1,
                    "weekday": 2,
                    "variation": "nearest",
                },
            ),
            ("FY5253", {"weekday": 0, "startingMonth": 2, "variation": "nearest"}),
            ("WeekOfMonth", {"weekday": 2, "week": 2}),
            "Easter",
            ("DateOffset", {"day": 4}),
            ("DateOffset", {"month": 5}),
        ],
    )
    @pytest.mark.parametrize("normalize", [True, False])
    @pytest.mark.parametrize("n", [0, 5])
    @pytest.mark.parametrize("unit", ["s", "ms", "us", "ns"])
    @pytest.mark.parametrize("tz", [None, "US/Central"])
    def test_dt64arr_add_sub_DateOffsets(
        self, box_with_array, n, normalize, cls_and_kwargs, unit, tz
    ):
        # GH#10699
        # assert vectorized operation matches pointwise operations

        if isinstance(cls_and_kwargs, tuple):
            # If cls_name param is a tuple, then 2nd entry is kwargs for
            # the offset constructor
            cls_name, kwargs = cls_and_kwargs
        else:
            cls_name = cls_and_kwargs
            kwargs = {}

        if n == 0 and cls_name in [
            "WeekOfMonth",
            "LastWeekOfMonth",
            "FY5253Quarter",
            "FY5253",
        ]:
            # passing n = 0 is invalid for these offset classes
            return

        vec = (
            DatetimeIndex(
                [
                    Timestamp("2000-01-05 00:15:00"),
                    Timestamp("2000-01-31 00:23:00"),
                    Timestamp("2000-01-01"),
                    Timestamp("2000-03-31"),
                    Timestamp("2000-02-29"),
                    Timestamp("2000-12-31"),
                    Timestamp("2000-05-15"),
                    Timestamp("2001-06-15"),
                ]
            )
            .as_unit(unit)
            .tz_localize(tz)
        )
        vec = tm.box_expected(vec, box_with_array)
        vec_items = vec.iloc[0] if box_with_array is pd.DataFrame else vec

        offset_cls = getattr(pd.offsets, cls_name)
        offset = offset_cls(n, normalize=normalize, **kwargs)

        # TODO(GH#55564): as_unit will be unnecessary
        expected = DatetimeIndex([x + offset for x in vec_items]).as_unit(unit)
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(expected, vec + offset)
        tm.assert_equal(expected, offset + vec)

        expected = DatetimeIndex([x - offset for x in vec_items]).as_unit(unit)
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(expected, vec - offset)

        expected = DatetimeIndex([offset + x for x in vec_items]).as_unit(unit)
        expected = tm.box_expected(expected, box_with_array)
        tm.assert_equal(expected, offset + vec)
        msg = "(bad|unsupported) operand type for unary"
        with pytest.raises(TypeError, match=msg):
            offset - vec

    @pytest.mark.parametrize(
        "other",
        [
            np.array([pd.offsets.MonthEnd(), pd.offsets.Day(n=2)]),
            np.array([pd.offsets.DateOffset(years=1), pd.offsets.MonthEnd()]),
            np.array(  # matching offsets
                [pd.offsets.DateOffset(years=1), pd.offsets.DateOffset(years=1)]
            ),
        ],
    )
    @pytest.mark.parametrize("op", [operator.add, roperator.radd, operator.sub])
    def test_dt64arr_add_sub_offset_array(
        self, tz_naive_fixture, box_with_array, op, other
    ):
        # GH#18849
        # GH#10699 array of offsets

        tz = tz_naive_fixture
        dti = date_range("2017-01-01", periods=2, tz=tz)
        dtarr = tm.box_expected(dti, box_with_array)

        expected = DatetimeIndex([op(dti[n], other[n]) for n in range(len(dti))])
        expected = tm.box_expected(expected, box_with_array).astype(object)

        with tm.assert_produces_warning(PerformanceWarning):
            res = op(dtarr, other)
        tm.assert_equal(res, expected)

        # Same thing but boxing other
        other = tm.box_expected(other, box_with_array)
        if box_with_array is pd.array and op is roperator.radd:
            # We expect a NumpyExtensionArray, not ndarray[object] here
            expected = pd.array(expected, dtype=object)
        with tm.assert_produces_warning(PerformanceWarning):
            res = op(dtarr, other)
        tm.assert_equal(res, expected)

    @pytest.mark.parametrize(
        "op, offset, exp, exp_freq",
        [
            (
                "__add__",
                DateOffset(months=3, days=10),
                [
                    Timestamp("2014-04-11"),
                    Timestamp("2015-04-11"),
                    Timestamp("2016-04-11"),
                    Timestamp("2017-04-11"),
                ],
                None,
            ),
            (
                "__add__",
                DateOffset(months=3),
                [
                    Timestamp("2014-04-01"),
                    Timestamp("2015-04-01"),
                    Timestamp("2016-04-01"),
                    Timestamp("2017-04-01"),
                ],
                "YS-APR",
            ),
            (
                "__sub__",
                DateOffset(months=3, days=10),
                [
                    Timestamp("2013-09-21"),
                    Timestamp("2014-09-21"),
                    Timestamp("2015-09-21"),
                    Timestamp("2016-09-21"),
                ],
                None,
            ),
            (
                "__sub__",
                DateOffset(months=3),
                [
                    Timestamp("2013-10-01"),
                    Timestamp("2014-10-01"),
                    Timestamp("2015-10-01"),
                    Timestamp("2016-10-01"),
                ],
                "YS-OCT",
            ),
        ],
    )
    def test_dti_add_sub_nonzero_mth_offset(
        self, op, offset, exp, exp_freq, tz_aware_fixture, box_with_array
    ):
        # GH 26258
        tz = tz_aware_fixture
        date = date_range(start="01 Jan 2014", end="01 Jan 2017", freq="YS", tz=tz)
        date = tm.box_expected(date, box_with_array, False)
        mth = getattr(date, op)
        result = mth(offset)

        expected = DatetimeIndex(exp, tz=tz).as_unit("ns")
        expected = tm.box_expected(expected, box_with_array, False)
        tm.assert_equal(result, expected)

    def test_dt64arr_series_add_DateOffset_with_milli(self):
        # GH 57529
        dti = DatetimeIndex(
            [
                "2000-01-01 00:00:00.012345678",
                "2000-01-31 00:00:00.012345678",
                "2000-02-29 00:00:00.012345678",
            ],
            dtype="datetime64[ns]",
        )
        result = dti + DateOffset(milliseconds=4)
        expected = DatetimeIndex(
            [
                "2000-01-01 00:00:00.016345678",
                "2000-01-31 00:00:00.016345678",
                "2000-02-29 00:00:00.016345678",
            ],
            dtype="datetime64[ns]",
        )
        tm.assert_index_equal(result, expected)

        result = dti + DateOffset(days=1, milliseconds=4)
        expected = DatetimeIndex(
            [
                "2000-01-02 00:00:00.016345678",
                "2000-02-01 00:00:00.016345678",
                "2000-03-01 00:00:00.016345678",
            ],
            dtype="datetime64[ns]",
        )
        tm.assert_index_equal(result, expected)


class TestDatetime64OverflowHandling:
    # TODO: box + de-duplicate

    def test_dt64_overflow_masking(self, box_with_array):
        # GH#25317
        left = Series([Timestamp("1969-12-31")], dtype="M8[ns]")
        right = Series([NaT])

        left = tm.box_expected(left, box_with_array)
        right = tm.box_expected(right, box_with_array)

        expected = TimedeltaIndex([NaT], dtype="m8[ns]")
        expected = tm.box_expected(expected, box_with_array)

        result = left - right
        tm.assert_equal(result, expected)

    def test_dt64_series_arith_overflow(self):
        # GH#12534, fixed by GH#19024
        dt = Timestamp("1700-01-31")
        td = Timedelta("20000 Days")
        dti = date_range("1949-09-30", freq="100YE", periods=4)
        ser = Series(dti)
        msg = "Overflow in int64 addition"
        with pytest.raises(OverflowError, match=msg):
            ser - dt
        with pytest.raises(OverflowError, match=msg):
            dt - ser
        with pytest.raises(OverflowError, match=msg):
            ser + td
        with pytest.raises(OverflowError, match=msg):
            td + ser

        ser.iloc[-1] = NaT
        expected = Series(
            ["2004-10-03", "2104-10-04", "2204-10-04", "NaT"], dtype="datetime64[ns]"
        )
        res = ser + td
        tm.assert_series_equal(res, expected)
        res = td + ser
        tm.assert_series_equal(res, expected)

        ser.iloc[1:] = NaT
        expected = Series(["91279 Days", "NaT", "NaT", "NaT"], dtype="timedelta64[ns]")
        res = ser - dt
        tm.assert_series_equal(res, expected)
        res = dt - ser
        tm.assert_series_equal(res, -expected)

    def test_datetimeindex_sub_timestamp_overflow(self):
        dtimax = pd.to_datetime(["2021-12-28 17:19", Timestamp.max]).as_unit("ns")
        dtimin = pd.to_datetime(["2021-12-28 17:19", Timestamp.min]).as_unit("ns")

        tsneg = Timestamp("1950-01-01").as_unit("ns")
        ts_neg_variants = [
            tsneg,
            tsneg.to_pydatetime(),
            tsneg.to_datetime64().astype("datetime64[ns]"),
            tsneg.to_datetime64().astype("datetime64[D]"),
        ]

        tspos = Timestamp("1980-01-01").as_unit("ns")
        ts_pos_variants = [
            tspos,
            tspos.to_pydatetime(),
            tspos.to_datetime64().astype("datetime64[ns]"),
            tspos.to_datetime64().astype("datetime64[D]"),
        ]
        msg = "Overflow in int64 addition"
        for variant in ts_neg_variants:
            with pytest.raises(OverflowError, match=msg):
                dtimax - variant

        expected = Timestamp.max._value - tspos._value
        for variant in ts_pos_variants:
            res = dtimax - variant
            assert res[1]._value == expected

        expected = Timestamp.min._value - tsneg._value
        for variant in ts_neg_variants:
            res = dtimin - variant
            assert res[1]._value == expected

        for variant in ts_pos_variants:
            with pytest.raises(OverflowError, match=msg):
                dtimin - variant

    def test_datetimeindex_sub_datetimeindex_overflow(self):
        # GH#22492, GH#22508
        dtimax = pd.to_datetime(["2021-12-28 17:19", Timestamp.max]).as_unit("ns")
        dtimin = pd.to_datetime(["2021-12-28 17:19", Timestamp.min]).as_unit("ns")

        ts_neg = pd.to_datetime(["1950-01-01", "1950-01-01"]).as_unit("ns")
        ts_pos = pd.to_datetime(["1980-01-01", "1980-01-01"]).as_unit("ns")

        # General tests
        expected = Timestamp.max._value - ts_pos[1]._value
        result = dtimax - ts_pos
        assert result[1]._value == expected

        expected = Timestamp.min._value - ts_neg[1]._value
        result = dtimin - ts_neg
        assert result[1]._value == expected
        msg = "Overflow in int64 addition"
        with pytest.raises(OverflowError, match=msg):
            dtimax - ts_neg

        with pytest.raises(OverflowError, match=msg):
            dtimin - ts_pos

        # Edge cases
        tmin = pd.to_datetime([Timestamp.min])
        t1 = tmin + Timedelta.max + Timedelta("1us")
        with pytest.raises(OverflowError, match=msg):
            t1 - tmin

        tmax = pd.to_datetime([Timestamp.max])
        t2 = tmax + Timedelta.min - Timedelta("1us")
        with pytest.raises(OverflowError, match=msg):
            tmax - t2


class TestTimestampSeriesArithmetic:
    def test_empty_series_add_sub(self, box_with_array):
        # GH#13844
        a = Series(dtype="M8[ns]")
        b = Series(dtype="m8[ns]")
        a = box_with_array(a)
        b = box_with_array(b)
        tm.assert_equal(a, a + b)
        tm.assert_equal(a, a - b)
        tm.assert_equal(a, b + a)
        msg = "cannot subtract"
        with pytest.raises(TypeError, match=msg):
            b - a

    def test_operators_datetimelike(self):
        # ## timedelta64 ###
        td1 = Series([timedelta(minutes=5, seconds=3)] * 3)
        td1.iloc[2] = np.nan

        # ## datetime64 ###
        dt1 = Series(
            [
                Timestamp("20111230"),
                Timestamp("20120101"),
                Timestamp("20120103"),
            ]
        )
        dt1.iloc[2] = np.nan
        dt2 = Series(
            [
                Timestamp("20111231"),
                Timestamp("20120102"),
                Timestamp("20120104"),
            ]
        )
        dt1 - dt2
        dt2 - dt1

        # datetime64 with timetimedelta
        dt1 + td1
        td1 + dt1
        dt1 - td1

        # timetimedelta with datetime64
        td1 + dt1
        dt1 + td1

    def test_dt64ser_sub_datetime_dtype(self, unit):
        ts = Timestamp(datetime(1993, 1, 7, 13, 30, 00))
        dt = datetime(1993, 6, 22, 13, 30)
        ser = Series([ts], dtype=f"M8[{unit}]")
        result = ser - dt

        # the expected unit is the max of `unit` and the unit imputed to `dt`,
        #  which is "us"
        exp_unit = tm.get_finest_unit(unit, "us")
        assert result.dtype == f"timedelta64[{exp_unit}]"

    # -------------------------------------------------------------
    # TODO: This next block of tests came from tests.series.test_operators,
    # needs to be de-duplicated and parametrized over `box` classes

    @pytest.mark.parametrize(
        "left, right, op_fail",
        [
            [
                [Timestamp("20111230"), Timestamp("20120101"), NaT],
                [Timestamp("20111231"), Timestamp("20120102"), Timestamp("20120104")],
                ["__sub__", "__rsub__"],
            ],
            [
                [Timestamp("20111230"), Timestamp("20120101"), NaT],
                [timedelta(minutes=5, seconds=3), timedelta(minutes=5, seconds=3), NaT],
                ["__add__", "__radd__", "__sub__"],
            ],
            [
                [
                    Timestamp("20111230", tz="US/Eastern"),
                    Timestamp("20111230", tz="US/Eastern"),
                    NaT,
                ],
                [timedelta(minutes=5, seconds=3), NaT, timedelta(minutes=5, seconds=3)],
                ["__add__", "__radd__", "__sub__"],
            ],
        ],
    )
    def test_operators_datetimelike_invalid(
        self, left, right, op_fail, all_arithmetic_operators
    ):
        # these are all TypeError ops
        op_str = all_arithmetic_operators
        arg1 = Series(left)
        arg2 = Series(right)
        # check that we are getting a TypeError
        # with 'operate' (from core/ops.py) for the ops that are not
        # defined
        op = getattr(arg1, op_str, None)
        # Previously, _validate_for_numeric_binop in core/indexes/base.py
        # did this for us.
        if op_str not in op_fail:
            with pytest.raises(
                TypeError, match="operate|[cC]annot|unsupported operand"
            ):
                op(arg2)
        else:
            # Smoke test
            op(arg2)

    def test_sub_single_tz(self, unit):
        # GH#12290
        s1 = Series([Timestamp("2016-02-10", tz="America/Sao_Paulo")]).dt.as_unit(unit)
        s2 = Series([Timestamp("2016-02-08", tz="America/Sao_Paulo")]).dt.as_unit(unit)
        result = s1 - s2
        expected = Series([Timedelta("2days")]).dt.as_unit(unit)
        tm.assert_series_equal(result, expected)
        result = s2 - s1
        expected = Series([Timedelta("-2days")]).dt.as_unit(unit)
        tm.assert_series_equal(result, expected)

    def test_dt64tz_series_sub_dtitz(self):
        # GH#19071 subtracting tzaware DatetimeIndex from tzaware Series
        # (with same tz) raises, fixed by #19024
        dti = date_range("1999-09-30", periods=10, tz="US/Pacific")
        ser = Series(dti)
        expected = Series(TimedeltaIndex(["0days"] * 10))

        res = dti - ser
        tm.assert_series_equal(res, expected)
        res = ser - dti
        tm.assert_series_equal(res, expected)

    def test_sub_datetime_compat(self, unit):
        # see GH#14088
        ser = Series([datetime(2016, 8, 23, 12, tzinfo=pytz.utc), NaT]).dt.as_unit(unit)
        dt = datetime(2016, 8, 22, 12, tzinfo=pytz.utc)
        # The datetime object has "us" so we upcast lower units
        exp_unit = tm.get_finest_unit(unit, "us")
        exp = Series([Timedelta("1 days"), NaT]).dt.as_unit(exp_unit)
        result = ser - dt
        tm.assert_series_equal(result, exp)
        result2 = ser - Timestamp(dt)
        tm.assert_series_equal(result2, exp)

    def test_dt64_series_add_mixed_tick_DateOffset(self):
        # GH#4532
        # operate with pd.offsets
        s = Series([Timestamp("20130101 9:01"), Timestamp("20130101 9:02")])

        result = s + pd.offsets.Milli(5)
        result2 = pd.offsets.Milli(5) + s
        expected = Series(
            [Timestamp("20130101 9:01:00.005"), Timestamp("20130101 9:02:00.005")]
        )
        tm.assert_series_equal(result, expected)
        tm.assert_series_equal(result2, expected)

        result = s + pd.offsets.Minute(5) + pd.offsets.Milli(5)
        expected = Series(
            [Timestamp("20130101 9:06:00.005"), Timestamp("20130101 9:07:00.005")]
        )
        tm.assert_series_equal(result, expected)

    def test_datetime64_ops_nat(self, unit):
        # GH#11349
        datetime_series = Series([NaT, Timestamp("19900315")]).dt.as_unit(unit)
        nat_series_dtype_timestamp = Series([NaT, NaT], dtype=f"datetime64[{unit}]")
        single_nat_dtype_datetime = Series([NaT], dtype=f"datetime64[{unit}]")

        # subtraction
        tm.assert_series_equal(-NaT + datetime_series, nat_series_dtype_timestamp)
        msg = "bad operand type for unary -: 'DatetimeArray'"
        with pytest.raises(TypeError, match=msg):
            -single_nat_dtype_datetime + datetime_series

        tm.assert_series_equal(
            -NaT + nat_series_dtype_timestamp, nat_series_dtype_timestamp
        )
        with pytest.raises(TypeError, match=msg):
            -single_nat_dtype_datetime + nat_series_dtype_timestamp

        # addition
        tm.assert_series_equal(
            nat_series_dtype_timestamp + NaT, nat_series_dtype_timestamp
        )
        tm.assert_series_equal(
            NaT + nat_series_dtype_timestamp, nat_series_dtype_timestamp
        )

        tm.assert_series_equal(
            nat_series_dtype_timestamp + NaT, nat_series_dtype_timestamp
        )
        tm.assert_series_equal(
            NaT + nat_series_dtype_timestamp, nat_series_dtype_timestamp
        )

    # -------------------------------------------------------------
    # Timezone-Centric Tests

    def test_operators_datetimelike_with_timezones(self):
        tz = "US/Eastern"
        dt1 = Series(date_range("2000-01-01 09:00:00", periods=5, tz=tz), name="foo")
        dt2 = dt1.copy()
        dt2.iloc[2] = np.nan

        td1 = Series(pd.timedelta_range("1 days 1 min", periods=5, freq="h"))
        td2 = td1.copy()
        td2.iloc[1] = np.nan
        assert td2._values.freq is None

        result = dt1 + td1[0]
        exp = (dt1.dt.tz_localize(None) + td1[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = dt2 + td2[0]
        exp = (dt2.dt.tz_localize(None) + td2[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        # odd numpy behavior with scalar timedeltas
        result = td1[0] + dt1
        exp = (dt1.dt.tz_localize(None) + td1[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = td2[0] + dt2
        exp = (dt2.dt.tz_localize(None) + td2[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = dt1 - td1[0]
        exp = (dt1.dt.tz_localize(None) - td1[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)
        msg = "(bad|unsupported) operand type for unary"
        with pytest.raises(TypeError, match=msg):
            td1[0] - dt1

        result = dt2 - td2[0]
        exp = (dt2.dt.tz_localize(None) - td2[0]).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)
        with pytest.raises(TypeError, match=msg):
            td2[0] - dt2

        result = dt1 + td1
        exp = (dt1.dt.tz_localize(None) + td1).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = dt2 + td2
        exp = (dt2.dt.tz_localize(None) + td2).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = dt1 - td1
        exp = (dt1.dt.tz_localize(None) - td1).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)

        result = dt2 - td2
        exp = (dt2.dt.tz_localize(None) - td2).dt.tz_localize(tz)
        tm.assert_series_equal(result, exp)
        msg = "cannot (add|subtract)"
        with pytest.raises(TypeError, match=msg):
            td1 - dt1
        with pytest.raises(TypeError, match=msg):
            td2 - dt2


class TestDatetimeIndexArithmetic:
    # -------------------------------------------------------------
    # Binary operations DatetimeIndex and TimedeltaIndex/array

    def test_dti_add_tdi(self, tz_naive_fixture):
        # GH#17558
        tz = tz_naive_fixture
        dti = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10)
        tdi = pd.timedelta_range("0 days", periods=10)
        expected = date_range("2017-01-01", periods=10, tz=tz)
        expected = expected._with_freq(None)

        # add with TimedeltaIndex
        result = dti + tdi
        tm.assert_index_equal(result, expected)

        result = tdi + dti
        tm.assert_index_equal(result, expected)

        # add with timedelta64 array
        result = dti + tdi.values
        tm.assert_index_equal(result, expected)

        result = tdi.values + dti
        tm.assert_index_equal(result, expected)

    def test_dti_iadd_tdi(self, tz_naive_fixture):
        # GH#17558
        tz = tz_naive_fixture
        dti = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10)
        tdi = pd.timedelta_range("0 days", periods=10)
        expected = date_range("2017-01-01", periods=10, tz=tz)
        expected = expected._with_freq(None)

        # iadd with TimedeltaIndex
        result = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10)
        result += tdi
        tm.assert_index_equal(result, expected)

        result = pd.timedelta_range("0 days", periods=10)
        result += dti
        tm.assert_index_equal(result, expected)

        # iadd with timedelta64 array
        result = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10)
        result += tdi.values
        tm.assert_index_equal(result, expected)

        result = pd.timedelta_range("0 days", periods=10)
        result += dti
        tm.assert_index_equal(result, expected)

    def test_dti_sub_tdi(self, tz_naive_fixture):
        # GH#17558
        tz = tz_naive_fixture
        dti = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10)
        tdi = pd.timedelta_range("0 days", periods=10)
        expected = date_range("2017-01-01", periods=10, tz=tz, freq="-1D")
        expected = expected._with_freq(None)

        # sub with TimedeltaIndex
        result = dti - tdi
        tm.assert_index_equal(result, expected)

        msg = "cannot subtract .*TimedeltaArray"
        with pytest.raises(TypeError, match=msg):
            tdi - dti

        # sub with timedelta64 array
        result = dti - tdi.values
        tm.assert_index_equal(result, expected)

        msg = "cannot subtract a datelike from a TimedeltaArray"
        with pytest.raises(TypeError, match=msg):
            tdi.values - dti

    def test_dti_isub_tdi(self, tz_naive_fixture, unit):
        # GH#17558
        tz = tz_naive_fixture
        dti = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10).as_unit(unit)
        tdi = pd.timedelta_range("0 days", periods=10, unit=unit)
        expected = date_range("2017-01-01", periods=10, tz=tz, freq="-1D", unit=unit)
        expected = expected._with_freq(None)

        # isub with TimedeltaIndex
        result = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10).as_unit(unit)
        result -= tdi
        tm.assert_index_equal(result, expected)

        # DTA.__isub__ GH#43904
        dta = dti._data.copy()
        dta -= tdi
        tm.assert_datetime_array_equal(dta, expected._data)

        out = dti._data.copy()
        np.subtract(out, tdi, out=out)
        tm.assert_datetime_array_equal(out, expected._data)

        msg = "cannot subtract a datelike from a TimedeltaArray"
        with pytest.raises(TypeError, match=msg):
            tdi -= dti

        # isub with timedelta64 array
        result = DatetimeIndex([Timestamp("2017-01-01", tz=tz)] * 10).as_unit(unit)
        result -= tdi.values
        tm.assert_index_equal(result, expected)

        with pytest.raises(TypeError, match=msg):
            tdi.values -= dti

        with pytest.raises(TypeError, match=msg):
            tdi._values -= dti

    # -------------------------------------------------------------
    # Binary Operations DatetimeIndex and datetime-like
    # TODO: A couple other tests belong in this section.  Move them in
    # A PR where there isn't already a giant diff.

    # -------------------------------------------------------------

    def test_dta_add_sub_index(self, tz_naive_fixture):
        # Check that DatetimeArray defers to Index classes
        dti = date_range("20130101", periods=3, tz=tz_naive_fixture)
        dta = dti.array
        result = dta - dti
        expected = dti - dti
        tm.assert_index_equal(result, expected)

        tdi = result
        result = dta + tdi
        expected = dti + tdi
        tm.assert_index_equal(result, expected)

        result = dta - tdi
        expected = dti - tdi
        tm.assert_index_equal(result, expected)

    def test_sub_dti_dti(self, unit):
        # previously performed setop (deprecated in 0.16.0), now changed to
        # return subtraction -> TimeDeltaIndex (GH ...)

        dti = date_range("20130101", periods=3, unit=unit)
        dti_tz = date_range("20130101", periods=3, unit=unit).tz_localize("US/Eastern")
        expected = TimedeltaIndex([0, 0, 0]).as_unit(unit)

        result = dti - dti
        tm.assert_index_equal(result, expected)

        result = dti_tz - dti_tz
        tm.assert_index_equal(result, expected)
        msg = "Cannot subtract tz-naive and tz-aware datetime-like objects"
        with pytest.raises(TypeError, match=msg):
            dti_tz - dti

        with pytest.raises(TypeError, match=msg):
            dti - dti_tz

        # isub
        dti -= dti
        tm.assert_index_equal(dti, expected)

        # different length raises ValueError
        dti1 = date_range("20130101", periods=3, unit=unit)
        dti2 = date_range("20130101", periods=4, unit=unit)
        msg = "cannot add indices of unequal length"
        with pytest.raises(ValueError, match=msg):
            dti1 - dti2

        # NaN propagation
        dti1 = DatetimeIndex(["2012-01-01", np.nan, "2012-01-03"]).as_unit(unit)
        dti2 = DatetimeIndex(["2012-01-02", "2012-01-03", np.nan]).as_unit(unit)
        expected = TimedeltaIndex(["1 days", np.nan, np.nan]).as_unit(unit)
        result = dti2 - dti1
        tm.assert_index_equal(result, expected)

    # -------------------------------------------------------------------
    # TODO: Most of this block is moved from series or frame tests, needs
    # cleanup, box-parametrization, and de-duplication

    @pytest.mark.parametrize("op", [operator.add, operator.sub])
    def test_timedelta64_equal_timedelta_supported_ops(self, op, box_with_array):
        ser = Series(
            [
                Timestamp("20130301"),
                Timestamp("20130228 23:00:00"),
                Timestamp("20130228 22:00:00"),
                Timestamp("20130228 21:00:00"),
            ]
        )
        obj = box_with_array(ser)

        intervals = ["D", "h", "m", "s", "us"]

        def timedelta64(*args):
            # see casting notes in NumPy gh-12927
            return np.sum(list(starmap(np.timedelta64, zip(args, intervals))))

        for d, h, m, s, us in product(*([range(2)] * 5)):
            nptd = timedelta64(d, h, m, s, us)
            pytd = timedelta(days=d, hours=h, minutes=m, seconds=s, microseconds=us)
            lhs = op(obj, nptd)
            rhs = op(obj, pytd)

            tm.assert_equal(lhs, rhs)

    def test_ops_nat_mixed_datetime64_timedelta64(self):
        # GH#11349
        timedelta_series = Series([NaT, Timedelta("1s")])
        datetime_series = Series([NaT, Timestamp("19900315")])
        nat_series_dtype_timedelta = Series([NaT, NaT], dtype="timedelta64[ns]")
        nat_series_dtype_timestamp = Series([NaT, NaT], dtype="datetime64[ns]")
        single_nat_dtype_datetime = Series([NaT], dtype="datetime64[ns]")
        single_nat_dtype_timedelta = Series([NaT], dtype="timedelta64[ns]")

        # subtraction
        tm.assert_series_equal(
            datetime_series - single_nat_dtype_datetime, nat_series_dtype_timedelta
        )

        tm.assert_series_equal(
            datetime_series - single_nat_dtype_timedelta, nat_series_dtype_timestamp
        )
        tm.assert_series_equal(
            -single_nat_dtype_timedelta + datetime_series, nat_series_dtype_timestamp
        )

        # without a Series wrapping the NaT, it is ambiguous
        # whether it is a datetime64 or timedelta64
        # defaults to interpreting it as timedelta64
        tm.assert_series_equal(
            nat_series_dtype_timestamp - single_nat_dtype_datetime,
            nat_series_dtype_timedelta,
        )

        tm.assert_series_equal(
            nat_series_dtype_timestamp - single_nat_dtype_timedelta,
            nat_series_dtype_timestamp,
        )
        tm.assert_series_equal(
            -single_nat_dtype_timedelta + nat_series_dtype_timestamp,
            nat_series_dtype_timestamp,
        )
        msg = "cannot subtract a datelike"
        with pytest.raises(TypeError, match=msg):
            timedelta_series - single_nat_dtype_datetime

        # addition
        tm.assert_series_equal(
            nat_series_dtype_timestamp + single_nat_dtype_timedelta,
            nat_series_dtype_timestamp,
        )
        tm.assert_series_equal(
            single_nat_dtype_timedelta + nat_series_dtype_timestamp,
            nat_series_dtype_timestamp,
        )

        tm.assert_series_equal(
            nat_series_dtype_timestamp + single_nat_dtype_timedelta,
            nat_series_dtype_timestamp,
        )
        tm.assert_series_equal(
            single_nat_dtype_timedelta + nat_series_dtype_timestamp,
            nat_series_dtype_timestamp,
        )

        tm.assert_series_equal(
            nat_series_dtype_timedelta + single_nat_dtype_datetime,
            nat_series_dtype_timestamp,
        )
        tm.assert_series_equal(
            single_nat_dtype_datetime + nat_series_dtype_timedelta,
            nat_series_dtype_timestamp,
        )

    def test_ufunc_coercions(self, unit):
        idx = date_range("2011-01-01", periods=3, freq="2D", name="x", unit=unit)

        delta = np.timedelta64(1, "D")
        exp = date_range("2011-01-02", periods=3, freq="2D", name="x", unit=unit)
        for result in [idx + delta, np.add(idx, delta)]:
            assert isinstance(result, DatetimeIndex)
            tm.assert_index_equal(result, exp)
            assert result.freq == "2D"

        exp = date_range("2010-12-31", periods=3, freq="2D", name="x", unit=unit)

        for result in [idx - delta, np.subtract(idx, delta)]:
            assert isinstance(result, DatetimeIndex)
            tm.assert_index_equal(result, exp)
            assert result.freq == "2D"

        # When adding/subtracting an ndarray (which has no .freq), the result
        #  does not infer freq
        idx = idx._with_freq(None)
        delta = np.array(
            [np.timedelta64(1, "D"), np.timedelta64(2, "D"), np.timedelta64(3, "D")]
        )
        exp = DatetimeIndex(
            ["2011-01-02", "2011-01-05", "2011-01-08"], name="x"
        ).as_unit(unit)

        for result in [idx + delta, np.add(idx, delta)]:
            tm.assert_index_equal(result, exp)
            assert result.freq == exp.freq

        exp = DatetimeIndex(
            ["2010-12-31", "2011-01-01", "2011-01-02"], name="x"
        ).as_unit(unit)
        for result in [idx - delta, np.subtract(idx, delta)]:
            assert isinstance(result, DatetimeIndex)
            tm.assert_index_equal(result, exp)
            assert result.freq == exp.freq

    def test_dti_add_series(self, tz_naive_fixture, names):
        # GH#13905
        tz = tz_naive_fixture
        index = DatetimeIndex(
            ["2016-06-28 05:30", "2016-06-28 05:31"], tz=tz, name=names[0]
        ).as_unit("ns")
        ser = Series([Timedelta(seconds=5)] * 2, index=index, name=names[1])
        expected = Series(index + Timedelta(seconds=5), index=index, name=names[2])

        # passing name arg isn't enough when names[2] is None
        expected.name = names[2]
        assert expected.dtype == index.dtype
        result = ser + index
        tm.assert_series_equal(result, expected)
        result2 = index + ser
        tm.assert_series_equal(result2, expected)

        expected = index + Timedelta(seconds=5)
        result3 = ser.values + index
        tm.assert_index_equal(result3, expected)
        result4 = index + ser.values
        tm.assert_index_equal(result4, expected)

    @pytest.mark.parametrize("op", [operator.add, roperator.radd, operator.sub])
    def test_dti_addsub_offset_arraylike(
        self, tz_naive_fixture, names, op, index_or_series
    ):
        # GH#18849, GH#19744
        other_box = index_or_series

        tz = tz_naive_fixture
        dti = date_range("2017-01-01", periods=2, tz=tz, name=names[0])
        other = other_box([pd.offsets.MonthEnd(), pd.offsets.Day(n=2)], name=names[1])

        xbox = get_upcast_box(dti, other)

        with tm.assert_produces_warning(PerformanceWarning):
            res = op(dti, other)

        expected = DatetimeIndex(
            [op(dti[n], other[n]) for n in range(len(dti))], name=names[2], freq="infer"
        )
        expected = tm.box_expected(expected, xbox).astype(object)
        tm.assert_equal(res, expected)

    @pytest.mark.parametrize("other_box", [pd.Index, np.array])
    def test_dti_addsub_object_arraylike(
        self, tz_naive_fixture, box_with_array, other_box
    ):
        tz = tz_naive_fixture

        dti = date_range("2017-01-01", periods=2, tz=tz)
        dtarr = tm.box_expected(dti, box_with_array)
        other = other_box([pd.offsets.MonthEnd(), Timedelta(days=4)])
        xbox = get_upcast_box(dtarr, other)

        expected = DatetimeIndex(["2017-01-31", "2017-01-06"], tz=tz_naive_fixture)
        expected = tm.box_expected(expected, xbox).astype(object)

        with tm.assert_produces_warning(PerformanceWarning):
            result = dtarr + other
        tm.assert_equal(result, expected)

        expected = DatetimeIndex(["2016-12-31", "2016-12-29"], tz=tz_naive_fixture)
        expected = tm.box_expected(expected, xbox).astype(object)

        with tm.assert_produces_warning(PerformanceWarning):
            result = dtarr - other
        tm.assert_equal(result, expected)


@pytest.mark.parametrize("years", [-1, 0, 1])
@pytest.mark.parametrize("months", [-2, 0, 2])
@pytest.mark.parametrize("unit", ["s", "ms", "us", "ns"])
def test_shift_months(years, months, unit):
    dti = DatetimeIndex(
        [
            Timestamp("2000-01-05 00:15:00"),
            Timestamp("2000-01-31 00:23:00"),
            Timestamp("2000-01-01"),
            Timestamp("2000-02-29"),
            Timestamp("2000-12-31"),
        ]
    ).as_unit(unit)
    shifted = shift_months(dti.asi8, years * 12 + months, reso=dti._data._creso)
    shifted_dt64 = shifted.view(f"M8[{dti.unit}]")
    actual = DatetimeIndex(shifted_dt64)

    raw = [x + pd.offsets.DateOffset(years=years, months=months) for x in dti]
    expected = DatetimeIndex(raw).as_unit(dti.unit)
    tm.assert_index_equal(actual, expected)


def test_dt64arr_addsub_object_dtype_2d():
    # block-wise DataFrame operations will require operating on 2D
    #  DatetimeArray/TimedeltaArray, so check that specifically.
    dti = date_range("1994-02-13", freq="2W", periods=4)
    dta = dti._data.reshape((4, 1))

    other = np.array([[pd.offsets.Day(n)] for n in range(4)])
    assert other.shape == dta.shape

    with tm.assert_produces_warning(PerformanceWarning):
        result = dta + other
    with tm.assert_produces_warning(PerformanceWarning):
        expected = (dta[:, 0] + other[:, 0]).reshape(-1, 1)

    tm.assert_numpy_array_equal(result, expected)

    with tm.assert_produces_warning(PerformanceWarning):
        # Case where we expect to get a TimedeltaArray back
        result2 = dta - dta.astype(object)

    assert result2.shape == (4, 1)
    assert all(td._value == 0 for td in result2.ravel())


def test_non_nano_dt64_addsub_np_nat_scalars():
    # GH 52295
    ser = Series([1233242342344, 232432434324, 332434242344], dtype="datetime64[ms]")
    result = ser - np.datetime64("nat", "ms")
    expected = Series([NaT] * 3, dtype="timedelta64[ms]")
    tm.assert_series_equal(result, expected)

    result = ser + np.timedelta64("nat", "ms")
    expected = Series([NaT] * 3, dtype="datetime64[ms]")
    tm.assert_series_equal(result, expected)


def test_non_nano_dt64_addsub_np_nat_scalars_unitless():
    # GH 52295
    # TODO: Can we default to the ser unit?
    ser = Series([1233242342344, 232432434324, 332434242344], dtype="datetime64[ms]")
    result = ser - np.datetime64("nat")
    expected = Series([NaT] * 3, dtype="timedelta64[ns]")
    tm.assert_series_equal(result, expected)

    result = ser + np.timedelta64("nat")
    expected = Series([NaT] * 3, dtype="datetime64[ns]")
    tm.assert_series_equal(result, expected)


def test_non_nano_dt64_addsub_np_nat_scalars_unsupported_unit():
    # GH 52295
    ser = Series([12332, 23243, 33243], dtype="datetime64[s]")
    result = ser - np.datetime64("nat", "D")
    expected = Series([NaT] * 3, dtype="timedelta64[s]")
    tm.assert_series_equal(result, expected)

    result = ser + np.timedelta64("nat", "D")
    expected = Series([NaT] * 3, dtype="datetime64[s]")
    tm.assert_series_equal(result, expected)

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\numpy\f2py\_backends\_distutils.py ===
import os
import shutil
import sys
import warnings

from numpy.distutils.core import Extension, setup
from numpy.distutils.misc_util import dict_append
from numpy.distutils.system_info import get_info
from numpy.exceptions import VisibleDeprecationWarning

from ._backend import Backend


class DistutilsBackend(Backend):
    def __init__(sef, *args, **kwargs):
        warnings.warn(
            "\ndistutils has been deprecated since NumPy 1.26.x\n"
            "Use the Meson backend instead, or generate wrappers"
            " without -c and use a custom build script",
            VisibleDeprecationWarning,
            stacklevel=2,
        )
        super().__init__(*args, **kwargs)

    def compile(self):
        num_info = {}
        if num_info:
            self.include_dirs.extend(num_info.get("include_dirs", []))
        ext_args = {
            "name": self.modulename,
            "sources": self.sources,
            "include_dirs": self.include_dirs,
            "library_dirs": self.library_dirs,
            "libraries": self.libraries,
            "define_macros": self.define_macros,
            "undef_macros": self.undef_macros,
            "extra_objects": self.extra_objects,
            "f2py_options": self.f2py_flags,
        }

        if self.sysinfo_flags:
            for n in self.sysinfo_flags:
                i = get_info(n)
                if not i:
                    print(
                        f"No {n!r} resources found"
                        "in system (try `f2py --help-link`)"
                    )
                dict_append(ext_args, **i)

        ext = Extension(**ext_args)

        sys.argv = [sys.argv[0]] + self.setup_flags
        sys.argv.extend(
            [
                "build",
                "--build-temp",
                self.build_dir,
                "--build-base",
                self.build_dir,
                "--build-platlib",
                ".",
                "--disable-optimization",
            ]
        )

        if self.fc_flags:
            sys.argv.extend(["config_fc"] + self.fc_flags)
        if self.flib_flags:
            sys.argv.extend(["build_ext"] + self.flib_flags)

        setup(ext_modules=[ext])

        if self.remove_build_dir and os.path.exists(self.build_dir):
            print(f"Removing build directory {self.build_dir}")
            shutil.rmtree(self.build_dir)

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\transformers\quantize_helper.py ===
# -------------------------------------------------------------------------
# Copyright (c) Microsoft Corporation.  All rights reserved.
# Licensed under the MIT License.  See License.txt in the project root for
# license information.
# --------------------------------------------------------------------------

import logging
import os

import onnx
import torch
from transformers.modeling_utils import Conv1D

logger = logging.getLogger(__name__)


def _conv1d_to_linear(module):
    in_size, out_size = module.weight.shape
    linear = torch.nn.Linear(in_size, out_size)
    linear.weight.data = module.weight.data.T.contiguous()
    linear.bias.data = module.bias.data
    return linear


def conv1d_to_linear(model):
    """in-place
    This is for Dynamic Quantization, as Conv1D is not recognized by PyTorch, convert it to nn.Linear
    """
    logger.debug("replace Conv1D with Linear")
    for name in list(model._modules):
        module = model._modules[name]
        if isinstance(module, Conv1D):
            linear = _conv1d_to_linear(module)
            model._modules[name] = linear
        else:
            conv1d_to_linear(module)


def _get_size_of_pytorch_model(model):
    torch.save(model.state_dict(), "temp.p")
    size = os.path.getsize("temp.p") / (1024 * 1024)
    os.remove("temp.p")
    return size


class QuantizeHelper:
    @staticmethod
    def quantize_torch_model(model, dtype=torch.qint8):
        """
        Usage: model = quantize_model(model)

        TODO: mix of in-place and return, but results are different
        """
        conv1d_to_linear(model)
        quantized_model = torch.quantization.quantize_dynamic(model, {torch.nn.Linear}, dtype=dtype)
        logger.info(f"Size of full precision Torch model(MB):{_get_size_of_pytorch_model(model)}")
        logger.info(f"Size of quantized Torch model(MB):{_get_size_of_pytorch_model(quantized_model)}")
        return quantized_model

    @staticmethod
    def quantize_onnx_model(onnx_model_path, quantized_model_path, use_external_data_format=False):
        from pathlib import Path

        from onnxruntime.quantization import quantize_dynamic

        Path(quantized_model_path).parent.mkdir(parents=True, exist_ok=True)
        logger.info(f"Size of full precision ONNX model(MB):{os.path.getsize(onnx_model_path) / (1024 * 1024)}")
        quantize_dynamic(
            onnx_model_path,
            quantized_model_path,
            use_external_data_format=use_external_data_format,
            extra_options={"DefaultTensorType": onnx.TensorProto.FLOAT},
        )
        logger.info(f"quantized model saved to:{quantized_model_path}")
        # TODO: inlcude external data in total model size.
        logger.info(f"Size of quantized ONNX model(MB):{os.path.getsize(quantized_model_path) / (1024 * 1024)}")

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\transformers\models\longformer\longformer_helper.py ===
# -------------------------------------------------------------------------
# Copyright (c) Microsoft Corporation.  All rights reserved.
# Licensed under the MIT License.  See License.txt in the project root for
# license information.
# --------------------------------------------------------------------------
# This script helps creating dummy inputs for Longformer model.

import logging

import numpy
import torch

logger = logging.getLogger(__name__)

PRETRAINED_LONGFORMER_MODELS = {
    "longformer-base-4096": "allenai/longformer-base-4096",
    "longformer-large-4096": "allenai/longformer-large-4096",
    "longformer-random-tiny": "patrickvonplaten/longformer-random-tiny",  # A tiny model for debugging
}


class LongformerInputs:
    def __init__(self, input_ids, attention_mask, global_attention_mask):
        self.input_ids: torch.LongTensor = input_ids
        self.attention_mask: torch.FloatTensor | torch.HalfTensor = attention_mask
        self.global_attention_mask: torch.FloatTensor | torch.HalfTensor = global_attention_mask

    def to_list(self) -> list:
        return [v for v in [self.input_ids, self.attention_mask, self.global_attention_mask] if v is not None]

    def to_tuple(self) -> tuple:
        return tuple(v for v in self.to_list())

    def get_ort_inputs(self) -> dict:
        return {
            "input_ids": numpy.ascontiguousarray(self.input_ids.cpu().numpy()),
            "attention_mask": numpy.ascontiguousarray(self.attention_mask.cpu().numpy()),
            "global_attention_mask": numpy.ascontiguousarray(self.global_attention_mask.cpu().numpy()),
        }


class LongformerHelper:
    """A helper class for Longformer model conversion, inference and verification."""

    @staticmethod
    def get_dummy_inputs(
        batch_size: int,
        sequence_length: int,
        num_global_tokens: int,
        device: torch.device,
        vocab_size: int = 100,
    ) -> LongformerInputs:
        """Create random inputs for Longformer model.
        Returns torch tensors of input_ids, attention_mask and global_attention_mask tensors.
        """

        input_ids = torch.randint(
            low=0,
            high=vocab_size - 1,
            size=(batch_size, sequence_length),
            dtype=torch.long,
            device=device,
        )
        attention_mask = torch.ones(input_ids.shape, dtype=torch.long, device=device)
        global_attention_mask = torch.zeros(input_ids.shape, dtype=torch.long, device=device)
        global_token_index = list(range(num_global_tokens))
        global_attention_mask[:, global_token_index] = 1
        return LongformerInputs(input_ids, attention_mask, global_attention_mask)

    @staticmethod
    def get_output_shapes(batch_size: int, sequence_length: int, hidden_size: int) -> dict[str, list[int]]:
        """Returns a dictionary with output name as key, and shape as value."""
        return {
            "last_state": [batch_size, sequence_length, hidden_size],
            "pooler": [batch_size, sequence_length],
        }

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\openpyxl\chart\title.py ===
# Copyright (c) 2010-2024 openpyxl

from openpyxl.descriptors.serialisable import Serialisable
from openpyxl.descriptors import (
    Typed,
    Alias,
)

from openpyxl.descriptors.excel import ExtensionList
from openpyxl.descriptors.nested import NestedBool

from .text import Text, RichText
from .layout import Layout
from .shapes import GraphicalProperties

from openpyxl.drawing.text import (
    Paragraph,
    RegularTextRun,
    LineBreak,
    ParagraphProperties,
    CharacterProperties,
)


class Title(Serialisable):
    tagname = "title"

    tx = Typed(expected_type=Text, allow_none=True)
    text = Alias('tx')
    layout = Typed(expected_type=Layout, allow_none=True)
    overlay = NestedBool(allow_none=True)
    spPr = Typed(expected_type=GraphicalProperties, allow_none=True)
    graphicalProperties = Alias('spPr')
    txPr = Typed(expected_type=RichText, allow_none=True)
    body = Alias('txPr')
    extLst = Typed(expected_type=ExtensionList, allow_none=True)

    __elements__ = ('tx', 'layout', 'overlay', 'spPr', 'txPr')

    def __init__(self,
                 tx=None,
                 layout=None,
                 overlay=None,
                 spPr=None,
                 txPr=None,
                 extLst=None,
                ):
        if tx is None:
            tx = Text()
        self.tx = tx
        self.layout = layout
        self.overlay = overlay
        self.spPr = spPr
        self.txPr = txPr



def title_maker(text):
    title = Title()
    paraprops = ParagraphProperties()
    paraprops.defRPr = CharacterProperties()
    paras = [Paragraph(r=[RegularTextRun(t=s)], pPr=paraprops) for s in text.split("\n")]

    title.tx.rich.paragraphs = paras
    return title


class TitleDescriptor(Typed):

    expected_type = Title
    allow_none = True

    def __set__(self, instance, value):
        if isinstance(value, str):
            value = title_maker(value)
        super().__set__(instance, value)

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pip\_vendor\rich\_export_format.py ===
CONSOLE_HTML_FORMAT = """\
<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<style>
{stylesheet}
body {{
    color: {foreground};
    background-color: {background};
}}
</style>
</head>
<body>
    <pre style="font-family:Menlo,'DejaVu Sans Mono',consolas,'Courier New',monospace"><code style="font-family:inherit">{code}</code></pre>
</body>
</html>
"""

CONSOLE_SVG_FORMAT = """\
<svg class="rich-terminal" viewBox="0 0 {width} {height}" xmlns="http://www.w3.org/2000/svg">
    <!-- Generated with Rich https://www.textualize.io -->
    <style>

    @font-face {{
        font-family: "Fira Code";
        src: local("FiraCode-Regular"),
                url("https://cdnjs.cloudflare.com/ajax/libs/firacode/6.2.0/woff2/FiraCode-Regular.woff2") format("woff2"),
                url("https://cdnjs.cloudflare.com/ajax/libs/firacode/6.2.0/woff/FiraCode-Regular.woff") format("woff");
        font-style: normal;
        font-weight: 400;
    }}
    @font-face {{
        font-family: "Fira Code";
        src: local("FiraCode-Bold"),
                url("https://cdnjs.cloudflare.com/ajax/libs/firacode/6.2.0/woff2/FiraCode-Bold.woff2") format("woff2"),
                url("https://cdnjs.cloudflare.com/ajax/libs/firacode/6.2.0/woff/FiraCode-Bold.woff") format("woff");
        font-style: bold;
        font-weight: 700;
    }}

    .{unique_id}-matrix {{
        font-family: Fira Code, monospace;
        font-size: {char_height}px;
        line-height: {line_height}px;
        font-variant-east-asian: full-width;
    }}

    .{unique_id}-title {{
        font-size: 18px;
        font-weight: bold;
        font-family: arial;
    }}

    {styles}
    </style>

    <defs>
    <clipPath id="{unique_id}-clip-terminal">
      <rect x="0" y="0" width="{terminal_width}" height="{terminal_height}" />
    </clipPath>
    {lines}
    </defs>

    {chrome}
    <g transform="translate({terminal_x}, {terminal_y})" clip-path="url(#{unique_id}-clip-terminal)">
    {backgrounds}
    <g class="{unique_id}-matrix">
    {matrix}
    </g>
    </g>
</svg>
"""

_SVG_FONT_FAMILY = "Rich Fira Code"
_SVG_CLASSES_PREFIX = "rich-svg"

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\selenium\webdriver\common\devtools\v135\inspector.py ===
# DO NOT EDIT THIS FILE!
#
# This file is generated from the CDP specification. If you need to make
# changes, edit the generator and regenerate all of the modules.
#
# CDP domain: Inspector (experimental)
from __future__ import annotations
from .util import event_class, T_JSON_DICT
from dataclasses import dataclass
import enum
import typing

def disable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Disables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.disable',
    }
    json = yield cmd_dict


def enable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Enables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.enable',
    }
    json = yield cmd_dict


@event_class('Inspector.detached')
@dataclass
class Detached:
    '''
    Fired when remote debugging connection is about to be terminated. Contains detach reason.
    '''
    #: The reason why connection has been terminated.
    reason: str

    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> Detached:
        return cls(
            reason=str(json['reason'])
        )


@event_class('Inspector.targetCrashed')
@dataclass
class TargetCrashed:
    '''
    Fired when debugging target has crashed
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetCrashed:
        return cls(

        )


@event_class('Inspector.targetReloadedAfterCrash')
@dataclass
class TargetReloadedAfterCrash:
    '''
    Fired when debugging target has reloaded after crash
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetReloadedAfterCrash:
        return cls(

        )

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\selenium\webdriver\common\devtools\v136\inspector.py ===
# DO NOT EDIT THIS FILE!
#
# This file is generated from the CDP specification. If you need to make
# changes, edit the generator and regenerate all of the modules.
#
# CDP domain: Inspector (experimental)
from __future__ import annotations
from .util import event_class, T_JSON_DICT
from dataclasses import dataclass
import enum
import typing

def disable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Disables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.disable',
    }
    json = yield cmd_dict


def enable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Enables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.enable',
    }
    json = yield cmd_dict


@event_class('Inspector.detached')
@dataclass
class Detached:
    '''
    Fired when remote debugging connection is about to be terminated. Contains detach reason.
    '''
    #: The reason why connection has been terminated.
    reason: str

    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> Detached:
        return cls(
            reason=str(json['reason'])
        )


@event_class('Inspector.targetCrashed')
@dataclass
class TargetCrashed:
    '''
    Fired when debugging target has crashed
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetCrashed:
        return cls(

        )


@event_class('Inspector.targetReloadedAfterCrash')
@dataclass
class TargetReloadedAfterCrash:
    '''
    Fired when debugging target has reloaded after crash
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetReloadedAfterCrash:
        return cls(

        )

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\selenium\webdriver\common\devtools\v137\inspector.py ===
# DO NOT EDIT THIS FILE!
#
# This file is generated from the CDP specification. If you need to make
# changes, edit the generator and regenerate all of the modules.
#
# CDP domain: Inspector (experimental)
from __future__ import annotations
from .util import event_class, T_JSON_DICT
from dataclasses import dataclass
import enum
import typing

def disable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Disables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.disable',
    }
    json = yield cmd_dict


def enable() -> typing.Generator[T_JSON_DICT,T_JSON_DICT,None]:
    '''
    Enables inspector domain notifications.
    '''
    cmd_dict: T_JSON_DICT = {
        'method': 'Inspector.enable',
    }
    json = yield cmd_dict


@event_class('Inspector.detached')
@dataclass
class Detached:
    '''
    Fired when remote debugging connection is about to be terminated. Contains detach reason.
    '''
    #: The reason why connection has been terminated.
    reason: str

    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> Detached:
        return cls(
            reason=str(json['reason'])
        )


@event_class('Inspector.targetCrashed')
@dataclass
class TargetCrashed:
    '''
    Fired when debugging target has crashed
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetCrashed:
        return cls(

        )


@event_class('Inspector.targetReloadedAfterCrash')
@dataclass
class TargetReloadedAfterCrash:
    '''
    Fired when debugging target has reloaded after crash
    '''


    @classmethod
    def from_json(cls, json: T_JSON_DICT) -> TargetReloadedAfterCrash:
        return cls(

        )

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sqlalchemy\types.py ===
# types.py
# Copyright (C) 2005-2025 the SQLAlchemy authors and contributors
# <see AUTHORS file>
#
# This module is part of SQLAlchemy and is released under
# the MIT License: https://www.opensource.org/licenses/mit-license.php

"""Compatibility namespace for sqlalchemy.sql.types.

"""


from __future__ import annotations

from .sql.sqltypes import _Binary as _Binary
from .sql.sqltypes import ARRAY as ARRAY
from .sql.sqltypes import BIGINT as BIGINT
from .sql.sqltypes import BigInteger as BigInteger
from .sql.sqltypes import BINARY as BINARY
from .sql.sqltypes import BLOB as BLOB
from .sql.sqltypes import BOOLEAN as BOOLEAN
from .sql.sqltypes import Boolean as Boolean
from .sql.sqltypes import CHAR as CHAR
from .sql.sqltypes import CLOB as CLOB
from .sql.sqltypes import Concatenable as Concatenable
from .sql.sqltypes import DATE as DATE
from .sql.sqltypes import Date as Date
from .sql.sqltypes import DATETIME as DATETIME
from .sql.sqltypes import DateTime as DateTime
from .sql.sqltypes import DECIMAL as DECIMAL
from .sql.sqltypes import DOUBLE as DOUBLE
from .sql.sqltypes import Double as Double
from .sql.sqltypes import DOUBLE_PRECISION as DOUBLE_PRECISION
from .sql.sqltypes import Enum as Enum
from .sql.sqltypes import FLOAT as FLOAT
from .sql.sqltypes import Float as Float
from .sql.sqltypes import Indexable as Indexable
from .sql.sqltypes import INT as INT
from .sql.sqltypes import INTEGER as INTEGER
from .sql.sqltypes import Integer as Integer
from .sql.sqltypes import Interval as Interval
from .sql.sqltypes import JSON as JSON
from .sql.sqltypes import LargeBinary as LargeBinary
from .sql.sqltypes import MatchType as MatchType
from .sql.sqltypes import NCHAR as NCHAR
from .sql.sqltypes import NULLTYPE as NULLTYPE
from .sql.sqltypes import NullType as NullType
from .sql.sqltypes import NUMERIC as NUMERIC
from .sql.sqltypes import Numeric as Numeric
from .sql.sqltypes import NVARCHAR as NVARCHAR
from .sql.sqltypes import PickleType as PickleType
from .sql.sqltypes import REAL as REAL
from .sql.sqltypes import SchemaType as SchemaType
from .sql.sqltypes import SMALLINT as SMALLINT
from .sql.sqltypes import SmallInteger as SmallInteger
from .sql.sqltypes import String as String
from .sql.sqltypes import STRINGTYPE as STRINGTYPE
from .sql.sqltypes import TEXT as TEXT
from .sql.sqltypes import Text as Text
from .sql.sqltypes import TIME as TIME
from .sql.sqltypes import Time as Time
from .sql.sqltypes import TIMESTAMP as TIMESTAMP
from .sql.sqltypes import TupleType as TupleType
from .sql.sqltypes import Unicode as Unicode
from .sql.sqltypes import UnicodeText as UnicodeText
from .sql.sqltypes import UUID as UUID
from .sql.sqltypes import Uuid as Uuid
from .sql.sqltypes import VARBINARY as VARBINARY
from .sql.sqltypes import VARCHAR as VARCHAR
from .sql.type_api import adapt_type as adapt_type
from .sql.type_api import ExternalType as ExternalType
from .sql.type_api import to_instance as to_instance
from .sql.type_api import TypeDecorator as TypeDecorator
from .sql.type_api import TypeEngine as TypeEngine
from .sql.type_api import UserDefinedType as UserDefinedType
from .sql.type_api import Variant as Variant

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sympy\utilities\memoization.py ===
from functools import wraps


def recurrence_memo(initial):
    """
    Memo decorator for sequences defined by recurrence

    Examples
    ========

    >>> from sympy.utilities.memoization import recurrence_memo
    >>> @recurrence_memo([1]) # 0! = 1
    ... def factorial(n, prev):
    ...     return n * prev[-1]
    >>> factorial(4)
    24
    >>> factorial(3) # use cache values
    6
    >>> factorial.cache_length() # cache length can be obtained
    5
    >>> factorial.fetch_item(slice(2, 4))
    [2, 6]

    """
    cache = initial

    def decorator(f):
        @wraps(f)
        def g(n):
            L = len(cache)
            if n < L:
                return cache[n]
            for i in range(L, n + 1):
                cache.append(f(i, cache))
            return cache[-1]
        g.cache_length = lambda: len(cache)
        g.fetch_item = lambda x: cache[x]
        return g
    return decorator


def assoc_recurrence_memo(base_seq):
    """
    Memo decorator for associated sequences defined by recurrence starting from base

    base_seq(n) -- callable to get base sequence elements

    XXX works only for Pn0 = base_seq(0) cases
    XXX works only for m <= n cases
    """

    cache = []

    def decorator(f):
        @wraps(f)
        def g(n, m):
            L = len(cache)
            if n < L:
                return cache[n][m]

            for i in range(L, n + 1):
                # get base sequence
                F_i0 = base_seq(i)
                F_i_cache = [F_i0]
                cache.append(F_i_cache)

                # XXX only works for m <= n cases
                # generate assoc sequence
                for j in range(1, i + 1):
                    F_ij = f(i, j, cache)
                    F_i_cache.append(F_ij)

            return cache[n][m]

        return g
    return decorator

# === atelier-kyo-manager/project-root\scripts\fetch_warehouses.py ===
import asyncio, logging, re, sys
from pathlib import Path
import yaml
from playwright.async_api import async_playwright

# ────────── 設定 ──────────
HEADLESS = True
TIMEOUT  = 60_000
ROOT     = Path(__file__).resolve().parents[1]
CONF     = ROOT / "configs"
STORAGE  = CONF / "storage.json"
OUT      = CONF / "warehouses.yaml"
CONF.mkdir(exist_ok=True)

URL_LIST = "https://www.buyandship.co.jp/account/v2020/warehouse/"
STOPPED  = {"UK", "CA"}         # 転送停止倉庫（必要なら編集）

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s | %(levelname)-8s | %(message)s")

# ────────── レスポンスフィルタ ──────────
def is_warehouse_api(resp):
    return ("warehouse" in resp.url) and resp.headers.get("content-type","").startswith("application/json")

# ────────── メイン ──────────
async def main():
    if not STORAGE.exists():
        sys.exit("storage.json がありません。先に codegen で作成してください")

    async with async_playwright() as p:
        br  = await p.chromium.launch(headless=HEADLESS)
        ctx = await br.new_context(storage_state=str(STORAGE))
        page = await ctx.new_page()

        # ページを開き、倉庫一覧 API をキャッチ
        logging.info("倉庫一覧ページをロード中…")
        fut = page.expect_response(is_warehouse_api, timeout=TIMEOUT)
        await page.goto(URL_LIST, wait_until="domcontentloaded", timeout=TIMEOUT)

        try:
            async with fut as resp:                    # API レスポンスを取得
        data = await resp.json()
        except asyncio.TimeoutError:
            sys.exit("❌ 住所 API を検出できませんでした。storage.json が失効している可能性")
        await br.close()

    # レスポンス構造が配列 or dict どちらでも吸収
    if isinstance(data, list):
        rows = {row.get("countryCode",""): row for row in data}
    else:
        rows = data

    result = {}
    for code, row in rows.items():
        if code in STOPPED:
            result[code] = {"disabled": True}
            continue

        result[code] = {
            "name"    : row.get("name", ""),
            "address1": row.get("address1", ""),
            "address2": row.get("address2", ""),
            "city"    : row.get("city", ""),
            "state"   : row.get("state", ""),
            "zip"     : row.get("zip", ""),
            "phone"   : row.get("phone", ""),
            "disabled": False,
        }

    with OUT.open("w", encoding="utf-8") as f:
        yaml.safe_dump({"buyandship": result}, f, allow_unicode=True, sort_keys=False)
    logging.info("✅ 完了 → %s", OUT)

# ──────────
if __name__ == "__main__":
    asyncio.run(main())

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\dateutil\zoneinfo\rebuild.py ===
import logging
import os
import tempfile
import shutil
import json
from subprocess import check_call, check_output
from tarfile import TarFile

from dateutil.zoneinfo import METADATA_FN, ZONEFILENAME


def rebuild(filename, tag=None, format="gz", zonegroups=[], metadata=None):
    """Rebuild the internal timezone info in dateutil/zoneinfo/zoneinfo*tar*

    filename is the timezone tarball from ``ftp.iana.org/tz``.

    """
    tmpdir = tempfile.mkdtemp()
    zonedir = os.path.join(tmpdir, "zoneinfo")
    moduledir = os.path.dirname(__file__)
    try:
        with TarFile.open(filename) as tf:
            for name in zonegroups:
                tf.extract(name, tmpdir)
            filepaths = [os.path.join(tmpdir, n) for n in zonegroups]

            _run_zic(zonedir, filepaths)

        # write metadata file
        with open(os.path.join(zonedir, METADATA_FN), 'w') as f:
            json.dump(metadata, f, indent=4, sort_keys=True)
        target = os.path.join(moduledir, ZONEFILENAME)
        with TarFile.open(target, "w:%s" % format) as tf:
            for entry in os.listdir(zonedir):
                entrypath = os.path.join(zonedir, entry)
                tf.add(entrypath, entry)
    finally:
        shutil.rmtree(tmpdir)


def _run_zic(zonedir, filepaths):
    """Calls the ``zic`` compiler in a compatible way to get a "fat" binary.

    Recent versions of ``zic`` default to ``-b slim``, while older versions
    don't even have the ``-b`` option (but default to "fat" binaries). The
    current version of dateutil does not support Version 2+ TZif files, which
    causes problems when used in conjunction with "slim" binaries, so this
    function is used to ensure that we always get a "fat" binary.
    """

    try:
        help_text = check_output(["zic", "--help"])
    except OSError as e:
        _print_on_nosuchfile(e)
        raise

    if b"-b " in help_text:
        bloat_args = ["-b", "fat"]
    else:
        bloat_args = []

    check_call(["zic"] + bloat_args + ["-d", zonedir] + filepaths)


def _print_on_nosuchfile(e):
    """Print helpful troubleshooting message

    e is an exception raised by subprocess.check_call()

    """
    if e.errno == 2:
        logging.error(
            "Could not find zic. Perhaps you need to install "
            "libc-bin or some other package that provides it, "
            "or it's not in your PATH?")

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\flask_wtf\recaptcha\validators.py ===
import json
from urllib import request as http
from urllib.parse import urlencode

from flask import current_app
from flask import request
from wtforms import ValidationError

RECAPTCHA_VERIFY_SERVER_DEFAULT = "https://www.google.com/recaptcha/api/siteverify"
RECAPTCHA_ERROR_CODES = {
    "missing-input-secret": "The secret parameter is missing.",
    "invalid-input-secret": "The secret parameter is invalid or malformed.",
    "missing-input-response": "The response parameter is missing.",
    "invalid-input-response": "The response parameter is invalid or malformed.",
}


__all__ = ["Recaptcha"]


class Recaptcha:
    """Validates a ReCaptcha."""

    def __init__(self, message=None):
        if message is None:
            message = RECAPTCHA_ERROR_CODES["missing-input-response"]
        self.message = message

    def __call__(self, form, field):
        if current_app.testing:
            return True

        if request.is_json:
            response = request.json.get("g-recaptcha-response", "")
        else:
            response = request.form.get("g-recaptcha-response", "")
        remote_ip = request.remote_addr

        if not response:
            raise ValidationError(field.gettext(self.message))

        if not self._validate_recaptcha(response, remote_ip):
            field.recaptcha_error = "incorrect-captcha-sol"
            raise ValidationError(field.gettext(self.message))

    def _validate_recaptcha(self, response, remote_addr):
        """Performs the actual validation."""
        try:
            private_key = current_app.config["RECAPTCHA_PRIVATE_KEY"]
        except KeyError:
            raise RuntimeError("No RECAPTCHA_PRIVATE_KEY config set") from None

        verify_server = current_app.config.get("RECAPTCHA_VERIFY_SERVER")
        if not verify_server:
            verify_server = RECAPTCHA_VERIFY_SERVER_DEFAULT

        data = urlencode(
            {"secret": private_key, "remoteip": remote_addr, "response": response}
        )

        http_response = http.urlopen(verify_server, data.encode("utf-8"))

        if http_response.code != 200:
            return False

        json_resp = json.loads(http_response.read())

        if json_resp["success"]:
            return True

        for error in json_resp.get("error-codes", []):
            if error in RECAPTCHA_ERROR_CODES:
                raise ValidationError(RECAPTCHA_ERROR_CODES[error])

        return False

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\openpyxl\chart\legend.py ===
# Copyright (c) 2010-2024 openpyxl

from openpyxl.descriptors.serialisable import Serialisable
from openpyxl.descriptors import (
    Typed,
    Integer,
    Alias,
    Sequence,
)
from openpyxl.descriptors.excel import ExtensionList
from openpyxl.descriptors.nested import (
    NestedBool,
    NestedSet,
    NestedInteger
)

from .layout import Layout
from .shapes import GraphicalProperties
from .text import RichText


class LegendEntry(Serialisable):

    tagname = "legendEntry"

    idx = NestedInteger()
    delete = NestedBool()
    txPr = Typed(expected_type=RichText, allow_none=True)
    extLst = Typed(expected_type=ExtensionList, allow_none=True)

    __elements__ = ('idx', 'delete', 'txPr')

    def __init__(self,
                 idx=0,
                 delete=False,
                 txPr=None,
                 extLst=None,
                ):
        self.idx = idx
        self.delete = delete
        self.txPr = txPr


class Legend(Serialisable):

    tagname = "legend"

    legendPos = NestedSet(values=(['b', 'tr', 'l', 'r', 't']))
    position = Alias('legendPos')
    legendEntry = Sequence(expected_type=LegendEntry)
    layout = Typed(expected_type=Layout, allow_none=True)
    overlay = NestedBool(allow_none=True)
    spPr = Typed(expected_type=GraphicalProperties, allow_none=True)
    graphicalProperties = Alias('spPr')
    txPr = Typed(expected_type=RichText, allow_none=True)
    textProperties = Alias('txPr')
    extLst = Typed(expected_type=ExtensionList, allow_none=True)

    __elements__ = ('legendPos', 'legendEntry', 'layout', 'overlay', 'spPr', 'txPr',)

    def __init__(self,
                 legendPos="r",
                 legendEntry=(),
                 layout=None,
                 overlay=None,
                 spPr=None,
                 txPr=None,
                 extLst=None,
                ):
        self.legendPos = legendPos
        self.legendEntry = legendEntry
        self.layout = layout
        self.overlay = overlay
        self.spPr = spPr
        self.txPr = txPr

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pip\_vendor\cachecontrol\cache.py ===
# SPDX-FileCopyrightText: 2015 Eric Larson
#
# SPDX-License-Identifier: Apache-2.0

"""
The cache object API for implementing caches. The default is a thread
safe in-memory dictionary.
"""

from __future__ import annotations

from threading import Lock
from typing import IO, TYPE_CHECKING, MutableMapping

if TYPE_CHECKING:
    from datetime import datetime


class BaseCache:
    def get(self, key: str) -> bytes | None:
        raise NotImplementedError()

    def set(
        self, key: str, value: bytes, expires: int | datetime | None = None
    ) -> None:
        raise NotImplementedError()

    def delete(self, key: str) -> None:
        raise NotImplementedError()

    def close(self) -> None:
        pass


class DictCache(BaseCache):
    def __init__(self, init_dict: MutableMapping[str, bytes] | None = None) -> None:
        self.lock = Lock()
        self.data = init_dict or {}

    def get(self, key: str) -> bytes | None:
        return self.data.get(key, None)

    def set(
        self, key: str, value: bytes, expires: int | datetime | None = None
    ) -> None:
        with self.lock:
            self.data.update({key: value})

    def delete(self, key: str) -> None:
        with self.lock:
            if key in self.data:
                self.data.pop(key)


class SeparateBodyBaseCache(BaseCache):
    """
    In this variant, the body is not stored mixed in with the metadata, but is
    passed in (as a bytes-like object) in a separate call to ``set_body()``.

    That is, the expected interaction pattern is::

        cache.set(key, serialized_metadata)
        cache.set_body(key)

    Similarly, the body should be loaded separately via ``get_body()``.
    """

    def set_body(self, key: str, body: bytes) -> None:
        raise NotImplementedError()

    def get_body(self, key: str) -> IO[bytes] | None:
        """
        Return the body as file-like object.
        """
        raise NotImplementedError()

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pyreadline3\clipboard\get_clipboard_text_and_convert.py ===
# -*- coding: utf-8 -*-
# *****************************************************************************
#     Copyright (C) 2006-2020 Jorgen Stenarson. <jorgen.stenarson@bostream.nu>
#     Copyright (C) 2020 Bassem Girgis. <brgirgis@gmail.com>
#
#  Distributed under the terms of the BSD License.  The full license is in
#  the file COPYING, distributed as part of this software.
# *****************************************************************************

from typing import Any, List, Tuple

from .api import get_clipboard_text


def _make_num(x: Any) -> Any:
    try:
        return int(x)
    except ValueError:
        try:
            return float(x)
        except ValueError:
            try:
                return complex(x)
            except ValueError:
                return x


def _make_list_of_list(txt: str) -> Tuple[
    List[List[Any]],
    bool,
]:

    ut = []
    flag = False

    for line in [x for x in txt.split("\r\n") if x != ""]:
        words = [_make_num(x) for x in line.split("\t")]

        if str in list(map(type, words)):
            flag = True

        ut.append(words)

    return (
        ut,
        flag,
    )


def get_clipboard_text_and_convert(paste_list: bool = False) -> str:
    """Get txt from clipboard. if paste_list==True the convert tab separated
    data to list of lists. Enclose list of list in array() if all elements are
    numeric"""

    txt = get_clipboard_text()

    if not txt:
        return ""

    if not paste_list:
        return txt

    if "\t" not in txt:
        return txt

    array, flag = _make_list_of_list(txt)

    if flag:
        txt = repr(array)
    else:
        txt = "array(%s)" % repr(array)

    txt = "".join([c for c in txt if c not in " \t\r\n"])

    return txt

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pyreadline3\logger\control.py ===
# -*- coding: utf-8 -*-
# *****************************************************************************
#     Copyright (C) 2006-2020 Jorgen Stenarson. <jorgen.stenarson@bostream.nu>
#     Copyright (C) 2020 Bassem Girgis. <brgirgis@gmail.com>
#
#  Distributed under the terms of the BSD License.  The full license is in
#  the file COPYING, distributed as part of this software.
# *****************************************************************************

import os
from logging import FileHandler, Formatter, StreamHandler
from logging.handlers import DEFAULT_TCP_LOGGING_PORT
from typing import Optional

from .logger import LOGGER
from .socket_stream import SocketStream

_default_formatter_str = os.environ.get("PYREADLINE_FORMATTER", "%(message)s")

SOCKET_HANDLER: Optional["StreamHandler[SocketStream]"] = None
FILE_HANDLER: Optional[FileHandler] = None


def start_socket_log(
    host: str = "localhost",
    port: int = DEFAULT_TCP_LOGGING_PORT,
    formatter_str: str = _default_formatter_str,
) -> None:
    global SOCKET_HANDLER

    if SOCKET_HANDLER is not None:
        return

    SOCKET_HANDLER = StreamHandler(SocketStream(host, port))
    SOCKET_HANDLER.setFormatter(Formatter(formatter_str))

    LOGGER.addHandler(SOCKET_HANDLER)


def stop_socket_log() -> None:
    global SOCKET_HANDLER

    if SOCKET_HANDLER is None:
        return

    LOGGER.removeHandler(SOCKET_HANDLER)

    SOCKET_HANDLER = None


def start_file_log(filename: str) -> None:
    global FILE_HANDLER

    if FILE_HANDLER is not None:
        return

    FILE_HANDLER = FileHandler(filename, "w")
    LOGGER.addHandler(FILE_HANDLER)


def stop_file_log() -> None:
    global FILE_HANDLER

    if FILE_HANDLER is None:
        return

    LOGGER.removeHandler(FILE_HANDLER)

    FILE_HANDLER.close()
    FILE_HANDLER = None


def stop_logging() -> None:
    stop_file_log()
    stop_socket_log()

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sqlalchemy\sql\_py_util.py ===
# sql/_py_util.py
# Copyright (C) 2005-2025 the SQLAlchemy authors and contributors
# <see AUTHORS file>
#
# This module is part of SQLAlchemy and is released under
# the MIT License: https://www.opensource.org/licenses/mit-license.php

from __future__ import annotations

import typing
from typing import Any
from typing import Dict
from typing import Tuple
from typing import Union

from ..util.typing import Literal

if typing.TYPE_CHECKING:
    from .cache_key import CacheConst


class prefix_anon_map(Dict[str, str]):
    """A map that creates new keys for missing key access.

    Considers keys of the form "<ident> <name>" to produce
    new symbols "<name>_<index>", where "index" is an incrementing integer
    corresponding to <name>.

    Inlines the approach taken by :class:`sqlalchemy.util.PopulateDict` which
    is otherwise usually used for this type of operation.

    """

    def __missing__(self, key: str) -> str:
        (ident, derived) = key.split(" ", 1)
        anonymous_counter = self.get(derived, 1)
        self[derived] = anonymous_counter + 1  # type: ignore
        value = f"{derived}_{anonymous_counter}"
        self[key] = value
        return value


class cache_anon_map(
    Dict[Union[int, "Literal[CacheConst.NO_CACHE]"], Union[Literal[True], str]]
):
    """A map that creates new keys for missing key access.

    Produces an incrementing sequence given a series of unique keys.

    This is similar to the compiler prefix_anon_map class although simpler.

    Inlines the approach taken by :class:`sqlalchemy.util.PopulateDict` which
    is otherwise usually used for this type of operation.

    """

    _index = 0

    def get_anon(self, object_: Any) -> Tuple[str, bool]:
        idself = id(object_)
        if idself in self:
            s_val = self[idself]
            assert s_val is not True
            return s_val, True
        else:
            # inline of __missing__
            self[idself] = id_ = str(self._index)
            self._index += 1

            return id_, False

    def __missing__(self, key: int) -> str:
        self[key] = val = str(self._index)
        self._index += 1
        return val

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sympy\solvers\__init__.py ===
"""A module for solving all kinds of equations.

    Examples
    ========

    >>> from sympy.solvers import solve
    >>> from sympy.abc import x
    >>> solve(((x + 1)**5).expand(), x)
    [-1]
"""
from sympy.core.assumptions import check_assumptions, failing_assumptions

from .solvers import solve, solve_linear_system, solve_linear_system_LU, \
    solve_undetermined_coeffs, nsolve, solve_linear, checksol, \
    det_quick, inv_quick

from sympy.solvers.diophantine.diophantine import diophantine

from .recurr import rsolve, rsolve_poly, rsolve_ratio, rsolve_hyper

from .ode import checkodesol, classify_ode, dsolve, \
    homogeneous_order

from .polysys import solve_poly_system, solve_triangulated, factor_system

from .pde import pde_separate, pde_separate_add, pde_separate_mul, \
    pdsolve, classify_pde, checkpdesol

from .deutils import ode_order

from .inequalities import reduce_inequalities, reduce_abs_inequality, \
    reduce_abs_inequalities, solve_poly_inequality, solve_rational_inequalities, solve_univariate_inequality

from .decompogen import decompogen

from .solveset import solveset, linsolve, linear_eq_to_matrix, nonlinsolve, substitution

from .simplex import lpmin, lpmax, linprog

# This is here instead of sympy/sets/__init__.py to avoid circular import issues
from ..core.singleton import S
Complexes = S.Complexes

__all__ = [
    'solve', 'solve_linear_system', 'solve_linear_system_LU',
    'solve_undetermined_coeffs', 'nsolve', 'solve_linear', 'checksol',
    'det_quick', 'inv_quick', 'check_assumptions', 'failing_assumptions',

    'diophantine',

    'rsolve', 'rsolve_poly', 'rsolve_ratio', 'rsolve_hyper',

    'checkodesol', 'classify_ode', 'dsolve', 'homogeneous_order',

    'solve_poly_system', 'solve_triangulated', 'factor_system',

    'pde_separate', 'pde_separate_add', 'pde_separate_mul', 'pdsolve',
    'classify_pde', 'checkpdesol',

    'ode_order',

    'reduce_inequalities', 'reduce_abs_inequality', 'reduce_abs_inequalities',
    'solve_poly_inequality', 'solve_rational_inequalities',
    'solve_univariate_inequality',

    'decompogen',

    'solveset', 'linsolve', 'linear_eq_to_matrix', 'nonlinsolve',
    'substitution',

    # This is here instead of sympy/sets/__init__.py to avoid circular import issues
    'Complexes',

    'lpmin', 'lpmax', 'linprog'
]

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sympy\utilities\timeutils.py ===
"""Simple tools for timing functions' execution, when IPython is not available. """


import timeit
import math


_scales = [1e0, 1e3, 1e6, 1e9]
_units = ['s', 'ms', '\N{GREEK SMALL LETTER MU}s', 'ns']


def timed(func, setup="pass", limit=None):
    """Adaptively measure execution time of a function. """
    timer = timeit.Timer(func, setup=setup)
    repeat, number = 3, 1

    for i in range(1, 10):
        if timer.timeit(number) >= 0.2:
            break
        elif limit is not None and number >= limit:
            break
        else:
            number *= 10

    time = min(timer.repeat(repeat, number)) / number

    if time > 0.0:
        order = min(-int(math.floor(math.log10(time)) // 3), 3)
    else:
        order = 3

    return (number, time, time*_scales[order], _units[order])


# Code for doing inline timings of recursive algorithms.

def __do_timings():
    import os
    res = os.getenv('SYMPY_TIMINGS', '')
    res = [x.strip() for x in res.split(',')]
    return set(res)

_do_timings = __do_timings()
_timestack = None


def _print_timestack(stack, level=1):
    print('-'*level, '%.2f %s%s' % (stack[2], stack[0], stack[3]))
    for s in stack[1]:
        _print_timestack(s, level + 1)


def timethis(name):
    def decorator(func):
        if name not in _do_timings:
            return func

        def wrapper(*args, **kwargs):
            from time import time
            global _timestack
            oldtimestack = _timestack
            _timestack = [func.func_name, [], 0, args]
            t1 = time()
            r = func(*args, **kwargs)
            t2 = time()
            _timestack[2] = t2 - t1
            if oldtimestack is not None:
                oldtimestack[1].append(_timestack)
                _timestack = oldtimestack
            else:
                _print_timestack(_timestack)
                _timestack = None
            return r
        return wrapper
    return decorator

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\trio\_core\_wakeup_socketpair.py ===
from __future__ import annotations

import contextlib
import signal
import socket
import warnings

from .. import _core
from .._util import is_main_thread


class WakeupSocketpair:
    def __init__(self) -> None:
        # explicitly typed to please `pyright --verifytypes` without `--ignoreexternal`
        self.wakeup_sock: socket.socket
        self.write_sock: socket.socket

        self.wakeup_sock, self.write_sock = socket.socketpair()
        self.wakeup_sock.setblocking(False)
        self.write_sock.setblocking(False)
        # This somewhat reduces the amount of memory wasted queueing up data
        # for wakeups. With these settings, maximum number of 1-byte sends
        # before getting BlockingIOError:
        #   Linux 4.8: 6
        #   macOS (darwin 15.5): 1
        #   Windows 10: 525347
        # Windows you're weird. (And on Windows setting SNDBUF to 0 makes send
        # blocking, even on non-blocking sockets, so don't do that.)
        self.wakeup_sock.setsockopt(socket.SOL_SOCKET, socket.SO_RCVBUF, 1)
        self.write_sock.setsockopt(socket.SOL_SOCKET, socket.SO_SNDBUF, 1)
        # On Windows this is a TCP socket so this might matter. On other
        # platforms this fails b/c AF_UNIX sockets aren't actually TCP.
        with contextlib.suppress(OSError):
            self.write_sock.setsockopt(socket.IPPROTO_TCP, socket.TCP_NODELAY, 1)
        self.old_wakeup_fd: int | None = None

    def wakeup_thread_and_signal_safe(self) -> None:
        with contextlib.suppress(BlockingIOError):
            self.write_sock.send(b"\x00")

    async def wait_woken(self) -> None:
        await _core.wait_readable(self.wakeup_sock)
        self.drain()

    def drain(self) -> None:
        try:
            while True:
                self.wakeup_sock.recv(2**16)
        except BlockingIOError:
            pass

    def wakeup_on_signals(self) -> None:
        assert self.old_wakeup_fd is None
        if not is_main_thread():
            return
        fd = self.write_sock.fileno()
        self.old_wakeup_fd = signal.set_wakeup_fd(fd, warn_on_full_buffer=False)
        if self.old_wakeup_fd != -1:
            warnings.warn(
                RuntimeWarning(
                    "It looks like Trio's signal handling code might have "
                    "collided with another library you're using. If you're "
                    "running Trio in guest mode, then this might mean you "
                    "should set host_uses_signal_set_wakeup_fd=True. "
                    "Otherwise, file a bug on Trio and we'll help you figure "
                    "out what's going on.",
                ),
                stacklevel=1,
            )

    def close(self) -> None:
        self.wakeup_sock.close()
        self.write_sock.close()
        if self.old_wakeup_fd is not None:
            signal.set_wakeup_fd(self.old_wakeup_fd)

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\webdriver_manager\microsoft.py ===
import os
from typing import Optional

from webdriver_manager.core.download_manager import DownloadManager
from webdriver_manager.core.driver_cache import DriverCacheManager
from webdriver_manager.core.os_manager import OperationSystemManager
from webdriver_manager.drivers.edge import EdgeChromiumDriver
from webdriver_manager.drivers.ie import IEDriver
from webdriver_manager.core.manager import DriverManager


class IEDriverManager(DriverManager):
    def __init__(
            self,
            version: Optional[str] = None,
            name: str = "IEDriverServer",
            url: str = "https://github.com/seleniumhq/selenium/releases/download",
            latest_release_url: str = "https://api.github.com/repos/seleniumhq/selenium/releases",
            ie_release_tag: str = "https://api.github.com/repos/seleniumhq/selenium/releases/tags/selenium-{0}",
            download_manager: Optional[DownloadManager] = None,
            cache_manager: Optional[DriverCacheManager] = None,
            os_system_manager: Optional[OperationSystemManager] = None
    ):
        super().__init__(
            download_manager=download_manager,
            cache_manager=cache_manager
        )

        self.driver = IEDriver(
            driver_version=version,
            name=name,
            url=url,
            latest_release_url=latest_release_url,
            ie_release_tag=ie_release_tag,
            http_client=self.http_client,
            os_system_manager=os_system_manager
        )

    def install(self) -> str:
        return self._get_driver_binary_path(self.driver)

    def get_os_type(self):
        return "x64" if self._os_system_manager.get_os_type() == "win64" else "Win32"


class EdgeChromiumDriverManager(DriverManager):
    def __init__(
            self,
            version: Optional[str] = None,
            name: str = "edgedriver",
            url: str = "https://msedgedriver.azureedge.net",
            latest_release_url: str = "https://msedgedriver.azureedge.net/LATEST_RELEASE",
            download_manager: Optional[DownloadManager] = None,
            cache_manager: Optional[DriverCacheManager] = None,
            os_system_manager: Optional[OperationSystemManager] = None
    ):
        super().__init__(
            download_manager=download_manager,
            cache_manager=cache_manager,
            os_system_manager=os_system_manager
        )

        self.driver = EdgeChromiumDriver(
            driver_version=version,
            name=name,
            url=url,
            latest_release_url=latest_release_url,
            http_client=self.http_client,
            os_system_manager=os_system_manager
        )

    def install(self) -> str:
        driver_path = self._get_driver_binary_path(self.driver)
        os.chmod(driver_path, 0o755)
        return driver_path

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\websocket\_cookiejar.py ===
import http.cookies
from typing import Optional

"""
_cookiejar.py
websocket - WebSocket client library for Python

Copyright 2024 engn33r

Licensed under the Apache License, Version 2.0 (the "License");
you may not use this file except in compliance with the License.
You may obtain a copy of the License at

    http://www.apache.org/licenses/LICENSE-2.0

Unless required by applicable law or agreed to in writing, software
distributed under the License is distributed on an "AS IS" BASIS,
WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
See the License for the specific language governing permissions and
limitations under the License.
"""


class SimpleCookieJar:
    def __init__(self) -> None:
        self.jar: dict = {}

    def add(self, set_cookie: Optional[str]) -> None:
        if set_cookie:
            simple_cookie = http.cookies.SimpleCookie(set_cookie)

            for v in simple_cookie.values():
                if domain := v.get("domain"):
                    if not domain.startswith("."):
                        domain = f".{domain}"
                    cookie = (
                        self.jar.get(domain)
                        if self.jar.get(domain)
                        else http.cookies.SimpleCookie()
                    )
                    cookie.update(simple_cookie)
                    self.jar[domain.lower()] = cookie

    def set(self, set_cookie: str) -> None:
        if set_cookie:
            simple_cookie = http.cookies.SimpleCookie(set_cookie)

            for v in simple_cookie.values():
                if domain := v.get("domain"):
                    if not domain.startswith("."):
                        domain = f".{domain}"
                    self.jar[domain.lower()] = simple_cookie

    def get(self, host: str) -> str:
        if not host:
            return ""

        cookies = []
        for domain, _ in self.jar.items():
            host = host.lower()
            if host.endswith(domain) or host == domain[1:]:
                cookies.append(self.jar.get(domain))

        return "; ".join(
            filter(
                None,
                sorted(
                    [
                        f"{k}={v.value}"
                        for cookie in filter(None, cookies)
                        for k, v in cookie.items()
                    ]
                ),
            )
        )

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\numpy\linalg\tests\test_linalg.py ===
""" Test functions for linalg module

"""
import itertools
import os
import subprocess
import sys
import textwrap
import threading
import traceback

import pytest

import numpy as np
from numpy import (
    array,
    asarray,
    atleast_2d,
    cdouble,
    csingle,
    dot,
    double,
    identity,
    inf,
    linalg,
    matmul,
    multiply,
    single,
)
from numpy._core import swapaxes
from numpy.exceptions import AxisError
from numpy.linalg import LinAlgError, matrix_power, matrix_rank, multi_dot, norm
from numpy.linalg._linalg import _multi_dot_matrix_chain_order
from numpy.testing import (
    HAS_LAPACK64,
    IS_WASM,
    NOGIL_BUILD,
    assert_,
    assert_allclose,
    assert_almost_equal,
    assert_array_equal,
    assert_equal,
    assert_raises,
    assert_raises_regex,
    suppress_warnings,
)

try:
    import numpy.linalg.lapack_lite
except ImportError:
    # May be broken when numpy was built without BLAS/LAPACK present
    # If so, ensure we don't break the whole test suite - the `lapack_lite`
    # submodule should be removed, it's only used in two tests in this file.
    pass


def consistent_subclass(out, in_):
    # For ndarray subclass input, our output should have the same subclass
    # (non-ndarray input gets converted to ndarray).
    return type(out) is (type(in_) if isinstance(in_, np.ndarray)
                         else np.ndarray)


old_assert_almost_equal = assert_almost_equal


def assert_almost_equal(a, b, single_decimal=6, double_decimal=12, **kw):
    if asarray(a).dtype.type in (single, csingle):
        decimal = single_decimal
    else:
        decimal = double_decimal
    old_assert_almost_equal(a, b, decimal=decimal, **kw)


def get_real_dtype(dtype):
    return {single: single, double: double,
            csingle: single, cdouble: double}[dtype]


def get_complex_dtype(dtype):
    return {single: csingle, double: cdouble,
            csingle: csingle, cdouble: cdouble}[dtype]


def get_rtol(dtype):
    # Choose a safe rtol
    if dtype in (single, csingle):
        return 1e-5
    else:
        return 1e-11


# used to categorize tests
all_tags = {
  'square', 'nonsquare', 'hermitian',  # mutually exclusive
  'generalized', 'size-0', 'strided'  # optional additions
}


class LinalgCase:
    def __init__(self, name, a, b, tags=set()):
        """
        A bundle of arguments to be passed to a test case, with an identifying
        name, the operands a and b, and a set of tags to filter the tests
        """
        assert_(isinstance(name, str))
        self.name = name
        self.a = a
        self.b = b
        self.tags = frozenset(tags)  # prevent shared tags

    def check(self, do):
        """
        Run the function `do` on this test case, expanding arguments
        """
        do(self.a, self.b, tags=self.tags)

    def __repr__(self):
        return f'<LinalgCase: {self.name}>'


def apply_tag(tag, cases):
    """
    Add the given tag (a string) to each of the cases (a list of LinalgCase
    objects)
    """
    assert tag in all_tags, "Invalid tag"
    for case in cases:
        case.tags = case.tags | {tag}
    return cases


#
# Base test cases
#

np.random.seed(1234)

CASES = []

# square test cases
CASES += apply_tag('square', [
    LinalgCase("single",
               array([[1., 2.], [3., 4.]], dtype=single),
               array([2., 1.], dtype=single)),
    LinalgCase("double",
               array([[1., 2.], [3., 4.]], dtype=double),
               array([2., 1.], dtype=double)),
    LinalgCase("double_2",
               array([[1., 2.], [3., 4.]], dtype=double),
               array([[2., 1., 4.], [3., 4., 6.]], dtype=double)),
    LinalgCase("csingle",
               array([[1. + 2j, 2 + 3j], [3 + 4j, 4 + 5j]], dtype=csingle),
               array([2. + 1j, 1. + 2j], dtype=csingle)),
    LinalgCase("cdouble",
               array([[1. + 2j, 2 + 3j], [3 + 4j, 4 + 5j]], dtype=cdouble),
               array([2. + 1j, 1. + 2j], dtype=cdouble)),
    LinalgCase("cdouble_2",
               array([[1. + 2j, 2 + 3j], [3 + 4j, 4 + 5j]], dtype=cdouble),
               array([[2. + 1j, 1. + 2j, 1 + 3j], [1 - 2j, 1 - 3j, 1 - 6j]], dtype=cdouble)),
    LinalgCase("0x0",
               np.empty((0, 0), dtype=double),
               np.empty((0,), dtype=double),
               tags={'size-0'}),
    LinalgCase("8x8",
               np.random.rand(8, 8),
               np.random.rand(8)),
    LinalgCase("1x1",
               np.random.rand(1, 1),
               np.random.rand(1)),
    LinalgCase("nonarray",
               [[1, 2], [3, 4]],
               [2, 1]),
])

# non-square test-cases
CASES += apply_tag('nonsquare', [
    LinalgCase("single_nsq_1",
               array([[1., 2., 3.], [3., 4., 6.]], dtype=single),
               array([2., 1.], dtype=single)),
    LinalgCase("single_nsq_2",
               array([[1., 2.], [3., 4.], [5., 6.]], dtype=single),
               array([2., 1., 3.], dtype=single)),
    LinalgCase("double_nsq_1",
               array([[1., 2., 3.], [3., 4., 6.]], dtype=double),
               array([2., 1.], dtype=double)),
    LinalgCase("double_nsq_2",
               array([[1., 2.], [3., 4.], [5., 6.]], dtype=double),
               array([2., 1., 3.], dtype=double)),
    LinalgCase("csingle_nsq_1",
               array(
                   [[1. + 1j, 2. + 2j, 3. - 3j], [3. - 5j, 4. + 9j, 6. + 2j]], dtype=csingle),
               array([2. + 1j, 1. + 2j], dtype=csingle)),
    LinalgCase("csingle_nsq_2",
               array(
                   [[1. + 1j, 2. + 2j], [3. - 3j, 4. - 9j], [5. - 4j, 6. + 8j]], dtype=csingle),
               array([2. + 1j, 1. + 2j, 3. - 3j], dtype=csingle)),
    LinalgCase("cdouble_nsq_1",
               array(
                   [[1. + 1j, 2. + 2j, 3. - 3j], [3. - 5j, 4. + 9j, 6. + 2j]], dtype=cdouble),
               array([2. + 1j, 1. + 2j], dtype=cdouble)),
    LinalgCase("cdouble_nsq_2",
               array(
                   [[1. + 1j, 2. + 2j], [3. - 3j, 4. - 9j], [5. - 4j, 6. + 8j]], dtype=cdouble),
               array([2. + 1j, 1. + 2j, 3. - 3j], dtype=cdouble)),
    LinalgCase("cdouble_nsq_1_2",
               array(
                   [[1. + 1j, 2. + 2j, 3. - 3j], [3. - 5j, 4. + 9j, 6. + 2j]], dtype=cdouble),
               array([[2. + 1j, 1. + 2j], [1 - 1j, 2 - 2j]], dtype=cdouble)),
    LinalgCase("cdouble_nsq_2_2",
               array(
                   [[1. + 1j, 2. + 2j], [3. - 3j, 4. - 9j], [5. - 4j, 6. + 8j]], dtype=cdouble),
               array([[2. + 1j, 1. + 2j], [1 - 1j, 2 - 2j], [1 - 1j, 2 - 2j]], dtype=cdouble)),
    LinalgCase("8x11",
               np.random.rand(8, 11),
               np.random.rand(8)),
    LinalgCase("1x5",
               np.random.rand(1, 5),
               np.random.rand(1)),
    LinalgCase("5x1",
               np.random.rand(5, 1),
               np.random.rand(5)),
    LinalgCase("0x4",
               np.random.rand(0, 4),
               np.random.rand(0),
               tags={'size-0'}),
    LinalgCase("4x0",
               np.random.rand(4, 0),
               np.random.rand(4),
               tags={'size-0'}),
])

# hermitian test-cases
CASES += apply_tag('hermitian', [
    LinalgCase("hsingle",
               array([[1., 2.], [2., 1.]], dtype=single),
               None),
    LinalgCase("hdouble",
               array([[1., 2.], [2., 1.]], dtype=double),
               None),
    LinalgCase("hcsingle",
               array([[1., 2 + 3j], [2 - 3j, 1]], dtype=csingle),
               None),
    LinalgCase("hcdouble",
               array([[1., 2 + 3j], [2 - 3j, 1]], dtype=cdouble),
               None),
    LinalgCase("hempty",
               np.empty((0, 0), dtype=double),
               None,
               tags={'size-0'}),
    LinalgCase("hnonarray",
               [[1, 2], [2, 1]],
               None),
    LinalgCase("matrix_b_only",
               array([[1., 2.], [2., 1.]]),
               None),
    LinalgCase("hmatrix_1x1",
               np.random.rand(1, 1),
               None),
])


#
# Gufunc test cases
#
def _make_generalized_cases():
    new_cases = []

    for case in CASES:
        if not isinstance(case.a, np.ndarray):
            continue

        a = np.array([case.a, 2 * case.a, 3 * case.a])
        if case.b is None:
            b = None
        elif case.b.ndim == 1:
            b = case.b
        else:
            b = np.array([case.b, 7 * case.b, 6 * case.b])
        new_case = LinalgCase(case.name + "_tile3", a, b,
                              tags=case.tags | {'generalized'})
        new_cases.append(new_case)

        a = np.array([case.a] * 2 * 3).reshape((3, 2) + case.a.shape)
        if case.b is None:
            b = None
        elif case.b.ndim == 1:
            b = np.array([case.b] * 2 * 3 * a.shape[-1])\
                  .reshape((3, 2) + case.a.shape[-2:])
        else:
            b = np.array([case.b] * 2 * 3).reshape((3, 2) + case.b.shape)
        new_case = LinalgCase(case.name + "_tile213", a, b,
                              tags=case.tags | {'generalized'})
        new_cases.append(new_case)

    return new_cases


CASES += _make_generalized_cases()


#
# Generate stride combination variations of the above
#
def _stride_comb_iter(x):
    """
    Generate cartesian product of strides for all axes
    """

    if not isinstance(x, np.ndarray):
        yield x, "nop"
        return

    stride_set = [(1,)] * x.ndim
    stride_set[-1] = (1, 3, -4)
    if x.ndim > 1:
        stride_set[-2] = (1, 3, -4)
    if x.ndim > 2:
        stride_set[-3] = (1, -4)

    for repeats in itertools.product(*tuple(stride_set)):
        new_shape = [abs(a * b) for a, b in zip(x.shape, repeats)]
        slices = tuple(slice(None, None, repeat) for repeat in repeats)

        # new array with different strides, but same data
        xi = np.empty(new_shape, dtype=x.dtype)
        xi.view(np.uint32).fill(0xdeadbeef)
        xi = xi[slices]
        xi[...] = x
        xi = xi.view(x.__class__)
        assert_(np.all(xi == x))
        yield xi, "stride_" + "_".join(["%+d" % j for j in repeats])

        # generate also zero strides if possible
        if x.ndim >= 1 and x.shape[-1] == 1:
            s = list(x.strides)
            s[-1] = 0
            xi = np.lib.stride_tricks.as_strided(x, strides=s)
            yield xi, "stride_xxx_0"
        if x.ndim >= 2 and x.shape[-2] == 1:
            s = list(x.strides)
            s[-2] = 0
            xi = np.lib.stride_tricks.as_strided(x, strides=s)
            yield xi, "stride_xxx_0_x"
        if x.ndim >= 2 and x.shape[:-2] == (1, 1):
            s = list(x.strides)
            s[-1] = 0
            s[-2] = 0
            xi = np.lib.stride_tricks.as_strided(x, strides=s)
            yield xi, "stride_xxx_0_0"


def _make_strided_cases():
    new_cases = []
    for case in CASES:
        for a, a_label in _stride_comb_iter(case.a):
            for b, b_label in _stride_comb_iter(case.b):
                new_case = LinalgCase(case.name + "_" + a_label + "_" + b_label, a, b,
                                      tags=case.tags | {'strided'})
                new_cases.append(new_case)
    return new_cases


CASES += _make_strided_cases()


#
# Test different routines against the above cases
#
class LinalgTestCase:
    TEST_CASES = CASES

    def check_cases(self, require=set(), exclude=set()):
        """
        Run func on each of the cases with all of the tags in require, and none
        of the tags in exclude
        """
        for case in self.TEST_CASES:
            # filter by require and exclude
            if case.tags & require != require:
                continue
            if case.tags & exclude:
                continue

            try:
                case.check(self.do)
            except Exception as e:
                msg = f'In test case: {case!r}\n\n'
                msg += traceback.format_exc()
                raise AssertionError(msg) from e


class LinalgSquareTestCase(LinalgTestCase):

    def test_sq_cases(self):
        self.check_cases(require={'square'},
                         exclude={'generalized', 'size-0'})

    def test_empty_sq_cases(self):
        self.check_cases(require={'square', 'size-0'},
                         exclude={'generalized'})


class LinalgNonsquareTestCase(LinalgTestCase):

    def test_nonsq_cases(self):
        self.check_cases(require={'nonsquare'},
                         exclude={'generalized', 'size-0'})

    def test_empty_nonsq_cases(self):
        self.check_cases(require={'nonsquare', 'size-0'},
                         exclude={'generalized'})


class HermitianTestCase(LinalgTestCase):

    def test_herm_cases(self):
        self.check_cases(require={'hermitian'},
                         exclude={'generalized', 'size-0'})

    def test_empty_herm_cases(self):
        self.check_cases(require={'hermitian', 'size-0'},
                         exclude={'generalized'})


class LinalgGeneralizedSquareTestCase(LinalgTestCase):

    @pytest.mark.slow
    def test_generalized_sq_cases(self):
        self.check_cases(require={'generalized', 'square'},
                         exclude={'size-0'})

    @pytest.mark.slow
    def test_generalized_empty_sq_cases(self):
        self.check_cases(require={'generalized', 'square', 'size-0'})


class LinalgGeneralizedNonsquareTestCase(LinalgTestCase):

    @pytest.mark.slow
    def test_generalized_nonsq_cases(self):
        self.check_cases(require={'generalized', 'nonsquare'},
                         exclude={'size-0'})

    @pytest.mark.slow
    def test_generalized_empty_nonsq_cases(self):
        self.check_cases(require={'generalized', 'nonsquare', 'size-0'})


class HermitianGeneralizedTestCase(LinalgTestCase):

    @pytest.mark.slow
    def test_generalized_herm_cases(self):
        self.check_cases(require={'generalized', 'hermitian'},
                         exclude={'size-0'})

    @pytest.mark.slow
    def test_generalized_empty_herm_cases(self):
        self.check_cases(require={'generalized', 'hermitian', 'size-0'},
                         exclude={'none'})


def identity_like_generalized(a):
    a = asarray(a)
    if a.ndim >= 3:
        r = np.empty(a.shape, dtype=a.dtype)
        r[...] = identity(a.shape[-2])
        return r
    else:
        return identity(a.shape[0])


class SolveCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):
    # kept apart from TestSolve for use for testing with matrices.
    def do(self, a, b, tags):
        x = linalg.solve(a, b)
        if np.array(b).ndim == 1:
            # When a is (..., M, M) and b is (M,), it is the same as when b is
            # (M, 1), except the result has shape (..., M)
            adotx = matmul(a, x[..., None])[..., 0]
            assert_almost_equal(np.broadcast_to(b, adotx.shape), adotx)
        else:
            adotx = matmul(a, x)
            assert_almost_equal(b, adotx)
        assert_(consistent_subclass(x, b))


class TestSolve(SolveCases):
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        assert_equal(linalg.solve(x, x).dtype, dtype)

    def test_1_d(self):
        class ArraySubclass(np.ndarray):
            pass
        a = np.arange(8).reshape(2, 2, 2)
        b = np.arange(2).view(ArraySubclass)
        result = linalg.solve(a, b)
        assert result.shape == (2, 2)

        # If b is anything other than 1-D it should be treated as a stack of
        # matrices
        b = np.arange(4).reshape(2, 2).view(ArraySubclass)
        result = linalg.solve(a, b)
        assert result.shape == (2, 2, 2)

        b = np.arange(2).reshape(1, 2).view(ArraySubclass)
        assert_raises(ValueError, linalg.solve, a, b)

    def test_0_size(self):
        class ArraySubclass(np.ndarray):
            pass
        # Test system of 0x0 matrices
        a = np.arange(8).reshape(2, 2, 2)
        b = np.arange(6).reshape(1, 2, 3).view(ArraySubclass)

        expected = linalg.solve(a, b)[:, 0:0, :]
        result = linalg.solve(a[:, 0:0, 0:0], b[:, 0:0, :])
        assert_array_equal(result, expected)
        assert_(isinstance(result, ArraySubclass))

        # Test errors for non-square and only b's dimension being 0
        assert_raises(linalg.LinAlgError, linalg.solve, a[:, 0:0, 0:1], b)
        assert_raises(ValueError, linalg.solve, a, b[:, 0:0, :])

        # Test broadcasting error
        b = np.arange(6).reshape(1, 3, 2)  # broadcasting error
        assert_raises(ValueError, linalg.solve, a, b)
        assert_raises(ValueError, linalg.solve, a[0:0], b[0:0])

        # Test zero "single equations" with 0x0 matrices.
        b = np.arange(2).view(ArraySubclass)
        expected = linalg.solve(a, b)[:, 0:0]
        result = linalg.solve(a[:, 0:0, 0:0], b[0:0])
        assert_array_equal(result, expected)
        assert_(isinstance(result, ArraySubclass))

        b = np.arange(3).reshape(1, 3)
        assert_raises(ValueError, linalg.solve, a, b)
        assert_raises(ValueError, linalg.solve, a[0:0], b[0:0])
        assert_raises(ValueError, linalg.solve, a[:, 0:0, 0:0], b)

    def test_0_size_k(self):
        # test zero multiple equation (K=0) case.
        class ArraySubclass(np.ndarray):
            pass
        a = np.arange(4).reshape(1, 2, 2)
        b = np.arange(6).reshape(3, 2, 1).view(ArraySubclass)

        expected = linalg.solve(a, b)[:, :, 0:0]
        result = linalg.solve(a, b[:, :, 0:0])
        assert_array_equal(result, expected)
        assert_(isinstance(result, ArraySubclass))

        # test both zero.
        expected = linalg.solve(a, b)[:, 0:0, 0:0]
        result = linalg.solve(a[:, 0:0, 0:0], b[:, 0:0, 0:0])
        assert_array_equal(result, expected)
        assert_(isinstance(result, ArraySubclass))


class InvCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):

    def do(self, a, b, tags):
        a_inv = linalg.inv(a)
        assert_almost_equal(matmul(a, a_inv),
                            identity_like_generalized(a))
        assert_(consistent_subclass(a_inv, a))


class TestInv(InvCases):
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        assert_equal(linalg.inv(x).dtype, dtype)

    def test_0_size(self):
        # Check that all kinds of 0-sized arrays work
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res = linalg.inv(a)
        assert_(res.dtype.type is np.float64)
        assert_equal(a.shape, res.shape)
        assert_(isinstance(res, ArraySubclass))

        a = np.zeros((0, 0), dtype=np.complex64).view(ArraySubclass)
        res = linalg.inv(a)
        assert_(res.dtype.type is np.complex64)
        assert_equal(a.shape, res.shape)
        assert_(isinstance(res, ArraySubclass))


class EigvalsCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):

    def do(self, a, b, tags):
        ev = linalg.eigvals(a)
        evalues, evectors = linalg.eig(a)
        assert_almost_equal(ev, evalues)


class TestEigvals(EigvalsCases):
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        assert_equal(linalg.eigvals(x).dtype, dtype)
        x = np.array([[1, 0.5], [-1, 1]], dtype=dtype)
        assert_equal(linalg.eigvals(x).dtype, get_complex_dtype(dtype))

    def test_0_size(self):
        # Check that all kinds of 0-sized arrays work
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res = linalg.eigvals(a)
        assert_(res.dtype.type is np.float64)
        assert_equal((0, 1), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(res, np.ndarray))

        a = np.zeros((0, 0), dtype=np.complex64).view(ArraySubclass)
        res = linalg.eigvals(a)
        assert_(res.dtype.type is np.complex64)
        assert_equal((0,), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(res, np.ndarray))


class EigCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):

    def do(self, a, b, tags):
        res = linalg.eig(a)
        eigenvalues, eigenvectors = res.eigenvalues, res.eigenvectors
        assert_allclose(matmul(a, eigenvectors),
                        np.asarray(eigenvectors) * np.asarray(eigenvalues)[..., None, :],
                        rtol=get_rtol(eigenvalues.dtype))
        assert_(consistent_subclass(eigenvectors, a))


class TestEig(EigCases):
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        w, v = np.linalg.eig(x)
        assert_equal(w.dtype, dtype)
        assert_equal(v.dtype, dtype)

        x = np.array([[1, 0.5], [-1, 1]], dtype=dtype)
        w, v = np.linalg.eig(x)
        assert_equal(w.dtype, get_complex_dtype(dtype))
        assert_equal(v.dtype, get_complex_dtype(dtype))

    def test_0_size(self):
        # Check that all kinds of 0-sized arrays work
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res, res_v = linalg.eig(a)
        assert_(res_v.dtype.type is np.float64)
        assert_(res.dtype.type is np.float64)
        assert_equal(a.shape, res_v.shape)
        assert_equal((0, 1), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(a, np.ndarray))

        a = np.zeros((0, 0), dtype=np.complex64).view(ArraySubclass)
        res, res_v = linalg.eig(a)
        assert_(res_v.dtype.type is np.complex64)
        assert_(res.dtype.type is np.complex64)
        assert_equal(a.shape, res_v.shape)
        assert_equal((0,), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(a, np.ndarray))


class SVDBaseTests:
    hermitian = False

    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        res = linalg.svd(x)
        U, S, Vh = res.U, res.S, res.Vh
        assert_equal(U.dtype, dtype)
        assert_equal(S.dtype, get_real_dtype(dtype))
        assert_equal(Vh.dtype, dtype)
        s = linalg.svd(x, compute_uv=False, hermitian=self.hermitian)
        assert_equal(s.dtype, get_real_dtype(dtype))


class SVDCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):

    def do(self, a, b, tags):
        u, s, vt = linalg.svd(a, False)
        assert_allclose(a, matmul(np.asarray(u) * np.asarray(s)[..., None, :],
                                           np.asarray(vt)),
                        rtol=get_rtol(u.dtype))
        assert_(consistent_subclass(u, a))
        assert_(consistent_subclass(vt, a))


class TestSVD(SVDCases, SVDBaseTests):
    def test_empty_identity(self):
        """ Empty input should put an identity matrix in u or vh """
        x = np.empty((4, 0))
        u, s, vh = linalg.svd(x, compute_uv=True, hermitian=self.hermitian)
        assert_equal(u.shape, (4, 4))
        assert_equal(vh.shape, (0, 0))
        assert_equal(u, np.eye(4))

        x = np.empty((0, 4))
        u, s, vh = linalg.svd(x, compute_uv=True, hermitian=self.hermitian)
        assert_equal(u.shape, (0, 0))
        assert_equal(vh.shape, (4, 4))
        assert_equal(vh, np.eye(4))

    def test_svdvals(self):
        x = np.array([[1, 0.5], [0.5, 1]])
        s_from_svd = linalg.svd(x, compute_uv=False, hermitian=self.hermitian)
        s_from_svdvals = linalg.svdvals(x)
        assert_almost_equal(s_from_svd, s_from_svdvals)


class SVDHermitianCases(HermitianTestCase, HermitianGeneralizedTestCase):

    def do(self, a, b, tags):
        u, s, vt = linalg.svd(a, False, hermitian=True)
        assert_allclose(a, matmul(np.asarray(u) * np.asarray(s)[..., None, :],
                                           np.asarray(vt)),
                        rtol=get_rtol(u.dtype))

        def hermitian(mat):
            axes = list(range(mat.ndim))
            axes[-1], axes[-2] = axes[-2], axes[-1]
            return np.conj(np.transpose(mat, axes=axes))

        assert_almost_equal(np.matmul(u, hermitian(u)), np.broadcast_to(np.eye(u.shape[-1]), u.shape))
        assert_almost_equal(np.matmul(vt, hermitian(vt)), np.broadcast_to(np.eye(vt.shape[-1]), vt.shape))
        assert_equal(np.sort(s)[..., ::-1], s)
        assert_(consistent_subclass(u, a))
        assert_(consistent_subclass(vt, a))


class TestSVDHermitian(SVDHermitianCases, SVDBaseTests):
    hermitian = True


class CondCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):
    # cond(x, p) for p in (None, 2, -2)

    def do(self, a, b, tags):
        c = asarray(a)  # a might be a matrix
        if 'size-0' in tags:
            assert_raises(LinAlgError, linalg.cond, c)
            return

        # +-2 norms
        s = linalg.svd(c, compute_uv=False)
        assert_almost_equal(
            linalg.cond(a), s[..., 0] / s[..., -1],
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, 2), s[..., 0] / s[..., -1],
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, -2), s[..., -1] / s[..., 0],
            single_decimal=5, double_decimal=11)

        # Other norms
        cinv = np.linalg.inv(c)
        assert_almost_equal(
            linalg.cond(a, 1),
            abs(c).sum(-2).max(-1) * abs(cinv).sum(-2).max(-1),
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, -1),
            abs(c).sum(-2).min(-1) * abs(cinv).sum(-2).min(-1),
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, np.inf),
            abs(c).sum(-1).max(-1) * abs(cinv).sum(-1).max(-1),
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, -np.inf),
            abs(c).sum(-1).min(-1) * abs(cinv).sum(-1).min(-1),
            single_decimal=5, double_decimal=11)
        assert_almost_equal(
            linalg.cond(a, 'fro'),
            np.sqrt((abs(c)**2).sum(-1).sum(-1)
                    * (abs(cinv)**2).sum(-1).sum(-1)),
            single_decimal=5, double_decimal=11)


class TestCond(CondCases):
    def test_basic_nonsvd(self):
        # Smoketest the non-svd norms
        A = array([[1., 0, 1], [0, -2., 0], [0, 0, 3.]])
        assert_almost_equal(linalg.cond(A, inf), 4)
        assert_almost_equal(linalg.cond(A, -inf), 2 / 3)
        assert_almost_equal(linalg.cond(A, 1), 4)
        assert_almost_equal(linalg.cond(A, -1), 0.5)
        assert_almost_equal(linalg.cond(A, 'fro'), np.sqrt(265 / 12))

    def test_singular(self):
        # Singular matrices have infinite condition number for
        # positive norms, and negative norms shouldn't raise
        # exceptions
        As = [np.zeros((2, 2)), np.ones((2, 2))]
        p_pos = [None, 1, 2, 'fro']
        p_neg = [-1, -2]
        for A, p in itertools.product(As, p_pos):
            # Inversion may not hit exact infinity, so just check the
            # number is large
            assert_(linalg.cond(A, p) > 1e15)
        for A, p in itertools.product(As, p_neg):
            linalg.cond(A, p)

    @pytest.mark.xfail(True, run=False,
                       reason="Platform/LAPACK-dependent failure, "
                              "see gh-18914")
    def test_nan(self):
        # nans should be passed through, not converted to infs
        ps = [None, 1, -1, 2, -2, 'fro']
        p_pos = [None, 1, 2, 'fro']

        A = np.ones((2, 2))
        A[0, 1] = np.nan
        for p in ps:
            c = linalg.cond(A, p)
            assert_(isinstance(c, np.float64))
            assert_(np.isnan(c))

        A = np.ones((3, 2, 2))
        A[1, 0, 1] = np.nan
        for p in ps:
            c = linalg.cond(A, p)
            assert_(np.isnan(c[1]))
            if p in p_pos:
                assert_(c[0] > 1e15)
                assert_(c[2] > 1e15)
            else:
                assert_(not np.isnan(c[0]))
                assert_(not np.isnan(c[2]))

    def test_stacked_singular(self):
        # Check behavior when only some of the stacked matrices are
        # singular
        np.random.seed(1234)
        A = np.random.rand(2, 2, 2, 2)
        A[0, 0] = 0
        A[1, 1] = 0

        for p in (None, 1, 2, 'fro', -1, -2):
            c = linalg.cond(A, p)
            assert_equal(c[0, 0], np.inf)
            assert_equal(c[1, 1], np.inf)
            assert_(np.isfinite(c[0, 1]))
            assert_(np.isfinite(c[1, 0]))


class PinvCases(LinalgSquareTestCase,
                LinalgNonsquareTestCase,
                LinalgGeneralizedSquareTestCase,
                LinalgGeneralizedNonsquareTestCase):

    def do(self, a, b, tags):
        a_ginv = linalg.pinv(a)
        # `a @ a_ginv == I` does not hold if a is singular
        dot = matmul
        assert_almost_equal(dot(dot(a, a_ginv), a), a, single_decimal=5, double_decimal=11)
        assert_(consistent_subclass(a_ginv, a))


class TestPinv(PinvCases):
    pass


class PinvHermitianCases(HermitianTestCase, HermitianGeneralizedTestCase):

    def do(self, a, b, tags):
        a_ginv = linalg.pinv(a, hermitian=True)
        # `a @ a_ginv == I` does not hold if a is singular
        dot = matmul
        assert_almost_equal(dot(dot(a, a_ginv), a), a, single_decimal=5, double_decimal=11)
        assert_(consistent_subclass(a_ginv, a))


class TestPinvHermitian(PinvHermitianCases):
    pass


def test_pinv_rtol_arg():
    a = np.array([[1, 2, 3], [4, 1, 1], [2, 3, 1]])

    assert_almost_equal(
        np.linalg.pinv(a, rcond=0.5),
        np.linalg.pinv(a, rtol=0.5),
    )

    with pytest.raises(
        ValueError, match=r"`rtol` and `rcond` can't be both set."
    ):
        np.linalg.pinv(a, rcond=0.5, rtol=0.5)


class DetCases(LinalgSquareTestCase, LinalgGeneralizedSquareTestCase):

    def do(self, a, b, tags):
        d = linalg.det(a)
        res = linalg.slogdet(a)
        s, ld = res.sign, res.logabsdet
        if asarray(a).dtype.type in (single, double):
            ad = asarray(a).astype(double)
        else:
            ad = asarray(a).astype(cdouble)
        ev = linalg.eigvals(ad)
        assert_almost_equal(d, multiply.reduce(ev, axis=-1))
        assert_almost_equal(s * np.exp(ld), multiply.reduce(ev, axis=-1))

        s = np.atleast_1d(s)
        ld = np.atleast_1d(ld)
        m = (s != 0)
        assert_almost_equal(np.abs(s[m]), 1)
        assert_equal(ld[~m], -inf)


class TestDet(DetCases):
    def test_zero(self):
        assert_equal(linalg.det([[0.0]]), 0.0)
        assert_equal(type(linalg.det([[0.0]])), double)
        assert_equal(linalg.det([[0.0j]]), 0.0)
        assert_equal(type(linalg.det([[0.0j]])), cdouble)

        assert_equal(linalg.slogdet([[0.0]]), (0.0, -inf))
        assert_equal(type(linalg.slogdet([[0.0]])[0]), double)
        assert_equal(type(linalg.slogdet([[0.0]])[1]), double)
        assert_equal(linalg.slogdet([[0.0j]]), (0.0j, -inf))
        assert_equal(type(linalg.slogdet([[0.0j]])[0]), cdouble)
        assert_equal(type(linalg.slogdet([[0.0j]])[1]), double)

    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        assert_equal(np.linalg.det(x).dtype, dtype)
        ph, s = np.linalg.slogdet(x)
        assert_equal(s.dtype, get_real_dtype(dtype))
        assert_equal(ph.dtype, dtype)

    def test_0_size(self):
        a = np.zeros((0, 0), dtype=np.complex64)
        res = linalg.det(a)
        assert_equal(res, 1.)
        assert_(res.dtype.type is np.complex64)
        res = linalg.slogdet(a)
        assert_equal(res, (1, 0))
        assert_(res[0].dtype.type is np.complex64)
        assert_(res[1].dtype.type is np.float32)

        a = np.zeros((0, 0), dtype=np.float64)
        res = linalg.det(a)
        assert_equal(res, 1.)
        assert_(res.dtype.type is np.float64)
        res = linalg.slogdet(a)
        assert_equal(res, (1, 0))
        assert_(res[0].dtype.type is np.float64)
        assert_(res[1].dtype.type is np.float64)


class LstsqCases(LinalgSquareTestCase, LinalgNonsquareTestCase):

    def do(self, a, b, tags):
        arr = np.asarray(a)
        m, n = arr.shape
        u, s, vt = linalg.svd(a, False)
        x, residuals, rank, sv = linalg.lstsq(a, b, rcond=-1)
        if m == 0:
            assert_((x == 0).all())
        if m <= n:
            assert_almost_equal(b, dot(a, x))
            assert_equal(rank, m)
        else:
            assert_equal(rank, n)
        assert_almost_equal(sv, sv.__array_wrap__(s))
        if rank == n and m > n:
            expect_resids = (
                np.asarray(abs(np.dot(a, x) - b)) ** 2).sum(axis=0)
            expect_resids = np.asarray(expect_resids)
            if np.asarray(b).ndim == 1:
                expect_resids.shape = (1,)
                assert_equal(residuals.shape, expect_resids.shape)
        else:
            expect_resids = np.array([]).view(type(x))
        assert_almost_equal(residuals, expect_resids)
        assert_(np.issubdtype(residuals.dtype, np.floating))
        assert_(consistent_subclass(x, b))
        assert_(consistent_subclass(residuals, b))


class TestLstsq(LstsqCases):
    def test_rcond(self):
        a = np.array([[0., 1.,  0.,  1.,  2.,  0.],
                      [0., 2.,  0.,  0.,  1.,  0.],
                      [1., 0.,  1.,  0.,  0.,  4.],
                      [0., 0.,  0.,  2.,  3.,  0.]]).T

        b = np.array([1, 0, 0, 0, 0, 0])

        x, residuals, rank, s = linalg.lstsq(a, b, rcond=-1)
        assert_(rank == 4)
        x, residuals, rank, s = linalg.lstsq(a, b)
        assert_(rank == 3)
        x, residuals, rank, s = linalg.lstsq(a, b, rcond=None)
        assert_(rank == 3)

    @pytest.mark.parametrize(["m", "n", "n_rhs"], [
        (4, 2, 2),
        (0, 4, 1),
        (0, 4, 2),
        (4, 0, 1),
        (4, 0, 2),
        (4, 2, 0),
        (0, 0, 0)
    ])
    def test_empty_a_b(self, m, n, n_rhs):
        a = np.arange(m * n).reshape(m, n)
        b = np.ones((m, n_rhs))
        x, residuals, rank, s = linalg.lstsq(a, b, rcond=None)
        if m == 0:
            assert_((x == 0).all())
        assert_equal(x.shape, (n, n_rhs))
        assert_equal(residuals.shape, ((n_rhs,) if m > n else (0,)))
        if m > n and n_rhs > 0:
            # residuals are exactly the squared norms of b's columns
            r = b - np.dot(a, x)
            assert_almost_equal(residuals, (r * r).sum(axis=-2))
        assert_equal(rank, min(m, n))
        assert_equal(s.shape, (min(m, n),))

    def test_incompatible_dims(self):
        # use modified version of docstring example
        x = np.array([0, 1, 2, 3])
        y = np.array([-1, 0.2, 0.9, 2.1, 3.3])
        A = np.vstack([x, np.ones(len(x))]).T
        with assert_raises_regex(LinAlgError, "Incompatible dimensions"):
            linalg.lstsq(A, y, rcond=None)


@pytest.mark.parametrize('dt', [np.dtype(c) for c in '?bBhHiIqQefdgFDGO'])
class TestMatrixPower:

    rshft_0 = np.eye(4)
    rshft_1 = rshft_0[[3, 0, 1, 2]]
    rshft_2 = rshft_0[[2, 3, 0, 1]]
    rshft_3 = rshft_0[[1, 2, 3, 0]]
    rshft_all = [rshft_0, rshft_1, rshft_2, rshft_3]
    noninv = array([[1, 0], [0, 0]])
    stacked = np.block([[[rshft_0]]] * 2)
    # FIXME the 'e' dtype might work in future
    dtnoinv = [object, np.dtype('e'), np.dtype('g'), np.dtype('G')]

    def test_large_power(self, dt):
        rshft = self.rshft_1.astype(dt)
        assert_equal(
            matrix_power(rshft, 2**100 + 2**10 + 2**5 + 0), self.rshft_0)
        assert_equal(
            matrix_power(rshft, 2**100 + 2**10 + 2**5 + 1), self.rshft_1)
        assert_equal(
            matrix_power(rshft, 2**100 + 2**10 + 2**5 + 2), self.rshft_2)
        assert_equal(
            matrix_power(rshft, 2**100 + 2**10 + 2**5 + 3), self.rshft_3)

    def test_power_is_zero(self, dt):
        def tz(M):
            mz = matrix_power(M, 0)
            assert_equal(mz, identity_like_generalized(M))
            assert_equal(mz.dtype, M.dtype)

        for mat in self.rshft_all:
            tz(mat.astype(dt))
            if dt != object:
                tz(self.stacked.astype(dt))

    def test_power_is_one(self, dt):
        def tz(mat):
            mz = matrix_power(mat, 1)
            assert_equal(mz, mat)
            assert_equal(mz.dtype, mat.dtype)

        for mat in self.rshft_all:
            tz(mat.astype(dt))
            if dt != object:
                tz(self.stacked.astype(dt))

    def test_power_is_two(self, dt):
        def tz(mat):
            mz = matrix_power(mat, 2)
            mmul = matmul if mat.dtype != object else dot
            assert_equal(mz, mmul(mat, mat))
            assert_equal(mz.dtype, mat.dtype)

        for mat in self.rshft_all:
            tz(mat.astype(dt))
            if dt != object:
                tz(self.stacked.astype(dt))

    def test_power_is_minus_one(self, dt):
        def tz(mat):
            invmat = matrix_power(mat, -1)
            mmul = matmul if mat.dtype != object else dot
            assert_almost_equal(
                mmul(invmat, mat), identity_like_generalized(mat))

        for mat in self.rshft_all:
            if dt not in self.dtnoinv:
                tz(mat.astype(dt))

    def test_exceptions_bad_power(self, dt):
        mat = self.rshft_0.astype(dt)
        assert_raises(TypeError, matrix_power, mat, 1.5)
        assert_raises(TypeError, matrix_power, mat, [1])

    def test_exceptions_non_square(self, dt):
        assert_raises(LinAlgError, matrix_power, np.array([1], dt), 1)
        assert_raises(LinAlgError, matrix_power, np.array([[1], [2]], dt), 1)
        assert_raises(LinAlgError, matrix_power, np.ones((4, 3, 2), dt), 1)

    @pytest.mark.skipif(IS_WASM, reason="fp errors don't work in wasm")
    def test_exceptions_not_invertible(self, dt):
        if dt in self.dtnoinv:
            return
        mat = self.noninv.astype(dt)
        assert_raises(LinAlgError, matrix_power, mat, -1)


class TestEigvalshCases(HermitianTestCase, HermitianGeneralizedTestCase):

    def do(self, a, b, tags):
        # note that eigenvalue arrays returned by eig must be sorted since
        # their order isn't guaranteed.
        ev = linalg.eigvalsh(a, 'L')
        evalues, evectors = linalg.eig(a)
        evalues.sort(axis=-1)
        assert_allclose(ev, evalues, rtol=get_rtol(ev.dtype))

        ev2 = linalg.eigvalsh(a, 'U')
        assert_allclose(ev2, evalues, rtol=get_rtol(ev.dtype))


class TestEigvalsh:
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        w = np.linalg.eigvalsh(x)
        assert_equal(w.dtype, get_real_dtype(dtype))

    def test_invalid(self):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=np.float32)
        assert_raises(ValueError, np.linalg.eigvalsh, x, UPLO="lrong")
        assert_raises(ValueError, np.linalg.eigvalsh, x, "lower")
        assert_raises(ValueError, np.linalg.eigvalsh, x, "upper")

    def test_UPLO(self):
        Klo = np.array([[0, 0], [1, 0]], dtype=np.double)
        Kup = np.array([[0, 1], [0, 0]], dtype=np.double)
        tgt = np.array([-1, 1], dtype=np.double)
        rtol = get_rtol(np.double)

        # Check default is 'L'
        w = np.linalg.eigvalsh(Klo)
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'L'
        w = np.linalg.eigvalsh(Klo, UPLO='L')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'l'
        w = np.linalg.eigvalsh(Klo, UPLO='l')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'U'
        w = np.linalg.eigvalsh(Kup, UPLO='U')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'u'
        w = np.linalg.eigvalsh(Kup, UPLO='u')
        assert_allclose(w, tgt, rtol=rtol)

    def test_0_size(self):
        # Check that all kinds of 0-sized arrays work
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res = linalg.eigvalsh(a)
        assert_(res.dtype.type is np.float64)
        assert_equal((0, 1), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(res, np.ndarray))

        a = np.zeros((0, 0), dtype=np.complex64).view(ArraySubclass)
        res = linalg.eigvalsh(a)
        assert_(res.dtype.type is np.float32)
        assert_equal((0,), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(res, np.ndarray))


class TestEighCases(HermitianTestCase, HermitianGeneralizedTestCase):

    def do(self, a, b, tags):
        # note that eigenvalue arrays returned by eig must be sorted since
        # their order isn't guaranteed.
        res = linalg.eigh(a)
        ev, evc = res.eigenvalues, res.eigenvectors
        evalues, evectors = linalg.eig(a)
        evalues.sort(axis=-1)
        assert_almost_equal(ev, evalues)

        assert_allclose(matmul(a, evc),
                        np.asarray(ev)[..., None, :] * np.asarray(evc),
                        rtol=get_rtol(ev.dtype))

        ev2, evc2 = linalg.eigh(a, 'U')
        assert_almost_equal(ev2, evalues)

        assert_allclose(matmul(a, evc2),
                        np.asarray(ev2)[..., None, :] * np.asarray(evc2),
                        rtol=get_rtol(ev.dtype), err_msg=repr(a))


class TestEigh:
    @pytest.mark.parametrize('dtype', [single, double, csingle, cdouble])
    def test_types(self, dtype):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=dtype)
        w, v = np.linalg.eigh(x)
        assert_equal(w.dtype, get_real_dtype(dtype))
        assert_equal(v.dtype, dtype)

    def test_invalid(self):
        x = np.array([[1, 0.5], [0.5, 1]], dtype=np.float32)
        assert_raises(ValueError, np.linalg.eigh, x, UPLO="lrong")
        assert_raises(ValueError, np.linalg.eigh, x, "lower")
        assert_raises(ValueError, np.linalg.eigh, x, "upper")

    def test_UPLO(self):
        Klo = np.array([[0, 0], [1, 0]], dtype=np.double)
        Kup = np.array([[0, 1], [0, 0]], dtype=np.double)
        tgt = np.array([-1, 1], dtype=np.double)
        rtol = get_rtol(np.double)

        # Check default is 'L'
        w, v = np.linalg.eigh(Klo)
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'L'
        w, v = np.linalg.eigh(Klo, UPLO='L')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'l'
        w, v = np.linalg.eigh(Klo, UPLO='l')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'U'
        w, v = np.linalg.eigh(Kup, UPLO='U')
        assert_allclose(w, tgt, rtol=rtol)
        # Check 'u'
        w, v = np.linalg.eigh(Kup, UPLO='u')
        assert_allclose(w, tgt, rtol=rtol)

    def test_0_size(self):
        # Check that all kinds of 0-sized arrays work
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res, res_v = linalg.eigh(a)
        assert_(res_v.dtype.type is np.float64)
        assert_(res.dtype.type is np.float64)
        assert_equal(a.shape, res_v.shape)
        assert_equal((0, 1), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(a, np.ndarray))

        a = np.zeros((0, 0), dtype=np.complex64).view(ArraySubclass)
        res, res_v = linalg.eigh(a)
        assert_(res_v.dtype.type is np.complex64)
        assert_(res.dtype.type is np.float32)
        assert_equal(a.shape, res_v.shape)
        assert_equal((0,), res.shape)
        # This is just for documentation, it might make sense to change:
        assert_(isinstance(a, np.ndarray))


class _TestNormBase:
    dt = None
    dec = None

    @staticmethod
    def check_dtype(x, res):
        if issubclass(x.dtype.type, np.inexact):
            assert_equal(res.dtype, x.real.dtype)
        else:
            # For integer input, don't have to test float precision of output.
            assert_(issubclass(res.dtype.type, np.floating))


class _TestNormGeneral(_TestNormBase):

    def test_empty(self):
        assert_equal(norm([]), 0.0)
        assert_equal(norm(array([], dtype=self.dt)), 0.0)
        assert_equal(norm(atleast_2d(array([], dtype=self.dt))), 0.0)

    def test_vector_return_type(self):
        a = np.array([1, 0, 1])

        exact_types = np.typecodes['AllInteger']
        inexact_types = np.typecodes['AllFloat']

        all_types = exact_types + inexact_types

        for each_type in all_types:
            at = a.astype(each_type)

            an = norm(at, -np.inf)
            self.check_dtype(at, an)
            assert_almost_equal(an, 0.0)

            with suppress_warnings() as sup:
                sup.filter(RuntimeWarning, "divide by zero encountered")
                an = norm(at, -1)
                self.check_dtype(at, an)
                assert_almost_equal(an, 0.0)

            an = norm(at, 0)
            self.check_dtype(at, an)
            assert_almost_equal(an, 2)

            an = norm(at, 1)
            self.check_dtype(at, an)
            assert_almost_equal(an, 2.0)

            an = norm(at, 2)
            self.check_dtype(at, an)
            assert_almost_equal(an, an.dtype.type(2.0)**an.dtype.type(1.0 / 2.0))

            an = norm(at, 4)
            self.check_dtype(at, an)
            assert_almost_equal(an, an.dtype.type(2.0)**an.dtype.type(1.0 / 4.0))

            an = norm(at, np.inf)
            self.check_dtype(at, an)
            assert_almost_equal(an, 1.0)

    def test_vector(self):
        a = [1, 2, 3, 4]
        b = [-1, -2, -3, -4]
        c = [-1, 2, -3, 4]

        def _test(v):
            np.testing.assert_almost_equal(norm(v), 30 ** 0.5,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, inf), 4.0,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, -inf), 1.0,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, 1), 10.0,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, -1), 12.0 / 25,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, 2), 30 ** 0.5,
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, -2), ((205. / 144) ** -0.5),
                                           decimal=self.dec)
            np.testing.assert_almost_equal(norm(v, 0), 4,
                                           decimal=self.dec)

        for v in (a, b, c,):
            _test(v)

        for v in (array(a, dtype=self.dt), array(b, dtype=self.dt),
                  array(c, dtype=self.dt)):
            _test(v)

    def test_axis(self):
        # Vector norms.
        # Compare the use of `axis` with computing the norm of each row
        # or column separately.
        A = array([[1, 2, 3], [4, 5, 6]], dtype=self.dt)
        for order in [None, -1, 0, 1, 2, 3, np.inf, -np.inf]:
            expected0 = [norm(A[:, k], ord=order) for k in range(A.shape[1])]
            assert_almost_equal(norm(A, ord=order, axis=0), expected0)
            expected1 = [norm(A[k, :], ord=order) for k in range(A.shape[0])]
            assert_almost_equal(norm(A, ord=order, axis=1), expected1)

        # Matrix norms.
        B = np.arange(1, 25, dtype=self.dt).reshape(2, 3, 4)
        nd = B.ndim
        for order in [None, -2, 2, -1, 1, np.inf, -np.inf, 'fro']:
            for axis in itertools.combinations(range(-nd, nd), 2):
                row_axis, col_axis = axis
                if row_axis < 0:
                    row_axis += nd
                if col_axis < 0:
                    col_axis += nd
                if row_axis == col_axis:
                    assert_raises(ValueError, norm, B, ord=order, axis=axis)
                else:
                    n = norm(B, ord=order, axis=axis)

                    # The logic using k_index only works for nd = 3.
                    # This has to be changed if nd is increased.
                    k_index = nd - (row_axis + col_axis)
                    if row_axis < col_axis:
                        expected = [norm(B[:].take(k, axis=k_index), ord=order)
                                    for k in range(B.shape[k_index])]
                    else:
                        expected = [norm(B[:].take(k, axis=k_index).T, ord=order)
                                    for k in range(B.shape[k_index])]
                    assert_almost_equal(n, expected)

    def test_keepdims(self):
        A = np.arange(1, 25, dtype=self.dt).reshape(2, 3, 4)

        allclose_err = 'order {0}, axis = {1}'
        shape_err = 'Shape mismatch found {0}, expected {1}, order={2}, axis={3}'

        # check the order=None, axis=None case
        expected = norm(A, ord=None, axis=None)
        found = norm(A, ord=None, axis=None, keepdims=True)
        assert_allclose(np.squeeze(found), expected,
                        err_msg=allclose_err.format(None, None))
        expected_shape = (1, 1, 1)
        assert_(found.shape == expected_shape,
                shape_err.format(found.shape, expected_shape, None, None))

        # Vector norms.
        for order in [None, -1, 0, 1, 2, 3, np.inf, -np.inf]:
            for k in range(A.ndim):
                expected = norm(A, ord=order, axis=k)
                found = norm(A, ord=order, axis=k, keepdims=True)
                assert_allclose(np.squeeze(found), expected,
                                err_msg=allclose_err.format(order, k))
                expected_shape = list(A.shape)
                expected_shape[k] = 1
                expected_shape = tuple(expected_shape)
                assert_(found.shape == expected_shape,
                        shape_err.format(found.shape, expected_shape, order, k))

        # Matrix norms.
        for order in [None, -2, 2, -1, 1, np.inf, -np.inf, 'fro', 'nuc']:
            for k in itertools.permutations(range(A.ndim), 2):
                expected = norm(A, ord=order, axis=k)
                found = norm(A, ord=order, axis=k, keepdims=True)
                assert_allclose(np.squeeze(found), expected,
                                err_msg=allclose_err.format(order, k))
                expected_shape = list(A.shape)
                expected_shape[k[0]] = 1
                expected_shape[k[1]] = 1
                expected_shape = tuple(expected_shape)
                assert_(found.shape == expected_shape,
                        shape_err.format(found.shape, expected_shape, order, k))


class _TestNorm2D(_TestNormBase):
    # Define the part for 2d arrays separately, so we can subclass this
    # and run the tests using np.matrix in matrixlib.tests.test_matrix_linalg.
    array = np.array

    def test_matrix_empty(self):
        assert_equal(norm(self.array([[]], dtype=self.dt)), 0.0)

    def test_matrix_return_type(self):
        a = self.array([[1, 0, 1], [0, 1, 1]])

        exact_types = np.typecodes['AllInteger']

        # float32, complex64, float64, complex128 types are the only types
        # allowed by `linalg`, which performs the matrix operations used
        # within `norm`.
        inexact_types = 'fdFD'

        all_types = exact_types + inexact_types

        for each_type in all_types:
            at = a.astype(each_type)

            an = norm(at, -np.inf)
            self.check_dtype(at, an)
            assert_almost_equal(an, 2.0)

            with suppress_warnings() as sup:
                sup.filter(RuntimeWarning, "divide by zero encountered")
                an = norm(at, -1)
                self.check_dtype(at, an)
                assert_almost_equal(an, 1.0)

            an = norm(at, 1)
            self.check_dtype(at, an)
            assert_almost_equal(an, 2.0)

            an = norm(at, 2)
            self.check_dtype(at, an)
            assert_almost_equal(an, 3.0**(1.0 / 2.0))

            an = norm(at, -2)
            self.check_dtype(at, an)
            assert_almost_equal(an, 1.0)

            an = norm(at, np.inf)
            self.check_dtype(at, an)
            assert_almost_equal(an, 2.0)

            an = norm(at, 'fro')
            self.check_dtype(at, an)
            assert_almost_equal(an, 2.0)

            an = norm(at, 'nuc')
            self.check_dtype(at, an)
            # Lower bar needed to support low precision floats.
            # They end up being off by 1 in the 7th place.
            np.testing.assert_almost_equal(an, 2.7320508075688772, decimal=6)

    def test_matrix_2x2(self):
        A = self.array([[1, 3], [5, 7]], dtype=self.dt)
        assert_almost_equal(norm(A), 84 ** 0.5)
        assert_almost_equal(norm(A, 'fro'), 84 ** 0.5)
        assert_almost_equal(norm(A, 'nuc'), 10.0)
        assert_almost_equal(norm(A, inf), 12.0)
        assert_almost_equal(norm(A, -inf), 4.0)
        assert_almost_equal(norm(A, 1), 10.0)
        assert_almost_equal(norm(A, -1), 6.0)
        assert_almost_equal(norm(A, 2), 9.1231056256176615)
        assert_almost_equal(norm(A, -2), 0.87689437438234041)

        assert_raises(ValueError, norm, A, 'nofro')
        assert_raises(ValueError, norm, A, -3)
        assert_raises(ValueError, norm, A, 0)

    def test_matrix_3x3(self):
        # This test has been added because the 2x2 example
        # happened to have equal nuclear norm and induced 1-norm.
        # The 1/10 scaling factor accommodates the absolute tolerance
        # used in assert_almost_equal.
        A = (1 / 10) * \
            self.array([[1, 2, 3], [6, 0, 5], [3, 2, 1]], dtype=self.dt)
        assert_almost_equal(norm(A), (1 / 10) * 89 ** 0.5)
        assert_almost_equal(norm(A, 'fro'), (1 / 10) * 89 ** 0.5)
        assert_almost_equal(norm(A, 'nuc'), 1.3366836911774836)
        assert_almost_equal(norm(A, inf), 1.1)
        assert_almost_equal(norm(A, -inf), 0.6)
        assert_almost_equal(norm(A, 1), 1.0)
        assert_almost_equal(norm(A, -1), 0.4)
        assert_almost_equal(norm(A, 2), 0.88722940323461277)
        assert_almost_equal(norm(A, -2), 0.19456584790481812)

    def test_bad_args(self):
        # Check that bad arguments raise the appropriate exceptions.

        A = self.array([[1, 2, 3], [4, 5, 6]], dtype=self.dt)
        B = np.arange(1, 25, dtype=self.dt).reshape(2, 3, 4)

        # Using `axis=<integer>` or passing in a 1-D array implies vector
        # norms are being computed, so also using `ord='fro'`
        # or `ord='nuc'` or any other string raises a ValueError.
        assert_raises(ValueError, norm, A, 'fro', 0)
        assert_raises(ValueError, norm, A, 'nuc', 0)
        assert_raises(ValueError, norm, [3, 4], 'fro', None)
        assert_raises(ValueError, norm, [3, 4], 'nuc', None)
        assert_raises(ValueError, norm, [3, 4], 'test', None)

        # Similarly, norm should raise an exception when ord is any finite
        # number other than 1, 2, -1 or -2 when computing matrix norms.
        for order in [0, 3]:
            assert_raises(ValueError, norm, A, order, None)
            assert_raises(ValueError, norm, A, order, (0, 1))
            assert_raises(ValueError, norm, B, order, (1, 2))

        # Invalid axis
        assert_raises(AxisError, norm, B, None, 3)
        assert_raises(AxisError, norm, B, None, (2, 3))
        assert_raises(ValueError, norm, B, None, (0, 1, 2))


class _TestNorm(_TestNorm2D, _TestNormGeneral):
    pass


class TestNorm_NonSystematic:

    def test_longdouble_norm(self):
        # Non-regression test: p-norm of longdouble would previously raise
        # UnboundLocalError.
        x = np.arange(10, dtype=np.longdouble)
        old_assert_almost_equal(norm(x, ord=3), 12.65, decimal=2)

    def test_intmin(self):
        # Non-regression test: p-norm of signed integer would previously do
        # float cast and abs in the wrong order.
        x = np.array([-2 ** 31], dtype=np.int32)
        old_assert_almost_equal(norm(x, ord=3), 2 ** 31, decimal=5)

    def test_complex_high_ord(self):
        # gh-4156
        d = np.empty((2,), dtype=np.clongdouble)
        d[0] = 6 + 7j
        d[1] = -6 + 7j
        res = 11.615898132184
        old_assert_almost_equal(np.linalg.norm(d, ord=3), res, decimal=10)
        d = d.astype(np.complex128)
        old_assert_almost_equal(np.linalg.norm(d, ord=3), res, decimal=9)
        d = d.astype(np.complex64)
        old_assert_almost_equal(np.linalg.norm(d, ord=3), res, decimal=5)


# Separate definitions so we can use them for matrix tests.
class _TestNormDoubleBase(_TestNormBase):
    dt = np.double
    dec = 12


class _TestNormSingleBase(_TestNormBase):
    dt = np.float32
    dec = 6


class _TestNormInt64Base(_TestNormBase):
    dt = np.int64
    dec = 12


class TestNormDouble(_TestNorm, _TestNormDoubleBase):
    pass


class TestNormSingle(_TestNorm, _TestNormSingleBase):
    pass


class TestNormInt64(_TestNorm, _TestNormInt64Base):
    pass


class TestMatrixRank:

    def test_matrix_rank(self):
        # Full rank matrix
        assert_equal(4, matrix_rank(np.eye(4)))
        # rank deficient matrix
        I = np.eye(4)
        I[-1, -1] = 0.
        assert_equal(matrix_rank(I), 3)
        # All zeros - zero rank
        assert_equal(matrix_rank(np.zeros((4, 4))), 0)
        # 1 dimension - rank 1 unless all 0
        assert_equal(matrix_rank([1, 0, 0, 0]), 1)
        assert_equal(matrix_rank(np.zeros((4,))), 0)
        # accepts array-like
        assert_equal(matrix_rank([1]), 1)
        # greater than 2 dimensions treated as stacked matrices
        ms = np.array([I, np.eye(4), np.zeros((4, 4))])
        assert_equal(matrix_rank(ms), np.array([3, 4, 0]))
        # works on scalar
        assert_equal(matrix_rank(1), 1)

        with assert_raises_regex(
            ValueError, "`tol` and `rtol` can\'t be both set."
        ):
            matrix_rank(I, tol=0.01, rtol=0.01)

    def test_symmetric_rank(self):
        assert_equal(4, matrix_rank(np.eye(4), hermitian=True))
        assert_equal(1, matrix_rank(np.ones((4, 4)), hermitian=True))
        assert_equal(0, matrix_rank(np.zeros((4, 4)), hermitian=True))
        # rank deficient matrix
        I = np.eye(4)
        I[-1, -1] = 0.
        assert_equal(3, matrix_rank(I, hermitian=True))
        # manually supplied tolerance
        I[-1, -1] = 1e-8
        assert_equal(4, matrix_rank(I, hermitian=True, tol=0.99e-8))
        assert_equal(3, matrix_rank(I, hermitian=True, tol=1.01e-8))


def test_reduced_rank():
    # Test matrices with reduced rank
    rng = np.random.RandomState(20120714)
    for i in range(100):
        # Make a rank deficient matrix
        X = rng.normal(size=(40, 10))
        X[:, 0] = X[:, 1] + X[:, 2]
        # Assert that matrix_rank detected deficiency
        assert_equal(matrix_rank(X), 9)
        X[:, 3] = X[:, 4] + X[:, 5]
        assert_equal(matrix_rank(X), 8)


class TestQR:
    # Define the array class here, so run this on matrices elsewhere.
    array = np.array

    def check_qr(self, a):
        # This test expects the argument `a` to be an ndarray or
        # a subclass of an ndarray of inexact type.
        a_type = type(a)
        a_dtype = a.dtype
        m, n = a.shape
        k = min(m, n)

        # mode == 'complete'
        res = linalg.qr(a, mode='complete')
        Q, R = res.Q, res.R
        assert_(Q.dtype == a_dtype)
        assert_(R.dtype == a_dtype)
        assert_(isinstance(Q, a_type))
        assert_(isinstance(R, a_type))
        assert_(Q.shape == (m, m))
        assert_(R.shape == (m, n))
        assert_almost_equal(dot(Q, R), a)
        assert_almost_equal(dot(Q.T.conj(), Q), np.eye(m))
        assert_almost_equal(np.triu(R), R)

        # mode == 'reduced'
        q1, r1 = linalg.qr(a, mode='reduced')
        assert_(q1.dtype == a_dtype)
        assert_(r1.dtype == a_dtype)
        assert_(isinstance(q1, a_type))
        assert_(isinstance(r1, a_type))
        assert_(q1.shape == (m, k))
        assert_(r1.shape == (k, n))
        assert_almost_equal(dot(q1, r1), a)
        assert_almost_equal(dot(q1.T.conj(), q1), np.eye(k))
        assert_almost_equal(np.triu(r1), r1)

        # mode == 'r'
        r2 = linalg.qr(a, mode='r')
        assert_(r2.dtype == a_dtype)
        assert_(isinstance(r2, a_type))
        assert_almost_equal(r2, r1)

    @pytest.mark.parametrize(["m", "n"], [
        (3, 0),
        (0, 3),
        (0, 0)
    ])
    def test_qr_empty(self, m, n):
        k = min(m, n)
        a = np.empty((m, n))

        self.check_qr(a)

        h, tau = np.linalg.qr(a, mode='raw')
        assert_equal(h.dtype, np.double)
        assert_equal(tau.dtype, np.double)
        assert_equal(h.shape, (n, m))
        assert_equal(tau.shape, (k,))

    def test_mode_raw(self):
        # The factorization is not unique and varies between libraries,
        # so it is not possible to check against known values. Functional
        # testing is a possibility, but awaits the exposure of more
        # of the functions in lapack_lite. Consequently, this test is
        # very limited in scope. Note that the results are in FORTRAN
        # order, hence the h arrays are transposed.
        a = self.array([[1, 2], [3, 4], [5, 6]], dtype=np.double)

        # Test double
        h, tau = linalg.qr(a, mode='raw')
        assert_(h.dtype == np.double)
        assert_(tau.dtype == np.double)
        assert_(h.shape == (2, 3))
        assert_(tau.shape == (2,))

        h, tau = linalg.qr(a.T, mode='raw')
        assert_(h.dtype == np.double)
        assert_(tau.dtype == np.double)
        assert_(h.shape == (3, 2))
        assert_(tau.shape == (2,))

    def test_mode_all_but_economic(self):
        a = self.array([[1, 2], [3, 4]])
        b = self.array([[1, 2], [3, 4], [5, 6]])
        for dt in "fd":
            m1 = a.astype(dt)
            m2 = b.astype(dt)
            self.check_qr(m1)
            self.check_qr(m2)
            self.check_qr(m2.T)

        for dt in "fd":
            m1 = 1 + 1j * a.astype(dt)
            m2 = 1 + 1j * b.astype(dt)
            self.check_qr(m1)
            self.check_qr(m2)
            self.check_qr(m2.T)

    def check_qr_stacked(self, a):
        # This test expects the argument `a` to be an ndarray or
        # a subclass of an ndarray of inexact type.
        a_type = type(a)
        a_dtype = a.dtype
        m, n = a.shape[-2:]
        k = min(m, n)

        # mode == 'complete'
        q, r = linalg.qr(a, mode='complete')
        assert_(q.dtype == a_dtype)
        assert_(r.dtype == a_dtype)
        assert_(isinstance(q, a_type))
        assert_(isinstance(r, a_type))
        assert_(q.shape[-2:] == (m, m))
        assert_(r.shape[-2:] == (m, n))
        assert_almost_equal(matmul(q, r), a)
        I_mat = np.identity(q.shape[-1])
        stack_I_mat = np.broadcast_to(I_mat,
                        q.shape[:-2] + (q.shape[-1],) * 2)
        assert_almost_equal(matmul(swapaxes(q, -1, -2).conj(), q), stack_I_mat)
        assert_almost_equal(np.triu(r[..., :, :]), r)

        # mode == 'reduced'
        q1, r1 = linalg.qr(a, mode='reduced')
        assert_(q1.dtype == a_dtype)
        assert_(r1.dtype == a_dtype)
        assert_(isinstance(q1, a_type))
        assert_(isinstance(r1, a_type))
        assert_(q1.shape[-2:] == (m, k))
        assert_(r1.shape[-2:] == (k, n))
        assert_almost_equal(matmul(q1, r1), a)
        I_mat = np.identity(q1.shape[-1])
        stack_I_mat = np.broadcast_to(I_mat,
                        q1.shape[:-2] + (q1.shape[-1],) * 2)
        assert_almost_equal(matmul(swapaxes(q1, -1, -2).conj(), q1),
                            stack_I_mat)
        assert_almost_equal(np.triu(r1[..., :, :]), r1)

        # mode == 'r'
        r2 = linalg.qr(a, mode='r')
        assert_(r2.dtype == a_dtype)
        assert_(isinstance(r2, a_type))
        assert_almost_equal(r2, r1)

    @pytest.mark.parametrize("size", [
        (3, 4), (4, 3), (4, 4),
        (3, 0), (0, 3)])
    @pytest.mark.parametrize("outer_size", [
        (2, 2), (2,), (2, 3, 4)])
    @pytest.mark.parametrize("dt", [
        np.single, np.double,
        np.csingle, np.cdouble])
    def test_stacked_inputs(self, outer_size, size, dt):

        rng = np.random.default_rng(123)
        A = rng.normal(size=outer_size + size).astype(dt)
        B = rng.normal(size=outer_size + size).astype(dt)
        self.check_qr_stacked(A)
        self.check_qr_stacked(A + 1.j * B)


class TestCholesky:

    @pytest.mark.parametrize(
        'shape', [(1, 1), (2, 2), (3, 3), (50, 50), (3, 10, 10)]
    )
    @pytest.mark.parametrize(
        'dtype', (np.float32, np.float64, np.complex64, np.complex128)
    )
    @pytest.mark.parametrize(
        'upper', [False, True])
    def test_basic_property(self, shape, dtype, upper):
        np.random.seed(1)
        a = np.random.randn(*shape)
        if np.issubdtype(dtype, np.complexfloating):
            a = a + 1j * np.random.randn(*shape)

        t = list(range(len(shape)))
        t[-2:] = -1, -2

        a = np.matmul(a.transpose(t).conj(), a)
        a = np.asarray(a, dtype=dtype)

        c = np.linalg.cholesky(a, upper=upper)

        # Check A = L L^H or A = U^H U
        if upper:
            b = np.matmul(c.transpose(t).conj(), c)
        else:
            b = np.matmul(c, c.transpose(t).conj())

        atol = 500 * a.shape[0] * np.finfo(dtype).eps
        assert_allclose(b, a, atol=atol, err_msg=f'{shape} {dtype}\n{a}\n{c}')

        # Check diag(L or U) is real and positive
        d = np.diagonal(c, axis1=-2, axis2=-1)
        assert_(np.all(np.isreal(d)))
        assert_(np.all(d >= 0))

    def test_0_size(self):
        class ArraySubclass(np.ndarray):
            pass
        a = np.zeros((0, 1, 1), dtype=np.int_).view(ArraySubclass)
        res = linalg.cholesky(a)
        assert_equal(a.shape, res.shape)
        assert_(res.dtype.type is np.float64)
        # for documentation purpose:
        assert_(isinstance(res, np.ndarray))

        a = np.zeros((1, 0, 0), dtype=np.complex64).view(ArraySubclass)
        res = linalg.cholesky(a)
        assert_equal(a.shape, res.shape)
        assert_(res.dtype.type is np.complex64)
        assert_(isinstance(res, np.ndarray))

    def test_upper_lower_arg(self):
        # Explicit test of upper argument that also checks the default.
        a = np.array([[1 + 0j, 0 - 2j], [0 + 2j, 5 + 0j]])

        assert_equal(linalg.cholesky(a), linalg.cholesky(a, upper=False))

        assert_equal(
            linalg.cholesky(a, upper=True),
            linalg.cholesky(a).T.conj()
        )


class TestOuter:
    arr1 = np.arange(3)
    arr2 = np.arange(3)
    expected = np.array(
        [[0, 0, 0],
         [0, 1, 2],
         [0, 2, 4]]
    )

    assert_array_equal(np.linalg.outer(arr1, arr2), expected)

    with assert_raises_regex(
        ValueError, "Input arrays must be one-dimensional"
    ):
        np.linalg.outer(arr1[:, np.newaxis], arr2)


def test_byteorder_check():
    # Byte order check should pass for native order
    if sys.byteorder == 'little':
        native = '<'
    else:
        native = '>'

    for dtt in (np.float32, np.float64):
        arr = np.eye(4, dtype=dtt)
        n_arr = arr.view(arr.dtype.newbyteorder(native))
        sw_arr = arr.view(arr.dtype.newbyteorder("S")).byteswap()
        assert_equal(arr.dtype.byteorder, '=')
        for routine in (linalg.inv, linalg.det, linalg.pinv):
            # Normal call
            res = routine(arr)
            # Native but not '='
            assert_array_equal(res, routine(n_arr))
            # Swapped
            assert_array_equal(res, routine(sw_arr))


@pytest.mark.skipif(IS_WASM, reason="fp errors don't work in wasm")
def test_generalized_raise_multiloop():
    # It should raise an error even if the error doesn't occur in the
    # last iteration of the ufunc inner loop

    invertible = np.array([[1, 2], [3, 4]])
    non_invertible = np.array([[1, 1], [1, 1]])

    x = np.zeros([4, 4, 2, 2])[1::2]
    x[...] = invertible
    x[0, 0] = non_invertible

    assert_raises(np.linalg.LinAlgError, np.linalg.inv, x)


@pytest.mark.skipif(
    threading.active_count() > 1,
    reason="skipping test that uses fork because there are multiple threads")
@pytest.mark.skipif(
    NOGIL_BUILD,
    reason="Cannot safely use fork in tests on the free-threaded build")
def test_xerbla_override():
    # Check that our xerbla has been successfully linked in. If it is not,
    # the default xerbla routine is called, which prints a message to stdout
    # and may, or may not, abort the process depending on the LAPACK package.

    XERBLA_OK = 255

    try:
        pid = os.fork()
    except (OSError, AttributeError):
        # fork failed, or not running on POSIX
        pytest.skip("Not POSIX or fork failed.")

    if pid == 0:
        # child; close i/o file handles
        os.close(1)
        os.close(0)
        # Avoid producing core files.
        import resource
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
        # These calls may abort.
        try:
            np.linalg.lapack_lite.xerbla()
        except ValueError:
            pass
        except Exception:
            os._exit(os.EX_CONFIG)

        try:
            a = np.array([[1.]])
            np.linalg.lapack_lite.dorgqr(
                1, 1, 1, a,
                0,  # <- invalid value
                a, a, 0, 0)
        except ValueError as e:
            if "DORGQR parameter number 5" in str(e):
                # success, reuse error code to mark success as
                # FORTRAN STOP returns as success.
                os._exit(XERBLA_OK)

        # Did not abort, but our xerbla was not linked in.
        os._exit(os.EX_CONFIG)
    else:
        # parent
        pid, status = os.wait()
        if os.WEXITSTATUS(status) != XERBLA_OK:
            pytest.skip('Numpy xerbla not linked in.')


@pytest.mark.skipif(IS_WASM, reason="Cannot start subprocess")
@pytest.mark.slow
def test_sdot_bug_8577():
    # Regression test that loading certain other libraries does not
    # result to wrong results in float32 linear algebra.
    #
    # There's a bug gh-8577 on OSX that can trigger this, and perhaps
    # there are also other situations in which it occurs.
    #
    # Do the check in a separate process.

    bad_libs = ['PyQt5.QtWidgets', 'IPython']

    template = textwrap.dedent("""
    import sys
    {before}
    try:
        import {bad_lib}
    except ImportError:
        sys.exit(0)
    {after}
    x = np.ones(2, dtype=np.float32)
    sys.exit(0 if np.allclose(x.dot(x), 2.0) else 1)
    """)

    for bad_lib in bad_libs:
        code = template.format(before="import numpy as np", after="",
                               bad_lib=bad_lib)
        subprocess.check_call([sys.executable, "-c", code])

        # Swapped import order
        code = template.format(after="import numpy as np", before="",
                               bad_lib=bad_lib)
        subprocess.check_call([sys.executable, "-c", code])


class TestMultiDot:

    def test_basic_function_with_three_arguments(self):
        # multi_dot with three arguments uses a fast hand coded algorithm to
        # determine the optimal order. Therefore test it separately.
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))

        assert_almost_equal(multi_dot([A, B, C]), A.dot(B).dot(C))
        assert_almost_equal(multi_dot([A, B, C]), np.dot(A, np.dot(B, C)))

    def test_basic_function_with_two_arguments(self):
        # separate code path with two arguments
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))

        assert_almost_equal(multi_dot([A, B]), A.dot(B))
        assert_almost_equal(multi_dot([A, B]), np.dot(A, B))

    def test_basic_function_with_dynamic_programming_optimization(self):
        # multi_dot with four or more arguments uses the dynamic programming
        # optimization and therefore deserve a separate
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))
        D = np.random.random((2, 1))
        assert_almost_equal(multi_dot([A, B, C, D]), A.dot(B).dot(C).dot(D))

    def test_vector_as_first_argument(self):
        # The first argument can be 1-D
        A1d = np.random.random(2)  # 1-D
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))
        D = np.random.random((2, 2))

        # the result should be 1-D
        assert_equal(multi_dot([A1d, B, C, D]).shape, (2,))

    def test_vector_as_last_argument(self):
        # The last argument can be 1-D
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))
        D1d = np.random.random(2)  # 1-D

        # the result should be 1-D
        assert_equal(multi_dot([A, B, C, D1d]).shape, (6,))

    def test_vector_as_first_and_last_argument(self):
        # The first and last arguments can be 1-D
        A1d = np.random.random(2)  # 1-D
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))
        D1d = np.random.random(2)  # 1-D

        # the result should be a scalar
        assert_equal(multi_dot([A1d, B, C, D1d]).shape, ())

    def test_three_arguments_and_out(self):
        # multi_dot with three arguments uses a fast hand coded algorithm to
        # determine the optimal order. Therefore test it separately.
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))

        out = np.zeros((6, 2))
        ret = multi_dot([A, B, C], out=out)
        assert out is ret
        assert_almost_equal(out, A.dot(B).dot(C))
        assert_almost_equal(out, np.dot(A, np.dot(B, C)))

    def test_two_arguments_and_out(self):
        # separate code path with two arguments
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        out = np.zeros((6, 6))
        ret = multi_dot([A, B], out=out)
        assert out is ret
        assert_almost_equal(out, A.dot(B))
        assert_almost_equal(out, np.dot(A, B))

    def test_dynamic_programming_optimization_and_out(self):
        # multi_dot with four or more arguments uses the dynamic programming
        # optimization and therefore deserve a separate test
        A = np.random.random((6, 2))
        B = np.random.random((2, 6))
        C = np.random.random((6, 2))
        D = np.random.random((2, 1))
        out = np.zeros((6, 1))
        ret = multi_dot([A, B, C, D], out=out)
        assert out is ret
        assert_almost_equal(out, A.dot(B).dot(C).dot(D))

    def test_dynamic_programming_logic(self):
        # Test for the dynamic programming part
        # This test is directly taken from Cormen page 376.
        arrays = [np.random.random((30, 35)),
                  np.random.random((35, 15)),
                  np.random.random((15, 5)),
                  np.random.random((5, 10)),
                  np.random.random((10, 20)),
                  np.random.random((20, 25))]
        m_expected = np.array([[0., 15750., 7875., 9375., 11875., 15125.],
                               [0.,     0., 2625., 4375.,  7125., 10500.],
                               [0.,     0.,    0.,  750.,  2500.,  5375.],
                               [0.,     0.,    0.,    0.,  1000.,  3500.],
                               [0.,     0.,    0.,    0.,     0.,  5000.],
                               [0.,     0.,    0.,    0.,     0.,     0.]])
        s_expected = np.array([[0,  1,  1,  3,  3,  3],
                               [0,  0,  2,  3,  3,  3],
                               [0,  0,  0,  3,  3,  3],
                               [0,  0,  0,  0,  4,  5],
                               [0,  0,  0,  0,  0,  5],
                               [0,  0,  0,  0,  0,  0]], dtype=int)
        s_expected -= 1  # Cormen uses 1-based index, python does not.

        s, m = _multi_dot_matrix_chain_order(arrays, return_costs=True)

        # Only the upper triangular part (without the diagonal) is interesting.
        assert_almost_equal(np.triu(s[:-1, 1:]),
                            np.triu(s_expected[:-1, 1:]))
        assert_almost_equal(np.triu(m), np.triu(m_expected))

    def test_too_few_input_arrays(self):
        assert_raises(ValueError, multi_dot, [])
        assert_raises(ValueError, multi_dot, [np.random.random((3, 3))])


class TestTensorinv:

    @pytest.mark.parametrize("arr, ind", [
        (np.ones((4, 6, 8, 2)), 2),
        (np.ones((3, 3, 2)), 1),
        ])
    def test_non_square_handling(self, arr, ind):
        with assert_raises(LinAlgError):
            linalg.tensorinv(arr, ind=ind)

    @pytest.mark.parametrize("shape, ind", [
        # examples from docstring
        ((4, 6, 8, 3), 2),
        ((24, 8, 3), 1),
        ])
    def test_tensorinv_shape(self, shape, ind):
        a = np.eye(24)
        a.shape = shape
        ainv = linalg.tensorinv(a=a, ind=ind)
        expected = a.shape[ind:] + a.shape[:ind]
        actual = ainv.shape
        assert_equal(actual, expected)

    @pytest.mark.parametrize("ind", [
        0, -2,
        ])
    def test_tensorinv_ind_limit(self, ind):
        a = np.eye(24)
        a.shape = (4, 6, 8, 3)
        with assert_raises(ValueError):
            linalg.tensorinv(a=a, ind=ind)

    def test_tensorinv_result(self):
        # mimic a docstring example
        a = np.eye(24)
        a.shape = (24, 8, 3)
        ainv = linalg.tensorinv(a, ind=1)
        b = np.ones(24)
        assert_allclose(np.tensordot(ainv, b, 1), np.linalg.tensorsolve(a, b))


class TestTensorsolve:

    @pytest.mark.parametrize("a, axes", [
        (np.ones((4, 6, 8, 2)), None),
        (np.ones((3, 3, 2)), (0, 2)),
        ])
    def test_non_square_handling(self, a, axes):
        with assert_raises(LinAlgError):
            b = np.ones(a.shape[:2])
            linalg.tensorsolve(a, b, axes=axes)

    @pytest.mark.parametrize("shape",
        [(2, 3, 6), (3, 4, 4, 3), (0, 3, 3, 0)],
    )
    def test_tensorsolve_result(self, shape):
        a = np.random.randn(*shape)
        b = np.ones(a.shape[:2])
        x = np.linalg.tensorsolve(a, b)
        assert_allclose(np.tensordot(a, x, axes=len(x.shape)), b)


def test_unsupported_commontype():
    # linalg gracefully handles unsupported type
    arr = np.array([[1, -2], [2, 5]], dtype='float16')
    with assert_raises_regex(TypeError, "unsupported in linalg"):
        linalg.cholesky(arr)


#@pytest.mark.slow
#@pytest.mark.xfail(not HAS_LAPACK64, run=False,
#                   reason="Numpy not compiled with 64-bit BLAS/LAPACK")
#@requires_memory(free_bytes=16e9)
@pytest.mark.skip(reason="Bad memory reports lead to OOM in ci testing")
def test_blas64_dot():
    n = 2**32
    a = np.zeros([1, n], dtype=np.float32)
    b = np.ones([1, 1], dtype=np.float32)
    a[0, -1] = 1
    c = np.dot(b, a)
    assert_equal(c[0, -1], 1)


@pytest.mark.xfail(not HAS_LAPACK64,
                   reason="Numpy not compiled with 64-bit BLAS/LAPACK")
def test_blas64_geqrf_lwork_smoketest():
    # Smoke test LAPACK geqrf lwork call with 64-bit integers
    dtype = np.float64
    lapack_routine = np.linalg.lapack_lite.dgeqrf

    m = 2**32 + 1
    n = 2**32 + 1
    lda = m

    # Dummy arrays, not referenced by the lapack routine, so don't
    # need to be of the right size
    a = np.zeros([1, 1], dtype=dtype)
    work = np.zeros([1], dtype=dtype)
    tau = np.zeros([1], dtype=dtype)

    # Size query
    results = lapack_routine(m, n, a, lda, tau, work, -1, 0)
    assert_equal(results['info'], 0)
    assert_equal(results['m'], m)
    assert_equal(results['n'], m)

    # Should result to an integer of a reasonable size
    lwork = int(work.item())
    assert_(2**32 < lwork < 2**42)


def test_diagonal():
    # Here we only test if selected axes are compatible
    # with Array API (last two). Core implementation
    # of `diagonal` is tested in `test_multiarray.py`.
    x = np.arange(60).reshape((3, 4, 5))
    actual = np.linalg.diagonal(x)
    expected = np.array(
        [
            [0,  6, 12, 18],
            [20, 26, 32, 38],
            [40, 46, 52, 58],
        ]
    )
    assert_equal(actual, expected)


def test_trace():
    # Here we only test if selected axes are compatible
    # with Array API (last two). Core implementation
    # of `trace` is tested in `test_multiarray.py`.
    x = np.arange(60).reshape((3, 4, 5))
    actual = np.linalg.trace(x)
    expected = np.array([36, 116, 196])

    assert_equal(actual, expected)


def test_cross():
    x = np.arange(9).reshape((3, 3))
    actual = np.linalg.cross(x, x + 1)
    expected = np.array([
        [-1, 2, -1],
        [-1, 2, -1],
        [-1, 2, -1],
    ])

    assert_equal(actual, expected)

    # We test that lists are converted to arrays.
    u = [1, 2, 3]
    v = [4, 5, 6]
    actual = np.linalg.cross(u, v)
    expected = array([-3,  6, -3])

    assert_equal(actual, expected)

    with assert_raises_regex(
        ValueError,
        r"input arrays must be \(arrays of\) 3-dimensional vectors"
    ):
        x_2dim = x[:, 1:]
        np.linalg.cross(x_2dim, x_2dim)


def test_tensordot():
    # np.linalg.tensordot is just an alias for np.tensordot
    x = np.arange(6).reshape((2, 3))

    assert np.linalg.tensordot(x, x) == 55
    assert np.linalg.tensordot(x, x, axes=[(0, 1), (0, 1)]) == 55


def test_matmul():
    # np.linalg.matmul and np.matmul only differs in the number
    # of arguments in the signature
    x = np.arange(6).reshape((2, 3))
    actual = np.linalg.matmul(x, x.T)
    expected = np.array([[5, 14], [14, 50]])

    assert_equal(actual, expected)


def test_matrix_transpose():
    x = np.arange(6).reshape((2, 3))
    actual = np.linalg.matrix_transpose(x)
    expected = x.T

    assert_equal(actual, expected)

    with assert_raises_regex(
        ValueError, "array must be at least 2-dimensional"
    ):
        np.linalg.matrix_transpose(x[:, 0])


def test_matrix_norm():
    x = np.arange(9).reshape((3, 3))
    actual = np.linalg.matrix_norm(x)

    assert_almost_equal(actual, np.float64(14.2828), double_decimal=3)

    actual = np.linalg.matrix_norm(x, keepdims=True)

    assert_almost_equal(actual, np.array([[14.2828]]), double_decimal=3)


def test_matrix_norm_empty():
    for shape in [(0, 2), (2, 0), (0, 0)]:
        for dtype in [np.float64, np.float32, np.int32]:
            x = np.zeros(shape, dtype)
            assert_equal(np.linalg.matrix_norm(x, ord="fro"), 0)
            assert_equal(np.linalg.matrix_norm(x, ord="nuc"), 0)
            assert_equal(np.linalg.matrix_norm(x, ord=1), 0)
            assert_equal(np.linalg.matrix_norm(x, ord=2), 0)
            assert_equal(np.linalg.matrix_norm(x, ord=np.inf), 0)

def test_vector_norm():
    x = np.arange(9).reshape((3, 3))
    actual = np.linalg.vector_norm(x)

    assert_almost_equal(actual, np.float64(14.2828), double_decimal=3)

    actual = np.linalg.vector_norm(x, axis=0)

    assert_almost_equal(
        actual, np.array([6.7082, 8.124, 9.6436]), double_decimal=3
    )

    actual = np.linalg.vector_norm(x, keepdims=True)
    expected = np.full((1, 1), 14.2828, dtype='float64')
    assert_equal(actual.shape, expected.shape)
    assert_almost_equal(actual, expected, double_decimal=3)


def test_vector_norm_empty():
    for dtype in [np.float64, np.float32, np.int32]:
        x = np.zeros(0, dtype)
        assert_equal(np.linalg.vector_norm(x, ord=1), 0)
        assert_equal(np.linalg.vector_norm(x, ord=2), 0)
        assert_equal(np.linalg.vector_norm(x, ord=np.inf), 0)

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\quantization\operators\softmax.py ===
import onnx
import onnx.helper

from ..quant_utils import TENSOR_NAME_QUANT_SUFFIX, QuantizedValue, QuantizedValueType, attribute_to_kwarg, ms_domain
from .base_operator import QuantOperatorBase


class QLinearSoftmax(QuantOperatorBase):
    def quantize(self):
        node = self.node
        # set limitations for softmax output scale and zp, because the output of softmax is always 0-1
        if self.quantizer.activation_qType == onnx.onnx_pb.TensorProto.UINT8:
            out_scale = 1 / 256.0
            out_zero_point = 0
        else:
            out_scale = 1 / 256.0
            out_zero_point = -128
        # only try to quantize when given quantization parameters for it
        (
            data_found,
            output_scale_name,
            output_zp_name,
            _,
            _,
        ) = self.quantizer._get_quantization_params(node.output[0], out_scale, out_zero_point)

        # get quantized input tensor names, quantize input if needed
        (
            quantized_input_names,
            input_zero_point_names,
            input_scale_names,
            nodes,
        ) = self.quantizer.quantize_activation(node, [0])

        if not data_found or quantized_input_names is None:
            return super().quantize()

        # Create an entry for output quantized value.
        qlinear_output_name = node.output[0] + TENSOR_NAME_QUANT_SUFFIX
        quantized_output_value = QuantizedValue(
            node.output[0],
            qlinear_output_name,
            output_scale_name,
            output_zp_name,
            QuantizedValueType.Input,
        )
        self.quantizer.quantized_value_map[node.output[0]] = quantized_output_value

        # Create qlinear softmax node for given type
        kwargs = {}
        for attribute in node.attribute:
            kwargs.update(attribute_to_kwarg(attribute))
        kwargs["domain"] = ms_domain
        # make qlinearsoft has the real opset_version, its default SinceVersion would be 1
        kwargs["opset"] = self.quantizer.opset_version
        qlinear_node_name = node.name + "_quant" if node.name else ""
        qnode = onnx.helper.make_node(
            "QLinear" + node.op_type,
            [
                quantized_input_names[0],
                input_scale_names[0],
                input_zero_point_names[0],
                output_scale_name,
                output_zp_name,
            ],
            [qlinear_output_name],
            qlinear_node_name,
            **kwargs,
        )

        # add all newly created nodes
        nodes.append(qnode)
        self.quantizer.new_nodes += nodes
        return None

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\transformers\fusion_quickgelu.py ===
# -------------------------------------------------------------------------
# Copyright (c) Microsoft Corporation.  All rights reserved.
# Licensed under the MIT License.
# --------------------------------------------------------------------------

import logging

from fusion_base import Fusion
from onnx import helper
from onnx_model import OnnxModel

logger = logging.getLogger(__name__)


class FusionQuickGelu(Fusion):
    def __init__(self, model: OnnxModel):
        super().__init__(model, "QuickGelu", ["Mul"])

    def fuse(self, node, input_name_to_nodes, output_name_to_node):
        # Fuse the following subgraph to `QuickGelu`
        #
        #       root_input
        #      /          \
        #     |           Mul       ----+
        #     |       (B = ~1.702)      |
        #      \           |            |
        #       \       Sigmoid         |---- `QuickGelu`
        #        \     /                |
        #         \   /                 |
        #          Mul              ----+
        #           |
        #       root_output

        if node.op_type != "Mul":
            logger.debug("fuse_quickgelu: failed to match second Mul node")
            return

        second_mul_node = node
        root_input = second_mul_node.input[0]

        sigmoid_node = self.model.match_parent_path(second_mul_node, ["Sigmoid"], [1])
        if sigmoid_node is None:
            logger.debug("fuse_quickgelu: failed to match Sigmoid node")
            return
        sigmoid_node = sigmoid_node[0]

        first_mul_node = self.model.match_parent_path(sigmoid_node, ["Mul"], [0])
        if first_mul_node is None:
            logger.debug("fuse_quickgelu: failed to match first Mul node")
            return
        first_mul_node = first_mul_node[0]

        approximation_value = self.model.get_constant_value(first_mul_node.input[1]).item()
        if abs(approximation_value - 1.7021484375) >= 1e-3:
            logger.debug("fuse_quickgelu: failed to match approximation value")
            return

        if first_mul_node.input[0] != root_input:
            logger.debug("fuse_quickgelu: failed to match root input with first Mul node's input")
            return

        new_node = helper.make_node(
            "QuickGelu",
            inputs=[root_input],
            outputs=[second_mul_node.output[0]],
            name=self.model.create_node_name("QuickGelu"),
        )
        new_node.domain = "com.microsoft"
        new_node.attribute.extend([helper.make_attribute("alpha", approximation_value)])

        self.nodes_to_remove.extend([first_mul_node, sigmoid_node, second_mul_node])
        self.nodes_to_add.append(new_node)
        self.node_name_to_graph_name[new_node.name] = self.this_graph_name
        self.increase_counter("QuickGelu")

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\transformers\huggingface_models.py ===
# -------------------------------------------------------------------------
# Copyright (c) Microsoft Corporation.  All rights reserved.
# Licensed under the MIT License.  See License.txt in the project root for
# license information.
# --------------------------------------------------------------------------

# Maps model class name to a tuple of model class
MODEL_CLASSES = [
    "AutoModel",
    "AutoModelWithLMHead",
    "AutoModelForSequenceClassification",
    "AutoModelForQuestionAnswering",
    "AutoModelForCausalLM",
]

# Pretrained model name to a tuple of input names, opset_version, use_external_data_format, optimization model type
# Some models like GPT, T5, Bart etc has its own convert_to_onnx.py in models sub-directory, and they are excluded here.
MODELS = {
    # BERT
    "bert-base-cased": (["input_ids", "attention_mask", "token_type_ids"], 16, False, "bert"),
    "bert-large-cased": (["input_ids", "attention_mask", "token_type_ids"], 16, False, "bert"),
    # Transformer-XL (Models uses Einsum, which need opset version 16 or later.)
    "transfo-xl-wt103": (["input_ids", "mems"], 16, False, "bert"),
    # XLNet
    "xlnet-base-cased": (["input_ids"], 16, False, "bert"),
    "xlnet-large-cased": (["input_ids"], 16, False, "bert"),
    # XLM
    "xlm-mlm-en-2048": (["input_ids"], 16, True, "bert"),
    "xlm-mlm-ende-1024": (["input_ids"], 16, False, "bert"),
    "xlm-mlm-enfr-1024": (["input_ids"], 16, False, "bert"),
    # RoBERTa
    "roberta-base": (["input_ids", "attention_mask"], 16, False, "bert"),
    "roberta-large": (["input_ids", "attention_mask"], 16, False, "bert"),
    "roberta-large-mnli": (["input_ids", "attention_mask"], 16, False, "bert"),
    "deepset/roberta-base-squad2": (["input_ids", "attention_mask"], 16, False, "bert"),
    "distilroberta-base": (["input_ids", "attention_mask"], 16, False, "bert"),
    # DistilBERT
    "distilbert-base-uncased": (["input_ids", "attention_mask"], 16, False, "bert"),
    "distilbert-base-uncased-distilled-squad": (["input_ids", "attention_mask"], 16, False, "bert"),
    # CTRL
    "ctrl": (["input_ids"], 16, True, "bert"),
    # CamemBERT
    "camembert-base": (["input_ids"], 16, False, "bert"),
    # ALBERT
    "albert-base-v1": (["input_ids"], 16, False, "bert"),
    "albert-large-v1": (["input_ids"], 16, False, "bert"),
    "albert-xlarge-v1": (["input_ids"], 16, True, "bert"),
    # "albert-xxlarge-v1": (["input_ids"], 16, True, "bert"),
    "albert-base-v2": (["input_ids"], 16, False, "bert"),
    "albert-large-v2": (["input_ids"], 16, False, "bert"),
    "albert-xlarge-v2": (["input_ids"], 16, True, "bert"),
    # "albert-xxlarge-v2": (["input_ids"], 16, True, "bert"),
    # XLM-RoBERTa
    "xlm-roberta-base": (["input_ids"], 16, False, "bert"),
    "xlm-roberta-large": (["input_ids"], 16, True, "bert"),
    # FlauBERT
    "flaubert/flaubert_small_cased": (["input_ids"], 16, False, "bert"),
    "flaubert/flaubert_base_cased": (["input_ids"], 16, False, "bert"),
    # "flaubert/flaubert_large_cased": (["input_ids"], 16, False, "bert"),
    # Layoutlm
    "microsoft/layoutlm-base-uncased": (["input_ids"], 16, False, "bert"),
    "microsoft/layoutlm-large-uncased": (["input_ids"], 16, False, "bert"),
    # Squeezebert
    "squeezebert/squeezebert-uncased": (["input_ids"], 16, False, "bert"),
    "squeezebert/squeezebert-mnli": (["input_ids"], 16, False, "bert"),
    "squeezebert/squeezebert-mnli-headless": (["input_ids"], 16, False, "bert"),
    "unc-nlp/lxmert-base-uncased": (["input_ids", "visual_feats", "visual_pos"], 16, False, "bert"),
    # ViT
    "google/vit-base-patch16-224": (["pixel_values"], 16, False, "vit"),
    # Swin
    "microsoft/swin-base-patch4-window7-224": (["pixel_values"], 16, False, "swin"),
    "microsoft/swin-small-patch4-window7-224": (["pixel_values"], 16, False, "swin"),
    "microsoft/swin-tiny-patch4-window7-224": (["pixel_values"], 16, False, "swin"),
}

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\onnxruntime\transformers\torch_onnx_export_helper.py ===
# -------------------------------------------------------------------------
# Copyright (c) Microsoft Corporation.  All rights reserved.
# Licensed under the MIT License.
# --------------------------------------------------------------------------

import torch
from torch._C._onnx import OperatorExportTypes

TrainingMode = torch.onnx.TrainingMode
from packaging.version import Version  # noqa: E402


def torch_onnx_export(
    model,
    args,
    f,
    export_params=True,
    verbose=False,
    training=TrainingMode.EVAL,
    input_names=None,
    output_names=None,
    operator_export_type=OperatorExportTypes.ONNX,
    opset_version=None,
    _retain_param_name=None,
    do_constant_folding=True,
    example_outputs=None,
    strip_doc_string=None,
    dynamic_axes=None,
    keep_initializers_as_inputs=None,
    custom_opsets=None,
    enable_onnx_checker=None,
    use_external_data_format=None,
    export_modules_as_functions=False,
):
    if Version(torch.__version__) >= Version("1.11.0"):
        torch.onnx.export(
            model=model,
            args=args,
            f=f,
            export_params=export_params,
            verbose=verbose,
            training=training,
            input_names=input_names,
            output_names=output_names,
            operator_export_type=operator_export_type,
            opset_version=opset_version,
            do_constant_folding=do_constant_folding,
            dynamic_axes=dynamic_axes,
            keep_initializers_as_inputs=keep_initializers_as_inputs,
            custom_opsets=custom_opsets,
            export_modules_as_functions=export_modules_as_functions,
        )
    else:
        torch.onnx.export(
            model=model,
            args=args,
            f=f,
            export_params=export_params,
            verbose=verbose,
            training=training,
            input_names=input_names,
            output_names=output_names,
            operator_export_type=operator_export_type,
            opset_version=opset_version,
            _retain_param_name=_retain_param_name,
            do_constant_folding=do_constant_folding,
            example_outputs=example_outputs,
            strip_doc_string=strip_doc_string,
            dynamic_axes=dynamic_axes,
            keep_initializers_as_inputs=keep_initializers_as_inputs,
            custom_opsets=custom_opsets,
            enable_onnx_checker=enable_onnx_checker,
            use_external_data_format=use_external_data_format,
        )

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\openpyxl\chart\layout.py ===
# Copyright (c) 2010-2024 openpyxl

from openpyxl.descriptors.serialisable import Serialisable
from openpyxl.descriptors import (
    NoneSet,
    Float,
    Typed,
    Alias,
)

from openpyxl.descriptors.excel import ExtensionList
from openpyxl.descriptors.nested import (
    NestedNoneSet,
    NestedSet,
    NestedMinMax,
)

class ManualLayout(Serialisable):

    tagname = "manualLayout"

    layoutTarget = NestedNoneSet(values=(['inner', 'outer']))
    xMode = NestedNoneSet(values=(['edge', 'factor']))
    yMode = NestedNoneSet(values=(['edge', 'factor']))
    wMode = NestedSet(values=(['edge', 'factor']))
    hMode = NestedSet(values=(['edge', 'factor']))
    x = NestedMinMax(min=-1, max=1, allow_none=True)
    y = NestedMinMax(min=-1, max=1, allow_none=True)
    w = NestedMinMax(min=0, max=1, allow_none=True)
    width = Alias('w')
    h = NestedMinMax(min=0, max=1,  allow_none=True)
    height = Alias('h')
    extLst = Typed(expected_type=ExtensionList, allow_none=True)

    __elements__ = ('layoutTarget', 'xMode', 'yMode', 'wMode', 'hMode', 'x',
                    'y', 'w', 'h')

    def __init__(self,
                 layoutTarget=None,
                 xMode=None,
                 yMode=None,
                 wMode="factor",
                 hMode="factor",
                 x=None,
                 y=None,
                 w=None,
                 h=None,
                 extLst=None,
                ):
        self.layoutTarget = layoutTarget
        self.xMode = xMode
        self.yMode = yMode
        self.wMode = wMode
        self.hMode = hMode
        self.x = x
        self.y = y
        self.w = w
        self.h = h


class Layout(Serialisable):

    tagname = "layout"

    manualLayout = Typed(expected_type=ManualLayout, allow_none=True)
    extLst = Typed(expected_type=ExtensionList, allow_none=True)

    __elements__ = ('manualLayout',)

    def __init__(self,
                 manualLayout=None,
                 extLst=None,
                ):
        self.manualLayout = manualLayout

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pycparser\ply\ygen.py ===
# ply: ygen.py
#
# This is a support program that auto-generates different versions of the YACC parsing
# function with different features removed for the purposes of performance.
#
# Users should edit the method LParser.parsedebug() in yacc.py.   The source code 
# for that method is then used to create the other methods.   See the comments in
# yacc.py for further details.

import os.path
import shutil

def get_source_range(lines, tag):
    srclines = enumerate(lines)
    start_tag = '#--! %s-start' % tag
    end_tag = '#--! %s-end' % tag

    for start_index, line in srclines:
        if line.strip().startswith(start_tag):
            break

    for end_index, line in srclines:
        if line.strip().endswith(end_tag):
            break

    return (start_index + 1, end_index)

def filter_section(lines, tag):
    filtered_lines = []
    include = True
    tag_text = '#--! %s' % tag
    for line in lines:
        if line.strip().startswith(tag_text):
            include = not include
        elif include:
            filtered_lines.append(line)
    return filtered_lines

def main():
    dirname = os.path.dirname(__file__)
    shutil.copy2(os.path.join(dirname, 'yacc.py'), os.path.join(dirname, 'yacc.py.bak'))
    with open(os.path.join(dirname, 'yacc.py'), 'r') as f:
        lines = f.readlines()

    parse_start, parse_end = get_source_range(lines, 'parsedebug')
    parseopt_start, parseopt_end = get_source_range(lines, 'parseopt')
    parseopt_notrack_start, parseopt_notrack_end = get_source_range(lines, 'parseopt-notrack')

    # Get the original source
    orig_lines = lines[parse_start:parse_end]

    # Filter the DEBUG sections out
    parseopt_lines = filter_section(orig_lines, 'DEBUG')

    # Filter the TRACKING sections out
    parseopt_notrack_lines = filter_section(parseopt_lines, 'TRACKING')

    # Replace the parser source sections with updated versions
    lines[parseopt_notrack_start:parseopt_notrack_end] = parseopt_notrack_lines
    lines[parseopt_start:parseopt_end] = parseopt_lines

    lines = [line.rstrip()+'\n' for line in lines]
    with open(os.path.join(dirname, 'yacc.py'), 'w') as f:
        f.writelines(lines)

    print('Updated yacc.py')

if __name__ == '__main__':
    main()






# === atelier-kyo-manager/.venv_backup\Lib\site-packages\selenium\webdriver\common\bidi\webextension.py ===
# Licensed to the Software Freedom Conservancy (SFC) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The SFC licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

from typing import Dict, Union

from selenium.webdriver.common.bidi.common import command_builder


class WebExtension:
    """
    BiDi implementation of the webExtension module.
    """

    def __init__(self, conn):
        self.conn = conn

    def install(self, path=None, archive_path=None, base64_value=None) -> Dict:
        """Installs a web extension in the remote end.

        You must provide exactly one of the parameters.

        Parameters:
        -----------
            path: Path to an extension directory
            archive_path: Path to an extension archive file
            base64_value: Base64 encoded string of the extension archive

        Returns:
        -------
            Dict: A dictionary containing the extension ID.
        """
        if sum(x is not None for x in (path, archive_path, base64_value)) != 1:
            raise ValueError("Exactly one of path, archive_path, or base64_value must be provided")

        if path is not None:
            extension_data = {"type": "path", "path": path}
        elif archive_path is not None:
            extension_data = {"type": "archivePath", "path": archive_path}
        elif base64_value is not None:
            extension_data = {"type": "base64", "value": base64_value}

        params = {"extensionData": extension_data}
        result = self.conn.execute(command_builder("webExtension.install", params))
        return result

    def uninstall(self, extension_id_or_result: Union[str, Dict]) -> None:
        """Uninstalls a web extension from the remote end.

        Parameters:
        -----------
            extension_id_or_result: Either the extension ID as a string or the result dictionary
                                   from a previous install() call containing the extension ID.
        """
        if isinstance(extension_id_or_result, dict):
            extension_id = extension_id_or_result.get("extension")
        else:
            extension_id = extension_id_or_result

        params = {"extension": extension_id}
        self.conn.execute(command_builder("webExtension.uninstall", params))

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sortedcontainers\__init__.py ===
"""Sorted Containers -- Sorted List, Sorted Dict, Sorted Set

Sorted Containers is an Apache2 licensed containers library, written in
pure-Python, and fast as C-extensions.

Python's standard library is great until you need a sorted collections
type. Many will attest that you can get really far without one, but the moment
you **really need** a sorted list, dict, or set, you're faced with a dozen
different implementations, most using C-extensions without great documentation
and benchmarking.

In Python, we can do better. And we can do it in pure-Python!

::

    >>> from sortedcontainers import SortedList
    >>> sl = SortedList(['e', 'a', 'c', 'd', 'b'])
    >>> sl
    SortedList(['a', 'b', 'c', 'd', 'e'])
    >>> sl *= 1000000
    >>> sl.count('c')
    1000000
    >>> sl[-3:]
    ['e', 'e', 'e']
    >>> from sortedcontainers import SortedDict
    >>> sd = SortedDict({'c': 3, 'a': 1, 'b': 2})
    >>> sd
    SortedDict({'a': 1, 'b': 2, 'c': 3})
    >>> sd.popitem(index=-1)
    ('c', 3)
    >>> from sortedcontainers import SortedSet
    >>> ss = SortedSet('abracadabra')
    >>> ss
    SortedSet(['a', 'b', 'c', 'd', 'r'])
    >>> ss.bisect_left('c')
    2

Sorted Containers takes all of the work out of Python sorted types - making
your deployment and use of Python easy. There's no need to install a C compiler
or pre-build and distribute custom extensions. Performance is a feature and
testing has 100% coverage with unit tests and hours of stress.

:copyright: (c) 2014-2019 by Grant Jenks.
:license: Apache 2.0, see LICENSE for more details.

"""


from .sortedlist import SortedList, SortedKeyList, SortedListWithKey
from .sortedset import SortedSet
from .sorteddict import (
    SortedDict,
    SortedKeysView,
    SortedItemsView,
    SortedValuesView,
)

__all__ = [
    'SortedList',
    'SortedKeyList',
    'SortedListWithKey',
    'SortedDict',
    'SortedKeysView',
    'SortedItemsView',
    'SortedValuesView',
    'SortedSet',
]

__title__ = 'sortedcontainers'
__version__ = '2.4.0'
__build__ = 0x020400
__author__ = 'Grant Jenks'
__license__ = 'Apache 2.0'
__copyright__ = '2014-2019, Grant Jenks'

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\sqlalchemy\engine\_py_util.py ===
# engine/_py_util.py
# Copyright (C) 2005-2025 the SQLAlchemy authors and contributors
# <see AUTHORS file>
#
# This module is part of SQLAlchemy and is released under
# the MIT License: https://www.opensource.org/licenses/mit-license.php
from __future__ import annotations

import typing
from typing import Any
from typing import Mapping
from typing import Optional
from typing import Tuple

from .. import exc

if typing.TYPE_CHECKING:
    from .interfaces import _CoreAnyExecuteParams
    from .interfaces import _CoreMultiExecuteParams
    from .interfaces import _DBAPIAnyExecuteParams
    from .interfaces import _DBAPIMultiExecuteParams


_no_tuple: Tuple[Any, ...] = ()


def _distill_params_20(
    params: Optional[_CoreAnyExecuteParams],
) -> _CoreMultiExecuteParams:
    if params is None:
        return _no_tuple
    # Assume list is more likely than tuple
    elif isinstance(params, list) or isinstance(params, tuple):
        # collections_abc.MutableSequence): # avoid abc.__instancecheck__
        if params and not isinstance(params[0], (tuple, Mapping)):
            raise exc.ArgumentError(
                "List argument must consist only of tuples or dictionaries"
            )

        return params
    elif isinstance(params, dict) or isinstance(
        # only do immutabledict or abc.__instancecheck__ for Mapping after
        # we've checked for plain dictionaries and would otherwise raise
        params,
        Mapping,
    ):
        return [params]
    else:
        raise exc.ArgumentError("mapping or list expected for parameters")


def _distill_raw_params(
    params: Optional[_DBAPIAnyExecuteParams],
) -> _DBAPIMultiExecuteParams:
    if params is None:
        return _no_tuple
    elif isinstance(params, list):
        # collections_abc.MutableSequence): # avoid abc.__instancecheck__
        if params and not isinstance(params[0], (tuple, Mapping)):
            raise exc.ArgumentError(
                "List argument must consist only of tuples or dictionaries"
            )

        return params
    elif isinstance(params, (tuple, dict)) or isinstance(
        # only do abc.__instancecheck__ for Mapping after we've checked
        # for plain dictionaries and would otherwise raise
        params,
        Mapping,
    ):
        # cast("Union[List[Mapping[str, Any]], Tuple[Any, ...]]", [params])
        return [params]  # type: ignore
    else:
        raise exc.ArgumentError("mapping or sequence expected for parameters")

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\pandas\tests\io\test_stata.py ===
import bz2
import datetime as dt
from datetime import datetime
import gzip
import io
import os
import struct
import tarfile
import zipfile

import numpy as np
import pytest

import pandas.util._test_decorators as td

import pandas as pd
from pandas import CategoricalDtype
import pandas._testing as tm
from pandas.core.frame import (
    DataFrame,
    Series,
)

from pandas.io.parsers import read_csv
from pandas.io.stata import (
    CategoricalConversionWarning,
    InvalidColumnName,
    PossiblePrecisionLoss,
    StataMissingValue,
    StataReader,
    StataWriter,
    StataWriterUTF8,
    ValueLabelTypeMismatch,
    read_stata,
)


@pytest.fixture
def mixed_frame():
    return DataFrame(
        {
            "a": [1, 2, 3, 4],
            "b": [1.0, 3.0, 27.0, 81.0],
            "c": ["Atlanta", "Birmingham", "Cincinnati", "Detroit"],
        }
    )


@pytest.fixture
def parsed_114(datapath):
    dta14_114 = datapath("io", "data", "stata", "stata5_114.dta")
    parsed_114 = read_stata(dta14_114, convert_dates=True)
    parsed_114.index.name = "index"
    return parsed_114


class TestStata:
    def read_dta(self, file):
        # Legacy default reader configuration
        return read_stata(file, convert_dates=True)

    def read_csv(self, file):
        return read_csv(file, parse_dates=True)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_read_empty_dta(self, version):
        empty_ds = DataFrame(columns=["unit"])
        # GH 7369, make sure can read a 0-obs dta file
        with tm.ensure_clean() as path:
            empty_ds.to_stata(path, write_index=False, version=version)
            empty_ds2 = read_stata(path)
            tm.assert_frame_equal(empty_ds, empty_ds2)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_read_empty_dta_with_dtypes(self, version):
        # GH 46240
        # Fixing above bug revealed that types are not correctly preserved when
        # writing empty DataFrames
        empty_df_typed = DataFrame(
            {
                "i8": np.array([0], dtype=np.int8),
                "i16": np.array([0], dtype=np.int16),
                "i32": np.array([0], dtype=np.int32),
                "i64": np.array([0], dtype=np.int64),
                "u8": np.array([0], dtype=np.uint8),
                "u16": np.array([0], dtype=np.uint16),
                "u32": np.array([0], dtype=np.uint32),
                "u64": np.array([0], dtype=np.uint64),
                "f32": np.array([0], dtype=np.float32),
                "f64": np.array([0], dtype=np.float64),
            }
        )
        expected = empty_df_typed.copy()
        # No uint# support. Downcast since values in range for int#
        expected["u8"] = expected["u8"].astype(np.int8)
        expected["u16"] = expected["u16"].astype(np.int16)
        expected["u32"] = expected["u32"].astype(np.int32)
        # No int64 supported at all. Downcast since values in range for int32
        expected["u64"] = expected["u64"].astype(np.int32)
        expected["i64"] = expected["i64"].astype(np.int32)

        # GH 7369, make sure can read a 0-obs dta file
        with tm.ensure_clean() as path:
            empty_df_typed.to_stata(path, write_index=False, version=version)
            empty_reread = read_stata(path)
            tm.assert_frame_equal(expected, empty_reread)
            tm.assert_series_equal(expected.dtypes, empty_reread.dtypes)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_read_index_col_none(self, version):
        df = DataFrame({"a": range(5), "b": ["b1", "b2", "b3", "b4", "b5"]})
        # GH 7369, make sure can read a 0-obs dta file
        with tm.ensure_clean() as path:
            df.to_stata(path, write_index=False, version=version)
            read_df = read_stata(path)

        assert isinstance(read_df.index, pd.RangeIndex)
        expected = df.copy()
        expected["a"] = expected["a"].astype(np.int32)
        tm.assert_frame_equal(read_df, expected, check_index_type=True)

    @pytest.mark.parametrize("file", ["stata1_114", "stata1_117"])
    def test_read_dta1(self, file, datapath):
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = self.read_dta(file)

        # Pandas uses np.nan as missing value.
        # Thus, all columns will be of type float, regardless of their name.
        expected = DataFrame(
            [(np.nan, np.nan, np.nan, np.nan, np.nan)],
            columns=["float_miss", "double_miss", "byte_miss", "int_miss", "long_miss"],
        )

        # this is an oddity as really the nan should be float64, but
        # the casting doesn't fail so need to match stata here
        expected["float_miss"] = expected["float_miss"].astype(np.float32)

        tm.assert_frame_equal(parsed, expected)

    def test_read_dta2(self, datapath):
        expected = DataFrame.from_records(
            [
                (
                    datetime(2006, 11, 19, 23, 13, 20),
                    1479596223000,
                    datetime(2010, 1, 20),
                    datetime(2010, 1, 8),
                    datetime(2010, 1, 1),
                    datetime(1974, 7, 1),
                    datetime(2010, 1, 1),
                    datetime(2010, 1, 1),
                ),
                (
                    datetime(1959, 12, 31, 20, 3, 20),
                    -1479590,
                    datetime(1953, 10, 2),
                    datetime(1948, 6, 10),
                    datetime(1955, 1, 1),
                    datetime(1955, 7, 1),
                    datetime(1955, 1, 1),
                    datetime(2, 1, 1),
                ),
                (pd.NaT, pd.NaT, pd.NaT, pd.NaT, pd.NaT, pd.NaT, pd.NaT, pd.NaT),
            ],
            columns=[
                "datetime_c",
                "datetime_big_c",
                "date",
                "weekly_date",
                "monthly_date",
                "quarterly_date",
                "half_yearly_date",
                "yearly_date",
            ],
        )
        expected["yearly_date"] = expected["yearly_date"].astype("O")

        path1 = datapath("io", "data", "stata", "stata2_114.dta")
        path2 = datapath("io", "data", "stata", "stata2_115.dta")
        path3 = datapath("io", "data", "stata", "stata2_117.dta")

        with tm.assert_produces_warning(UserWarning):
            parsed_114 = self.read_dta(path1)
        with tm.assert_produces_warning(UserWarning):
            parsed_115 = self.read_dta(path2)
        with tm.assert_produces_warning(UserWarning):
            parsed_117 = self.read_dta(path3)
            # FIXME: don't leave commented-out
            # 113 is buggy due to limits of date format support in Stata
            # parsed_113 = self.read_dta(
            # datapath("io", "data", "stata", "stata2_113.dta")
            # )

        # FIXME: don't leave commented-out
        # buggy test because of the NaT comparison on certain platforms
        # Format 113 test fails since it does not support tc and tC formats
        # tm.assert_frame_equal(parsed_113, expected)
        tm.assert_frame_equal(parsed_114, expected, check_datetimelike_compat=True)
        tm.assert_frame_equal(parsed_115, expected, check_datetimelike_compat=True)
        tm.assert_frame_equal(parsed_117, expected, check_datetimelike_compat=True)

    @pytest.mark.parametrize(
        "file", ["stata3_113", "stata3_114", "stata3_115", "stata3_117"]
    )
    def test_read_dta3(self, file, datapath):
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = self.read_dta(file)

        # match stata here
        expected = self.read_csv(datapath("io", "data", "stata", "stata3.csv"))
        expected = expected.astype(np.float32)
        expected["year"] = expected["year"].astype(np.int16)
        expected["quarter"] = expected["quarter"].astype(np.int8)

        tm.assert_frame_equal(parsed, expected)

    @pytest.mark.parametrize(
        "file", ["stata4_113", "stata4_114", "stata4_115", "stata4_117"]
    )
    def test_read_dta4(self, file, datapath):
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = self.read_dta(file)

        expected = DataFrame.from_records(
            [
                ["one", "ten", "one", "one", "one"],
                ["two", "nine", "two", "two", "two"],
                ["three", "eight", "three", "three", "three"],
                ["four", "seven", 4, "four", "four"],
                ["five", "six", 5, np.nan, "five"],
                ["six", "five", 6, np.nan, "six"],
                ["seven", "four", 7, np.nan, "seven"],
                ["eight", "three", 8, np.nan, "eight"],
                ["nine", "two", 9, np.nan, "nine"],
                ["ten", "one", "ten", np.nan, "ten"],
            ],
            columns=[
                "fully_labeled",
                "fully_labeled2",
                "incompletely_labeled",
                "labeled_with_missings",
                "float_labelled",
            ],
        )

        # these are all categoricals
        for col in expected:
            orig = expected[col].copy()

            categories = np.asarray(expected["fully_labeled"][orig.notna()])
            if col == "incompletely_labeled":
                categories = orig

            cat = orig.astype("category")._values
            cat = cat.set_categories(categories, ordered=True)
            cat.categories.rename(None, inplace=True)

            expected[col] = cat

        # stata doesn't save .category metadata
        tm.assert_frame_equal(parsed, expected)

    # File containing strls
    def test_read_dta12(self, datapath):
        parsed_117 = self.read_dta(datapath("io", "data", "stata", "stata12_117.dta"))
        expected = DataFrame.from_records(
            [
                [1, "abc", "abcdefghi"],
                [3, "cba", "qwertywertyqwerty"],
                [93, "", "strl"],
            ],
            columns=["x", "y", "z"],
        )

        tm.assert_frame_equal(parsed_117, expected, check_dtype=False)

    def test_read_dta18(self, datapath):
        parsed_118 = self.read_dta(datapath("io", "data", "stata", "stata14_118.dta"))
        parsed_118["Bytes"] = parsed_118["Bytes"].astype("O")
        expected = DataFrame.from_records(
            [
                ["Cat", "Bogota", "Bogotá", 1, 1.0, "option b Ünicode", 1.0],
                ["Dog", "Boston", "Uzunköprü", np.nan, np.nan, np.nan, np.nan],
                ["Plane", "Rome", "Tromsø", 0, 0.0, "option a", 0.0],
                ["Potato", "Tokyo", "Elâzığ", -4, 4.0, 4, 4],  # noqa: RUF001
                ["", "", "", 0, 0.3332999, "option a", 1 / 3.0],
            ],
            columns=[
                "Things",
                "Cities",
                "Unicode_Cities_Strl",
                "Ints",
                "Floats",
                "Bytes",
                "Longs",
            ],
        )
        expected["Floats"] = expected["Floats"].astype(np.float32)
        for col in parsed_118.columns:
            tm.assert_almost_equal(parsed_118[col], expected[col])

        with StataReader(datapath("io", "data", "stata", "stata14_118.dta")) as rdr:
            vl = rdr.variable_labels()
            vl_expected = {
                "Unicode_Cities_Strl": "Here are some strls with Ünicode chars",
                "Longs": "long data",
                "Things": "Here are some things",
                "Bytes": "byte data",
                "Ints": "int data",
                "Cities": "Here are some cities",
                "Floats": "float data",
            }
            tm.assert_dict_equal(vl, vl_expected)

            assert rdr.data_label == "This is a  Ünicode data label"

    def test_read_write_dta5(self):
        original = DataFrame(
            [(np.nan, np.nan, np.nan, np.nan, np.nan)],
            columns=["float_miss", "double_miss", "byte_miss", "int_miss", "long_miss"],
        )
        original.index.name = "index"

        with tm.ensure_clean() as path:
            original.to_stata(path, convert_dates=None)
            written_and_read_again = self.read_dta(path)

        expected = original.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    def test_write_dta6(self, datapath):
        original = self.read_csv(datapath("io", "data", "stata", "stata3.csv"))
        original.index.name = "index"
        original.index = original.index.astype(np.int32)
        original["year"] = original["year"].astype(np.int32)
        original["quarter"] = original["quarter"].astype(np.int32)

        with tm.ensure_clean() as path:
            original.to_stata(path, convert_dates=None)
            written_and_read_again = self.read_dta(path)
            tm.assert_frame_equal(
                written_and_read_again.set_index("index"),
                original,
                check_index_type=False,
            )

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_read_write_dta10(self, version, using_infer_string):
        original = DataFrame(
            data=[["string", "object", 1, 1.1, np.datetime64("2003-12-25")]],
            columns=["string", "object", "integer", "floating", "datetime"],
        )
        original["object"] = Series(original["object"], dtype=object)
        original.index.name = "index"
        original.index = original.index.astype(np.int32)
        original["integer"] = original["integer"].astype(np.int32)

        with tm.ensure_clean() as path:
            original.to_stata(path, convert_dates={"datetime": "tc"}, version=version)
            written_and_read_again = self.read_dta(path)

        expected = original.copy()
        if using_infer_string:
            expected["object"] = expected["object"].astype("str")

        # original.index is np.int32, read index is np.int64
        tm.assert_frame_equal(
            written_and_read_again.set_index("index"),
            expected,
            check_index_type=False,
        )

    def test_stata_doc_examples(self):
        with tm.ensure_clean() as path:
            df = DataFrame(
                np.random.default_rng(2).standard_normal((10, 2)), columns=list("AB")
            )
            df.to_stata(path)

    def test_write_preserves_original(self):
        # 9795

        df = DataFrame(
            np.random.default_rng(2).standard_normal((5, 4)), columns=list("abcd")
        )
        df.loc[2, "a":"c"] = np.nan
        df_copy = df.copy()
        with tm.ensure_clean() as path:
            df.to_stata(path, write_index=False)
        tm.assert_frame_equal(df, df_copy)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_encoding(self, version, datapath):
        # GH 4626, proper encoding handling
        raw = read_stata(datapath("io", "data", "stata", "stata1_encoding.dta"))
        encoded = read_stata(datapath("io", "data", "stata", "stata1_encoding.dta"))
        result = encoded.kreis1849[0]

        expected = raw.kreis1849[0]
        assert result == expected
        assert isinstance(result, str)

        with tm.ensure_clean() as path:
            encoded.to_stata(path, write_index=False, version=version)
            reread_encoded = read_stata(path)
            tm.assert_frame_equal(encoded, reread_encoded)

    def test_read_write_dta11(self):
        original = DataFrame(
            [(1, 2, 3, 4)],
            columns=[
                "good",
                "b\u00E4d",
                "8number",
                "astringwithmorethan32characters______",
            ],
        )
        formatted = DataFrame(
            [(1, 2, 3, 4)],
            columns=["good", "b_d", "_8number", "astringwithmorethan32characters_"],
        )
        formatted.index.name = "index"
        formatted = formatted.astype(np.int32)

        with tm.ensure_clean() as path:
            with tm.assert_produces_warning(InvalidColumnName):
                original.to_stata(path, convert_dates=None)

            written_and_read_again = self.read_dta(path)

        expected = formatted.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_read_write_dta12(self, version):
        original = DataFrame(
            [(1, 2, 3, 4, 5, 6)],
            columns=[
                "astringwithmorethan32characters_1",
                "astringwithmorethan32characters_2",
                "+",
                "-",
                "short",
                "delete",
            ],
        )
        formatted = DataFrame(
            [(1, 2, 3, 4, 5, 6)],
            columns=[
                "astringwithmorethan32characters_",
                "_0astringwithmorethan32character",
                "_",
                "_1_",
                "_short",
                "_delete",
            ],
        )
        formatted.index.name = "index"
        formatted = formatted.astype(np.int32)

        with tm.ensure_clean() as path:
            with tm.assert_produces_warning(InvalidColumnName):
                original.to_stata(path, convert_dates=None, version=version)
                # should get a warning for that format.

            written_and_read_again = self.read_dta(path)

        expected = formatted.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    def test_read_write_dta13(self):
        s1 = Series(2**9, dtype=np.int16)
        s2 = Series(2**17, dtype=np.int32)
        s3 = Series(2**33, dtype=np.int64)
        original = DataFrame({"int16": s1, "int32": s2, "int64": s3})
        original.index.name = "index"

        formatted = original
        formatted["int64"] = formatted["int64"].astype(np.float64)

        with tm.ensure_clean() as path:
            original.to_stata(path)
            written_and_read_again = self.read_dta(path)

        expected = formatted.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    @pytest.mark.parametrize(
        "file", ["stata5_113", "stata5_114", "stata5_115", "stata5_117"]
    )
    def test_read_write_reread_dta14(self, file, parsed_114, version, datapath):
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = self.read_dta(file)
        parsed.index.name = "index"

        tm.assert_frame_equal(parsed_114, parsed)

        with tm.ensure_clean() as path:
            parsed_114.to_stata(path, convert_dates={"date_td": "td"}, version=version)
            written_and_read_again = self.read_dta(path)

        expected = parsed_114.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    @pytest.mark.parametrize(
        "file", ["stata6_113", "stata6_114", "stata6_115", "stata6_117"]
    )
    def test_read_write_reread_dta15(self, file, datapath):
        expected = self.read_csv(datapath("io", "data", "stata", "stata6.csv"))
        expected["byte_"] = expected["byte_"].astype(np.int8)
        expected["int_"] = expected["int_"].astype(np.int16)
        expected["long_"] = expected["long_"].astype(np.int32)
        expected["float_"] = expected["float_"].astype(np.float32)
        expected["double_"] = expected["double_"].astype(np.float64)
        expected["date_td"] = expected["date_td"].apply(
            datetime.strptime, args=("%Y-%m-%d",)
        )

        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = self.read_dta(file)

        tm.assert_frame_equal(expected, parsed)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_timestamp_and_label(self, version):
        original = DataFrame([(1,)], columns=["variable"])
        time_stamp = datetime(2000, 2, 29, 14, 21)
        data_label = "This is a data file."
        with tm.ensure_clean() as path:
            original.to_stata(
                path, time_stamp=time_stamp, data_label=data_label, version=version
            )

            with StataReader(path) as reader:
                assert reader.time_stamp == "29 Feb 2000 14:21"
                assert reader.data_label == data_label

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_invalid_timestamp(self, version):
        original = DataFrame([(1,)], columns=["variable"])
        time_stamp = "01 Jan 2000, 00:00:00"
        with tm.ensure_clean() as path:
            msg = "time_stamp should be datetime type"
            with pytest.raises(ValueError, match=msg):
                original.to_stata(path, time_stamp=time_stamp, version=version)
            assert not os.path.isfile(path)

    def test_numeric_column_names(self):
        original = DataFrame(np.reshape(np.arange(25.0), (5, 5)))
        original.index.name = "index"
        with tm.ensure_clean() as path:
            # should get a warning for that format.
            with tm.assert_produces_warning(InvalidColumnName):
                original.to_stata(path)

            written_and_read_again = self.read_dta(path)

        written_and_read_again = written_and_read_again.set_index("index")
        columns = list(written_and_read_again.columns)
        convert_col_name = lambda x: int(x[1])
        written_and_read_again.columns = map(convert_col_name, columns)

        expected = original.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(expected, written_and_read_again)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_nan_to_missing_value(self, version):
        s1 = Series(np.arange(4.0), dtype=np.float32)
        s2 = Series(np.arange(4.0), dtype=np.float64)
        s1[::2] = np.nan
        s2[1::2] = np.nan
        original = DataFrame({"s1": s1, "s2": s2})
        original.index.name = "index"

        with tm.ensure_clean() as path:
            original.to_stata(path, version=version)
            written_and_read_again = self.read_dta(path)

        written_and_read_again = written_and_read_again.set_index("index")
        expected = original.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again, expected)

    def test_no_index(self):
        columns = ["x", "y"]
        original = DataFrame(np.reshape(np.arange(10.0), (5, 2)), columns=columns)
        original.index.name = "index_not_written"
        with tm.ensure_clean() as path:
            original.to_stata(path, write_index=False)
            written_and_read_again = self.read_dta(path)
            with pytest.raises(KeyError, match=original.index.name):
                written_and_read_again["index_not_written"]

    def test_string_no_dates(self):
        s1 = Series(["a", "A longer string"])
        s2 = Series([1.0, 2.0], dtype=np.float64)
        original = DataFrame({"s1": s1, "s2": s2})
        original.index.name = "index"
        with tm.ensure_clean() as path:
            original.to_stata(path)
            written_and_read_again = self.read_dta(path)

        expected = original.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    def test_large_value_conversion(self):
        s0 = Series([1, 99], dtype=np.int8)
        s1 = Series([1, 127], dtype=np.int8)
        s2 = Series([1, 2**15 - 1], dtype=np.int16)
        s3 = Series([1, 2**63 - 1], dtype=np.int64)
        original = DataFrame({"s0": s0, "s1": s1, "s2": s2, "s3": s3})
        original.index.name = "index"
        with tm.ensure_clean() as path:
            with tm.assert_produces_warning(PossiblePrecisionLoss):
                original.to_stata(path)

            written_and_read_again = self.read_dta(path)

        modified = original.copy()
        modified["s1"] = Series(modified["s1"], dtype=np.int16)
        modified["s2"] = Series(modified["s2"], dtype=np.int32)
        modified["s3"] = Series(modified["s3"], dtype=np.float64)
        modified.index = original.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), modified)

    def test_dates_invalid_column(self):
        original = DataFrame([datetime(2006, 11, 19, 23, 13, 20)])
        original.index.name = "index"
        with tm.ensure_clean() as path:
            with tm.assert_produces_warning(InvalidColumnName):
                original.to_stata(path, convert_dates={0: "tc"})

            written_and_read_again = self.read_dta(path)

        modified = original.copy()
        modified.columns = ["_0"]
        modified.index = original.index.astype(np.int32)
        tm.assert_frame_equal(written_and_read_again.set_index("index"), modified)

    def test_105(self, datapath):
        # Data obtained from:
        # http://go.worldbank.org/ZXY29PVJ21
        dpath = datapath("io", "data", "stata", "S4_EDUC1.dta")
        df = read_stata(dpath)
        df0 = [[1, 1, 3, -2], [2, 1, 2, -2], [4, 1, 1, -2]]
        df0 = DataFrame(df0)
        df0.columns = ["clustnum", "pri_schl", "psch_num", "psch_dis"]
        df0["clustnum"] = df0["clustnum"].astype(np.int16)
        df0["pri_schl"] = df0["pri_schl"].astype(np.int8)
        df0["psch_num"] = df0["psch_num"].astype(np.int8)
        df0["psch_dis"] = df0["psch_dis"].astype(np.float32)
        tm.assert_frame_equal(df.head(3), df0)

    def test_value_labels_old_format(self, datapath):
        # GH 19417
        #
        # Test that value_labels() returns an empty dict if the file format
        # predates supporting value labels.
        dpath = datapath("io", "data", "stata", "S4_EDUC1.dta")
        with StataReader(dpath) as reader:
            assert reader.value_labels() == {}

    def test_date_export_formats(self):
        columns = ["tc", "td", "tw", "tm", "tq", "th", "ty"]
        conversions = {c: c for c in columns}
        data = [datetime(2006, 11, 20, 23, 13, 20)] * len(columns)
        original = DataFrame([data], columns=columns)
        original.index.name = "index"
        expected_values = [
            datetime(2006, 11, 20, 23, 13, 20),  # Time
            datetime(2006, 11, 20),  # Day
            datetime(2006, 11, 19),  # Week
            datetime(2006, 11, 1),  # Month
            datetime(2006, 10, 1),  # Quarter year
            datetime(2006, 7, 1),  # Half year
            datetime(2006, 1, 1),
        ]  # Year

        expected = DataFrame(
            [expected_values],
            index=pd.Index([0], dtype=np.int32, name="index"),
            columns=columns,
        )

        with tm.ensure_clean() as path:
            original.to_stata(path, convert_dates=conversions)
            written_and_read_again = self.read_dta(path)

        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    def test_write_missing_strings(self):
        original = DataFrame([["1"], [None]], columns=["foo"])

        expected = DataFrame(
            [["1"], [""]],
            index=pd.Index([0, 1], dtype=np.int32, name="index"),
            columns=["foo"],
        )

        with tm.ensure_clean() as path:
            original.to_stata(path)
            written_and_read_again = self.read_dta(path)

        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    @pytest.mark.parametrize("byteorder", [">", "<"])
    def test_bool_uint(self, byteorder, version):
        s0 = Series([0, 1, True], dtype=np.bool_)
        s1 = Series([0, 1, 100], dtype=np.uint8)
        s2 = Series([0, 1, 255], dtype=np.uint8)
        s3 = Series([0, 1, 2**15 - 100], dtype=np.uint16)
        s4 = Series([0, 1, 2**16 - 1], dtype=np.uint16)
        s5 = Series([0, 1, 2**31 - 100], dtype=np.uint32)
        s6 = Series([0, 1, 2**32 - 1], dtype=np.uint32)

        original = DataFrame(
            {"s0": s0, "s1": s1, "s2": s2, "s3": s3, "s4": s4, "s5": s5, "s6": s6}
        )
        original.index.name = "index"
        expected = original.copy()
        expected.index = original.index.astype(np.int32)
        expected_types = (
            np.int8,
            np.int8,
            np.int16,
            np.int16,
            np.int32,
            np.int32,
            np.float64,
        )
        for c, t in zip(expected.columns, expected_types):
            expected[c] = expected[c].astype(t)

        with tm.ensure_clean() as path:
            original.to_stata(path, byteorder=byteorder, version=version)
            written_and_read_again = self.read_dta(path)

        written_and_read_again = written_and_read_again.set_index("index")
        tm.assert_frame_equal(written_and_read_again, expected)

    def test_variable_labels(self, datapath):
        with StataReader(datapath("io", "data", "stata", "stata7_115.dta")) as rdr:
            sr_115 = rdr.variable_labels()
        with StataReader(datapath("io", "data", "stata", "stata7_117.dta")) as rdr:
            sr_117 = rdr.variable_labels()
        keys = ("var1", "var2", "var3")
        labels = ("label1", "label2", "label3")
        for k, v in sr_115.items():
            assert k in sr_117
            assert v == sr_117[k]
            assert k in keys
            assert v in labels

    def test_minimal_size_col(self):
        str_lens = (1, 100, 244)
        s = {}
        for str_len in str_lens:
            s["s" + str(str_len)] = Series(
                ["a" * str_len, "b" * str_len, "c" * str_len]
            )
        original = DataFrame(s)
        with tm.ensure_clean() as path:
            original.to_stata(path, write_index=False)

            with StataReader(path) as sr:
                sr._ensure_open()  # The `_*list` variables are initialized here
                for variable, fmt, typ in zip(sr._varlist, sr._fmtlist, sr._typlist):
                    assert int(variable[1:]) == int(fmt[1:-1])
                    assert int(variable[1:]) == typ

    def test_excessively_long_string(self):
        str_lens = (1, 244, 500)
        s = {}
        for str_len in str_lens:
            s["s" + str(str_len)] = Series(
                ["a" * str_len, "b" * str_len, "c" * str_len]
            )
        original = DataFrame(s)
        msg = (
            r"Fixed width strings in Stata \.dta files are limited to 244 "
            r"\(or fewer\)\ncharacters\.  Column 's500' does not satisfy "
            r"this restriction\. Use the\n'version=117' parameter to write "
            r"the newer \(Stata 13 and later\) format\."
        )
        with pytest.raises(ValueError, match=msg):
            with tm.ensure_clean() as path:
                original.to_stata(path)

    def test_missing_value_generator(self):
        types = ("b", "h", "l")
        df = DataFrame([[0.0]], columns=["float_"])
        with tm.ensure_clean() as path:
            df.to_stata(path)
            with StataReader(path) as rdr:
                valid_range = rdr.VALID_RANGE
        expected_values = ["." + chr(97 + i) for i in range(26)]
        expected_values.insert(0, ".")
        for t in types:
            offset = valid_range[t][1]
            for i in range(27):
                val = StataMissingValue(offset + 1 + i)
                assert val.string == expected_values[i]

        # Test extremes for floats
        val = StataMissingValue(struct.unpack("<f", b"\x00\x00\x00\x7f")[0])
        assert val.string == "."
        val = StataMissingValue(struct.unpack("<f", b"\x00\xd0\x00\x7f")[0])
        assert val.string == ".z"

        # Test extremes for floats
        val = StataMissingValue(
            struct.unpack("<d", b"\x00\x00\x00\x00\x00\x00\xe0\x7f")[0]
        )
        assert val.string == "."
        val = StataMissingValue(
            struct.unpack("<d", b"\x00\x00\x00\x00\x00\x1a\xe0\x7f")[0]
        )
        assert val.string == ".z"

    @pytest.mark.parametrize("file", ["stata8_113", "stata8_115", "stata8_117"])
    def test_missing_value_conversion(self, file, datapath):
        columns = ["int8_", "int16_", "int32_", "float32_", "float64_"]
        smv = StataMissingValue(101)
        keys = sorted(smv.MISSING_VALUES.keys())
        data = []
        for i in range(27):
            row = [StataMissingValue(keys[i + (j * 27)]) for j in range(5)]
            data.append(row)
        expected = DataFrame(data, columns=columns)

        parsed = read_stata(
            datapath("io", "data", "stata", f"{file}.dta"), convert_missing=True
        )
        tm.assert_frame_equal(parsed, expected)

    def test_big_dates(self, datapath):
        yr = [1960, 2000, 9999, 100, 2262, 1677]
        mo = [1, 1, 12, 1, 4, 9]
        dd = [1, 1, 31, 1, 22, 23]
        hr = [0, 0, 23, 0, 0, 0]
        mm = [0, 0, 59, 0, 0, 0]
        ss = [0, 0, 59, 0, 0, 0]
        expected = []
        for year, month, day, hour, minute, second in zip(yr, mo, dd, hr, mm, ss):
            row = []
            for j in range(7):
                if j == 0:
                    row.append(datetime(year, month, day, hour, minute, second))
                elif j == 6:
                    row.append(datetime(year, 1, 1))
                else:
                    row.append(datetime(year, month, day))
            expected.append(row)
        expected.append([pd.NaT] * 7)
        columns = [
            "date_tc",
            "date_td",
            "date_tw",
            "date_tm",
            "date_tq",
            "date_th",
            "date_ty",
        ]

        # Fixes for weekly, quarterly,half,year
        expected[2][2] = datetime(9999, 12, 24)
        expected[2][3] = datetime(9999, 12, 1)
        expected[2][4] = datetime(9999, 10, 1)
        expected[2][5] = datetime(9999, 7, 1)
        expected[4][2] = datetime(2262, 4, 16)
        expected[4][3] = expected[4][4] = datetime(2262, 4, 1)
        expected[4][5] = expected[4][6] = datetime(2262, 1, 1)
        expected[5][2] = expected[5][3] = expected[5][4] = datetime(1677, 10, 1)
        expected[5][5] = expected[5][6] = datetime(1678, 1, 1)

        expected = DataFrame(expected, columns=columns, dtype=object)
        parsed_115 = read_stata(datapath("io", "data", "stata", "stata9_115.dta"))
        parsed_117 = read_stata(datapath("io", "data", "stata", "stata9_117.dta"))
        tm.assert_frame_equal(expected, parsed_115, check_datetimelike_compat=True)
        tm.assert_frame_equal(expected, parsed_117, check_datetimelike_compat=True)

        date_conversion = {c: c[-2:] for c in columns}
        # {c : c[-2:] for c in columns}
        with tm.ensure_clean() as path:
            expected.index.name = "index"
            expected.to_stata(path, convert_dates=date_conversion)
            written_and_read_again = self.read_dta(path)

        tm.assert_frame_equal(
            written_and_read_again.set_index("index"),
            expected.set_index(expected.index.astype(np.int32)),
            check_datetimelike_compat=True,
        )

    def test_dtype_conversion(self, datapath):
        expected = self.read_csv(datapath("io", "data", "stata", "stata6.csv"))
        expected["byte_"] = expected["byte_"].astype(np.int8)
        expected["int_"] = expected["int_"].astype(np.int16)
        expected["long_"] = expected["long_"].astype(np.int32)
        expected["float_"] = expected["float_"].astype(np.float32)
        expected["double_"] = expected["double_"].astype(np.float64)
        expected["date_td"] = expected["date_td"].apply(
            datetime.strptime, args=("%Y-%m-%d",)
        )

        no_conversion = read_stata(
            datapath("io", "data", "stata", "stata6_117.dta"), convert_dates=True
        )
        tm.assert_frame_equal(expected, no_conversion)

        conversion = read_stata(
            datapath("io", "data", "stata", "stata6_117.dta"),
            convert_dates=True,
            preserve_dtypes=False,
        )

        # read_csv types are the same
        expected = self.read_csv(datapath("io", "data", "stata", "stata6.csv"))
        expected["date_td"] = expected["date_td"].apply(
            datetime.strptime, args=("%Y-%m-%d",)
        )

        tm.assert_frame_equal(expected, conversion)

    def test_drop_column(self, datapath):
        expected = self.read_csv(datapath("io", "data", "stata", "stata6.csv"))
        expected["byte_"] = expected["byte_"].astype(np.int8)
        expected["int_"] = expected["int_"].astype(np.int16)
        expected["long_"] = expected["long_"].astype(np.int32)
        expected["float_"] = expected["float_"].astype(np.float32)
        expected["double_"] = expected["double_"].astype(np.float64)
        expected["date_td"] = expected["date_td"].apply(
            datetime.strptime, args=("%Y-%m-%d",)
        )

        columns = ["byte_", "int_", "long_"]
        expected = expected[columns]
        dropped = read_stata(
            datapath("io", "data", "stata", "stata6_117.dta"),
            convert_dates=True,
            columns=columns,
        )

        tm.assert_frame_equal(expected, dropped)

        # See PR 10757
        columns = ["int_", "long_", "byte_"]
        expected = expected[columns]
        reordered = read_stata(
            datapath("io", "data", "stata", "stata6_117.dta"),
            convert_dates=True,
            columns=columns,
        )
        tm.assert_frame_equal(expected, reordered)

        msg = "columns contains duplicate entries"
        with pytest.raises(ValueError, match=msg):
            columns = ["byte_", "byte_"]
            read_stata(
                datapath("io", "data", "stata", "stata6_117.dta"),
                convert_dates=True,
                columns=columns,
            )

        msg = "The following columns were not found in the Stata data set: not_found"
        with pytest.raises(ValueError, match=msg):
            columns = ["byte_", "int_", "long_", "not_found"]
            read_stata(
                datapath("io", "data", "stata", "stata6_117.dta"),
                convert_dates=True,
                columns=columns,
            )

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    @pytest.mark.filterwarnings(
        "ignore:\\nStata value:pandas.io.stata.ValueLabelTypeMismatch"
    )
    def test_categorical_writing(self, version):
        original = DataFrame.from_records(
            [
                ["one", "ten", "one", "one", "one", 1],
                ["two", "nine", "two", "two", "two", 2],
                ["three", "eight", "three", "three", "three", 3],
                ["four", "seven", 4, "four", "four", 4],
                ["five", "six", 5, np.nan, "five", 5],
                ["six", "five", 6, np.nan, "six", 6],
                ["seven", "four", 7, np.nan, "seven", 7],
                ["eight", "three", 8, np.nan, "eight", 8],
                ["nine", "two", 9, np.nan, "nine", 9],
                ["ten", "one", "ten", np.nan, "ten", 10],
            ],
            columns=[
                "fully_labeled",
                "fully_labeled2",
                "incompletely_labeled",
                "labeled_with_missings",
                "float_labelled",
                "unlabeled",
            ],
        )
        expected = original.copy()

        # these are all categoricals
        original = pd.concat(
            [original[col].astype("category") for col in original], axis=1
        )
        expected.index = expected.index.set_names("index").astype(np.int32)

        expected["incompletely_labeled"] = expected["incompletely_labeled"].apply(str)
        expected["unlabeled"] = expected["unlabeled"].apply(str)
        for col in expected:
            orig = expected[col].copy()

            cat = orig.astype("category")._values
            cat = cat.as_ordered()
            if col == "unlabeled":
                cat = cat.set_categories(orig, ordered=True)

            cat.categories.rename(None, inplace=True)

            expected[col] = cat

        with tm.ensure_clean() as path:
            original.to_stata(path, version=version)
            written_and_read_again = self.read_dta(path)

        res = written_and_read_again.set_index("index")
        tm.assert_frame_equal(res, expected)

    def test_categorical_warnings_and_errors(self):
        # Warning for non-string labels
        # Error for labels too long
        original = DataFrame.from_records(
            [["a" * 10000], ["b" * 10000], ["c" * 10000], ["d" * 10000]],
            columns=["Too_long"],
        )

        original = pd.concat(
            [original[col].astype("category") for col in original], axis=1
        )
        with tm.ensure_clean() as path:
            msg = (
                "Stata value labels for a single variable must have "
                r"a combined length less than 32,000 characters\."
            )
            with pytest.raises(ValueError, match=msg):
                original.to_stata(path)

        original = DataFrame.from_records(
            [["a"], ["b"], ["c"], ["d"], [1]], columns=["Too_long"]
        )
        original = pd.concat(
            [original[col].astype("category") for col in original], axis=1
        )

        with tm.assert_produces_warning(ValueLabelTypeMismatch):
            original.to_stata(path)
            # should get a warning for mixed content

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_categorical_with_stata_missing_values(self, version):
        values = [["a" + str(i)] for i in range(120)]
        values.append([np.nan])
        original = DataFrame.from_records(values, columns=["many_labels"])
        original = pd.concat(
            [original[col].astype("category") for col in original], axis=1
        )
        original.index.name = "index"
        with tm.ensure_clean() as path:
            original.to_stata(path, version=version)
            written_and_read_again = self.read_dta(path)

        res = written_and_read_again.set_index("index")

        expected = original.copy()
        for col in expected:
            cat = expected[col]._values
            new_cats = cat.remove_unused_categories().categories
            cat = cat.set_categories(new_cats, ordered=True)
            expected[col] = cat
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(res, expected)

    @pytest.mark.parametrize("file", ["stata10_115", "stata10_117"])
    def test_categorical_order(self, file, datapath):
        # Directly construct using expected codes
        # Format is is_cat, col_name, labels (in order), underlying data
        expected = [
            (True, "ordered", ["a", "b", "c", "d", "e"], np.arange(5)),
            (True, "reverse", ["a", "b", "c", "d", "e"], np.arange(5)[::-1]),
            (True, "noorder", ["a", "b", "c", "d", "e"], np.array([2, 1, 4, 0, 3])),
            (True, "floating", ["a", "b", "c", "d", "e"], np.arange(0, 5)),
            (True, "float_missing", ["a", "d", "e"], np.array([0, 1, 2, -1, -1])),
            (False, "nolabel", [1.0, 2.0, 3.0, 4.0, 5.0], np.arange(5)),
            (True, "int32_mixed", ["d", 2, "e", "b", "a"], np.arange(5)),
        ]
        cols = []
        for is_cat, col, labels, codes in expected:
            if is_cat:
                cols.append(
                    (col, pd.Categorical.from_codes(codes, labels, ordered=True))
                )
            else:
                cols.append((col, Series(labels, dtype=np.float32)))
        expected = DataFrame.from_dict(dict(cols))

        # Read with and with out categoricals, ensure order is identical
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = read_stata(file)
        tm.assert_frame_equal(expected, parsed)

        # Check identity of codes
        for col in expected:
            if isinstance(expected[col].dtype, CategoricalDtype):
                tm.assert_series_equal(expected[col].cat.codes, parsed[col].cat.codes)
                tm.assert_index_equal(
                    expected[col].cat.categories, parsed[col].cat.categories
                )

    @pytest.mark.parametrize("file", ["stata11_115", "stata11_117"])
    def test_categorical_sorting(self, file, datapath):
        parsed = read_stata(datapath("io", "data", "stata", f"{file}.dta"))

        # Sort based on codes, not strings
        parsed = parsed.sort_values("srh", na_position="first")

        # Don't sort index
        parsed.index = pd.RangeIndex(len(parsed))
        codes = [-1, -1, 0, 1, 1, 1, 2, 2, 3, 4]
        categories = ["Poor", "Fair", "Good", "Very good", "Excellent"]
        cat = pd.Categorical.from_codes(
            codes=codes, categories=categories, ordered=True
        )
        expected = Series(cat, name="srh")
        tm.assert_series_equal(expected, parsed["srh"])

    @pytest.mark.parametrize("file", ["stata10_115", "stata10_117"])
    def test_categorical_ordering(self, file, datapath):
        file = datapath("io", "data", "stata", f"{file}.dta")
        parsed = read_stata(file)

        parsed_unordered = read_stata(file, order_categoricals=False)
        for col in parsed:
            if not isinstance(parsed[col].dtype, CategoricalDtype):
                continue
            assert parsed[col].cat.ordered
            assert not parsed_unordered[col].cat.ordered

    @pytest.mark.filterwarnings("ignore::UserWarning")
    @pytest.mark.parametrize(
        "file",
        [
            "stata1_117",
            "stata2_117",
            "stata3_117",
            "stata4_117",
            "stata5_117",
            "stata6_117",
            "stata7_117",
            "stata8_117",
            "stata9_117",
            "stata10_117",
            "stata11_117",
        ],
    )
    @pytest.mark.parametrize("chunksize", [1, 2])
    @pytest.mark.parametrize("convert_categoricals", [False, True])
    @pytest.mark.parametrize("convert_dates", [False, True])
    def test_read_chunks_117(
        self, file, chunksize, convert_categoricals, convert_dates, datapath
    ):
        fname = datapath("io", "data", "stata", f"{file}.dta")

        parsed = read_stata(
            fname,
            convert_categoricals=convert_categoricals,
            convert_dates=convert_dates,
        )
        with read_stata(
            fname,
            iterator=True,
            convert_categoricals=convert_categoricals,
            convert_dates=convert_dates,
        ) as itr:
            pos = 0
            for j in range(5):
                try:
                    chunk = itr.read(chunksize)
                except StopIteration:
                    break
                from_frame = parsed.iloc[pos : pos + chunksize, :].copy()
                from_frame = self._convert_categorical(from_frame)
                tm.assert_frame_equal(
                    from_frame, chunk, check_dtype=False, check_datetimelike_compat=True
                )
                pos += chunksize

    @staticmethod
    def _convert_categorical(from_frame: DataFrame) -> DataFrame:
        """
        Emulate the categorical casting behavior we expect from roundtripping.
        """
        for col in from_frame:
            ser = from_frame[col]
            if isinstance(ser.dtype, CategoricalDtype):
                cat = ser._values.remove_unused_categories()
                if cat.categories.dtype == object:
                    categories = pd.Index._with_infer(cat.categories._values)
                    cat = cat.set_categories(categories)
                elif cat.categories.dtype == "string" and len(cat.categories) == 0:
                    # if the read categories are empty, it comes back as object dtype
                    categories = cat.categories.astype(object)
                    cat = cat.set_categories(categories)
                from_frame[col] = cat
        return from_frame

    def test_iterator(self, datapath):
        fname = datapath("io", "data", "stata", "stata3_117.dta")

        parsed = read_stata(fname)

        with read_stata(fname, iterator=True) as itr:
            chunk = itr.read(5)
            tm.assert_frame_equal(parsed.iloc[0:5, :], chunk)

        with read_stata(fname, chunksize=5) as itr:
            chunk = list(itr)
            tm.assert_frame_equal(parsed.iloc[0:5, :], chunk[0])

        with read_stata(fname, iterator=True) as itr:
            chunk = itr.get_chunk(5)
            tm.assert_frame_equal(parsed.iloc[0:5, :], chunk)

        with read_stata(fname, chunksize=5) as itr:
            chunk = itr.get_chunk()
            tm.assert_frame_equal(parsed.iloc[0:5, :], chunk)

        # GH12153
        with read_stata(fname, chunksize=4) as itr:
            from_chunks = pd.concat(itr)
        tm.assert_frame_equal(parsed, from_chunks)

    @pytest.mark.filterwarnings("ignore::UserWarning")
    @pytest.mark.parametrize(
        "file",
        [
            "stata2_115",
            "stata3_115",
            "stata4_115",
            "stata5_115",
            "stata6_115",
            "stata7_115",
            "stata8_115",
            "stata9_115",
            "stata10_115",
            "stata11_115",
        ],
    )
    @pytest.mark.parametrize("chunksize", [1, 2])
    @pytest.mark.parametrize("convert_categoricals", [False, True])
    @pytest.mark.parametrize("convert_dates", [False, True])
    def test_read_chunks_115(
        self, file, chunksize, convert_categoricals, convert_dates, datapath
    ):
        fname = datapath("io", "data", "stata", f"{file}.dta")

        # Read the whole file
        parsed = read_stata(
            fname,
            convert_categoricals=convert_categoricals,
            convert_dates=convert_dates,
        )

        # Compare to what we get when reading by chunk
        with read_stata(
            fname,
            iterator=True,
            convert_dates=convert_dates,
            convert_categoricals=convert_categoricals,
        ) as itr:
            pos = 0
            for j in range(5):
                try:
                    chunk = itr.read(chunksize)
                except StopIteration:
                    break
                from_frame = parsed.iloc[pos : pos + chunksize, :].copy()
                from_frame = self._convert_categorical(from_frame)
                tm.assert_frame_equal(
                    from_frame, chunk, check_dtype=False, check_datetimelike_compat=True
                )
                pos += chunksize

    def test_read_chunks_columns(self, datapath):
        fname = datapath("io", "data", "stata", "stata3_117.dta")
        columns = ["quarter", "cpi", "m1"]
        chunksize = 2

        parsed = read_stata(fname, columns=columns)
        with read_stata(fname, iterator=True) as itr:
            pos = 0
            for j in range(5):
                chunk = itr.read(chunksize, columns=columns)
                if chunk is None:
                    break
                from_frame = parsed.iloc[pos : pos + chunksize, :]
                tm.assert_frame_equal(from_frame, chunk, check_dtype=False)
                pos += chunksize

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_write_variable_labels(self, version, mixed_frame):
        # GH 13631, add support for writing variable labels
        mixed_frame.index.name = "index"
        variable_labels = {"a": "City Rank", "b": "City Exponent", "c": "City"}
        with tm.ensure_clean() as path:
            mixed_frame.to_stata(path, variable_labels=variable_labels, version=version)
            with StataReader(path) as sr:
                read_labels = sr.variable_labels()
            expected_labels = {
                "index": "",
                "a": "City Rank",
                "b": "City Exponent",
                "c": "City",
            }
            assert read_labels == expected_labels

        variable_labels["index"] = "The Index"
        with tm.ensure_clean() as path:
            mixed_frame.to_stata(path, variable_labels=variable_labels, version=version)
            with StataReader(path) as sr:
                read_labels = sr.variable_labels()
            assert read_labels == variable_labels

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_invalid_variable_labels(self, version, mixed_frame):
        mixed_frame.index.name = "index"
        variable_labels = {"a": "very long" * 10, "b": "City Exponent", "c": "City"}
        with tm.ensure_clean() as path:
            msg = "Variable labels must be 80 characters or fewer"
            with pytest.raises(ValueError, match=msg):
                mixed_frame.to_stata(
                    path, variable_labels=variable_labels, version=version
                )

    @pytest.mark.parametrize("version", [114, 117])
    def test_invalid_variable_label_encoding(self, version, mixed_frame):
        mixed_frame.index.name = "index"
        variable_labels = {"a": "very long" * 10, "b": "City Exponent", "c": "City"}
        variable_labels["a"] = "invalid character Œ"
        with tm.ensure_clean() as path:
            with pytest.raises(
                ValueError, match="Variable labels must contain only characters"
            ):
                mixed_frame.to_stata(
                    path, variable_labels=variable_labels, version=version
                )

    def test_write_variable_label_errors(self, mixed_frame):
        values = ["\u03A1", "\u0391", "\u039D", "\u0394", "\u0391", "\u03A3"]

        variable_labels_utf8 = {
            "a": "City Rank",
            "b": "City Exponent",
            "c": "".join(values),
        }

        msg = (
            "Variable labels must contain only characters that can be "
            "encoded in Latin-1"
        )
        with pytest.raises(ValueError, match=msg):
            with tm.ensure_clean() as path:
                mixed_frame.to_stata(path, variable_labels=variable_labels_utf8)

        variable_labels_long = {
            "a": "City Rank",
            "b": "City Exponent",
            "c": "A very, very, very long variable label "
            "that is too long for Stata which means "
            "that it has more than 80 characters",
        }

        msg = "Variable labels must be 80 characters or fewer"
        with pytest.raises(ValueError, match=msg):
            with tm.ensure_clean() as path:
                mixed_frame.to_stata(path, variable_labels=variable_labels_long)

    def test_default_date_conversion(self):
        # GH 12259
        dates = [
            dt.datetime(1999, 12, 31, 12, 12, 12, 12000),
            dt.datetime(2012, 12, 21, 12, 21, 12, 21000),
            dt.datetime(1776, 7, 4, 7, 4, 7, 4000),
        ]
        original = DataFrame(
            {
                "nums": [1.0, 2.0, 3.0],
                "strs": ["apple", "banana", "cherry"],
                "dates": dates,
            }
        )

        with tm.ensure_clean() as path:
            original.to_stata(path, write_index=False)
            reread = read_stata(path, convert_dates=True)
            tm.assert_frame_equal(original, reread)

            original.to_stata(path, write_index=False, convert_dates={"dates": "tc"})
            direct = read_stata(path, convert_dates=True)
            tm.assert_frame_equal(reread, direct)

            dates_idx = original.columns.tolist().index("dates")
            original.to_stata(path, write_index=False, convert_dates={dates_idx: "tc"})
            direct = read_stata(path, convert_dates=True)
            tm.assert_frame_equal(reread, direct)

    def test_unsupported_type(self):
        original = DataFrame({"a": [1 + 2j, 2 + 4j]})

        msg = "Data type complex128 not supported"
        with pytest.raises(NotImplementedError, match=msg):
            with tm.ensure_clean() as path:
                original.to_stata(path)

    def test_unsupported_datetype(self):
        dates = [
            dt.datetime(1999, 12, 31, 12, 12, 12, 12000),
            dt.datetime(2012, 12, 21, 12, 21, 12, 21000),
            dt.datetime(1776, 7, 4, 7, 4, 7, 4000),
        ]
        original = DataFrame(
            {
                "nums": [1.0, 2.0, 3.0],
                "strs": ["apple", "banana", "cherry"],
                "dates": dates,
            }
        )

        msg = "Format %tC not implemented"
        with pytest.raises(NotImplementedError, match=msg):
            with tm.ensure_clean() as path:
                original.to_stata(path, convert_dates={"dates": "tC"})

        dates = pd.date_range("1-1-1990", periods=3, tz="Asia/Hong_Kong")
        original = DataFrame(
            {
                "nums": [1.0, 2.0, 3.0],
                "strs": ["apple", "banana", "cherry"],
                "dates": dates,
            }
        )
        with pytest.raises(NotImplementedError, match="Data type datetime64"):
            with tm.ensure_clean() as path:
                original.to_stata(path)

    def test_repeated_column_labels(self, datapath):
        # GH 13923, 25772
        msg = """
Value labels for column ethnicsn are not unique. These cannot be converted to
pandas categoricals.

Either read the file with `convert_categoricals` set to False or use the
low level interface in `StataReader` to separately read the values and the
value_labels.

The repeated labels are:\n-+\nwolof
"""
        with pytest.raises(ValueError, match=msg):
            read_stata(
                datapath("io", "data", "stata", "stata15.dta"),
                convert_categoricals=True,
            )

    def test_stata_111(self, datapath):
        # 111 is an old version but still used by current versions of
        # SAS when exporting to Stata format. We do not know of any
        # on-line documentation for this version.
        df = read_stata(datapath("io", "data", "stata", "stata7_111.dta"))
        original = DataFrame(
            {
                "y": [1, 1, 1, 1, 1, 0, 0, np.nan, 0, 0],
                "x": [1, 2, 1, 3, np.nan, 4, 3, 5, 1, 6],
                "w": [2, np.nan, 5, 2, 4, 4, 3, 1, 2, 3],
                "z": ["a", "b", "c", "d", "e", "", "g", "h", "i", "j"],
            }
        )
        original = original[["y", "x", "w", "z"]]
        tm.assert_frame_equal(original, df)

    def test_out_of_range_double(self):
        # GH 14618
        df = DataFrame(
            {
                "ColumnOk": [0.0, np.finfo(np.double).eps, 4.49423283715579e307],
                "ColumnTooBig": [0.0, np.finfo(np.double).eps, np.finfo(np.double).max],
            }
        )
        msg = (
            r"Column ColumnTooBig has a maximum value \(.+\) outside the range "
            r"supported by Stata \(.+\)"
        )
        with pytest.raises(ValueError, match=msg):
            with tm.ensure_clean() as path:
                df.to_stata(path)

    def test_out_of_range_float(self):
        original = DataFrame(
            {
                "ColumnOk": [
                    0.0,
                    np.finfo(np.float32).eps,
                    np.finfo(np.float32).max / 10.0,
                ],
                "ColumnTooBig": [
                    0.0,
                    np.finfo(np.float32).eps,
                    np.finfo(np.float32).max,
                ],
            }
        )
        original.index.name = "index"
        for col in original:
            original[col] = original[col].astype(np.float32)

        with tm.ensure_clean() as path:
            original.to_stata(path)
            reread = read_stata(path)

        original["ColumnTooBig"] = original["ColumnTooBig"].astype(np.float64)
        expected = original.copy()
        expected.index = expected.index.astype(np.int32)
        tm.assert_frame_equal(reread.set_index("index"), expected)

    @pytest.mark.parametrize("infval", [np.inf, -np.inf])
    def test_inf(self, infval):
        # GH 45350
        df = DataFrame({"WithoutInf": [0.0, 1.0], "WithInf": [2.0, infval]})
        msg = (
            "Column WithInf contains infinity or -infinity"
            "which is outside the range supported by Stata."
        )
        with pytest.raises(ValueError, match=msg):
            with tm.ensure_clean() as path:
                df.to_stata(path)

    def test_path_pathlib(self):
        df = DataFrame(
            1.1 * np.arange(120).reshape((30, 4)),
            columns=pd.Index(list("ABCD")),
            index=pd.Index([f"i-{i}" for i in range(30)]),
        )
        df.index.name = "index"
        reader = lambda x: read_stata(x).set_index("index")
        result = tm.round_trip_pathlib(df.to_stata, reader)
        tm.assert_frame_equal(df, result)

    def test_pickle_path_localpath(self):
        df = DataFrame(
            1.1 * np.arange(120).reshape((30, 4)),
            columns=pd.Index(list("ABCD")),
            index=pd.Index([f"i-{i}" for i in range(30)]),
        )
        df.index.name = "index"
        reader = lambda x: read_stata(x).set_index("index")
        result = tm.round_trip_localpath(df.to_stata, reader)
        tm.assert_frame_equal(df, result)

    @pytest.mark.parametrize("write_index", [True, False])
    def test_value_labels_iterator(self, write_index):
        # GH 16923
        d = {"A": ["B", "E", "C", "A", "E"]}
        df = DataFrame(data=d)
        df["A"] = df["A"].astype("category")
        with tm.ensure_clean() as path:
            df.to_stata(path, write_index=write_index)

            with read_stata(path, iterator=True) as dta_iter:
                value_labels = dta_iter.value_labels()
        assert value_labels == {"A": {0: "A", 1: "B", 2: "C", 3: "E"}}

    def test_set_index(self):
        # GH 17328
        df = DataFrame(
            1.1 * np.arange(120).reshape((30, 4)),
            columns=pd.Index(list("ABCD")),
            index=pd.Index([f"i-{i}" for i in range(30)]),
        )
        df.index.name = "index"
        with tm.ensure_clean() as path:
            df.to_stata(path)
            reread = read_stata(path, index_col="index")
        tm.assert_frame_equal(df, reread)

    @pytest.mark.parametrize(
        "column", ["ms", "day", "week", "month", "qtr", "half", "yr"]
    )
    def test_date_parsing_ignores_format_details(self, column, datapath):
        # GH 17797
        #
        # Test that display formats are ignored when determining if a numeric
        # column is a date value.
        #
        # All date types are stored as numbers and format associated with the
        # column denotes both the type of the date and the display format.
        #
        # STATA supports 9 date types which each have distinct units. We test 7
        # of the 9 types, ignoring %tC and %tb. %tC is a variant of %tc that
        # accounts for leap seconds and %tb relies on STATAs business calendar.
        df = read_stata(datapath("io", "data", "stata", "stata13_dates.dta"))
        unformatted = df.loc[0, column]
        formatted = df.loc[0, column + "_fmt"]
        assert unformatted == formatted

    def test_writer_117(self, using_infer_string):
        original = DataFrame(
            data=[
                [
                    "string",
                    "object",
                    1,
                    1,
                    1,
                    1.1,
                    1.1,
                    np.datetime64("2003-12-25"),
                    "a",
                    "a" * 2045,
                    "a" * 5000,
                    "a",
                ],
                [
                    "string-1",
                    "object-1",
                    1,
                    1,
                    1,
                    1.1,
                    1.1,
                    np.datetime64("2003-12-26"),
                    "b",
                    "b" * 2045,
                    "",
                    "",
                ],
            ],
            columns=[
                "string",
                "object",
                "int8",
                "int16",
                "int32",
                "float32",
                "float64",
                "datetime",
                "s1",
                "s2045",
                "srtl",
                "forced_strl",
            ],
        )
        original["object"] = Series(original["object"], dtype=object)
        original["int8"] = Series(original["int8"], dtype=np.int8)
        original["int16"] = Series(original["int16"], dtype=np.int16)
        original["int32"] = original["int32"].astype(np.int32)
        original["float32"] = Series(original["float32"], dtype=np.float32)
        original.index.name = "index"
        original.index = original.index.astype(np.int32)
        copy = original.copy()
        with tm.ensure_clean() as path:
            original.to_stata(
                path,
                convert_dates={"datetime": "tc"},
                convert_strl=["forced_strl"],
                version=117,
            )
            written_and_read_again = self.read_dta(path)

        expected = original[:]
        if using_infer_string:
            # object dtype (with only strings/None) comes back as string dtype
            expected["object"] = expected["object"].astype("str")

        tm.assert_frame_equal(
            written_and_read_again.set_index("index"),
            expected,
        )
        tm.assert_frame_equal(original, copy)

    def test_convert_strl_name_swap(self):
        original = DataFrame(
            [["a" * 3000, "A", "apple"], ["b" * 1000, "B", "banana"]],
            columns=["long1" * 10, "long", 1],
        )
        original.index.name = "index"

        with tm.assert_produces_warning(InvalidColumnName):
            with tm.ensure_clean() as path:
                original.to_stata(path, convert_strl=["long", 1], version=117)
                reread = self.read_dta(path)
                reread = reread.set_index("index")
                reread.columns = original.columns
                tm.assert_frame_equal(reread, original, check_index_type=False)

    def test_invalid_date_conversion(self):
        # GH 12259
        dates = [
            dt.datetime(1999, 12, 31, 12, 12, 12, 12000),
            dt.datetime(2012, 12, 21, 12, 21, 12, 21000),
            dt.datetime(1776, 7, 4, 7, 4, 7, 4000),
        ]
        original = DataFrame(
            {
                "nums": [1.0, 2.0, 3.0],
                "strs": ["apple", "banana", "cherry"],
                "dates": dates,
            }
        )

        with tm.ensure_clean() as path:
            msg = "convert_dates key must be a column or an integer"
            with pytest.raises(ValueError, match=msg):
                original.to_stata(path, convert_dates={"wrong_name": "tc"})

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_nonfile_writing(self, version):
        # GH 21041
        bio = io.BytesIO()
        df = DataFrame(
            1.1 * np.arange(120).reshape((30, 4)),
            columns=pd.Index(list("ABCD")),
            index=pd.Index([f"i-{i}" for i in range(30)]),
        )
        df.index.name = "index"
        with tm.ensure_clean() as path:
            df.to_stata(bio, version=version)
            bio.seek(0)
            with open(path, "wb") as dta:
                dta.write(bio.read())
            reread = read_stata(path, index_col="index")
        tm.assert_frame_equal(df, reread)

    def test_gzip_writing(self):
        # writing version 117 requires seek and cannot be used with gzip
        df = DataFrame(
            1.1 * np.arange(120).reshape((30, 4)),
            columns=pd.Index(list("ABCD")),
            index=pd.Index([f"i-{i}" for i in range(30)]),
        )
        df.index.name = "index"
        with tm.ensure_clean() as path:
            with gzip.GzipFile(path, "wb") as gz:
                df.to_stata(gz, version=114)
            with gzip.GzipFile(path, "rb") as gz:
                reread = read_stata(gz, index_col="index")
        tm.assert_frame_equal(df, reread)

    def test_unicode_dta_118(self, datapath):
        unicode_df = self.read_dta(datapath("io", "data", "stata", "stata16_118.dta"))

        columns = ["utf8", "latin1", "ascii", "utf8_strl", "ascii_strl"]
        values = [
            ["ραηδας", "PÄNDÄS", "p", "ραηδας", "p"],
            ["ƤĀńĐąŜ", "Ö", "a", "ƤĀńĐąŜ", "a"],
            ["ᴘᴀᴎᴅᴀS", "Ü", "n", "ᴘᴀᴎᴅᴀS", "n"],
            ["      ", "      ", "d", "      ", "d"],
            [" ", "", "a", " ", "a"],
            ["", "", "s", "", "s"],
            ["", "", " ", "", " "],
        ]
        expected = DataFrame(values, columns=columns)

        tm.assert_frame_equal(unicode_df, expected)

    def test_mixed_string_strl(self, using_infer_string):
        # GH 23633
        output = [{"mixed": "string" * 500, "number": 0}, {"mixed": None, "number": 1}]
        output = DataFrame(output)
        output.number = output.number.astype("int32")

        with tm.ensure_clean() as path:
            output.to_stata(path, write_index=False, version=117)
            reread = read_stata(path)
            expected = output.fillna("")
            tm.assert_frame_equal(reread, expected)

            # Check strl supports all None (null)
            output["mixed"] = None
            output.to_stata(
                path, write_index=False, convert_strl=["mixed"], version=117
            )
            reread = read_stata(path)
            expected = output.copy()
            if using_infer_string:
                expected["mixed"] = expected["mixed"].astype("str")
            expected = expected.fillna("")
            tm.assert_frame_equal(reread, expected)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_all_none_exception(self, version):
        output = [{"none": "none", "number": 0}, {"none": None, "number": 1}]
        output = DataFrame(output)
        output["none"] = None
        with tm.ensure_clean() as path:
            with pytest.raises(ValueError, match="Column `none` cannot be exported"):
                output.to_stata(path, version=version)

    @pytest.mark.parametrize("version", [114, 117, 118, 119, None])
    def test_invalid_file_not_written(self, version):
        content = "Here is one __�__ Another one __·__ Another one __½__"
        df = DataFrame([content], columns=["invalid"])
        with tm.ensure_clean() as path:
            msg1 = (
                r"'latin-1' codec can't encode character '\\ufffd' "
                r"in position 14: ordinal not in range\(256\)"
            )
            msg2 = (
                "'ascii' codec can't decode byte 0xef in position 14: "
                r"ordinal not in range\(128\)"
            )
            with pytest.raises(UnicodeEncodeError, match=f"{msg1}|{msg2}"):
                df.to_stata(path)

    def test_strl_latin1(self):
        # GH 23573, correct GSO data to reflect correct size
        output = DataFrame(
            [["pandas"] * 2, ["þâÑÐÅ§"] * 2], columns=["var_str", "var_strl"]
        )

        with tm.ensure_clean() as path:
            output.to_stata(path, version=117, convert_strl=["var_strl"])
            with open(path, "rb") as reread:
                content = reread.read()
                expected = "þâÑÐÅ§"
                assert expected.encode("latin-1") in content
                assert expected.encode("utf-8") in content
                gsos = content.split(b"strls")[1][1:-2]
                for gso in gsos.split(b"GSO")[1:]:
                    val = gso.split(b"\x00")[-2]
                    size = gso[gso.find(b"\x82") + 1]
                    assert len(val) == size - 1

    def test_encoding_latin1_118(self, datapath):
        # GH 25960
        msg = """
One or more strings in the dta file could not be decoded using utf-8, and
so the fallback encoding of latin-1 is being used.  This can happen when a file
has been incorrectly encoded by Stata or some other software. You should verify
the string values returned are correct."""
        # Move path outside of read_stata, or else assert_produces_warning
        # will block pytests skip mechanism from triggering (failing the test)
        # if the path is not present
        path = datapath("io", "data", "stata", "stata1_encoding_118.dta")
        with tm.assert_produces_warning(UnicodeWarning, filter_level="once") as w:
            encoded = read_stata(path)
            # with filter_level="always", produces 151 warnings which can be slow
            assert len(w) == 1
            assert w[0].message.args[0] == msg

        expected = DataFrame([["Düsseldorf"]] * 151, columns=["kreis1849"])
        tm.assert_frame_equal(encoded, expected)

    @pytest.mark.slow
    def test_stata_119(self, datapath):
        # Gzipped since contains 32,999 variables and uncompressed is 20MiB
        # Just validate that the reader reports correct number of variables
        # to avoid high peak memory
        with gzip.open(
            datapath("io", "data", "stata", "stata1_119.dta.gz"), "rb"
        ) as gz:
            with StataReader(gz) as reader:
                reader._ensure_open()
                assert reader._nvar == 32999

    @pytest.mark.filterwarnings("ignore:Downcasting behavior:FutureWarning")
    @pytest.mark.parametrize("version", [118, 119, None])
    def test_utf8_writer(self, version):
        cat = pd.Categorical(["a", "β", "ĉ"], ordered=True)
        data = DataFrame(
            [
                [1.0, 1, "ᴬ", "ᴀ relatively long ŝtring"],
                [2.0, 2, "ᴮ", ""],
                [3.0, 3, "ᴰ", None],
            ],
            columns=["Å", "β", "ĉ", "strls"],
        )
        data["ᴐᴬᵀ"] = cat
        variable_labels = {
            "Å": "apple",
            "β": "ᵈᵉᵊ",
            "ĉ": "ᴎტჄႲႳႴႶႺ",
            "strls": "Long Strings",
            "ᴐᴬᵀ": "",
        }
        data_label = "ᴅaᵀa-label"
        value_labels = {"β": {1: "label", 2: "æøå", 3: "ŋot valid latin-1"}}
        data["β"] = data["β"].astype(np.int32)
        with tm.ensure_clean() as path:
            writer = StataWriterUTF8(
                path,
                data,
                data_label=data_label,
                convert_strl=["strls"],
                variable_labels=variable_labels,
                write_index=False,
                version=version,
                value_labels=value_labels,
            )
            writer.write_file()
            reread_encoded = read_stata(path)
            # Missing is intentionally converted to empty strl
            data["strls"] = data["strls"].fillna("")
            # Variable with value labels is reread as categorical
            data["β"] = (
                data["β"].replace(value_labels["β"]).astype("category").cat.as_ordered()
            )
            tm.assert_frame_equal(data, reread_encoded)
            with StataReader(path) as reader:
                assert reader.data_label == data_label
                assert reader.variable_labels() == variable_labels

            data.to_stata(path, version=version, write_index=False)
            reread_to_stata = read_stata(path)
            tm.assert_frame_equal(data, reread_to_stata)

    def test_writer_118_exceptions(self):
        df = DataFrame(np.zeros((1, 33000), dtype=np.int8))
        with tm.ensure_clean() as path:
            with pytest.raises(ValueError, match="version must be either 118 or 119."):
                StataWriterUTF8(path, df, version=117)
        with tm.ensure_clean() as path:
            with pytest.raises(ValueError, match="You must use version 119"):
                StataWriterUTF8(path, df, version=118)

    @pytest.mark.parametrize(
        "dtype_backend",
        ["numpy_nullable", pytest.param("pyarrow", marks=td.skip_if_no("pyarrow"))],
    )
    def test_read_write_ea_dtypes(self, dtype_backend):
        df = DataFrame(
            {
                "a": [1, 2, None],
                "b": ["a", "b", "c"],
                "c": [True, False, None],
                "d": [1.5, 2.5, 3.5],
                "e": pd.date_range("2020-12-31", periods=3, freq="D"),
            },
            index=pd.Index([0, 1, 2], name="index"),
        )
        df = df.convert_dtypes(dtype_backend=dtype_backend)
        df.to_stata("test_stata.dta", version=118)

        with tm.ensure_clean() as path:
            df.to_stata(path)
            written_and_read_again = self.read_dta(path)

        expected = DataFrame(
            {
                "a": [1, 2, np.nan],
                "b": ["a", "b", "c"],
                "c": [1.0, 0, np.nan],
                "d": [1.5, 2.5, 3.5],
                "e": pd.date_range("2020-12-31", periods=3, freq="D"),
            },
            index=pd.Index([0, 1, 2], name="index", dtype=np.int32),
        )

        tm.assert_frame_equal(written_and_read_again.set_index("index"), expected)


@pytest.mark.parametrize("version", [105, 108, 111, 113, 114])
def test_backward_compat(version, datapath):
    data_base = datapath("io", "data", "stata")
    ref = os.path.join(data_base, "stata-compat-118.dta")
    old = os.path.join(data_base, f"stata-compat-{version}.dta")
    expected = read_stata(ref)
    old_dta = read_stata(old)
    tm.assert_frame_equal(old_dta, expected, check_dtype=False)


def test_direct_read(datapath, monkeypatch):
    file_path = datapath("io", "data", "stata", "stata-compat-118.dta")

    # Test that opening a file path doesn't buffer the file.
    with StataReader(file_path) as reader:
        # Must not have been buffered to memory
        assert not reader.read().empty
        assert not isinstance(reader._path_or_buf, io.BytesIO)

    # Test that we use a given fp exactly, if possible.
    with open(file_path, "rb") as fp:
        with StataReader(fp) as reader:
            assert not reader.read().empty
            assert reader._path_or_buf is fp

    # Test that we use a given BytesIO exactly, if possible.
    with open(file_path, "rb") as fp:
        with io.BytesIO(fp.read()) as bio:
            with StataReader(bio) as reader:
                assert not reader.read().empty
                assert reader._path_or_buf is bio


def test_statareader_warns_when_used_without_context(datapath):
    file_path = datapath("io", "data", "stata", "stata-compat-118.dta")
    with tm.assert_produces_warning(
        ResourceWarning,
        match="without using a context manager",
    ):
        sr = StataReader(file_path)
        sr.read()
    with tm.assert_produces_warning(
        FutureWarning,
        match="is not part of the public API",
    ):
        sr.close()


@pytest.mark.parametrize("version", [114, 117, 118, 119, None])
@pytest.mark.parametrize("use_dict", [True, False])
@pytest.mark.parametrize("infer", [True, False])
def test_compression(compression, version, use_dict, infer, compression_to_extension):
    file_name = "dta_inferred_compression.dta"
    if compression:
        if use_dict:
            file_ext = compression
        else:
            file_ext = compression_to_extension[compression]
        file_name += f".{file_ext}"
    compression_arg = compression
    if infer:
        compression_arg = "infer"
    if use_dict:
        compression_arg = {"method": compression}

    df = DataFrame(
        np.random.default_rng(2).standard_normal((10, 2)), columns=list("AB")
    )
    df.index.name = "index"
    with tm.ensure_clean(file_name) as path:
        df.to_stata(path, version=version, compression=compression_arg)
        if compression == "gzip":
            with gzip.open(path, "rb") as comp:
                fp = io.BytesIO(comp.read())
        elif compression == "zip":
            with zipfile.ZipFile(path, "r") as comp:
                fp = io.BytesIO(comp.read(comp.filelist[0]))
        elif compression == "tar":
            with tarfile.open(path) as tar:
                fp = io.BytesIO(tar.extractfile(tar.getnames()[0]).read())
        elif compression == "bz2":
            with bz2.open(path, "rb") as comp:
                fp = io.BytesIO(comp.read())
        elif compression == "zstd":
            zstd = pytest.importorskip("zstandard")
            with zstd.open(path, "rb") as comp:
                fp = io.BytesIO(comp.read())
        elif compression == "xz":
            lzma = pytest.importorskip("lzma")
            with lzma.open(path, "rb") as comp:
                fp = io.BytesIO(comp.read())
        elif compression is None:
            fp = path
        reread = read_stata(fp, index_col="index")

    expected = df.copy()
    expected.index = expected.index.astype(np.int32)
    tm.assert_frame_equal(reread, expected)


@pytest.mark.parametrize("method", ["zip", "infer"])
@pytest.mark.parametrize("file_ext", [None, "dta", "zip"])
def test_compression_dict(method, file_ext):
    file_name = f"test.{file_ext}"
    archive_name = "test.dta"
    df = DataFrame(
        np.random.default_rng(2).standard_normal((10, 2)), columns=list("AB")
    )
    df.index.name = "index"
    with tm.ensure_clean(file_name) as path:
        compression = {"method": method, "archive_name": archive_name}
        df.to_stata(path, compression=compression)
        if method == "zip" or file_ext == "zip":
            with zipfile.ZipFile(path, "r") as zp:
                assert len(zp.filelist) == 1
                assert zp.filelist[0].filename == archive_name
                fp = io.BytesIO(zp.read(zp.filelist[0]))
        else:
            fp = path
        reread = read_stata(fp, index_col="index")

    expected = df.copy()
    expected.index = expected.index.astype(np.int32)
    tm.assert_frame_equal(reread, expected)


@pytest.mark.parametrize("version", [114, 117, 118, 119, None])
def test_chunked_categorical(version):
    df = DataFrame({"cats": Series(["a", "b", "a", "b", "c"], dtype="category")})
    df.index.name = "index"

    expected = df.copy()
    expected.index = expected.index.astype(np.int32)

    with tm.ensure_clean() as path:
        df.to_stata(path, version=version)
        with StataReader(path, chunksize=2, order_categoricals=False) as reader:
            for i, block in enumerate(reader):
                block = block.set_index("index")
                assert "cats" in block
                tm.assert_series_equal(
                    block.cats, expected.cats.iloc[2 * i : 2 * (i + 1)]
                )


def test_chunked_categorical_partial(datapath):
    dta_file = datapath("io", "data", "stata", "stata-dta-partially-labeled.dta")
    values = ["a", "b", "a", "b", 3.0]
    with StataReader(dta_file, chunksize=2) as reader:
        with tm.assert_produces_warning(CategoricalConversionWarning):
            for i, block in enumerate(reader):
                assert list(block.cats) == values[2 * i : 2 * (i + 1)]
                if i < 2:
                    idx = pd.Index(["a", "b"])
                else:
                    idx = pd.Index([3.0], dtype="float64")
                tm.assert_index_equal(block.cats.cat.categories, idx)
    with tm.assert_produces_warning(CategoricalConversionWarning):
        with StataReader(dta_file, chunksize=5) as reader:
            large_chunk = reader.__next__()
    direct = read_stata(dta_file)
    tm.assert_frame_equal(direct, large_chunk)


@pytest.mark.parametrize("chunksize", (-1, 0, "apple"))
def test_iterator_errors(datapath, chunksize):
    dta_file = datapath("io", "data", "stata", "stata-dta-partially-labeled.dta")
    with pytest.raises(ValueError, match="chunksize must be a positive"):
        with StataReader(dta_file, chunksize=chunksize):
            pass


def test_iterator_value_labels():
    # GH 31544
    values = ["c_label", "b_label"] + ["a_label"] * 500
    df = DataFrame({f"col{k}": pd.Categorical(values, ordered=True) for k in range(2)})
    with tm.ensure_clean() as path:
        df.to_stata(path, write_index=False)
        expected = pd.Index(["a_label", "b_label", "c_label"])
        with read_stata(path, chunksize=100) as reader:
            for j, chunk in enumerate(reader):
                for i in range(2):
                    tm.assert_index_equal(chunk.dtypes.iloc[i].categories, expected)
                tm.assert_frame_equal(chunk, df.iloc[j * 100 : (j + 1) * 100])


def test_precision_loss():
    df = DataFrame(
        [[sum(2**i for i in range(60)), sum(2**i for i in range(52))]],
        columns=["big", "little"],
    )
    with tm.ensure_clean() as path:
        with tm.assert_produces_warning(
            PossiblePrecisionLoss, match="Column converted from int64 to float64"
        ):
            df.to_stata(path, write_index=False)
        reread = read_stata(path)
        expected_dt = Series([np.float64, np.float64], index=["big", "little"])
        tm.assert_series_equal(reread.dtypes, expected_dt)
        assert reread.loc[0, "little"] == df.loc[0, "little"]
        assert reread.loc[0, "big"] == float(df.loc[0, "big"])


def test_compression_roundtrip(compression):
    df = DataFrame(
        [[0.123456, 0.234567, 0.567567], [12.32112, 123123.2, 321321.2]],
        index=["A", "B"],
        columns=["X", "Y", "Z"],
    )
    df.index.name = "index"

    with tm.ensure_clean() as path:
        df.to_stata(path, compression=compression)
        reread = read_stata(path, compression=compression, index_col="index")
        tm.assert_frame_equal(df, reread)

        # explicitly ensure file was compressed.
        with tm.decompress_file(path, compression) as fh:
            contents = io.BytesIO(fh.read())
        reread = read_stata(contents, index_col="index")
        tm.assert_frame_equal(df, reread)


@pytest.mark.parametrize("to_infer", [True, False])
@pytest.mark.parametrize("read_infer", [True, False])
def test_stata_compression(
    compression_only, read_infer, to_infer, compression_to_extension
):
    compression = compression_only

    ext = compression_to_extension[compression]
    filename = f"test.{ext}"

    df = DataFrame(
        [[0.123456, 0.234567, 0.567567], [12.32112, 123123.2, 321321.2]],
        index=["A", "B"],
        columns=["X", "Y", "Z"],
    )
    df.index.name = "index"

    to_compression = "infer" if to_infer else compression
    read_compression = "infer" if read_infer else compression

    with tm.ensure_clean(filename) as path:
        df.to_stata(path, compression=to_compression)
        result = read_stata(path, compression=read_compression, index_col="index")
        tm.assert_frame_equal(result, df)


def test_non_categorical_value_labels():
    data = DataFrame(
        {
            "fully_labelled": [1, 2, 3, 3, 1],
            "partially_labelled": [1.0, 2.0, np.nan, 9.0, np.nan],
            "Y": [7, 7, 9, 8, 10],
            "Z": pd.Categorical(["j", "k", "l", "k", "j"]),
        }
    )

    with tm.ensure_clean() as path:
        value_labels = {
            "fully_labelled": {1: "one", 2: "two", 3: "three"},
            "partially_labelled": {1.0: "one", 2.0: "two"},
        }
        expected = {**value_labels, "Z": {0: "j", 1: "k", 2: "l"}}

        writer = StataWriter(path, data, value_labels=value_labels)
        writer.write_file()

        with StataReader(path) as reader:
            reader_value_labels = reader.value_labels()
            assert reader_value_labels == expected

        msg = "Can't create value labels for notY, it wasn't found in the dataset."
        with pytest.raises(KeyError, match=msg):
            value_labels = {"notY": {7: "label1", 8: "label2"}}
            StataWriter(path, data, value_labels=value_labels)

        msg = (
            "Can't create value labels for Z, value labels "
            "can only be applied to numeric columns."
        )
        with pytest.raises(ValueError, match=msg):
            value_labels = {"Z": {1: "a", 2: "k", 3: "j", 4: "i"}}
            StataWriter(path, data, value_labels=value_labels)


def test_non_categorical_value_label_name_conversion():
    # Check conversion of invalid variable names
    data = DataFrame(
        {
            "invalid~!": [1, 1, 2, 3, 5, 8],  # Only alphanumeric and _
            "6_invalid": [1, 1, 2, 3, 5, 8],  # Must start with letter or _
            "invalid_name_longer_than_32_characters": [8, 8, 9, 9, 8, 8],  # Too long
            "aggregate": [2, 5, 5, 6, 6, 9],  # Reserved words
            (1, 2): [1, 2, 3, 4, 5, 6],  # Hashable non-string
        }
    )

    value_labels = {
        "invalid~!": {1: "label1", 2: "label2"},
        "6_invalid": {1: "label1", 2: "label2"},
        "invalid_name_longer_than_32_characters": {8: "eight", 9: "nine"},
        "aggregate": {5: "five"},
        (1, 2): {3: "three"},
    }

    expected = {
        "invalid__": {1: "label1", 2: "label2"},
        "_6_invalid": {1: "label1", 2: "label2"},
        "invalid_name_longer_than_32_char": {8: "eight", 9: "nine"},
        "_aggregate": {5: "five"},
        "_1__2_": {3: "three"},
    }

    with tm.ensure_clean() as path:
        with tm.assert_produces_warning(InvalidColumnName):
            data.to_stata(path, value_labels=value_labels)

        with StataReader(path) as reader:
            reader_value_labels = reader.value_labels()
            assert reader_value_labels == expected


def test_non_categorical_value_label_convert_categoricals_error():
    # Mapping more than one value to the same label is valid for Stata
    # labels, but can't be read with convert_categoricals=True
    value_labels = {
        "repeated_labels": {10: "Ten", 20: "More than ten", 40: "More than ten"}
    }

    data = DataFrame(
        {
            "repeated_labels": [10, 10, 20, 20, 40, 40],
        }
    )

    with tm.ensure_clean() as path:
        data.to_stata(path, value_labels=value_labels)

        with StataReader(path, convert_categoricals=False) as reader:
            reader_value_labels = reader.value_labels()
        assert reader_value_labels == value_labels

        col = "repeated_labels"
        repeats = "-" * 80 + "\n" + "\n".join(["More than ten"])

        msg = f"""
Value labels for column {col} are not unique. These cannot be converted to
pandas categoricals.

Either read the file with `convert_categoricals` set to False or use the
low level interface in `StataReader` to separately read the values and the
value_labels.

The repeated labels are:
{repeats}
"""
        with pytest.raises(ValueError, match=msg):
            read_stata(path, convert_categoricals=True)


@pytest.mark.parametrize("version", [114, 117, 118, 119, None])
@pytest.mark.parametrize(
    "dtype",
    [
        pd.BooleanDtype,
        pd.Int8Dtype,
        pd.Int16Dtype,
        pd.Int32Dtype,
        pd.Int64Dtype,
        pd.UInt8Dtype,
        pd.UInt16Dtype,
        pd.UInt32Dtype,
        pd.UInt64Dtype,
    ],
)
def test_nullable_support(dtype, version):
    df = DataFrame(
        {
            "a": Series([1.0, 2.0, 3.0]),
            "b": Series([1, pd.NA, pd.NA], dtype=dtype.name),
            "c": Series(["a", "b", None]),
        }
    )
    dtype_name = df.b.dtype.numpy_dtype.name
    # Only use supported names: no uint, bool or int64
    dtype_name = dtype_name.replace("u", "")
    if dtype_name == "int64":
        dtype_name = "int32"
    elif dtype_name == "bool":
        dtype_name = "int8"
    value = StataMissingValue.BASE_MISSING_VALUES[dtype_name]
    smv = StataMissingValue(value)
    expected_b = Series([1, smv, smv], dtype=object, name="b")
    expected_c = Series(["a", "b", ""], name="c")
    with tm.ensure_clean() as path:
        df.to_stata(path, write_index=False, version=version)
        reread = read_stata(path, convert_missing=True)
        tm.assert_series_equal(df.a, reread.a)
        tm.assert_series_equal(reread.b, expected_b)
        tm.assert_series_equal(reread.c, expected_c)


def test_empty_frame():
    # GH 46240
    # create an empty DataFrame with int64 and float64 dtypes
    df = DataFrame(data={"a": range(3), "b": [1.0, 2.0, 3.0]}).head(0)
    with tm.ensure_clean() as path:
        df.to_stata(path, write_index=False, version=117)
        # Read entire dataframe
        df2 = read_stata(path)
        assert "b" in df2
        # Dtypes don't match since no support for int32
        dtypes = Series({"a": np.dtype("int32"), "b": np.dtype("float64")})
        tm.assert_series_equal(df2.dtypes, dtypes)
        # read one column of empty .dta file
        df3 = read_stata(path, columns=["a"])
        assert "b" not in df3
        tm.assert_series_equal(df3.dtypes, dtypes.loc[["a"]])

# === atelier-kyo-manager/.venv_backup\Lib\site-packages\mpmath\tests\test_functions2.py ===
import math
import pytest
from mpmath import *

def test_bessel():
    mp.dps = 15
    assert j0(1).ae(0.765197686557966551)
    assert j0(pi).ae(-0.304242177644093864)
    assert j0(1000).ae(0.0247866861524201746)
    assert j0(-25).ae(0.0962667832759581162)
    assert j1(1).ae(0.440050585744933516)
    assert j1(pi).ae(0.284615343179752757)
    assert j1(1000).ae(0.00472831190708952392)
    assert j1(-25).ae(0.125350249580289905)
    assert besselj(5,1).ae(0.000249757730211234431)
    assert besselj(5+0j,1).ae(0.000249757730211234431)
    assert besselj(5,pi).ae(0.0521411843671184747)
    assert besselj(5,1000).ae(0.00502540694523318607)
    assert besselj(5,-25).ae(0.0660079953984229934)
    assert besselj(-3,2).ae(-0.128943249474402051)
    assert besselj(-4,2).ae(0.0339957198075684341)
    assert besselj(3,3+2j).ae(0.424718794929639595942 + 0.625665327745785804812j)
    assert besselj(0.25,4).ae(-0.374760630804249715)
    assert besselj(1+2j,3+4j).ae(0.319247428741872131 - 0.669557748880365678j)
    assert (besselj(3, 10**10) * 10**5).ae(0.76765081748139204023)
    assert bessely(-0.5, 0) == 0
    assert bessely(0.5, 0) == -inf
    assert bessely(1.5, 0) == -inf
    assert bessely(0,0) == -inf
    assert bessely(-0.4, 0) == -inf
    assert bessely(-0.6, 0) == inf
    assert bessely(-1, 0) == inf
    assert bessely(-1.4, 0) == inf
    assert bessely(-1.6, 0) == -inf
    assert bessely(-1, 0) == inf
    assert bessely(-2, 0) == -inf
    assert bessely(-3, 0) == inf
    assert bessely(0.5, 0) == -inf
    assert bessely(1, 0) == -inf
    assert bessely(1.5, 0) == -inf
    assert bessely(2, 0) == -inf
    assert bessely(2.5, 0) == -inf
    assert bessely(3, 0) == -inf
    assert bessely(0,0.5).ae(-0.44451873350670655715)
    assert bessely(1,0.5).ae(-1.4714723926702430692)
    assert bessely(-1,0.5).ae(1.4714723926702430692)
    assert bessely(3.5,0.5).ae(-138.86400867242488443)
    assert bessely(0,3+4j).ae(4.6047596915010138655-8.8110771408232264208j)
    assert bessely(0,j).ae(-0.26803248203398854876+1.26606587775200833560j)
    assert (bessely(3, 10**10) * 10**5).ae(0.21755917537013204058)
    assert besseli(0,0) == 1
    assert besseli(1,0) == 0
    assert besseli(2,0) == 0
    assert besseli(-1,0) == 0
    assert besseli(-2,0) == 0
    assert besseli(0,0.5).ae(1.0634833707413235193)
    assert besseli(1,0.5).ae(0.25789430539089631636)
    assert besseli(-1,0.5).ae(0.25789430539089631636)
    assert besseli(3.5,0.5).ae(0.00068103597085793815863)
    assert besseli(0,3+4j).ae(-3.3924877882755196097-1.3239458916287264815j)
    assert besseli(0,j).ae(besselj(0,1))
    assert (besseli(3, 10**10) * mpf(10)**(-4342944813)).ae(4.2996028505491271875)
    assert besselk(0,0) == inf
    assert besselk(1,0) == inf
    assert besselk(2,0) == inf
    assert besselk(-1,0) == inf
    assert besselk(-2,0) == inf
    assert besselk(0,0.5).ae(0.92441907122766586178)
    assert besselk(1,0.5).ae(1.6564411200033008937)
    assert besselk(-1,0.5).ae(1.6564411200033008937)
    assert besselk(3.5,0.5).ae(207.48418747548460607)
    assert besselk(0,3+4j).ae(-0.007239051213570155013+0.026510418350267677215j)
    assert besselk(0,j).ae(-0.13863371520405399968-1.20196971531720649914j)
    assert (besselk(3, 10**10) * mpf(10)**4342944824).ae(1.1628981033356187851)
    # test for issue 331, bug reported by Michael Hartmann
    for n in range(10,100,10):
        mp.dps = n
        assert besseli(91.5,24.7708).ae("4.00830632138673963619656140653537080438462342928377020695738635559218797348548092636896796324190271316137982810144874264e-41")

def test_bessel_zeros():
    mp.dps = 15
    assert besseljzero(0,1).ae(2.40482555769577276869)
    assert besseljzero(2,1).ae(5.1356223018406825563)
    assert besseljzero(1,50).ae(157.86265540193029781)
    assert besseljzero(10,1).ae(14.475500686554541220)
    assert besseljzero(0.5,3).ae(9.4247779607693797153)
    assert besseljzero(2,1,1).ae(3.0542369282271403228)
    assert besselyzero(0,1).ae(0.89357696627916752158)
    assert besselyzero(2,1).ae(3.3842417671495934727)
    assert besselyzero(1,50).ae(156.29183520147840108)
    assert besselyzero(10,1).ae(12.128927704415439387)
    assert besselyzero(0.5,3).ae(7.8539816339744830962)
    assert besselyzero(2,1,1).ae(5.0025829314460639452)

def test_hankel():
    mp.dps = 15
    assert hankel1(0,0.5).ae(0.93846980724081290423-0.44451873350670655715j)
    assert hankel1(1,0.5).ae(0.2422684576748738864-1.4714723926702430692j)
    assert hankel1(-1,0.5).ae(-0.2422684576748738864+1.4714723926702430692j)
    assert hankel1(1.5,0.5).ae(0.0917016996256513026-2.5214655504213378514j)
    assert hankel1(1.5,3+4j).ae(0.0066806866476728165382-0.0036684231610839127106j)
    assert hankel2(0,0.5).ae(0.93846980724081290423+0.44451873350670655715j)
    assert hankel2(1,0.5).ae(0.2422684576748738864+1.4714723926702430692j)
    assert hankel2(-1,0.5).ae(-0.2422684576748738864-1.4714723926702430692j)
    assert hankel2(1.5,0.5).ae(0.0917016996256513026+2.5214655504213378514j)
    assert hankel2(1.5,3+4j).ae(14.783528526098567526-7.397390270853446512j)

def test_struve():
    mp.dps = 15
    assert struveh(2,3).ae(0.74238666967748318564)
    assert struveh(-2.5,3).ae(0.41271003220971599344)
    assert struvel(2,3).ae(1.7476573277362782744)
    assert struvel(-2.5,3).ae(1.5153394466819651377)

def test_whittaker():
    mp.dps = 15
    assert whitm(2,3,4).ae(49.753745589025246591)
    assert whitw(2,3,4).ae(14.111656223052932215)

def test_kelvin():
    mp.dps = 15
    assert ber(2,3).ae(0.80836846563726819091)
    assert ber(3,4).ae(-0.28262680167242600233)
    assert ber(-3,2).ae(-0.085611448496796363669)
    assert bei(2,3).ae(-0.89102236377977331571)
    assert bei(-3,2).ae(-0.14420994155731828415)
    assert ker(2,3).ae(0.12839126695733458928)
    assert ker(-3,2).ae(-0.29802153400559142783)
    assert ker(0.5,3).ae(-0.085662378535217097524)
    assert kei(2,3).ae(0.036804426134164634000)
    assert kei(-3,2).ae(0.88682069845786731114)
    assert kei(0.5,3).ae(0.013633041571314302948)

def test_hyper_misc():
    mp.dps = 15
    assert hyp0f1(1,0) == 1
    assert hyp1f1(1,2,0) == 1
    assert hyp1f2(1,2,3,0) == 1
    assert hyp2f1(1,2,3,0) == 1
    assert hyp2f2(1,2,3,4,0) == 1
    assert hyp2f3(1,2,3,4,5,0) == 1
    # Degenerate case: 0F0
    assert hyper([],[],0) == 1
    assert hyper([],[],-2).ae(exp(-2))
    # Degenerate case: 1F0
    assert hyper([2],[],1.5) == 4
    #
    assert hyp2f1((1,3),(2,3),(5,6),mpf(27)/32).ae(1.6)
    assert hyp2f1((1,4),(1,2),(3,4),mpf(80)/81).ae(1.8)
    assert hyp2f1((2,3),(1,1),(3,2),(2+j)/3).ae(1.327531603558679093+0.439585080092769253j)
    mp.dps = 25
    v = mpc('1.2282306665029814734863026', '-0.1225033830118305184672133')
    assert hyper([(3,4),2+j,1],[1,5,j/3],mpf(1)/5+j/8).ae(v)
    mp.dps = 15

def test_elliptic_integrals():
    mp.dps = 15
    assert ellipk(0).ae(pi/2)
    assert ellipk(0.5).ae(gamma(0.25)**2/(4*sqrt(pi)))
    assert ellipk(1) == inf
    assert ellipk(1+0j) == inf
    assert ellipk(-1).ae('1.3110287771460599052')
    assert ellipk(-2).ae('1.1714200841467698589')
    assert isinstance(ellipk(-2), mpf)
    assert isinstance(ellipe(-2), mpf)
    assert ellipk(-50).ae('0.47103424540873331679')
    mp.dps = 30
    n1 = +fraction(99999,100000)
    n2 = +fraction(100001,100000)
    mp.dps = 15
    assert ellipk(n1).ae('7.1427724505817781901')
    assert ellipk(n2).ae(mpc('7.1427417367963090109', '-1.5707923998261688019'))
    assert ellipe(n1).ae('1.0000332138990829170')
    v = ellipe(n2)
    assert v.real.ae('0.999966786328145474069137')
    assert (v.imag*10**6).ae('7.853952181727432')
    assert ellipk(2).ae(mpc('1.3110287771460599052', '-1.3110287771460599052'))
    assert ellipk(50).ae(mpc('0.22326753950210985451', '-0.47434723226254522087'))
    assert ellipk(3+4j).ae(mpc('0.91119556380496500866', '0.63133428324134524388'))
    assert ellipk(3-4j).ae(mpc('0.91119556380496500866', '-0.63133428324134524388'))
    assert ellipk(-3+4j).ae(mpc('0.95357894880405122483', '0.23093044503746114444'))
    assert ellipk(-3-4j).ae(mpc('0.95357894880405122483', '-0.23093044503746114444'))
    assert isnan(ellipk(nan))
    assert isnan(ellipe(nan))
    assert ellipk(inf) == 0
    assert isinstance(ellipk(inf), mpc)
    assert ellipk(-inf) == 0
    assert ellipk(1+0j) == inf
    assert ellipe(0).ae(pi/2)
    assert ellipe(0.5).ae(pi**(mpf(3)/2)/gamma(0.25)**2 +gamma(0.25)**2/(8*sqrt(pi)))
    assert ellipe(1) == 1
    assert ellipe(1+0j) == 1
    assert ellipe(inf) == mpc(0,inf)
    assert ellipe(-inf) == inf
    assert ellipe(3+4j).ae(1.4995535209333469543-1.5778790079127582745j)
    assert ellipe(3-4j).ae(1.4995535209333469543+1.5778790079127582745j)
    assert ellipe(-3+4j).ae(2.5804237855343377803-0.8306096791000413778j)
    assert ellipe(-3-4j).ae(2.5804237855343377803+0.8306096791000413778j)
    assert ellipe(2).ae(0.59907011736779610372+0.59907011736779610372j)
    assert ellipe('1e-1000000000').ae(pi/2)
    assert ellipk('1e-1000000000').ae(pi/2)
    assert ellipe(-pi).ae(2.4535865983838923)
    mp.dps = 50
    assert ellipk(1/pi).ae('1.724756270009501831744438120951614673874904182624739673')
    assert ellipe(1/pi).ae('1.437129808135123030101542922290970050337425479058225712')
    assert ellipk(-10*pi).ae('0.5519067523886233967683646782286965823151896970015484512')
    assert ellipe(-10*pi).ae('5.926192483740483797854383268707108012328213431657645509')
    v = ellipk(pi)
    assert v.real.ae('0.973089521698042334840454592642137667227167622330325225')
    assert v.imag.ae('-1.156151296372835303836814390793087600271609993858798016')
    v = ellipe(pi)
    assert v.real.ae('0.4632848917264710404078033487934663562998345622611263332')
    assert v.imag.ae('1.0637961621753130852473300451583414489944099504180510966')
    mp.dps = 15

def test_exp_integrals():
    mp.dps = 15
    x = +e
    z = e + sqrt(3)*j
    assert ei(x).ae(8.21168165538361560)
    assert li(x).ae(1.89511781635593676)
    assert si(x).ae(1.82104026914756705)
    assert ci(x).ae(0.213958001340379779)
    assert shi(x).ae(4.11520706247846193)
    assert chi(x).ae(4.09647459290515367)
    assert fresnels(x).ae(0.437189718149787643)
    assert fresnelc(x).ae(0.401777759590243012)
    assert airyai(x).ae(0.0108502401568586681)
    assert airybi(x).ae(8.98245748585468627)
    assert ei(z).ae(3.72597969491314951 + 7.34213212314224421j)
    assert li(z).ae(2.28662658112562502 + 1.50427225297269364j)
    assert si(z).ae(2.48122029237669054 + 0.12684703275254834j)
    assert ci(z).ae(0.169255590269456633 - 0.892020751420780353j)
    assert shi(z).ae(1.85810366559344468 + 3.66435842914920263j)
    assert chi(z).ae(1.86787602931970484 + 3.67777369399304159j)
    assert fresnels(z/3).ae(0.034534397197008182 + 0.754859844188218737j)
    assert fresnelc(z/3).ae(1.261581645990027372 + 0.417949198775061893j)
    assert airyai(z).ae(-0.0162552579839056062 - 0.0018045715700210556j)
    assert airybi(z).ae(-4.98856113282883371 + 2.08558537872180623j)
    assert li(0) == 0.0
    assert li(1) == -inf
    assert li(inf) == inf
    assert isinstance(li(0.7), mpf)
    assert si(inf).ae(pi/2)
    assert si(-inf).ae(-pi/2)
    assert ci(inf) == 0
    assert ci(0) == -inf
    assert isinstance(ei(-0.7), mpf)
    assert airyai(inf) == 0
    assert airybi(inf) == inf
    assert airyai(-inf) == 0
    assert airybi(-inf) == 0
    assert fresnels(inf) == 0.5
    assert fresnelc(inf) == 0.5
    assert fresnels(-inf) == -0.5
    assert fresnelc(-inf) == -0.5
    assert shi(0) == 0
    assert shi(inf) == inf
    assert shi(-inf) == -inf
    assert chi(0) == -inf
    assert chi(inf) == inf

def test_ei():
    mp.dps = 15
    assert ei(0) == -inf
    assert ei(inf) == inf
    assert ei(-inf) == -0.0
    assert ei(20+70j).ae(6.1041351911152984397e6 - 2.7324109310519928872e6j)
    # tests for the asymptotic expansion
    # values checked with Mathematica ExpIntegralEi
    mp.dps = 50
    r = ei(20000)
    s = '3.8781962825045010930273870085501819470698476975019e+8681'
    assert str(r) == s
    r = ei(-200)
    s = '-6.8852261063076355977108174824557929738368086933303e-90'
    assert str(r) == s
    r =ei(20000 + 10*j)
    sre = '-3.255138234032069402493850638874410725961401274106e+8681'
    sim = '-2.1081929993474403520785942429469187647767369645423e+8681'
    assert str(r.real) == sre and str(r.imag) == sim
    mp.dps = 15
    # More asymptotic expansions
    assert chi(-10**6+100j).ae('1.3077239389562548386e+434288 + 7.6808956999707408158e+434287j')
    assert shi(-10**6+100j).ae('-1.3077239389562548386e+434288 - 7.6808956999707408158e+434287j')
    mp.dps = 15
    assert ei(10j).ae(-0.0454564330044553726+3.2291439210137706686j)
    assert ei(100j).ae(-0.0051488251426104921+3.1330217936839529126j)
    u = ei(fmul(10**20, j, exact=True))
    assert u.real.ae(-6.4525128526578084421345e-21, abs_eps=0, rel_eps=8*eps)
    assert u.imag.ae(pi)
    assert ei(-10j).ae(-0.0454564330044553726-3.2291439210137706686j)
    assert ei(-100j).ae(-0.0051488251426104921-3.1330217936839529126j)
    u = ei(fmul(-10**20, j, exact=True))
    assert u.real.ae(-6.4525128526578084421345e-21, abs_eps=0, rel_eps=8*eps)
    assert u.imag.ae(-pi)
    assert ei(10+10j).ae(-1576.1504265768517448+436.9192317011328140j)
    u = ei(-10+10j)
    assert u.real.ae(7.6698978415553488362543e-7, abs_eps=0, rel_eps=8*eps)
    assert u.imag.ae(3.141595611735621062025)

def test_e1():
    mp.dps = 15
    assert e1(0) == inf
    assert e1(inf) == 0
    assert e1(-inf) == mpc(-inf, -pi)
    assert e1(10j).ae(0.045456433004455372635 + 0.087551267423977430100j)
    assert e1(100j).ae(0.0051488251426104921444 - 0.0085708599058403258790j)
    assert e1(fmul(10**20, j, exact=True)).ae(6.4525128526578084421e-21 - 7.6397040444172830039e-21j, abs_eps=0, rel_eps=8*eps)
    assert e1(-10j).ae(0.045456433004455372635 - 0.087551267423977430100j)
    assert e1(-100j).ae(0.0051488251426104921444 + 0.0085708599058403258790j)
    assert e1(fmul(-10**20, j, exact=True)).ae(6.4525128526578084421e-21 + 7.6397040444172830039e-21j, abs_eps=0, rel_eps=8*eps)

def test_expint():
    mp.dps = 15
    assert expint(0,0) == inf
    assert expint(0,1).ae(1/e)
    assert expint(0,1.5).ae(2/exp(1.5)/3)
    assert expint(1,1).ae(-ei(-1))
    assert expint(2,0).ae(1)
    assert expint(3,0).ae(1/2.)
    assert expint(4,0).ae(1/3.)
    assert expint(-2, 0.5).ae(26/sqrt(e))
    assert expint(-1,-1) == 0
    assert expint(-2,-1).ae(-e)
    assert expint(5.5, 0).ae(2/9.)
    assert expint(2.00000001,0).ae(100000000./100000001)
    assert expint(2+3j,4-j).ae(0.0023461179581675065414+0.0020395540604713669262j)
    assert expint('1.01', '1e-1000').ae(99.9999999899412802)
    assert expint('1.000000000001', 3.5).ae(0.00697013985754701819446)
    assert expint(2,3).ae(3*ei(-3)+exp(-3))
    assert (expint(10,20)*10**10).ae(0.694439055541231353)
    assert expint(3,inf) == 0
    assert expint(3.2,inf) == 0
    assert expint(3.2+2j,inf) == 0
    assert expint(1,3j).ae(-0.11962978600800032763 + 0.27785620120457163717j)
    assert expint(1,3).ae(0.013048381094197037413)
    assert expint(1,-3).ae(-ei(3)-pi*j)
    #assert expint(3) == expint(1,3)
    assert expint(1,-20).ae(-25615652.66405658882 - 3.1415926535897932385j)
    assert expint(1000000,0).ae(1./999999)
    assert expint(0,2+3j).ae(-0.025019798357114678171 + 0.027980439405104419040j)
    assert expint(-1,2+3j).ae(-0.022411973626262070419 + 0.038058922011377716932j)
    assert expint(-1.5,0) == inf

def test_trig_integrals():
    mp.dps = 30
    assert si(mpf(1)/1000000).ae('0.000000999999999999944444444444446111')
    assert ci(mpf(1)/1000000).ae('-13.2382948930629912435014366276')
    assert si(10**10).ae('1.5707963267075846569685111517747537')
    assert ci(10**10).ae('-4.87506025174822653785729773959e-11')
    assert si(10**100).ae(pi/2)
    assert (ci(10**100)*10**100).ae('-0.372376123661276688262086695553')
    assert si(-3) == -si(3)
    assert ci(-3).ae(ci(3) + pi*j)
    # Test complex structure
    mp.dps = 15
    assert mp.ci(50).ae(-0.0056283863241163054402)
    assert mp.ci(50+2j).ae(-0.018378282946133067149+0.070352808023688336193j)
    assert mp.ci(20j).ae(1.28078263320282943611e7+1.5707963267949j)
    assert mp.ci(-2+20j).ae(-4.050116856873293505e6+1.207476188206989909e7j)
    assert mp.ci(-50+2j).ae(-0.0183782829461330671+3.0712398455661049023j)
    assert mp.ci(-50).ae(-0.0056283863241163054+3.1415926535897932385j)
    assert mp.ci(-50-2j).ae(-0.0183782829461330671-3.0712398455661049023j)
    assert mp.ci(-2-20j).ae(-4.050116856873293505e6-1.207476188206989909e7j)
    assert mp.ci(-20j).ae(1.28078263320282943611e7-1.5707963267949j)
    assert mp.ci(50-2j).ae(-0.018378282946133067149-0.070352808023688336193j)
    assert mp.si(50).ae(1.5516170724859358947)
    assert mp.si(50+2j).ae(1.497884414277228461-0.017515007378437448j)
    assert mp.si(20j).ae(1.2807826332028294459e7j)
    assert mp.si(-2+20j).ae(-1.20747603112735722103e7-4.050116856873293554e6j)
    assert mp.si(-50+2j).ae(-1.497884414277228461-0.017515007378437448j)
    assert mp.si(-50).ae(-1.5516170724859358947)
    assert mp.si(-50-2j).ae(-1.497884414277228461+0.017515007378437448j)
    assert mp.si(-2-20j).ae(-1.20747603112735722103e7+4.050116856873293554e6j)
    assert mp.si(-20j).ae(-1.2807826332028294459e7j)
    assert mp.si(50-2j).ae(1.497884414277228461+0.017515007378437448j)
    assert mp.chi(50j).ae(-0.0056283863241163054+1.5707963267948966192j)
    assert mp.chi(-2+50j).ae(-0.0183782829461330671+1.6411491348185849554j)
    assert mp.chi(-20).ae(1.28078263320282943611e7+3.1415926535898j)
    assert mp.chi(-20-2j).ae(-4.050116856873293505e6+1.20747571696809187053e7j)
    assert mp.chi(-2-50j).ae(-0.0183782829461330671-1.6411491348185849554j)
    assert mp.chi(-50j).ae(-0.0056283863241163054-1.5707963267948966192j)
    assert mp.chi(2-50j).ae(-0.0183782829461330671-1.500443518771208283j)
    assert mp.chi(20-2j).ae(-4.050116856873293505e6-1.20747603112735722951e7j)
    assert mp.chi(20).ae(1.2807826332028294361e7)
    assert mp.chi(2+50j).ae(-0.0183782829461330671+1.500443518771208283j)
    assert mp.shi(50j).ae(1.5516170724859358947j)
    assert mp.shi(-2+50j).ae(0.017515007378437448+1.497884414277228461j)
    assert mp.shi(-20).ae(-1.2807826332028294459e7)
    assert mp.shi(-20-2j).ae(4.050116856873293554e6-1.20747603112735722103e7j)
    assert mp.shi(-2-50j).ae(0.017515007378437448-1.497884414277228461j)
    assert mp.shi(-50j).ae(-1.5516170724859358947j)
    assert mp.shi(2-50j).ae(-0.017515007378437448-1.497884414277228461j)
    assert mp.shi(20-2j).ae(-4.050116856873293554e6-1.20747603112735722103e7j)
    assert mp.shi(20).ae(1.2807826332028294459e7)
    assert mp.shi(2+50j).ae(-0.017515007378437448+1.497884414277228461j)
    def ae(x,y,tol=1e-12):
        return abs(x-y) <= abs(y)*tol
    assert fp.ci(fp.inf) == 0
    assert ae(fp.ci(fp.ninf), fp.pi*1j)
    assert ae(fp.si(fp.inf), fp.pi/2)
    assert ae(fp.si(fp.ninf), -fp.pi/2)
    assert fp.si(0) == 0
    assert ae(fp.ci(50), -0.0056283863241163054402)
    assert ae(fp.ci(50+2j), -0.018378282946133067149+0.070352808023688336193j)
    assert ae(fp.ci(20j), 1.28078263320282943611e7+1.5707963267949j)
    assert ae(fp.ci(-2+20j), -4.050116856873293505e6+1.207476188206989909e7j)
    assert ae(fp.ci(-50+2j), -0.0183782829461330671+3.0712398455661049023j)
    assert ae(fp.ci(-50), -0.0056283863241163054+3.1415926535897932385j)
    assert ae(fp.ci(-50-2j), -0.0183782829461330671-3.0712398455661049023j)
    assert ae(fp.ci(-2-20j), -4.050116856873293505e6-1.207476188206989909e7j)
    assert ae(fp.ci(-20j), 1.28078263320282943611e7-1.5707963267949j)
    assert ae(fp.ci(50-2j), -0.018378282946133067149-0.070352808023688336193j)
    assert ae(fp.si(50), 1.5516170724859358947)
    assert ae(fp.si(50+2j), 1.497884414277228461-0.017515007378437448j)
    assert ae(fp.si(20j), 1.2807826332028294459e7j)
    assert ae(fp.si(-2+20j), -1.20747603112735722103e7-4.050116856873293554e6j)
    assert ae(fp.si(-50+2j), -1.497884414277228461-0.017515007378437448j)
    assert ae(fp.si(-50), -1.5516170724859358947)
    assert ae(fp.si(-50-2j), -1.497884414277228461+0.017515007378437448j)
    assert ae(fp.si(-2-20j), -1.20747603112735722103e7+4.050116856873293554e6j)
    assert ae(fp.si(-20j), -1.2807826332028294459e7j)
    assert ae(fp.si(50-2j), 1.497884414277228461+0.017515007378437448j)
    assert ae(fp.chi(50j), -0.0056283863241163054+1.5707963267948966192j)
    assert ae(fp.chi(-2+50j), -0.0183782829461330671+1.6411491348185849554j)
    assert ae(fp.chi(-20), 1.28078263320282943611e7+3.1415926535898j)
    assert ae(fp.chi(-20-2j), -4.050116856873293505e6+1.20747571696809187053e7j)
    assert ae(fp.chi(-2-50j), -0.0183782829461330671-1.6411491348185849554j)
    assert ae(fp.chi(-50j), -0.0056283863241163054-1.5707963267948966192j)
    assert ae(fp.chi(2-50j), -0.0183782829461330671-1.500443518771208283j)
    assert ae(fp.chi(20-2j), -4.050116856873293505e6-1.20747603112735722951e7j)
    assert ae(fp.chi(20), 1.2807826332028294361e7)
    assert ae(fp.chi(2+50j), -0.0183782829461330671+1.500443518771208283j)
    assert ae(fp.shi(50j), 1.5516170724859358947j)
    assert ae(fp.shi(-2+50j), 0.017515007378437448+1.497884414277228461j)
    assert ae(fp.shi(-20), -1.2807826332028294459e7)
    assert ae(fp.shi(-20-2j), 4.050116856873293554e6-1.20747603112735722103e7j)
    assert ae(fp.shi(-2-50j), 0.017515007378437448-1.497884414277228461j)
    assert ae(fp.shi(-50j), -1.5516170724859358947j)
    assert ae(fp.shi(2-50j), -0.017515007378437448-1.497884414277228461j)
    assert ae(fp.shi(20-2j), -4.050116856873293554e6-1.20747603112735722103e7j)
    assert ae(fp.shi(20), 1.2807826332028294459e7)
    assert ae(fp.shi(2+50j), -0.017515007378437448+1.497884414277228461j)

def test_airy():
    mp.dps = 15
    assert (airyai(10)*10**10).ae(1.1047532552898687)
    assert (airybi(10)/10**9).ae(0.45564115354822515)
    assert (airyai(1000)*10**9158).ae(9.306933063179556004)
    assert (airybi(1000)/10**9154).ae(5.4077118391949465477)
    assert airyai(-1000).ae(0.055971895773019918842)
    assert airybi(-1000).ae(-0.083264574117080633012)
    assert (airyai(100+100j)*10**188).ae(2.9099582462207032076 + 2.353013591706178756j)
    assert (airybi(100+100j)/10**185).ae(1.7086751714463652039 - 3.1416590020830804578j)

def test_hyper_0f1():
    mp.dps = 15
    v = 8.63911136507950465
    assert hyper([],[(1,3)],1.5).ae(v)
    assert hyper([],[1/3.],1.5).ae(v)
    assert hyp0f1(1/3.,1.5).ae(v)
    assert hyp0f1((1,3),1.5).ae(v)
    # Asymptotic expansion
    assert hyp0f1(3,1e9).ae('4.9679055380347771271e+27455')
    assert hyp0f1(3,1e9j).ae('-2.1222788784457702157e+19410 + 5.0840597555401854116e+19410j')

def test_hyper_1f1():
    mp.dps = 15
    v = 1.2917526488617656673
    assert hyper([(1,2)],[(3,2)],0.7).ae(v)
    assert hyper([(1,2)],[(3,2)],0.7+0j).ae(v)
    assert hyper([0.5],[(3,2)],0.7).ae(v)
    assert hyper([0.5],[1.5],0.7).ae(v)
    assert hyper([0.5],[(3,2)],0.7+0j).ae(v)
    assert hyper([0.5],[1.5],0.7+0j).ae(v)
    assert hyper([(1,2)],[1.5+0j],0.7).ae(v)
    assert hyper([0.5+0j],[1.5],0.7).ae(v)
    assert hyper([0.5+0j],[1.5+0j],0.7+0j).ae(v)
    assert hyp1f1(0.5,1.5,0.7).ae(v)
    assert hyp1f1((1,2),1.5,0.7).ae(v)
    # Asymptotic expansion
    assert hyp1f1(2,3,1e10).ae('2.1555012157015796988e+4342944809')
    assert (hyp1f1(2,3,1e10j)*10**10).ae(-0.97501205020039745852 - 1.7462392454512132074j)
    # Shouldn't use asymptotic expansion
    assert hyp1f1(-2, 1, 10000).ae(49980001)
    # Bug
    assert hyp1f1(1j,fraction(1,3),0.415-69.739j).ae(25.857588206024346592 + 15.738060264515292063j)

def test_hyper_2f1():
    mp.dps = 15
    v = 1.0652207633823291032
    assert hyper([(1,2), (3,4)], [2], 0.3).ae(v)
    assert hyper([(1,2), 0.75], [2], 0.3).ae(v)
    assert hyper([0.5, 0.75], [2.0], 0.3).ae(v)
    assert hyper([0.5, 0.75], [2.0], 0.3+0j).ae(v)
    assert hyper([0.5+0j, (3,4)], [2.0], 0.3+0j).ae(v)
    assert hyper([0.5+0j, (3,4)], [2.0], 0.3).ae(v)
    assert hyper([0.5, (3,4)], [2.0+0j], 0.3).ae(v)
    assert hyper([0.5+0j, 0.75+0j], [2.0+0j], 0.3+0j).ae(v)
    v = 1.09234681096223231717 + 0.18104859169479360380j
    assert hyper([(1,2),0.75+j], [2], 0.5).ae(v)
    assert hyper([0.5,0.75+j], [2.0], 0.5).ae(v)
    assert hyper([0.5,0.75+j], [2.0], 0.5+0j).ae(v)
    assert hyper([0.5,0.75+j], [2.0+0j], 0.5+0j).ae(v)
    v = 0.9625 - 0.125j
    assert hyper([(3,2),-1],[4], 0.1+j/3).ae(v)
    assert hyper([1.5,-1.0],[4], 0.1+j/3).ae(v)
    assert hyper([1.5,-1.0],[4+0j], 0.1+j/3).ae(v)
    assert hyper([1.5+0j,-1.0+0j],[4+0j], 0.1+j/3).ae(v)
    v = 1.02111069501693445001 - 0.50402252613466859521j
    assert hyper([(2,10),(3,10)],[(4,10)],1.5).ae(v)
    assert hyper([0.2,(3,10)],[0.4+0j],1.5).ae(v)
    assert hyper([0.2,(3,10)],[0.4+0j],1.5+0j).ae(v)
    v = 0.76922501362865848528 + 0.32640579593235886194j
    assert hyper([(2,10),(3,10)],[(4,10)],4+2j).ae(v)
    assert hyper([0.2,(3,10)],[0.4+0j],4+2j).ae(v)
    assert hyper([0.2,(3,10)],[(4,10)],4+2j).ae(v)

def test_hyper_2f1_hard():
    mp.dps = 15
    # Singular cases
    assert hyp2f1(2,-1,-1,3).ae(7)
    assert hyp2f1(2,-1,-1,3,eliminate_all=True).ae(0.25)
    assert hyp2f1(2,-2,-2,3).ae(34)
    assert hyp2f1(2,-2,-2,3,eliminate_all=True).ae(0.25)
    assert hyp2f1(2,-2,-3,3) == 14
    assert hyp2f1(2,-3,-2,3) == inf
    assert hyp2f1(2,-1.5,-1.5,3) == 0.25
    assert hyp2f1(1,2,3,0) == 1
    assert hyp2f1(0,1,0,0) == 1
    assert hyp2f1(0,0,0,0) == 1
    assert isnan(hyp2f1(1,1,0,0))
    assert hyp2f1(2,-1,-5, 0.25+0.25j).ae(1.1+0.1j)
    assert hyp2f1(2,-5,-5, 0.25+0.25j, eliminate=False).ae(163./128 + 125./128*j)
    assert hyp2f1(0.7235, -1, -5, 0.3).ae(1.04341)
    assert hyp2f1(0.7235, -5, -5, 0.3, eliminate=False).ae(1.2939225017815903812)
    assert hyp2f1(-1,-2,4,1) == 1.5
    assert hyp2f1(1,2,-3,1) == inf
    assert hyp2f1(-2,-2,1,1) == 6
    assert hyp2f1(1,-2,-4,1).ae(5./3)
    assert hyp2f1(0,-6,-4,1) == 1
    assert hyp2f1(0,-3,-4,1) == 1
    assert hyp2f1(0,0,0,1) == 1
    assert hyp2f1(1,0,0,1,eliminate=False) == 1
    assert hyp2f1(1,1,0,1) == inf
    assert hyp2f1(1,-6,-4,1) == inf
    assert hyp2f1(-7.2,-0.5,-4.5,1) == 0
    assert hyp2f1(-7.2,-1,-2,1).ae(-2.6)
    assert hyp2f1(1,-0.5,-4.5, 1) == inf
    assert hyp2f1(1,0.5,-4.5, 1) == -inf
    # Check evaluation on / close to unit circle
    z = exp(j*pi/3)
    w = (nthroot(2,3)+1)*exp(j*pi/12)/nthroot(3,4)**3
    assert hyp2f1('1/2','1/6','1/3', z).ae(w)
    assert hyp2f1('1/2','1/6','1/3', z.conjugate()).ae(w.conjugate())
    assert hyp2f1(0.25, (1,3), 2, '0.999').ae(1.06826449496030635)
    assert hyp2f1(0.25, (1,3), 2, '1.001').ae(1.06867299254830309446-0.00001446586793975874j)
    assert hyp2f1(0.25, (1,3), 2, -1).ae(0.96656584492524351673)
    assert hyp2f1(0.25, (1,3), 2, j).ae(0.99041766248982072266+0.03777135604180735522j)
    assert hyp2f1(2,3,5,'0.99').ae(27.699347904322690602)
    assert hyp2f1((3,2),-0.5,3,'0.99').ae(0.68403036843911661388)
    assert hyp2f1(2,3,5,1j).ae(0.37290667145974386127+0.59210004902748285917j)
    assert fsum([hyp2f1((7,10),(2,3),(-1,2), 0.95*exp(j*k)) for k in range(1,15)]).ae(52.851400204289452922+6.244285013912953225j)
    assert fsum([hyp2f1((7,10),(2,3),(-1,2), 1.05*exp(j*k)) for k in range(1,15)]).ae(54.506013786220655330-3.000118813413217097j)
    assert fsum([hyp2f1((7,10),(2,3),(-1,2), exp(j*k)) for k in range(1,15)]).ae(55.792077935955314887+1.731986485778500241j)
    assert hyp2f1(2,2.5,-3.25,0.999).ae(218373932801217082543180041.33)
    # Branches
    assert hyp2f1(1,1,2,1.01).ae(4.5595744415723676911-3.1104877758314784539j)
    assert hyp2f1(1,1,2,1.01+0.1j).ae(2.4149427480552782484+1.4148224796836938829j)
    assert hyp2f1(1,1,2,3+4j).ae(0.14576709331407297807+0.48379185417980360773j)
    assert hyp2f1(1,1,2,4).ae(-0.27465307216702742285 - 0.78539816339744830962j)
    assert hyp2f1(1,1,2,-4).ae(0.40235947810852509365)
    # Other:
    # Cancellation with a large parameter involved (bug reported on sage-devel)
    assert hyp2f1(112, (51,10), (-9,10), -0.99999).ae(-1.6241361047970862961e-24, abs_eps=0, rel_eps=eps*16)

def test_hyper_3f2_etc():
    assert hyper([1,2,3],[1.5,8],-1).ae(0.67108992351533333030)
    assert hyper([1,2,3,4],[5,6,7], -1).ae(0.90232988035425506008)
    assert hyper([1,2,3],[1.25,5], 1).ae(28.924181329701905701)
    assert hyper([1,2,3,4],[5,6,7],5).ae(1.5192307344006649499-1.1529845225075537461j)
    assert hyper([1,2,3,4,5],[6,7,8,9],-1).ae(0.96288759462882357253)
    assert hyper([1,2,3,4,5],[6,7,8,9],1).ae(1.0428697385885855841)
    assert hyper([1,2,3,4,5],[6,7,8,9],5).ae(1.33980653631074769423-0.07143405251029226699j)
    assert hyper([1,2.79,3.08,4.37],[5.2,6.1,7.3],5).ae(1.0996321464692607231-1.7748052293979985001j)
    assert hyper([1,1,1],[1,2],1) == inf
    assert hyper([1,1,1],[2,(101,100)],1).ae(100.01621213528313220)
    # slow -- covered by doctests
    #assert hyper([1,1,1],[2,3],0.9999).ae(1.2897972005319693905)

def test_hyper_u():
    mp.dps = 15
    assert hyperu(2,-3,0).ae(0.05)
    assert hyperu(2,-3.5,0).ae(4./99)
    assert hyperu(2,0,0) == 0.5
    assert hyperu(-5,1,0) == -120
    assert hyperu(-5,2,0) == inf
    assert hyperu(-5,-2,0) == 0
    assert hyperu(7,7,3).ae(0.00014681269365593503986)  #exp(3)*gammainc(-6,3)
    assert hyperu(2,-3,4).ae(0.011836478100271995559)
    assert hyperu(3,4,5).ae(1./125)
    assert hyperu(2,3,0.0625) == 256
    assert hyperu(-1,2,0.25+0.5j) == -1.75+0.5j
    assert hyperu(0.5,1.5,7.25).ae(2/sqrt(29))
    assert hyperu(2,6,pi).ae(0.55804439825913399130)
    assert (hyperu((3,2),8,100+201j)*10**4).ae(-0.3797318333856738798 - 2.9974928453561707782j)
    assert (hyperu((5,2),(-1,2),-5000)*10**10).ae(-5.6681877926881664678j)
    # XXX: fails because of undetected cancellation in low level series code
    # Alternatively: could use asymptotic series here, if convergence test
    # tweaked back to recognize this one
    #assert (hyperu((5,2),(-1,2),-500)*10**7).ae(-1.82526906001593252847j)

def test_hyper_2f0():
    mp.dps = 15
    assert hyper([1,2],[],3) == hyp2f0(1,2,3)
    assert hyp2f0(2,3,7).ae(0.0116108068639728714668 - 0.0073727413865865802130j)
    assert hyp2f0(2,3,0) == 1
    assert hyp2f0(0,0,0) == 1
    assert hyp2f0(-1,-1,1).ae(2)
    assert hyp2f0(-4,1,1.5).ae(62.5)
    assert hyp2f0(-4,1,50).ae(147029801)
    assert hyp2f0(-4,1,0.0001).ae(0.99960011997600240000)
    assert hyp2f0(0.5,0.25,0.001).ae(1.0001251174078538115)
    assert hyp2f0(0.5,0.25,3+4j).ae(0.85548875824755163518 + 0.21636041283392292973j)
    # Important: cancellation check
    assert hyp2f0((1,6),(5,6),-0.02371708245126284498).ae(0.996785723120804309)
    # Should be exact; polynomial case
    assert hyp2f0(-2,1,0.5+0.5j,zeroprec=200) == 0
    assert hyp2f0(1,-2,0.5+0.5j,zeroprec=200) == 0
    # There used to be a bug in thresholds that made one of the following hang
    for d in [15, 50, 80]:
        mp.dps = d
        assert hyp2f0(1.5, 0.5, 0.009).ae('1.006867007239309717945323585695344927904000945829843527398772456281301440034218290443367270629519483 + 1.238277162240704919639384945859073461954721356062919829456053965502443570466701567100438048602352623e-46j')

def test_hyper_1f2():
    mp.dps = 15
    assert hyper([1],[2,3],4) == hyp1f2(1,2,3,4)
    a1,b1,b2 = (1,10),(2,3),1./16
    assert hyp1f2(a1,b1,b2,10).ae(298.7482725554557568)
    assert hyp1f2(a1,b1,b2,100).ae(224128961.48602947604)
    assert hyp1f2(a1,b1,b2,1000).ae(1.1669528298622675109e+27)
    assert hyp1f2(a1,b1,b2,10000).ae(2.4780514622487212192e+86)
    assert hyp1f2(a1,b1,b2,100000).ae(1.3885391458871523997e+274)
    assert hyp1f2(a1,b1,b2,1000000).ae('9.8851796978960318255e+867')
    assert hyp1f2(a1,b1,b2,10**7).ae('1.1505659189516303646e+2746')
    assert hyp1f2(a1,b1,b2,10**8).ae('1.4672005404314334081e+8685')
    assert hyp1f2(a1,b1,b2,10**20).ae('3.6888217332150976493e+8685889636')
    assert hyp1f2(a1,b1,b2,10*j).ae(-16.163252524618572878 - 44.321567896480184312j)
    assert hyp1f2(a1,b1,b2,100*j).ae(61938.155294517848171 + 637349.45215942348739j)
    assert hyp1f2(a1,b1,b2,1000*j).ae(8455057657257695958.7 + 6261969266997571510.6j)
    assert hyp1f2(a1,b1,b2,10000*j).ae(-8.9771211184008593089e+60 + 4.6550528111731631456e+59j)
    assert hyp1f2(a1,b1,b2,100000*j).ae(2.6398091437239324225e+193 + 4.1658080666870618332e+193j)
    assert hyp1f2(a1,b1,b2,1000000*j).ae('3.5999042951925965458e+613 + 1.5026014707128947992e+613j')
    assert hyp1f2(a1,b1,b2,10**7*j).ae('-8.3208715051623234801e+1939 - 3.6752883490851869429e+1941j')
    assert hyp1f2(a1,b1,b2,10**8*j).ae('2.0724195707891484454e+6140 - 1.3276619482724266387e+6141j')
    assert hyp1f2(a1,b1,b2,10**20*j).ae('-1.1734497974795488504e+6141851462 + 1.1498106965385471542e+6141851462j')

def test_hyper_2f3():
    mp.dps = 15
    assert hyper([1,2],[3,4,5],6) == hyp2f3(1,2,3,4,5,6)
    a1,a2,b1,b2,b3 = (1,10),(2,3),(3,10), 2, 1./16
    # Check asymptotic expansion
    assert hyp2f3(a1,a2,b1,b2,b3,10).ae(128.98207160698659976)
    assert hyp2f3(a1,a2,b1,b2,b3,1000).ae(6.6309632883131273141e25)
    assert hyp2f3(a1,a2,b1,b2,b3,10000).ae(4.6863639362713340539e84)
    assert hyp2f3(a1,a2,b1,b2,b3,100000).ae(8.6632451236103084119e271)
    assert hyp2f3(a1,a2,b1,b2,b3,10**6).ae('2.0291718386574980641e865')
    assert hyp2f3(a1,a2,b1,b2,b3,10**7).ae('7.7639836665710030977e2742')
    assert hyp2f3(a1,a2,b1,b2,b3,10**8).ae('3.2537462584071268759e8681')
    assert hyp2f3(a1,a2,b1,b2,b3,10**20).ae('1.2966030542911614163e+8685889627')
    assert hyp2f3(a1,a2,b1,b2,b3,10*j).ae(-18.551602185587547854 - 13.348031097874113552j)
    assert hyp2f3(a1,a2,b1,b2,b3,100*j).ae(78634.359124504488695 + 74459.535945281973996j)
    assert hyp2f3(a1,a2,b1,b2,b3,1000*j).ae(597682550276527901.59 - 65136194809352613.078j)
    assert hyp2f3(a1,a2,b1,b2,b3,10000*j).ae(-1.1779696326238582496e+59 + 1.2297607505213133872e+59j)
    assert hyp2f3(a1,a2,b1,b2,b3,100000*j).ae(2.9844228969804380301e+191 + 7.5587163231490273296e+190j)
    assert hyp2f3(a1,a2,b1,b2,b3,1000000*j).ae('7.4859161049322370311e+610 - 2.8467477015940090189e+610j')
    assert hyp2f3(a1,a2,b1,b2,b3,10**7*j).ae('-1.7477645579418800826e+1938 - 1.7606522995808116405e+1938j')
    assert hyp2f3(a1,a2,b1,b2,b3,10**8*j).ae('-1.6932731942958401784e+6137 - 2.4521909113114629368e+6137j')
    assert hyp2f3(a1,a2,b1,b2,b3,10**20*j).ae('-2.0988815677627225449e+6141851451 + 5.7708223542739208681e+6141851452j')

def test_hyper_2f2():
    mp.dps = 15
    assert hyper([1,2],[3,4],5) == hyp2f2(1,2,3,4,5)
    a1,a2,b1,b2 = (3,10),4,(1,2),1./16
    assert hyp2f2(a1,a2,b1,b2,10).ae(448225936.3377556696)
    assert hyp2f2(a1,a2,b1,b2,10000).ae('1.2012553712966636711e+4358')
    assert hyp2f2(a1,a2,b1,b2,-20000).ae(-0.04182343755661214626)
    assert hyp2f2(a1,a2,b1,b2,10**20).ae('1.1148680024303263661e+43429448190325182840')

def test_orthpoly():
    mp.dps = 15
    assert jacobi(-4,2,3,0.7).ae(22800./4913)
    assert jacobi(3,2,4,5.5) == 4133.125
    assert jacobi(1.5,5/6.,4,0).ae(-1.0851951434075508417)
    assert jacobi(-2, 1, 2, 4).ae(-0.16)
    assert jacobi(2, -1, 2.5, 4).ae(34.59375)
    #assert jacobi(2, -1, 2, 4) == 28.5
    assert legendre(5, 7) == 129367
    assert legendre(0.5,0).ae(0.53935260118837935667)
    assert legendre(-1,-1) == 1
    assert legendre(0,-1) == 1
    assert legendre(0, 1) == 1
    assert legendre(1, -1) == -1
    assert legendre(7, 1) == 1
    assert legendre(7, -1) == -1
    assert legendre(8,1.5).ae(15457523./32768)
    assert legendre(j,-j).ae(2.4448182735671431011 + 0.6928881737669934843j)
    assert chebyu(5,1) == 6
    assert chebyt(3,2) == 26
    assert legendre(3.5,-1) == inf
    assert legendre(4.5,-1) == -inf
    assert legendre(3.5+1j,-1) == mpc(inf,inf)
    assert legendre(4.5+1j,-1) == mpc(-inf,-inf)
    assert laguerre(4, -2, 3).ae(-1.125)
    assert laguerre(3, 1+j, 0.5).ae(0.2291666666666666667 + 2.5416666666666666667j)

def test_hermite():
    mp.dps = 15
    assert hermite(-2, 0).ae(0.5)
    assert hermite(-1, 0).ae(0.88622692545275801365)
    assert hermite(0, 0).ae(1)
    assert hermite(1, 0) == 0
    assert hermite(2, 0).ae(-2)
    assert hermite(0, 2).ae(1)
    assert hermite(1, 2).ae(4)
    assert hermite(1, -2).ae(-4)
    assert hermite(2, -2).ae(14)
    assert hermite(0.5, 0).ae(0.69136733903629335053)
    assert hermite(9, 0) == 0
    assert hermite(4,4).ae(3340)
    assert hermite(3,4).ae(464)
    assert hermite(-4,4).ae(0.00018623860287512396181)
    assert hermite(-3,4).ae(0.0016540169879668766270)
    assert hermite(9, 2.5j).ae(13638725j)
    assert hermite(9, -2.5j).ae(-13638725j)
    assert hermite(9, 100).ae(511078883759363024000)
    assert hermite(9, -100).ae(-511078883759363024000)
    assert hermite(9, 100j).ae(512922083920643024000j)
    assert hermite(9, -100j).ae(-512922083920643024000j)
    assert hermite(-9.5, 2.5j).ae(-2.9004951258126778174e-6 + 1.7601372934039951100e-6j)
    assert hermite(-9.5, -2.5j).ae(-2.9004951258126778174e-6 - 1.7601372934039951100e-6j)
    assert hermite(-9.5, 100).ae(1.3776300722767084162e-22, abs_eps=0, rel_eps=eps)
    assert hermite(-9.5, -100).ae('1.3106082028470671626e4355')
    assert hermite(-9.5, 100j).ae(-9.7900218581864768430e-23 - 9.7900218581864768430e-23j, abs_eps=0, rel_eps=eps)
    assert hermite(-9.5, -100j).ae(-9.7900218581864768430e-23 + 9.7900218581864768430e-23j, abs_eps=0, rel_eps=eps)
    assert hermite(2+3j, -1-j).ae(851.3677063883687676 - 1496.4373467871007997j)

def test_gegenbauer():
    mp.dps = 15
    assert gegenbauer(1,2,3).ae(12)
    assert gegenbauer(2,3,4).ae(381)
    assert gegenbauer(0,0,0) == 0
    assert gegenbauer(2,-1,3) == 0
    assert gegenbauer(-7, 0.5, 3).ae(8989)
    assert gegenbauer(1, -0.5, 3).ae(-3)
    assert gegenbauer(1, -1.5, 3).ae(-9)
    assert gegenbauer(1, -0.5, 3).ae(-3)
    assert gegenbauer(-0.5, -0.5, 3).ae(-2.6383553159023906245)
    assert gegenbauer(2+3j, 1-j, 3+4j).ae(14.880536623203696780 + 20.022029711598032898j)
    #assert gegenbauer(-2, -0.5, 3).ae(-12)

def test_legenp():
    mp.dps = 15
    assert legenp(2,0,4) == legendre(2,4)
    assert legenp(-2, -1, 0.5).ae(0.43301270189221932338)
    assert legenp(-2, -1, 0.5, type=3).ae(0.43301270189221932338j)
    assert legenp(-2, 1, 0.5).ae(-0.86602540378443864676)
    assert legenp(2+j, 3+4j, -j).ae(134742.98773236786148 + 429782.72924463851745j)
    assert legenp(2+j, 3+4j, -j, type=3).ae(802.59463394152268507 - 251.62481308942906447j)
    assert legenp(2,4,3).ae(0)
    assert legenp(2,4,3,type=3).ae(0)
    assert legenp(2,1,0.5).ae(-1.2990381056766579701)
    assert legenp(2,1,0.5,type=3).ae(1.2990381056766579701j)
    assert legenp(3,2,3).ae(-360)
    assert legenp(3,3,3).ae(240j*2**0.5)
    assert legenp(3,4,3).ae(0)
    assert legenp(0,0.5,2).ae(0.52503756790433198939 - 0.52503756790433198939j)
    assert legenp(-1,-0.5,2).ae(0.60626116232846498110 + 0.60626116232846498110j)
    assert legenp(-2,0.5,2).ae(1.5751127037129959682 - 1.5751127037129959682j)
    assert legenp(-2,0.5,-0.5).ae(-0.85738275810499171286)

def test_legenq():
    mp.dps = 15
    f = legenq
    # Evaluation at poles
    assert isnan(f(3,2,1))
    assert isnan(f(3,2,-1))
    assert isnan(f(3,2,1,type=3))
    assert isnan(f(3,2,-1,type=3))
    # Evaluation at 0
    assert f(0,1,0,type=2).ae(-1)
    assert f(-2,2,0,type=2,zeroprec=200).ae(0)
    assert f(1.5,3,0,type=2).ae(-2.2239343475841951023)
    assert f(0,1,0,type=3).ae(j)
    assert f(-2,2,0,type=3,zeroprec=200).ae(0)
    assert f(1.5,3,0,type=3).ae(2.2239343475841951022*(1-1j))
    # Standard case, degree 0
    assert f(0,0,-1.5).ae(-0.8047189562170501873 + 1.5707963267948966192j)
    assert f(0,0,-0.5).ae(-0.54930614433405484570)
    assert f(0,0,0,zeroprec=200).ae(0)
    assert f(0,0,0.5).ae(0.54930614433405484570)
    assert f(0,0,1.5).ae(0.8047189562170501873 - 1.5707963267948966192j)
    assert f(0,0,-1.5,type=3).ae(-0.80471895621705018730)
    assert f(0,0,-0.5,type=3).ae(-0.5493061443340548457 - 1.5707963267948966192j)
    assert f(0,0,0,type=3).ae(-1.5707963267948966192j)
    assert f(0,0,0.5,type=3).ae(0.5493061443340548457 - 1.5707963267948966192j)
    assert f(0,0,1.5,type=3).ae(0.80471895621705018730)
    # Standard case, degree 1
    assert f(1,0,-1.5).ae(0.2070784343255752810 - 2.3561944901923449288j)
    assert f(1,0,-0.5).ae(-0.72534692783297257715)
    assert f(1,0,0).ae(-1)
    assert f(1,0,0.5).ae(-0.72534692783297257715)
    assert f(1,0,1.5).ae(0.2070784343255752810 - 2.3561944901923449288j)
    # Standard case, degree 2
    assert f(2,0,-1.5).ae(-0.0635669991240192885 + 4.5160394395353277803j)
    assert f(2,0,-0.5).ae(0.81866326804175685571)
    assert f(2,0,0,zeroprec=200).ae(0)
    assert f(2,0,0.5).ae(-0.81866326804175685571)
    assert f(2,0,1.5).ae(0.0635669991240192885 - 4.5160394395353277803j)
    # Misc orders and degrees
    assert f(2,3,1.5,type=2).ae(-5.7243340223994616228j)
    assert f(2,3,1.5,type=3).ae(-5.7243340223994616228)
    assert f(2,3,0.5,type=2).ae(-12.316805742712016310)
    assert f(2,3,0.5,type=3).ae(-12.316805742712016310j)
    assert f(2,3,-1.5,type=2).ae(-5.7243340223994616228j)
    assert f(2,3,-1.5,type=3).ae(5.7243340223994616228)
    assert f(2,3,-0.5,type=2).ae(-12.316805742712016310)
    assert f(2,3,-0.5,type=3).ae(-12.316805742712016310j)
    assert f(2+3j, 3+4j, 0.5, type=3).ae(0.0016119404873235186807 - 0.0005885900510718119836j)
    assert f(2+3j, 3+4j, -1.5, type=3).ae(0.008451400254138808670 + 0.020645193304593235298j)
    assert f(-2.5,1,-1.5).ae(3.9553395527435335749j)
    assert f(-2.5,1,-0.5).ae(1.9290561746445456908)
    assert f(-2.5,1,0).ae(1.2708196271909686299)
    assert f(-2.5,1,0.5).ae(-0.31584812990742202869)
    assert f(-2.5,1,1.5).ae(-3.9553395527435335742 + 0.2993235655044701706j)
    assert f(-2.5,1,-1.5,type=3).ae(0.29932356550447017254j)
    assert f(-2.5,1,-0.5,type=3).ae(-0.3158481299074220287 - 1.9290561746445456908j)
    assert f(-2.5,1,0,type=3).ae(1.2708196271909686292 - 1.2708196271909686299j)
    assert f(-2.5,1,0.5,type=3).ae(1.9290561746445456907 + 0.3158481299074220287j)
    assert f(-2.5,1,1.5,type=3).ae(-0.29932356550447017254)

def test_agm():
    mp.dps = 15
    assert agm(0,0) == 0
    assert agm(0,1) == 0
    assert agm(1,1) == 1
    assert agm(7,7) == 7
    assert agm(j,j) == j
    assert (1/agm(1,sqrt(2))).ae(0.834626841674073186)
    assert agm(1,2).ae(1.4567910310469068692)
    assert agm(1,3).ae(1.8636167832448965424)
    assert agm(1,j).ae(0.599070117367796104+0.599070117367796104j)
    assert agm(2) == agm(1,2)
    assert agm(-3,4).ae(0.63468509766550907+1.3443087080896272j)

def test_gammainc():
    mp.dps = 15
    assert gammainc(2,5).ae(6*exp(-5))
    assert gammainc(2,0,5).ae(1-6*exp(-5))
    assert gammainc(2,3,5).ae(-6*exp(-5)+4*exp(-3))
    assert gammainc(-2.5,-0.5).ae(-0.9453087204829418812-5.3164237738936178621j)
    assert gammainc(0,2,4).ae(0.045121158298212213088)
    assert gammainc(0,3).ae(0.013048381094197037413)
    assert gammainc(0,2+j,1-j).ae(0.00910653685850304839-0.22378752918074432574j)
    assert gammainc(0,1-j).ae(0.00028162445198141833+0.17932453503935894015j)
    assert gammainc(3,4,5,True).ae(0.11345128607046320253)
    assert gammainc(3.5,0,inf).ae(gamma(3.5))
    assert gammainc(-150.5,500).ae('6.9825435345798951153e-627')
    assert gammainc(-150.5,800).ae('4.6885137549474089431e-788')
    assert gammainc(-3.5, -20.5).ae(0.27008820585226911 - 1310.31447140574997636j)
    assert gammainc(-3.5, -200.5).ae(0.27008820585226911 - 5.3264597096208368435e76j) # XXX real part
    assert gammainc(0,0,2) == inf
    assert gammainc(1,b=1).ae(0.6321205588285576784)
    assert gammainc(3,2,2) == 0
    assert gammainc(2,3+j,3-j).ae(-0.28135485191849314194j)
    assert gammainc(4+0j,1).ae(5.8860710587430771455)
    # GH issue #301
    assert gammainc(-1,-1).ae(-0.8231640121031084799 + 3.1415926535897932385j)
    assert gammainc(-2,-1).ae(1.7707229202810768576 - 1.5707963267948966192j)
    assert gammainc(-3,-1).ae(-1.4963349162467073643 + 0.5235987755982988731j)
    assert gammainc(-4,-1).ae(1.05365418617643814992 - 0.13089969389957471827j)
    # Regularized upper gamma
    assert isnan(gammainc(0, 0, regularized=True))
    assert gammainc(-1, 0, regularized=True) == inf
    assert gammainc(1, 0, regularized=True) == 1
    assert gammainc(0, 5, regularized=True) == 0
    assert gammainc(0, 2+3j, regularized=True) == 0
    assert gammainc(0, 5000, regularized=True) == 0
    assert gammainc(0, 10**30, regularized=True) == 0
    assert gammainc(-1, 5, regularized=True) == 0
    assert gammainc(-1, 5000, regularized=True) == 0
    assert gammainc(-1, 10**30, regularized=True) == 0
    assert gammainc(-1, -5, regularized=True) == 0
    assert gammainc(-1, -5000, regularized=True) == 0
    assert gammainc(-1, -10**30, regularized=True) == 0
    assert gammainc(-1, 3+4j, regularized=True) == 0
    assert gammainc(1, 5, regularized=True).ae(exp(-5))
    assert gammainc(1, 5000, regularized=True).ae(exp(-5000))
    assert gammainc(1, 10**30, regularized=True).ae(exp(-10**30))
    assert gammainc(1, 3+4j, regularized=True).ae(exp(-3-4j))
    assert gammainc(-1000000,2).ae('1.3669297209397347754e-301037', abs_eps=0, rel_eps=8*eps)
    assert gammainc(-1000000,2,regularized=True) == 0
    assert gammainc(-1000000,3+4j).ae('-1.322575609404222361e-698979 - 4.9274570591854533273e-698978j', abs_eps=0, rel_eps=8*eps)
    assert gammainc(-1000000,3+4j,regularized=True) == 0
    assert gammainc(2+3j, 4+5j, regularized=True).ae(0.085422013530993285774-0.052595379150390078503j)
    assert gammainc(1000j, 1000j, regularized=True).ae(0.49702647628921131761 + 0.00297355675013575341j)
    # Generalized
    assert gammainc(3,4,2) == -gammainc(3,2,4)
    assert gammainc(4, 2, 3).ae(1.2593494302978947396)
    assert gammainc(4, 2, 3, regularized=True).ae(0.20989157171631578993)
    assert gammainc(0, 2, 3).ae(0.035852129613864082155)
    assert gammainc(0, 2, 3, regularized=True) == 0
    assert gammainc(-1, 2, 3).ae(0.015219822548487616132)
    assert gammainc(-1, 2, 3, regularized=True) == 0
    assert gammainc(0, 2, 3).ae(0.035852129613864082155)
    assert gammainc(0, 2, 3, regularized=True) == 0
    # Should use upper gammas
    assert gammainc(5, 10000, 12000).ae('1.1359381951461801687e-4327', abs_eps=0, rel_eps=8*eps)
    # Should use lower gammas
    assert gammainc(10000, 2, 3).ae('8.1244514125995785934e4765')
    # GH issue 306
    assert gammainc(3,-1-1j) == 0
    assert gammainc(3,-1+1j) == 0
    assert gammainc(2,-1) == 0
    assert gammainc(2,-1+0j) == 0
    assert gammainc(2+0j,-1) == 0

def test_gammainc_expint_n():
    # These tests are intended to check all cases of the low-level code
    # for upper gamma and expint with small integer index.
    # Need to cover positive/negative arguments; small/large/huge arguments
    # for both positive and negative indices, as well as indices 0 and 1
    # which may be special-cased
    mp.dps = 15
    assert expint(-3,3.5).ae(0.021456366563296693987)
    assert expint(-2,3.5).ae(0.014966633183073309405)
    assert expint(-1,3.5).ae(0.011092916359219041088)
    assert expint(0,3.5).ae(0.0086278238349481430685)
    assert expint(1,3.5).ae(0.0069701398575483929193)
    assert expint(2,3.5).ae(0.0058018939208991255223)
    assert expint(3,3.5).ae(0.0049453773495857807058)
    assert expint(-3,-3.5).ae(-4.6618170604073311319)
    assert expint(-2,-3.5).ae(-5.5996974157555515963)
    assert expint(-1,-3.5).ae(-6.7582555017739415818)
    assert expint(0,-3.5).ae(-9.4615577024835182145)
    assert expint(1,-3.5).ae(-13.925353995152335292 - 3.1415926535897932385j)
    assert expint(2,-3.5).ae(-15.62328702434085977 - 10.995574287564276335j)
    assert expint(3,-3.5).ae(-10.783026313250347722 - 19.242255003237483586j)
    assert expint(-3,350).ae(2.8614825451252838069e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(-2,350).ae(2.8532837224504675901e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(-1,350).ae(2.8451316155828634555e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(0,350).ae(2.8370258275042797989e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(1,350).ae(2.8289659656701459404e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(2,350).ae(2.8209516419468505006e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(3,350).ae(2.8129824725501272171e-155, abs_eps=0, rel_eps=8*eps)
    assert expint(-3,-350).ae(-2.8528796154044839443e+149)
    assert expint(-2,-350).ae(-2.8610072121701264351e+149)
    assert expint(-1,-350).ae(-2.8691813842677537647e+149)
    assert expint(0,-350).ae(-2.8774025343659421709e+149)
    u = expint(1,-350)
    assert u.ae(-2.8856710698020863568e+149)
    assert u.imag.ae(-3.1415926535897932385)
    u = expint(2,-350)
    assert u.ae(-2.8939874026504650534e+149)
    assert u.imag.ae(-1099.5574287564276335)
    u = expint(3,-350)
    assert u.ae(-2.9023519497915044349e+149)
    assert u.imag.ae(-192422.55003237483586)
    assert expint(-3,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(-2,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(-1,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(0,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(1,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(2,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(3,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert expint(-3,-350000000000000000000000).ae('-3.7805306852415755699e+152003068666138139677871')
    assert expint(-2,-350000000000000000000000).ae('-3.7805306852415755699e+152003068666138139677871')
    assert expint(-1,-350000000000000000000000).ae('-3.7805306852415755699e+152003068666138139677871')
    assert expint(0,-350000000000000000000000).ae('-3.7805306852415755699e+152003068666138139677871')
    u = expint(1,-350000000000000000000000)
    assert u.ae('-3.7805306852415755699e+152003068666138139677871')
    assert u.imag.ae(-3.1415926535897932385)
    u = expint(2,-350000000000000000000000)
    assert u.imag.ae(-1.0995574287564276335e+24)
    assert u.ae('-3.7805306852415755699e+152003068666138139677871')
    u = expint(3,-350000000000000000000000)
    assert u.imag.ae(-1.9242255003237483586e+47)
    assert u.ae('-3.7805306852415755699e+152003068666138139677871')
    # Small case; no branch cut
    assert gammainc(-3,3.5).ae(0.00010020262545203707109)
    assert gammainc(-2,3.5).ae(0.00040370427343557393517)
    assert gammainc(-1,3.5).ae(0.0016576839773997501492)
    assert gammainc(0,3.5).ae(0.0069701398575483929193)
    assert gammainc(1,3.5).ae(0.03019738342231850074)
    assert gammainc(2,3.5).ae(0.13588822540043325333)
    assert gammainc(3,3.5).ae(0.64169439772426814072)
    # Small case; with branch cut
    assert gammainc(-3,-3.5).ae(0.03595832954467563286 + 0.52359877559829887308j)
    assert gammainc(-2,-3.5).ae(-0.88024704597962022221 - 1.5707963267948966192j)
    assert gammainc(-1,-3.5).ae(4.4637962926688170771 + 3.1415926535897932385j)
    assert gammainc(0,-3.5).ae(-13.925353995152335292 - 3.1415926535897932385j)
    assert gammainc(1,-3.5).ae(33.115451958692313751)
    assert gammainc(2,-3.5).ae(-82.788629896730784377)
    assert gammainc(3,-3.5).ae(240.08702670051927469)
    # Asymptotic case; no branch cut
    assert gammainc(-3,350).ae(6.5424095113340358813e-163, abs_eps=0, rel_eps=8*eps)
    assert gammainc(-2,350).ae(2.296312222489899769e-160, abs_eps=0, rel_eps=8*eps)
    assert gammainc(-1,350).ae(8.059861834133858573e-158, abs_eps=0, rel_eps=8*eps)
    assert gammainc(0,350).ae(2.8289659656701459404e-155, abs_eps=0, rel_eps=8*eps)
    assert gammainc(1,350).ae(9.9295903962649792963e-153, abs_eps=0, rel_eps=8*eps)
    assert gammainc(2,350).ae(3.485286229089007733e-150, abs_eps=0, rel_eps=8*eps)
    assert gammainc(3,350).ae(1.2233453960006379793e-147, abs_eps=0, rel_eps=8*eps)
    # Asymptotic case; branch cut
    u = gammainc(-3,-350)
    assert u.ae(6.7889565783842895085e+141)
    assert u.imag.ae(0.52359877559829887308)
    u = gammainc(-2,-350)
    assert u.ae(-2.3692668977889832121e+144)
    assert u.imag.ae(-1.5707963267948966192)
    u = gammainc(-1,-350)
    assert u.ae(8.2685354361441858669e+146)
    assert u.imag.ae(3.1415926535897932385)
    u = gammainc(0,-350)
    assert u.ae(-2.8856710698020863568e+149)
    assert u.imag.ae(-3.1415926535897932385)
    u = gammainc(1,-350)
    assert u.ae(1.0070908870280797598e+152)
    assert u.imag == 0
    u = gammainc(2,-350)
    assert u.ae(-3.5147471957279983618e+154)
    assert u.imag == 0
    u = gammainc(3,-350)
    assert u.ae(1.2266568422179417091e+157)
    assert u.imag == 0
    # Extreme asymptotic case
    assert gammainc(-3,350000000000000000000000).ae('5.0362468738874738859e-152003068666138139677990', abs_eps=0, rel_eps=8*eps)
    assert gammainc(-2,350000000000000000000000).ae('1.7626864058606158601e-152003068666138139677966', abs_eps=0, rel_eps=8*eps)
    assert gammainc(-1,350000000000000000000000).ae('6.1694024205121555102e-152003068666138139677943', abs_eps=0, rel_eps=8*eps)
    assert gammainc(0,350000000000000000000000).ae('2.1592908471792544286e-152003068666138139677919', abs_eps=0, rel_eps=8*eps)
    assert gammainc(1,350000000000000000000000).ae('7.5575179651273905e-152003068666138139677896', abs_eps=0, rel_eps=8*eps)
    assert gammainc(2,350000000000000000000000).ae('2.645131287794586675e-152003068666138139677872', abs_eps=0, rel_eps=8*eps)
    assert gammainc(3,350000000000000000000000).ae('9.2579595072810533625e-152003068666138139677849', abs_eps=0, rel_eps=8*eps)
    u = gammainc(-3,-350000000000000000000000)
    assert u.ae('8.8175642804468234866e+152003068666138139677800')
    assert u.imag.ae(0.52359877559829887308)
    u = gammainc(-2,-350000000000000000000000)
    assert u.ae('-3.0861474981563882203e+152003068666138139677824')
    assert u.imag.ae(-1.5707963267948966192)
    u = gammainc(-1,-350000000000000000000000)
    assert u.ae('1.0801516243547358771e+152003068666138139677848')
    assert u.imag.ae(3.1415926535897932385)
    u = gammainc(0,-350000000000000000000000)
    assert u.ae('-3.7805306852415755699e+152003068666138139677871')
    assert u.imag.ae(-3.1415926535897932385)
    assert gammainc(1,-350000000000000000000000).ae('1.3231857398345514495e+152003068666138139677895')
    assert gammainc(2,-350000000000000000000000).ae('-4.6311500894209300731e+152003068666138139677918')
    assert gammainc(3,-350000000000000000000000).ae('1.6209025312973255256e+152003068666138139677942')

def test_incomplete_beta():
    mp.dps = 15
    assert betainc(-2,-3,0.5,0.75).ae(63.4305673311255413583969)
    assert betainc(4.5,0.5+2j,2.5,6).ae(0.2628801146130621387903065 + 0.5162565234467020592855378j)
    assert betainc(4,5,0,6).ae(90747.77142857142857142857)

def test_erf():
    mp.dps = 15
    assert erf(0) == 0
    assert erf(1).ae(0.84270079294971486934)
    assert erf(3+4j).ae(-120.186991395079444098 - 27.750337293623902498j)
    assert erf(-4-3j).ae(-0.99991066178539168236 + 0.00004972026054496604j)
    assert erf(pi).ae(0.99999112385363235839)
    assert erf(1j).ae(1.6504257587975428760j)
    assert erf(-1j).ae(-1.6504257587975428760j)
    assert isinstance(erf(1), mpf)
    assert isinstance(erf(-1), mpf)
    assert isinstance(erf(0), mpf)
    assert isinstance(erf(0j), mpc)
    assert erf(inf) == 1
    assert erf(-inf) == -1
    assert erfi(0) == 0
    assert erfi(1/pi).ae(0.371682698493894314)
    assert erfi(inf) == inf
    assert erfi(-inf) == -inf
    assert erf(1+0j) == erf(1)
    assert erfc(1+0j) == erfc(1)
    assert erf(0.2+0.5j).ae(1 - erfc(0.2+0.5j))
    assert erfc(0) == 1
    assert erfc(1).ae(1-erf(1))
    assert erfc(-1).ae(1-erf(-1))
    assert erfc(1/pi).ae(1-erf(1/pi))
    assert erfc(-10) == 2
    assert erfc(-1000000) == 2
    assert erfc(-inf) == 2
    assert erfc(inf) == 0
    assert isnan(erfc(nan))
    assert (erfc(10**4)*mpf(10)**43429453).ae('3.63998738656420')
    assert erf(8+9j).ae(-1072004.2525062051158 + 364149.91954310255423j)
    assert erfc(8+9j).ae(1072005.2525062051158 - 364149.91954310255423j)
    assert erfc(-8-9j).ae(-1072003.2525062051158 + 364149.91954310255423j)
    mp.dps = 50
    # This one does not use the asymptotic series
    assert (erfc(10)*10**45).ae('2.0884875837625447570007862949577886115608181193212')
    # This one does
    assert (erfc(50)*10**1088).ae('2.0709207788416560484484478751657887929322509209954')
    mp.dps = 15
    assert str(erfc(10**50)) == '3.66744826532555e-4342944819032518276511289189166050822943970058036665661144537831658646492088707747292249493384317534'
    assert erfinv(0) == 0
    assert erfinv(0.5).ae(0.47693627620446987338)
    assert erfinv(-0.5).ae(-0.47693627620446987338)
    assert erfinv(1) == inf
    assert erfinv(-1) == -inf
    assert erf(erfinv(0.95)).ae(0.95)
    assert erf(erfinv(0.999999999995)).ae(0.999999999995)
    assert erf(erfinv(-0.999999999995)).ae(-0.999999999995)
    mp.dps = 50
    assert erf(erfinv('0.99999999999999999999999999999995')).ae('0.99999999999999999999999999999995')
    assert erf(erfinv('0.999999999999999999999999999999995')).ae('0.999999999999999999999999999999995')
    assert erf(erfinv('-0.999999999999999999999999999999995')).ae('-0.999999999999999999999999999999995')
    mp.dps = 15
    # Complex asymptotic expansions
    v = erfc(50j)
    assert v.real == 1
    assert v.imag.ae('-6.1481820666053078736e+1083')
    assert erfc(-100+5j).ae(2)
    assert (erfc(100+5j)*10**4335).ae(2.3973567853824133572 - 3.9339259530609420597j)
    assert erfc(100+100j).ae(0.00065234366376857698698 - 0.0039357263629214118437j)

def test_pdf():
    mp.dps = 15
    assert npdf(-inf) == 0
    assert npdf(inf) == 0
    assert npdf(5,0,2).ae(npdf(5+4,4,2))
    assert quadts(lambda x: npdf(x,-0.5,0.8), [-inf, inf]) == 1
    assert ncdf(0) == 0.5
    assert ncdf(3,3) == 0.5
    assert ncdf(-inf) == 0
    assert ncdf(inf) == 1
    assert ncdf(10) == 1
    # Verify that this is computed accurately
    assert (ncdf(-10)*10**24).ae(7.619853024160526)

def test_lambertw():
    mp.dps = 15
    assert lambertw(0) == 0
    assert lambertw(0+0j) == 0
    assert lambertw(inf) == inf
    assert isnan(lambertw(nan))
    assert lambertw(inf,1).real == inf
    assert lambertw(inf,1).imag.ae(2*pi)
    assert lambertw(-inf,1).real == inf
    assert lambertw(-inf,1).imag.ae(3*pi)
    assert lambertw(0,-1) == -inf
    assert lambertw(0,1) == -inf
    assert lambertw(0,3) == -inf
    assert lambertw(e).ae(1)
    assert lambertw(1).ae(0.567143290409783873)
    assert lambertw(-pi/2).ae(j*pi/2)
    assert lambertw(-log(2)/2).ae(-log(2))
    assert lambertw(0.25).ae(0.203888354702240164)
    assert lambertw(-0.25).ae(-0.357402956181388903)
    assert lambertw(-1./10000,0).ae(-0.000100010001500266719)
    assert lambertw(-0.25,-1).ae(-2.15329236411034965)
    assert lambertw(0.25,-1).ae(-3.00899800997004620-4.07652978899159763j)
    assert lambertw(-0.25,-1).ae(-2.15329236411034965)
    assert lambertw(0.25,1).ae(-3.00899800997004620+4.07652978899159763j)
    assert lambertw(-0.25,1).ae(-3.48973228422959210+7.41405453009603664j)
    assert lambertw(-4).ae(0.67881197132094523+1.91195078174339937j)
    assert lambertw(-4,1).ae(-0.66743107129800988+7.76827456802783084j)
    assert lambertw(-4,-1).ae(0.67881197132094523-1.91195078174339937j)
    assert lambertw(1000).ae(5.24960285240159623)
    assert lambertw(1000,1).ae(4.91492239981054535+5.44652615979447070j)
    assert lambertw(1000,-1).ae(4.91492239981054535-5.44652615979447070j)
    assert lambertw(1000,5).ae(3.5010625305312892+29.9614548941181328j)
    assert lambertw(3+4j).ae(1.281561806123775878+0.533095222020971071j)
    assert lambertw(-0.4+0.4j).ae(-0.10396515323290657+0.61899273315171632j)
    assert lambertw(3+4j,1).ae(-0.11691092896595324+5.61888039871282334j)
    assert lambertw(3+4j,-1).ae(0.25856740686699742-3.85211668616143559j)
    assert lambertw(-0.5,-1).ae(-0.794023632344689368-0.770111750510379110j)
    assert lambertw(-1./10000,1).ae(-11.82350837248724344+6.80546081842002101j)
    assert lambertw(-1./10000,-1).ae(-11.6671145325663544)
    assert lambertw(-1./10000,-2).ae(-11.82350837248724344-6.80546081842002101j)
    assert lambertw(-1./100000,4).ae(-14.9186890769540539+26.1856750178782046j)
    assert lambertw(-1./100000,5).ae(-15.0931437726379218666+32.5525721210262290086j)
    assert lambertw((2+j)/10).ae(0.173704503762911669+0.071781336752835511j)
    assert lambertw((2+j)/10,1).ae(-3.21746028349820063+4.56175438896292539j)
    assert lambertw((2+j)/10,-1).ae(-3.03781405002993088-3.53946629633505737j)
    assert lambertw((2+j)/10,4).ae(-4.6878509692773249+23.8313630697683291j)
    assert lambertw(-(2+j)/10).ae(-0.226933772515757933-0.164986470020154580j)
    assert lambertw(-(2+j)/10,1).ae(-2.43569517046110001+0.76974067544756289j)
    assert lambertw(-(2+j)/10,-1).ae(-3.54858738151989450-6.91627921869943589j)
    assert lambertw(-(2+j)/10,4).ae(-4.5500846928118151+20.6672982215434637j)
    mp.dps = 50
    assert lambertw(pi).ae('1.073658194796149172092178407024821347547745350410314531')
    mp.dps = 15
    # Former bug in generated branch
    assert lambertw(-0.5+0.002j).ae(-0.78917138132659918344 + 0.76743539379990327749j)
    assert lambertw(-0.5-0.002j).ae(-0.78917138132659918344 - 0.76743539379990327749j)
    assert lambertw(-0.448+0.4j).ae(-0.11855133765652382241 + 0.66570534313583423116j)
    assert lambertw(-0.448-0.4j).ae(-0.11855133765652382241 - 0.66570534313583423116j)
    assert lambertw(-0.65475+0.0001j).ae(-0.61053421111385310898+1.0396534993944097723803j)
    # Huge branch index
    w = lambertw(1,10**20)
    assert w.real.ae(-47.889578926290259164)
    assert w.imag.ae(6.2831853071795864769e+20)

def test_lambertw_hard():
    def check(x,y):
        y = convert(y)
        type_ok = True
        if isinstance(y, mpf):
            type_ok = isinstance(x, mpf)
        real_ok = abs(x.real-y.real) <= abs(y.real)*8*eps
        imag_ok = abs(x.imag-y.imag) <= abs(y.imag)*8*eps
        #print x, y, abs(x.real-y.real), abs(x.imag-y.imag)
        return real_ok and imag_ok
    # Evaluation near 0
    mp.dps = 15
    assert check(lambertw(1e-10), 9.999999999000000000e-11)
    assert check(lambertw(-1e-10), -1.000000000100000000e-10)
    assert check(lambertw(1e-10j), 9.999999999999999999733e-21 + 9.99999999999999999985e-11j)
    assert check(lambertw(-1e-10j), 9.999999999999999999733e-21 - 9.99999999999999999985e-11j)
    assert check(lambertw(1e-10,1), -26.303186778379041559 + 3.265093911703828397j)
    assert check(lambertw(-1e-10,1), -26.326236166739163892 + 6.526183280686333315j)
    assert check(lambertw(1e-10j,1), -26.312931726911421551 + 4.896366881798013421j)
    assert check(lambertw(-1e-10j,1), -26.297238779529035066 + 1.632807161345576513j)
    assert check(lambertw(1e-10,-1), -26.303186778379041559 - 3.265093911703828397j)
    assert check(lambertw(-1e-10,-1), -26.295238819246925694)
    assert check(lambertw(1e-10j,-1), -26.297238779529035028 - 1.6328071613455765135j)
    assert check(lambertw(-1e-10j,-1), -26.312931726911421551 - 4.896366881798013421j)
    # Test evaluation very close to the branch point -1/e
    # on the -1, 0, and 1 branches
    add = lambda x, y: fadd(x,y,exact=True)
    sub = lambda x, y: fsub(x,y,exact=True)
    addj = lambda x, y: fadd(x,fmul(y,1j,exact=True),exact=True)
    subj = lambda x, y: fadd(x,fmul(y,-1j,exact=True),exact=True)
    mp.dps = 1500
    a = -1/e + 10*eps
    d3 = mpf('1e-3')
    d10 = mpf('1e-10')
    d20 = mpf('1e-20')
    d40 = mpf('1e-40')
    d80 = mpf('1e-80')
    d300 = mpf('1e-300')
    d1000 = mpf('1e-1000')
    mp.dps = 15
    # ---- Branch 0 ----
    # -1/e + eps
    assert check(lambertw(add(a,d3)), -0.92802015005456704876)
    assert check(lambertw(add(a,d10)), -0.99997668374140088071)
    assert check(lambertw(add(a,d20)), -0.99999999976683560186)
    assert lambertw(add(a,d40)) == -1
    assert lambertw(add(a,d80)) == -1
    assert lambertw(add(a,d300)) == -1
    assert lambertw(add(a,d1000)) == -1
    # -1/e - eps
    assert check(lambertw(sub(a,d3)), -0.99819016149860989001+0.07367191188934638577j)
    assert check(lambertw(sub(a,d10)), -0.9999999998187812114595992+0.0000233164398140346109194j)
    assert check(lambertw(sub(a,d20)), -0.99999999999999999998187+2.331643981597124203344e-10j)
    assert check(lambertw(sub(a,d40)), -1.0+2.33164398159712420336e-20j)
    assert check(lambertw(sub(a,d80)), -1.0+2.33164398159712420336e-40j)
    assert check(lambertw(sub(a,d300)), -1.0+2.33164398159712420336e-150j)
    assert check(lambertw(sub(a,d1000)), mpc(-1,'2.33164398159712420336e-500'))
    # -1/e + eps*j
    assert check(lambertw(addj(a,d3)), -0.94790387486938526634+0.05036819639190132490j)
    assert check(lambertw(addj(a,d10)), -0.9999835127872943680999899+0.0000164870314895821225256j)
    assert check(lambertw(addj(a,d20)), -0.999999999835127872929987+1.64872127051890935830e-10j)
    assert check(lambertw(addj(a,d40)), -0.9999999999999999999835+1.6487212707001281468305e-20j)
    assert check(lambertw(addj(a,d80)), -1.0 + 1.64872127070012814684865e-40j)
    assert check(lambertw(addj(a,d300)), -1.0 + 1.64872127070012814684865e-150j)
    assert check(lambertw(addj(a,d1000)), mpc(-1.0,'1.64872127070012814684865e-500'))
    # -1/e - eps*j
    assert check(lambertw(subj(a,d3)), -0.94790387486938526634-0.05036819639190132490j)
    assert check(lambertw(subj(a,d10)), -0.9999835127872943680999899-0.0000164870314895821225256j)
    assert check(lambertw(subj(a,d20)), -0.999999999835127872929987-1.64872127051890935830e-10j)
    assert check(lambertw(subj(a,d40)), -0.9999999999999999999835-1.6487212707001281468305e-20j)
    assert check(lambertw(subj(a,d80)), -1.0 - 1.64872127070012814684865e-40j)
    assert check(lambertw(subj(a,d300)), -1.0 - 1.64872127070012814684865e-150j)
    assert check(lambertw(subj(a,d1000)), mpc(-1.0,'-1.64872127070012814684865e-500'))
    # ---- Branch 1 ----
    assert check(lambertw(addj(a,d3),1), -3.088501303219933378005990 + 7.458676867597474813950098j)
    assert check(lambertw(addj(a,d80),1), -3.088843015613043855957087 + 7.461489285654254556906117j)
    assert check(lambertw(addj(a,d300),1), -3.088843015613043855957087 + 7.461489285654254556906117j)
    assert check(lambertw(addj(a,d1000),1), -3.088843015613043855957087 + 7.461489285654254556906117j)
    assert check(lambertw(subj(a,d3),1), -1.0520914180450129534365906 + 0.0539925638125450525673175j)
    assert check(lambertw(subj(a,d10),1), -1.0000164872127056318529390 + 0.000016487393927159250398333077j)
    assert check(lambertw(subj(a,d20),1), -1.0000000001648721270700128 + 1.64872127088134693542628e-10j)
    assert check(lambertw(subj(a,d40),1), -1.000000000000000000016487 + 1.64872127070012814686677e-20j)
    assert check(lambertw(subj(a,d80),1), -1.0 + 1.64872127070012814684865e-40j)
    assert check(lambertw(subj(a,d300),1), -1.0 + 1.64872127070012814684865e-150j)
    assert check(lambertw(subj(a,d1000),1), mpc(-1.0, '1.64872127070012814684865e-500'))
    # ---- Branch -1 ----
    # -1/e + eps
    assert check(lambertw(add(a,d3),-1), -1.075608941186624989414945)
    assert check(lambertw(add(a,d10),-1), -1.000023316621036696460620)
    assert check(lambertw(add(a,d20),-1), -1.000000000233164398177834)
    assert lambertw(add(a,d40),-1) == -1
    assert lambertw(add(a,d80),-1) == -1
    assert lambertw(add(a,d300),-1) == -1
    assert lambertw(add(a,d1000),-1) == -1
    # -1/e - eps
    assert check(lambertw(sub(a,d3),-1), -0.99819016149860989001-0.07367191188934638577j)
    assert check(lambertw(sub(a,d10),-1), -0.9999999998187812114595992-0.0000233164398140346109194j)
    assert check(lambertw(sub(a,d20),-1), -0.99999999999999999998187-2.331643981597124203344e-10j)
    assert check(lambertw(sub(a,d40),-1), -1.0-2.33164398159712420336e-20j)
    assert check(lambertw(sub(a,d80),-1), -1.0-2.33164398159712420336e-40j)
    assert check(lambertw(sub(a,d300),-1), -1.0-2.33164398159712420336e-150j)
    assert check(lambertw(sub(a,d1000),-1), mpc(-1,'-2.33164398159712420336e-500'))
    # -1/e + eps*j
    assert check(lambertw(addj(a,d3),-1), -1.0520914180450129534365906 - 0.0539925638125450525673175j)
    assert check(lambertw(addj(a,d10),-1), -1.0000164872127056318529390 - 0.0000164873939271592503983j)
    assert check(lambertw(addj(a,d20),-1), -1.0000000001648721270700 - 1.64872127088134693542628e-10j)
    assert check(lambertw(addj(a,d40),-1), -1.00000000000000000001648 - 1.6487212707001281468667726e-20j)
    assert check(lambertw(addj(a,d80),-1), -1.0 - 1.64872127070012814684865e-40j)
    assert check(lambertw(addj(a,d300),-1), -1.0 - 1.64872127070012814684865e-150j)
    assert check(lambertw(addj(a,d1000),-1), mpc(-1.0,'-1.64872127070012814684865e-500'))
    # -1/e - eps*j
    assert check(lambertw(subj(a,d3),-1), -3.088501303219933378005990-7.458676867597474813950098j)
    assert check(lambertw(subj(a,d10),-1), -3.088843015579260686911033-7.461489285372968780020716j)
    assert check(lambertw(subj(a,d20),-1), -3.088843015613043855953708-7.461489285654254556877988j)
    assert check(lambertw(subj(a,d40),-1), -3.088843015613043855957087-7.461489285654254556906117j)
    assert check(lambertw(subj(a,d80),-1), -3.088843015613043855957087 - 7.461489285654254556906117j)
    assert check(lambertw(subj(a,d300),-1), -3.088843015613043855957087 - 7.461489285654254556906117j)
    assert check(lambertw(subj(a,d1000),-1), -3.088843015613043855957087 - 7.461489285654254556906117j)
    # One more case, testing higher precision
    mp.dps = 500
    x = -1/e + mpf('1e-13')
    ans = "-0.99999926266961377166355784455394913638782494543377383"\
    "744978844374498153493943725364881490261187530235150668593869563"\
    "168276697689459394902153960200361935311512317183678882"
    mp.dps = 15
    assert lambertw(x).ae(ans)
    mp.dps = 50
    assert lambertw(x).ae(ans)
    mp.dps = 150
    assert lambertw(x).ae(ans)

def test_meijerg():
    mp.dps = 15
    assert meijerg([[2,3],[1]],[[0.5,2],[3,4]], 2.5).ae(4.2181028074787439386)
    assert meijerg([[],[1+j]],[[1],[1]], 3+4j).ae(271.46290321152464592 - 703.03330399954820169j)
    assert meijerg([[0.25],[1]],[[0.5],[2]],0) == 0
    assert meijerg([[0],[]],[[0,0,'1/3','2/3'], []], '2/27').ae(2.2019391389653314120)
    # Verify 1/z series being used
    assert meijerg([[-3],[-0.5]], [[-1],[-2.5]], -0.5).ae(-1.338096165935754898687431)
    assert meijerg([[1-(-1)],[1-(-2.5)]], [[1-(-3)],[1-(-0.5)]], -2.0).ae(-1.338096165935754898687431)
    assert meijerg([[-3],[-0.5]], [[-1],[-2.5]], -1).ae(-(pi+4)/(4*pi))
    a = 2.5
    b = 1.25
    for z in [mpf(0.25), mpf(2)]:
        x1 = hyp1f1(a,b,z)
        x2 = gamma(b)/gamma(a)*meijerg([[1-a],[]],[[0],[1-b]],-z)
        x3 = gamma(b)/gamma(a)*meijerg([[1-0],[1-(1-b)]],[[1-(1-a)],[]],-1/z)
        assert x1.ae(x2)
        assert x1.ae(x3)

def test_appellf1():
    mp.dps = 15
    assert appellf1(2,-2,1,1,2,3).ae(-1.75)
    assert appellf1(2,1,-2,1,2,3).ae(-8)
    assert appellf1(2,1,-2,1,0.5,0.25).ae(1.5)
    assert appellf1(-2,1,3,2,3,3).ae(19)
    assert appellf1(1,2,3,4,0.5,0.125).ae( 1.53843285792549786518)

def test_coulomb():
    # Note: most tests are doctests
    # Test for a bug:
    mp.dps = 15
    assert coulombg(mpc(-5,0),2,3).ae(20.087729487721430394)

def test_hyper_param_accuracy():
    mp.dps = 15
    As = [n+1e-10 for n in range(-5,-1)]
    Bs = [n+1e-10 for n in range(-12,-5)]
    assert hyper(As,Bs,10).ae(-381757055858.652671927)
    assert legenp(0.5, 100, 0.25).ae(-2.4124576567211311755e+144)
    assert (hyp1f1(1000,1,-100)*10**24).ae(5.2589445437370169113)
    assert (hyp2f1(10, -900, 10.5, 0.99)*10**24).ae(1.9185370579660768203)
    assert (hyp2f1(1000,1.5,-3.5,-1.5)*10**385).ae(-2.7367529051334000764)
    assert hyp2f1(-5, 10, 3, 0.5, zeroprec=500) == 0
    assert (hyp1f1(-10000, 1000, 100)*10**424).ae(-3.1046080515824859974)
    assert (hyp2f1(1000,1.5,-3.5,-0.75,maxterms=100000)*10**231).ae(-4.0534790813913998643)
    assert legenp(2, 3, 0.25) == 0
    pytest.raises(ValueError, lambda: hypercomb(lambda a: [([],[],[],[],[a],[-a],0.5)], [3]))
    assert hypercomb(lambda a: [([],[],[],[],[a],[-a],0.5)], [3], infprec=200) == inf
    assert meijerg([[],[]],[[0,0,0,0],[]],0.1).ae(1.5680822343832351418)
    assert (besselk(400,400)*10**94).ae(1.4387057277018550583)
    mp.dps = 5
    (hyp1f1(-5000.5, 1500, 100)*10**185).ae(8.5185229673381935522)
    (hyp1f1(-5000, 1500, 100)*10**185).ae(9.1501213424563944311)
    mp.dps = 15
    (hyp1f1(-5000.5, 1500, 100)*10**185).ae(8.5185229673381935522)
    (hyp1f1(-5000, 1500, 100)*10**185).ae(9.1501213424563944311)
    assert hyp0f1(fadd(-20,'1e-100',exact=True), 0.25).ae(1.85014429040102783e+49)
    assert hyp0f1((-20*10**100+1, 10**100), 0.25).ae(1.85014429040102783e+49)

def test_hypercomb_zero_pow():
    # check that 0^0 = 1
    assert hypercomb(lambda a: (([0],[a],[],[],[],[],0),), [0]) == 1
    assert meijerg([[-1.5],[]],[[0],[-0.75]],0).ae(1.4464090846320771425)

def test_spherharm():
    mp.dps = 15
    t = 0.5; r = 0.25
    assert spherharm(0,0,t,r).ae(0.28209479177387814347)
    assert spherharm(1,-1,t,r).ae(0.16048941205971996369 - 0.04097967481096344271j)
    assert spherharm(1,0,t,r).ae(0.42878904414183579379)
    assert spherharm(1,1,t,r).ae(-0.16048941205971996369 - 0.04097967481096344271j)
    assert spherharm(2,-2,t,r).ae(0.077915886919031181734 - 0.042565643022253962264j)
    assert spherharm(2,-1,t,r).ae(0.31493387233497459884 - 0.08041582001959297689j)
    assert spherharm(2,0,t,r).ae(0.41330596756220761898)
    assert spherharm(2,1,t,r).ae(-0.31493387233497459884 - 0.08041582001959297689j)
    assert spherharm(2,2,t,r).ae(0.077915886919031181734 + 0.042565643022253962264j)
    assert spherharm(3,-3,t,r).ae(0.033640236589690881646 - 0.031339125318637082197j)
    assert spherharm(3,-2,t,r).ae(0.18091018743101461963 - 0.09883168583167010241j)
    assert spherharm(3,-1,t,r).ae(0.42796713930907320351 - 0.10927795157064962317j)
    assert spherharm(3,0,t,r).ae(0.27861659336351639787)
    assert spherharm(3,1,t,r).ae(-0.42796713930907320351 - 0.10927795157064962317j)
    assert spherharm(3,2,t,r).ae(0.18091018743101461963 + 0.09883168583167010241j)
    assert spherharm(3,3,t,r).ae(-0.033640236589690881646 - 0.031339125318637082197j)
    assert spherharm(0,-1,t,r) == 0
    assert spherharm(0,-2,t,r) == 0
    assert spherharm(0,1,t,r) == 0
    assert spherharm(0,2,t,r) == 0
    assert spherharm(1,2,t,r) == 0
    assert spherharm(1,3,t,r) == 0
    assert spherharm(1,-2,t,r) == 0
    assert spherharm(1,-3,t,r) == 0
    assert spherharm(2,3,t,r) == 0
    assert spherharm(2,4,t,r) == 0
    assert spherharm(2,-3,t,r) == 0
    assert spherharm(2,-4,t,r) == 0
    assert spherharm(3,4.5,0.5,0.25).ae(-22.831053442240790148 + 10.910526059510013757j)
    assert spherharm(2+3j, 1-j, 1+j, 3+4j).ae(-2.6582752037810116935 - 1.0909214905642160211j)
    assert spherharm(-6,2.5,t,r).ae(0.39383644983851448178 + 0.28414687085358299021j)
    assert spherharm(-3.5, 3, 0.5, 0.25).ae(0.014516852987544698924 - 0.015582769591477628495j)
    assert spherharm(-3, 3, 0.5, 0.25) == 0
    assert spherharm(-6, 3, 0.5, 0.25).ae(-0.16544349818782275459 - 0.15412657723253924562j)
    assert spherharm(-6, 1.5, 0.5, 0.25).ae(0.032208193499767402477 + 0.012678000924063664921j)
    assert spherharm(3,0,0,1).ae(0.74635266518023078283)
    assert spherharm(3,-2,0,1) == 0
    assert spherharm(3,-2,1,1).ae(-0.16270707338254028971 - 0.35552144137546777097j)

def test_qfunctions():
    mp.dps = 15
    assert qp(2,3,100).ae('2.7291482267247332183e2391')

def test_issue_239():
    mp.prec = 150
    x = ldexp(2476979795053773,-52)
    assert betainc(206, 385, 0, 0.55, 1).ae('0.99999999999999999999996570910644857895771110649954')
    mp.dps = 15
    pytest.raises(ValueError, lambda: hyp2f1(-5,5,0.5,0.5))

# Extra stress testing for Bessel functions
# Reference zeros generated with the aid of scipy.special
# jn_zero, jnp_zero, yn_zero, ynp_zero

V = 15
M = 15

jn_small_zeros = \
[[2.4048255576957728,
  5.5200781102863106,
  8.6537279129110122,
  11.791534439014282,
  14.930917708487786,
  18.071063967910923,
  21.211636629879259,
  24.352471530749303,
  27.493479132040255,
  30.634606468431975,
  33.775820213573569,
  36.917098353664044,
  40.058425764628239,
  43.19979171317673,
  46.341188371661814],
 [3.8317059702075123,
  7.0155866698156188,
  10.173468135062722,
  13.323691936314223,
  16.470630050877633,
  19.615858510468242,
  22.760084380592772,
  25.903672087618383,
  29.046828534916855,
  32.189679910974404,
  35.332307550083865,
  38.474766234771615,
  41.617094212814451,
  44.759318997652822,
  47.901460887185447],
 [5.1356223018406826,
  8.4172441403998649,
  11.619841172149059,
  14.795951782351261,
  17.959819494987826,
  21.116997053021846,
  24.270112313573103,
  27.420573549984557,
  30.569204495516397,
  33.7165195092227,
  36.86285651128381,
  40.008446733478192,
  43.153453778371463,
  46.297996677236919,
  49.442164110416873],
 [6.3801618959239835,
  9.7610231299816697,
  13.015200721698434,
  16.223466160318768,
  19.409415226435012,
  22.582729593104442,
  25.748166699294978,
  28.908350780921758,
  32.064852407097709,
  35.218670738610115,
  38.370472434756944,
  41.520719670406776,
  44.669743116617253,
  47.817785691533302,
  50.965029906205183],
 [7.5883424345038044,
  11.064709488501185,
  14.37253667161759,
  17.615966049804833,
  20.826932956962388,
  24.01901952477111,
  27.199087765981251,
  30.371007667117247,
  33.537137711819223,
  36.699001128744649,
  39.857627302180889,
  43.01373772335443,
  46.167853512924375,
  49.320360686390272,
  52.471551398458023],
 [8.771483815959954,
  12.338604197466944,
  15.700174079711671,
  18.980133875179921,
  22.217799896561268,
  25.430341154222704,
  28.626618307291138,
  31.811716724047763,
  34.988781294559295,
  38.159868561967132,
  41.326383254047406,
  44.489319123219673,
  47.649399806697054,
  50.80716520300633,
  53.963026558378149],
 [9.9361095242176849,
  13.589290170541217,
  17.003819667816014,
  20.320789213566506,
  23.58608443558139,
  26.820151983411405,
  30.033722386570469,
  33.233041762847123,
  36.422019668258457,
  39.603239416075404,
  42.778481613199507,
  45.949015998042603,
  49.11577372476426,
  52.279453903601052,
  55.440592068853149],
 [11.086370019245084,
  14.821268727013171,
  18.287582832481726,
  21.641541019848401,
  24.934927887673022,
  28.191188459483199,
  31.42279419226558,
  34.637089352069324,
  37.838717382853611,
  41.030773691585537,
  44.21540850526126,
  47.394165755570512,
  50.568184679795566,
  53.738325371963291,
  56.905249991978781],
 [12.225092264004655,
  16.037774190887709,
  19.554536430997055,
  22.94517313187462,
  26.266814641176644,
  29.54565967099855,
  32.795800037341462,
  36.025615063869571,
  39.240447995178135,
  42.443887743273558,
  45.638444182199141,
  48.825930381553857,
  52.007691456686903,
  55.184747939289049,
  58.357889025269694],
 [13.354300477435331,
  17.241220382489128,
  20.807047789264107,
  24.233885257750552,
  27.583748963573006,
  30.885378967696675,
  34.154377923855096,
  37.400099977156589,
  40.628553718964528,
  43.843801420337347,
  47.048700737654032,
  50.245326955305383,
  53.435227157042058,
  56.619580266508436,
  59.799301630960228],
 [14.475500686554541,
  18.433463666966583,
  22.046985364697802,
  25.509450554182826,
  28.887375063530457,
  32.211856199712731,
  35.499909205373851,
  38.761807017881651,
  42.004190236671805,
  45.231574103535045,
  48.447151387269394,
  51.653251668165858,
  54.851619075963349,
  58.043587928232478,
  61.230197977292681],
 [15.589847884455485,
  19.61596690396692,
  23.275853726263409,
  26.773322545509539,
  30.17906117878486,
  33.526364075588624,
  36.833571341894905,
  40.111823270954241,
  43.368360947521711,
  46.608132676274944,
  49.834653510396724,
  53.050498959135054,
  56.257604715114484,
  59.457456908388002,
  62.651217388202912],
 [16.698249933848246,
  20.789906360078443,
  24.494885043881354,
  28.026709949973129,
  31.45996003531804,
  34.829986990290238,
  38.156377504681354,
  41.451092307939681,
  44.721943543191147,
  47.974293531269048,
  51.211967004101068,
  54.437776928325074,
  57.653844811906946,
  60.8618046824805,
  64.062937824850136],
 [17.801435153282442,
  21.95624406783631,
  25.705103053924724,
  29.270630441874802,
  32.731053310978403,
  36.123657666448762,
  39.469206825243883,
  42.780439265447158,
  46.06571091157561,
  49.330780096443524,
  52.579769064383396,
  55.815719876305778,
  59.040934037249271,
  62.257189393731728,
  65.465883797232125],
 [18.899997953174024,
  23.115778347252756,
  26.907368976182104,
  30.505950163896036,
  33.993184984781542,
  37.408185128639695,
  40.772827853501868,
  44.100590565798301,
  47.400347780543231,
  50.678236946479898,
  53.93866620912693,
  57.184898598119301,
  60.419409852130297,
  63.644117508962281,
  66.860533012260103]]

jnp_small_zeros = \
[[0.0,
  3.8317059702075123,
  7.0155866698156188,
  10.173468135062722,
  13.323691936314223,
  16.470630050877633,
  19.615858510468242,
  22.760084380592772,
  25.903672087618383,
  29.046828534916855,
  32.189679910974404,
  35.332307550083865,
  38.474766234771615,
  41.617094212814451,
  44.759318997652822],
 [1.8411837813406593,
  5.3314427735250326,
  8.5363163663462858,
  11.706004902592064,
  14.863588633909033,
  18.015527862681804,
  21.16436985918879,
  24.311326857210776,
  27.457050571059246,
  30.601922972669094,
  33.746182898667383,
  36.889987409236811,
  40.033444053350675,
  43.176628965448822,
  46.319597561173912],
 [3.0542369282271403,
  6.7061331941584591,
  9.9694678230875958,
  13.170370856016123,
  16.347522318321783,
  19.512912782488205,
  22.671581772477426,
  25.826037141785263,
  28.977672772993679,
  32.127327020443474,
  35.275535050674691,
  38.422654817555906,
  41.568934936074314,
  44.714553532819734,
  47.859641607992093],
 [4.2011889412105285,
  8.0152365983759522,
  11.345924310743006,
  14.585848286167028,
  17.78874786606647,
  20.9724769365377,
  24.144897432909265,
  27.310057930204349,
  30.470268806290424,
  33.626949182796679,
  36.781020675464386,
  39.933108623659488,
  43.083652662375079,
  46.232971081836478,
  49.381300092370349],
 [5.3175531260839944,
  9.2823962852416123,
  12.681908442638891,
  15.964107037731551,
  19.196028800048905,
  22.401032267689004,
  25.589759681386733,
  28.767836217666503,
  31.938539340972783,
  35.103916677346764,
  38.265316987088158,
  41.423666498500732,
  44.579623137359257,
  47.733667523865744,
  50.886159153182682],
 [6.4156163757002403,
  10.519860873772308,
  13.9871886301403,
  17.312842487884625,
  20.575514521386888,
  23.803581476593863,
  27.01030789777772,
  30.20284907898166,
  33.385443901010121,
  36.560777686880356,
  39.730640230067416,
  42.896273163494417,
  46.058566273567043,
  49.218174614666636,
  52.375591529563596],
 [7.501266144684147,
  11.734935953042708,
  15.268181461097873,
  18.637443009666202,
  21.931715017802236,
  25.183925599499626,
  28.409776362510085,
  31.617875716105035,
  34.81339298429743,
  37.999640897715301,
  41.178849474321413,
  44.352579199070217,
  47.521956905768113,
  50.687817781723741,
  53.85079463676896],
 [8.5778364897140741,
  12.932386237089576,
  16.529365884366944,
  19.941853366527342,
  23.268052926457571,
  26.545032061823576,
  29.790748583196614,
  33.015178641375142,
  36.224380548787162,
  39.422274578939259,
  42.611522172286684,
  45.793999658055002,
  48.971070951900596,
  52.143752969301988,
  55.312820330403446],
 [9.6474216519972168,
  14.115518907894618,
  17.774012366915256,
  21.229062622853124,
  24.587197486317681,
  27.889269427955092,
  31.155326556188325,
  34.39662855427218,
  37.620078044197086,
  40.830178681822041,
  44.030010337966153,
  47.221758471887113,
  50.407020967034367,
  53.586995435398319,
  56.762598475105272],
 [10.711433970699945,
  15.28673766733295,
  19.004593537946053,
  22.501398726777283,
  25.891277276839136,
  29.218563499936081,
  32.505247352375523,
  35.763792928808799,
  39.001902811514218,
  42.224638430753279,
  45.435483097475542,
  48.636922645305525,
  51.830783925834728,
  55.01844255063594,
  58.200955824859509],
 [11.770876674955582,
  16.447852748486498,
  20.223031412681701,
  23.760715860327448,
  27.182021527190532,
  30.534504754007074,
  33.841965775135715,
  37.118000423665604,
  40.371068905333891,
  43.606764901379516,
  46.828959446564562,
  50.040428970943456,
  53.243223214220535,
  56.438892058982552,
  59.628631306921512],
 [12.826491228033465,
  17.600266557468326,
  21.430854238060294,
  25.008518704644261,
  28.460857279654847,
  31.838424458616998,
  35.166714427392629,
  38.460388720328256,
  41.728625562624312,
  44.977526250903469,
  48.211333836373288,
  51.433105171422278,
  54.645106240447105,
  57.849056857839799,
  61.046288512821078],
 [13.878843069697276,
  18.745090916814406,
  22.629300302835503,
  26.246047773946584,
  29.72897816891134,
  33.131449953571661,
  36.480548302231658,
  39.791940718940855,
  43.075486800191012,
  46.337772104541405,
  49.583396417633095,
  52.815686826850452,
  56.037118687012179,
  59.249577075517968,
  62.454525995970462],
 [14.928374492964716,
  19.88322436109951,
  23.81938909003628,
  27.474339750968247,
  30.987394331665278,
  34.414545662167183,
  37.784378506209499,
  41.113512376883377,
  44.412454519229281,
  47.688252845993366,
  50.945849245830813,
  54.188831071035124,
  57.419876154678179,
  60.641030026538746,
  63.853885828967512],
 [15.975438807484321,
  21.015404934568315,
  25.001971500138194,
  28.694271223110755,
  32.236969407878118,
  35.688544091185301,
  39.078998185245057,
  42.425854432866141,
  45.740236776624833,
  49.029635055514276,
  52.299319390331728,
  55.553127779547459,
  58.793933759028134,
  62.02393848337554,
  65.244860767043859]]

yn_small_zeros = \
[[0.89357696627916752,
  3.9576784193148579,
  7.0860510603017727,
  10.222345043496417,
  13.361097473872763,
  16.500922441528091,
  19.64130970088794,
  22.782028047291559,
  25.922957653180923,
  29.064030252728398,
  32.205204116493281,
  35.346452305214321,
  38.487756653081537,
  41.629104466213808,
  44.770486607221993],
 [2.197141326031017,
  5.4296810407941351,
  8.5960058683311689,
  11.749154830839881,
  14.897442128336725,
  18.043402276727856,
  21.188068934142213,
  24.331942571356912,
  27.475294980449224,
  30.618286491641115,
  33.761017796109326,
  36.90355531614295,
  40.045944640266876,
  43.188218097393211,
  46.330399250701687],
 [3.3842417671495935,
  6.7938075132682675,
  10.023477979360038,
  13.209986710206416,
  16.378966558947457,
  19.539039990286384,
  22.69395593890929,
  25.845613720902269,
  28.995080395650151,
  32.143002257627551,
  35.289793869635804,
  38.435733485446343,
  41.581014867297885,
  44.725777117640461,
  47.870122696676504],
 [4.5270246611496439,
  8.0975537628604907,
  11.396466739595867,
  14.623077742393873,
  17.81845523294552,
  20.997284754187761,
  24.166235758581828,
  27.328799850405162,
  30.486989604098659,
  33.642049384702463,
  36.794791029185579,
  39.945767226378749,
  43.095367507846703,
  46.2438744334407,
  49.391498015725107],
 [5.6451478942208959,
  9.3616206152445429,
  12.730144474090465,
  15.999627085382479,
  19.22442895931681,
  22.424810599698521,
  25.610267054939328,
  28.785893657666548,
  31.954686680031668,
  35.118529525584828,
  38.278668089521758,
  41.435960629910073,
  44.591018225353424,
  47.744288086361052,
  50.896105199722123],
 [6.7471838248710219,
  10.597176726782031,
  14.033804104911233,
  17.347086393228382,
  20.602899017175335,
  23.826536030287532,
  27.030134937138834,
  30.220335654231385,
  33.401105611047908,
  36.574972486670962,
  39.743627733020277,
  42.908248189569535,
  46.069679073215439,
  49.228543693445843,
  52.385312123112282],
 [7.8377378223268716,
  11.811037107609447,
  15.313615118517857,
  18.670704965906724,
  21.958290897126571,
  25.206207715021249,
  28.429037095235496,
  31.634879502950644,
  34.828638524084437,
  38.013473399691765,
  41.19151880917741,
  44.364272633271975,
  47.53281875312084,
  50.697961822183806,
  53.860312300118388],
 [8.919605734873789,
  13.007711435388313,
  16.573915129085334,
  19.974342312352426,
  23.293972585596648,
  26.5667563757203,
  29.809531451608321,
  33.031769327150685,
  36.239265816598239,
  39.435790312675323,
  42.623910919472727,
  45.805442883111651,
  48.981708325514764,
  52.153694518185572,
  55.322154420959698],
 [9.9946283820824834,
  14.190361295800141,
  17.817887841179873,
  21.26093227125945,
  24.612576377421522,
  27.910524883974868,
  31.173701563441602,
  34.412862242025045,
  37.634648706110989,
  40.843415321050884,
  44.04214994542435,
  47.232978012841169,
  50.417456447370186,
  53.596753874948731,
  56.771765754432457],
 [11.064090256031013,
  15.361301343575925,
  19.047949646361388,
  22.532765416313869,
  25.91620496332662,
  29.2394205079349,
  32.523270869465881,
  35.779715464475261,
  39.016196664616095,
  42.237627509803703,
  45.4474001519274,
  48.647941127433196,
  51.841036928216499,
  55.028034667184916,
  58.209970905250097],
 [12.128927704415439,
  16.522284394784426,
  20.265984501212254,
  23.791669719454272,
  27.206568881574774,
  30.555020011020762,
  33.859683872746356,
  37.133649760307504,
  40.385117593813002,
  43.619533085646856,
  46.840676630553575,
  50.051265851897857,
  53.253310556711732,
  56.448332488918971,
  59.637507005589829],
 [13.189846995683845,
  17.674674253171487,
  21.473493977824902,
  25.03913093040942,
  28.485081336558058,
  31.858644293774859,
  35.184165245422787,
  38.475796636190897,
  41.742455848758449,
  44.990096293791186,
  48.222870660068338,
  51.443777308699826,
  54.655042589416311,
  57.858358441436511,
  61.055036135780528],
 [14.247395665073945,
  18.819555894710682,
  22.671697117872794,
  26.276375544903892,
  29.752925495549038,
  33.151412708998983,
  36.497763772987645,
  39.807134090704376,
  43.089121522203808,
  46.350163579538652,
  49.594769786270069,
  52.82620892320143,
  56.046916910756961,
  59.258751140598783,
  62.463155567737854],
 [15.30200785858925,
  19.957808654258601,
  23.861599172945054,
  27.504429642227545,
  31.011103429019229,
  34.434283425782942,
  37.801385632318459,
  41.128514139788358,
  44.425913324440663,
  47.700482714581842,
  50.957073905278458,
  54.199216028087261,
  57.429547607017405,
  60.65008661807661,
  63.862406280068586],
 [16.354034360047551,
  21.090156519983806,
  25.044040298785627,
  28.724161640881914,
  32.260472459522644,
  35.708083982611664,
  39.095820003878235,
  42.440684315990936,
  45.75353669045622,
  49.041718113283529,
  52.310408280968073,
  55.56338698149062,
  58.803488508906895,
  62.032886550960831,
  65.253280088312461]]

ynp_small_zeros = \
[[2.197141326031017,
  5.4296810407941351,
  8.5960058683311689,
  11.749154830839881,
  14.897442128336725,
  18.043402276727856,
  21.188068934142213,
  24.331942571356912,
  27.475294980449224,
  30.618286491641115,
  33.761017796109326,
  36.90355531614295,
  40.045944640266876,
  43.188218097393211,
  46.330399250701687],
 [3.6830228565851777,
  6.9414999536541757,
  10.123404655436613,
  13.285758156782854,
  16.440058007293282,
  19.590241756629495,
  22.738034717396327,
  25.884314618788867,
  29.029575819372535,
  32.174118233366201,
  35.318134458192094,
  38.461753870997549,
  41.605066618873108,
  44.74813744908079,
  47.891014070791065],
 [5.0025829314460639,
  8.3507247014130795,
  11.574195465217647,
  14.760909306207676,
  17.931285939466855,
  21.092894504412739,
  24.249231678519058,
  27.402145837145258,
  30.552708880564553,
  33.70158627151572,
  36.849213419846257,
  39.995887376143356,
  43.141817835750686,
  46.287157097544201,
  49.432018469138281],
 [6.2536332084598136,
  9.6987879841487711,
  12.972409052292216,
  16.19044719506921,
  19.38238844973613,
  22.559791857764261,
  25.728213194724094,
  28.890678419054777,
  32.048984005266337,
  35.204266606440635,
  38.357281675961019,
  41.508551443818436,
  44.658448731963676,
  47.807246956681162,
  50.95515126455207],
 [7.4649217367571329,
  11.005169149809189,
  14.3317235192331,
  17.58443601710272,
  20.801062338411128,
  23.997004122902644,
  27.179886689853435,
  30.353960608554323,
  33.521797098666792,
  36.685048382072301,
  39.844826969405863,
  43.001910515625288,
  46.15685955107263,
  49.310088614282257,
  52.461911043685864],
 [8.6495562436971983,
  12.280868725807848,
  15.660799304540377,
  18.949739756016503,
  22.192841809428241,
  25.409072788867674,
  28.608039283077593,
  31.795195353138159,
  34.973890634255288,
  38.14630522169358,
  41.313923188794905,
  44.477791768537617,
  47.638672065035628,
  50.797131066967842,
  53.953600129601663],
 [9.8147970120105779,
  13.532811875789828,
  16.965526446046053,
  20.291285512443867,
  23.56186260680065,
  26.799499736027237,
  30.015665481543419,
  33.216968050039509,
  36.407516858984748,
  39.590015243560459,
  42.766320595957378,
  45.937754257017323,
  49.105283450953203,
  52.269633324547373,
  55.431358715604255],
 [10.965152105242974,
  14.765687379508912,
  18.250123150217555,
  21.612750053384621,
  24.911310600813573,
  28.171051927637585,
  31.40518108895689,
  34.621401012564177,
  37.824552065973114,
  41.017847386464902,
  44.203512240871601,
  47.3831408366063,
  50.557907466622796,
  53.728697478957026,
  56.896191727313342],
 [12.103641941939539,
  15.982840905145284,
  19.517731005559611,
  22.916962141504605,
  26.243700855690533,
  29.525960140695407,
  32.778568197561124,
  36.010261572392516,
  39.226578757802172,
  42.43122493258747,
  45.626783824134354,
  48.815117837929515,
  51.997606404328863,
  55.175294723956816,
  58.348990221754937],
 [13.232403808592215,
  17.186756572616758,
  20.770762917490496,
  24.206152448722253,
  27.561059462697153,
  30.866053571250639,
  34.137476603379774,
  37.385039772270268,
  40.614946085165892,
  43.831373184731238,
  47.037251786726299,
  50.234705848765229,
  53.425316228549359,
  56.610286079882087,
  59.790548623216652],
 [14.35301374369987,
  18.379337301642568,
  22.011118775283494,
  25.482116178696707,
  28.865046588695164,
  32.192853922166294,
  35.483296655830277,
  38.747005493021857,
  41.990815194320955,
  45.219355876831731,
  48.435892856078888,
  51.642803925173029,
  54.84186659475857,
  58.034439083840155,
  61.221578745109862],
 [15.466672066554263,
  19.562077985759503,
  23.240325531101082,
  26.746322986645901,
  30.157042415639891,
  33.507642948240263,
  36.817212798512775,
  40.097251300178642,
  43.355193847719752,
  46.596103410173672,
  49.823567279972794,
  53.040208868780832,
  56.247996968470062,
  59.448441365714251,
  62.642721301357187],
 [16.574317035530872,
  20.73617763753932,
  24.459631728238804,
  27.999993668839644,
  31.438208790267783,
  34.811512070805535,
  38.140243708611251,
  41.436725143893739,
  44.708963264433333,
  47.962435051891027,
  51.201037321915983,
  54.427630745992975,
  57.644369734615238,
  60.852911791989989,
  64.054555435720397],
 [17.676697936439624,
  21.9026148697762,
  25.670073356263225,
  29.244155124266438,
  32.709534477396028,
  36.105399554497548,
  39.453272918267025,
  42.766255701958017,
  46.052899215578358,
  49.319076602061401,
  52.568982147952547,
  55.805705507386287,
  59.031580956740466,
  62.248409689597653,
  65.457606670836759],
 [18.774423978290318,
  23.06220035979272,
  26.872520985976736,
  30.479680663499762,
  33.971869047372436,
  37.390118854896324,
  40.757072537673599,
  44.086572292170345,
  47.387688809191869,
  50.66667461073936,
  53.928009929563275,
  57.175005343085052,
  60.410169281219877,
  63.635442539153021,
  66.85235358587768]]

@pytest.mark.slow
def test_bessel_zeros_extra():
    mp.dps = 15
    for v in range(V):
        for m in range(1,M+1):
            print(v, m, "of", V, M)
            # Twice to test cache (if used)
            assert besseljzero(v,m).ae(jn_small_zeros[v][m-1])
            assert besseljzero(v,m).ae(jn_small_zeros[v][m-1])
            assert besseljzero(v,m,1).ae(jnp_small_zeros[v][m-1])
            assert besseljzero(v,m,1).ae(jnp_small_zeros[v][m-1])
            assert besselyzero(v,m).ae(yn_small_zeros[v][m-1])
            assert besselyzero(v,m).ae(yn_small_zeros[v][m-1])
            assert besselyzero(v,m,1).ae(ynp_small_zeros[v][m-1])
            assert besselyzero(v,m,1).ae(ynp_small_zeros[v][m-1])